"""Shared tabular attachment transport for chat and source accounting.

This reader does not choose operations, norms, resources, or conversions.
"""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

logger = logging.getLogger(__name__)


def _compact_cell(value: Any, *, max_len: int | None = 240) -> str:
    text = " ".join(str(value if value is not None else "").replace("\u00a0", " ").split())
    if max_len is not None and len(text) > max_len:
        return text[: max_len - 1].rstrip() + "…"
    return text


def _transport_cell(value: Any, *, max_len: int | None) -> str:
    return _compact_cell(value, max_len=max_len).replace('\\', '\\\\').replace('|', '\\|')


def format_tabular_attachment_context(path: Path, original_name: str, *, max_chars: int | None, max_cell_chars: int | None = 240) -> tuple[str, bool] | None:
    """XLSX/CSV attachment → model-readable row context with sheet/row provenance.

    Это не ВОР-генератор и не сметный шаблон: только более честный транспорт таблицы
    в модель, чтобы она сама сделала спецификация → ВОР → сметный ход.
    """
    suffix = path.suffix.lower()
    parts: list[str] = [
        f"Файл: {original_name}",
        "Тип данных: таблица/спецификация. Строки ниже сохранены с номерами листов/строк.",
        "Формат ячеек: les.tabular.cells.v2",
        "",
    ]
    truncated = False
    current_chars = len("\n".join(parts))

    def _append(line: str) -> bool:
        nonlocal truncated, current_chars
        extra_chars = (1 if current_chars else 0) + len(line)
        if max_chars is not None and current_chars + extra_chars > max_chars:
            truncated = True
            return False
        parts.append(line)
        current_chars += extra_chars
        return True

    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as err:  # noqa: BLE001
            logger.warning("[ATTACH] xlsx structured read failed for %s: %s", original_name, err)
            return None
        try:
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                if not _append(f"## Лист: {sheet}"):
                    break
                nonempty = 0
                for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
                    vals = [_transport_cell(v, max_len=max_cell_chars) for v in row]
                    while vals and not vals[-1]:
                        vals.pop()
                    if not any(vals):
                        continue
                    nonempty += 1
                    if not _append(f"{sheet}!R{ri}: " + " | ".join(vals)):
                        break
                _append(f"Итого непустых строк на листе «{sheet}»: {nonempty}")
                if truncated:
                    break
        finally:
            wb.close()
    elif suffix == ".csv":
        import csv
        for enc in ("utf-8-sig", "utf-8", "cp1251", "latin-1"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
        sample = text.splitlines()[0] if text.splitlines() else ""
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
        reader = csv.reader(text.splitlines(), delimiter=delimiter)
        nonempty = 0
        _append("## CSV")
        for ri, row in enumerate(reader, 1):
            vals = [_transport_cell(v, max_len=max_cell_chars) for v in row]
            while vals and not vals[-1]:
                vals.pop()
            if not any(vals):
                continue
            nonempty += 1
            if not _append(f"CSV!R{ri}: " + " | ".join(vals)):
                break
        _append(f"Итого непустых строк CSV: {nonempty}")
    else:
        return None

    if truncated and parts[-1] != "[Табличный контекст усечён по лимиту; для полной сметы нужен файл как датасет или более узкий лист.]":
        parts.append("[Табличный контекст усечён по лимиту; для полной сметы нужен файл как датасет или более узкий лист.]")
    text = "\n".join(parts).strip()
    return (text, truncated) if text else None


def load_full_tabular_attachment_context(attachment_id: str) -> tuple[str, dict[str, Any]] | None:
    """Resolve verified server-owned table bytes before estimate portion planning.

    Upload text is only a bounded preview. This reads the existing original file;
    the caller must still fit whole rows through the normal context governor.
    Other attachment formats retain their existing text workflow.
    """
    from proxy.services.chat_attachment_service import resolve_read_attachment

    source, metadata = resolve_read_attachment(attachment_id)
    if source.suffix.lower() not in {'.xlsx', '.xlsm', '.csv'}:
        return None
    result = format_tabular_attachment_context(
        source, str(metadata.get('original_name') or source.name),
        max_chars=None, max_cell_chars=None,
    )
    if not result:
        raise ValueError('Cannot read the verified tabular attachment')
    return result[0], metadata
