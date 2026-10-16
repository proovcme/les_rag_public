"""Exact tabular source catalog for chat attachments.

The chat-facing text budget may still truncate a preview. This catalog is built
from the original server-owned file so ContextGovernor can pack real rows,
record omissions, and check model source_row/quantity/unit against the source.
It does not choose norms or rewrite model cells.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, Mapping, Sequence

from proxy.smeta_core.unit_contract import norm_measure

CATALOG_SCHEMA = "les.attachment-source-catalog.v2"
VOR_MODEL_AUDIT_SCHEMA = "les.vor-model-decisions.v1"
VOR_MODEL_AUDIT_HEADERS = ("Строка ВОР", "Фрагмент", "Авторские поля (JSON)")
_PDF_TEXT_BLOCK_CHARS = 4096
_PDF_REFERENCE_CONTEXT_CHARS = 8192
_WORK_UNITS = {
    "м", "м2", "м3", "шт", "кг", "т", "л", "компл", "комплект",
    "машч", "челч", "100м", "100м2", "100м3", "100шт", "1000шт",
}
_SUBTOTAL_MARKERS = ("итого", "всего", "subtotal", "total")
_SECTION_MARKERS = ("раздел", "глава", "этап", "section")
_HEADER_MARKERS = ("наименование", "единица", "ед изм", "количество", "объем", "объём")
# Header labels that, when exactly one column matches, unambiguously name
# the quantity column. Content-only; no sheet name or row number involved.
_QUANTITY_HEADER_TOKENS = (
    "кол-во", "кол", "количество", "объем", "объём", "qty", "quantity", "число",
)
_SUBJECT_HEADER_TOKENS = ("наименован", "описан", "назван", "name", "description")


def _compact_cell(value: Any, *, max_len: int | None = None) -> str:
    # This is an exact source catalog, not a preview. ContextGovernor owns the
    # downstream budget and records omissions; losing a cell here is invisible.
    text = " ".join(str(value if value is not None else "").replace("\u00a0", " ").split())
    if max_len is not None and len(text) > max_len:
        return text[: max_len - 1].rstrip() + "…"
    return text


def _parse_number(value: Any) -> int | float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)
    text = str(value or "").replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        parsed = float(text)
    except ValueError:
        return None
    return int(parsed) if parsed.is_integer() else parsed


def _numbers_in_values(values: Sequence[str]) -> tuple[float, ...]:
    found: list[float] = []
    for item in values:
        parsed = _parse_number(item)
        if parsed is not None:
            found.append(float(parsed))
    return tuple(found)


def quantities_close(left: float, right: float) -> bool:
    delta = abs(float(left) - float(right))
    return delta <= max(0.001, abs(float(right)) * 1e-6) or delta < 0.01


def normalize_unit(value: Any) -> str:
    text = (
        str(value or "")
        .casefold()
        .replace("м²", "м2")
        .replace("m²", "м2")
        .replace("м^2", "м2")
        .replace("м³", "м3")
        .replace("m³", "m3")
        .replace("м^3", "м3")
        .replace("шт.", "шт")
        .replace(".", "")
        .replace(" ", "")
    )
    text = text.strip()
    return {
        "пм": "м", "мп": "м", "погм": "м",
        "штука": "шт", "штуки": "шт", "штук": "шт",
        "комплект": "компл",
    }.get(text, text)


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _stable_id(*parts: object, prefix: str) -> str:
    raw = "\x1f".join(str(part) for part in parts).encode("utf-8")
    return f"{prefix}_{hashlib.sha256(raw).hexdigest()[:24]}"


def _column_letter(column: int) -> str:
    value = max(1, int(column))
    letters = ""
    while value:
        value, remainder = divmod(value - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _markdown_table_cells(line: str) -> tuple[str, ...]:
    """Split Markdown cells without treating an escaped pipe as a column."""
    cells: list[str] = []
    current: list[str] = []
    index = 0
    inner = line[1:-1]
    while index < len(inner):
        character = inner[index]
        if character == "\\" and index + 1 < len(inner) and inner[index + 1] in {"\\", "|"}:
            current.append(inner[index + 1])
            index += 2
            continue
        if character == "|":
            cells.append(_compact_cell("".join(current)))
            current = []
        else:
            current.append(character)
        index += 1
    cells.append(_compact_cell("".join(current)))
    return tuple(cells)


def _markdown_table_row_spans(text: str) -> list[tuple[tuple[str, ...], int, int]]:
    """Read converter-produced Markdown tables without interpreting their domain content."""

    rows: list[tuple[tuple[str, ...], int, int]] = []
    offset = 0
    for raw_line in str(text or "").splitlines(keepends=True):
        start, offset = offset, offset + len(raw_line)
        line = raw_line.strip()
        if not (line.startswith("|") and line.endswith("|")):
            continue
        values = _markdown_table_cells(line)
        if not values or all(not item for item in values):
            continue
        if all(item and set(item.replace(":", "")) <= {"-"} for item in values):
            continue
        rows.append((values, start, offset))
    return rows


def _markdown_table_rows(text: str) -> list[tuple[str, ...]]:
    return [values for values, _start, _end in _markdown_table_row_spans(text)]


def _is_numeric_column_legend(values: Sequence[str]) -> bool:
    parsed = [_parse_number(item) for item in values]
    return bool(parsed) and all(
        isinstance(value, int) and value == index
        for index, value in enumerate(parsed, 1)
    )


@dataclass(frozen=True)
class CatalogCell:
    coordinate: str
    column: int
    raw_value: Any
    display_value: str
    data_type: str
    number_format: str
    formula: str = ""
    cached_formula_value: Any = None
    merged_range: str = ""
    hidden_column: bool = False


@dataclass(frozen=True)
class CatalogRow:
    locator: str
    sheet: str
    excel_row: int
    values: tuple[str, ...]
    line: str
    numbers: tuple[float, ...]
    row_id: str = ""
    sheet_id: str = ""
    cells: tuple[CatalogCell, ...] = ()
    row_classification: str = "ambiguous"
    classification_reason: str = "legacy_numeric_heuristic"
    classification_confidence: float = 0.0
    ambiguities: tuple[str, ...] = ()
    section_context: str = ""
    hidden_row: bool = False
    hidden_sheet: bool = False
    header_quantity: float | None = None
    quantity_column_known: bool = False
    # ``(column, governing header label)`` for the cells of this row, kept as data
    # next to the labelled ``line``.  A consumer that must tell a subject-matter
    # column from a service column (totals, money, numbering) needs the label as a
    # field, not re-parsed out of the ``"<value> [<label>]"`` text.
    column_headers: tuple[tuple[int, str], ...] = ()
    # Zero-based, end-exclusive character offsets in the literal converter
    # text (legacy Markdown or explicitly identified native text), not original
    # file byte offsets or inferred professional associations.
    source_start: int | None = None
    source_end: int | None = None
    source_order_basis: str = ""

    def is_work_row(self) -> bool:
        # Ambiguous rows remain in the completeness denominator.  The catalog
        # must not silently erase a possible work row before the model/user can
        # inspect the exact cells.
        return self.row_classification in {"work", "ambiguous"}


def _looks_like_column_header_strip(values: Sequence[str]) -> bool:
    """Shape test for a column-header strip: several labels, none an expression.

    «no number anywhere» alone would not separate a header from a subtotal line,
    because the money cells of a сметная total are formulas («=SUBTOTAL(9,G9:G108)»)
    that resolve to no number either.  A header strip is a row of *names*: more than
    one cell, and no computed cell among them.  No word list is involved, so the
    test holds for any sheet layout rather than for one customer's captions.
    """

    cells = [str(text or "").strip() for text in values if str(text or "").strip()]
    if len(cells) < 2:
        return False
    return not any(cell.startswith("=") for cell in cells)


def _classify_row(
    values: Sequence[str],
    numbers: Sequence[float],
    *,
    excel_row: int,
) -> tuple[str, str, float, tuple[str, ...]]:
    folded = " ".join(values).casefold().replace("ё", "е")
    units = {normalize_unit(value) for value in values}
    has_unit = bool(units.intersection(_WORK_UNITS))
    has_text = any(any(character.isalpha() for character in value) for value in values)
    non_sequence_numbers = [
        number
        for number in numbers
        if not (float(number).is_integer() and int(number) == excel_row)
    ]
    # The column-header strip is recognised FIRST, before the subtotal rule, and by
    # its shape.  The order was the defect: a сметная sheet names its money columns
    # «Материал (итого), руб.» / «Стоимость (итого)…», so the header row itself
    # contains a subtotal word, and claiming it as a subtotal left the customer's
    # workbook without a single header row.  Nothing then bound: ``header_quantity``
    # stayed None and the quantity column stayed unknown, so every work row fell back
    # to "any number in the row" -- on this ВОР the item number («1.2», «1.3»), which
    # turned every correct model quantity into ``source_quantity_mismatch`` and made
    # the service-column filter unable to tell a money column from a subject one.
    # A genuine total line is not lost by looking first: it either carries a resolved
    # number or its cells are formulas, which the shape test rejects.
    if (
        not numbers
        and _looks_like_column_header_strip(values)
        and any(marker in folded for marker in _HEADER_MARKERS)
    ):
        return "header", "column_header_marker", 0.95, ()
    if any(marker in folded for marker in _SUBTOTAL_MARKERS):
        return "subtotal", "subtotal_marker", 0.98, ()
    if any(folded.startswith(marker) for marker in _SECTION_MARKERS):
        return "section", "section_marker", 0.95, ()
    if not numbers:
        return ("note", "text_without_quantity", 0.85, ()) if has_text else (
            "empty", "no_significant_values", 1.0, ()
        )
    if has_unit and non_sequence_numbers and has_text:
        return "work", "description_unit_quantity", 0.98, ()
    if len(values) == 1 and not has_text:
        return "note", "isolated_numeric_cell", 0.75, ()
    if has_text and non_sequence_numbers:
        return (
            "ambiguous",
            "quantity_without_unambiguous_unit",
            0.45,
            ("quantity_column_unconfirmed",),
        )
    return (
        "ambiguous",
        "numeric_structure_unresolved",
        0.25,
        ("work_row_unconfirmed",),
    )


def _normalize_header(text: str) -> str:
    return " ".join(str(text or "").casefold().replace("ё", "е").split())


def _header_labels_from_row(row: CatalogRow) -> dict[int, str]:
    """Column -> header label taken from a single header-classified row."""

    headers: dict[int, str] = {}
    for cell in row.cells:
        label = " ".join(str(cell.display_value or "").split())
        if label and cell.column not in headers:
            headers[cell.column] = label
    return headers


def _has_text_header(headers: Mapping[int, str]) -> bool:
    """A governing header must carry a letter; a numeric column-legend may not."""

    return any(any(ch.isalpha() for ch in label) for label in headers.values())


def _quantity_column(headers: Mapping[int, str]) -> int | None:
    """The quantity column when exactly one header names it, else None."""

    matches = [
        column
        for column, label in headers.items()
        if any(token in _normalize_header(label) for token in _QUANTITY_HEADER_TOKENS)
    ]
    return matches[0] if len(matches) == 1 else None


def _cell_number(cell: CatalogCell) -> float | None:
    """Numeric value of a cell, preferring a formula's cached computed result.

    A formula cell keeps its displayed expression as evidence, but its confirmed
    value is the cached result the workbook stored; a formula with no cached
    value yields None so the code never invents a quantity from an unresolved
    expression.
    """

    direct = _parse_number(cell.display_value)
    if direct is not None:
        return float(direct)
    if cell.formula and cell.cached_formula_value is not None:
        cached = _parse_number(str(cell.cached_formula_value))
        if cached is not None:
            return float(cached)
    return None


def _bind_column_headers(rows: Sequence[CatalogRow]) -> list[CatalogRow]:
    """Bind each row to its governing source header without dropping evidence.

    A sheet can stack several tables that repeat their headers or reorder
    columns, so the binding tracks the most recent header row ABOVE each data
    row, not the sheet's first header. For each data row the model-facing line
    is labelled per column and the quantity is bound to the governing header's
    quantity column:

      * usable value there -> a single header-bound candidate;
      * a known quantity column whose value is empty / erroneous / an unresolved
        formula -> an explicit "no confirmed quantity" (empty candidates), never
        a fallback to another number in the row;
      * no proven quantity column (missing or ambiguous header) -> the row stays
        ambiguous with the legacy candidate set.
    """

    headers: dict[int, str] = {}
    quantity_col: int | None = None
    bound: list[CatalogRow] = []
    for row in rows:
        if row.row_classification == "header":
            candidate = _header_labels_from_row(row)
            if _has_text_header(candidate):
                headers = candidate
                quantity_col = _quantity_column(headers)
            bound.append(row)
            continue
        if not headers:
            bound.append(row)
            continue
        parts: list[str] = []
        header_quantity: float | None = None
        labels: dict[int, str] = {}
        for cell in row.cells:
            label = headers.get(cell.column, "")
            number = _cell_number(cell)
            parts.append(
                f"{cell.display_value} [{label}]"
                if label and number is not None
                else cell.display_value
            )
            if cell.column not in labels:
                labels[cell.column] = label
            if quantity_col is not None and cell.column == quantity_col:
                header_quantity = number
        if not parts:
            bound.append(row)
            continue
        # A blank or formula-only quantity must not turn a populated subject
        # cell into a discarded note. Proven column labels establish only its
        # source scope; the model still owns its role and any missing quantity.
        subject_columns = {column for column, label in headers.items()
                           if any(token in _normalize_header(label) for token in _SUBJECT_HEADER_TOKENS)}
        subject_without_literal_quantity = (
            row.classification_reason == "text_without_quantity"
            and row.row_classification == "note"
            and quantity_col is not None
            and any(cell.column in subject_columns and cell.display_value for cell in row.cells)
        )
        bound.append(
            replace(
                row,
                line=f"{row.locator}: " + " | ".join(parts),
                header_quantity=header_quantity,
                quantity_column_known=quantity_col is not None,
                # The labels stay as data next to the labelled ``line`` so a
                # consumer can tell *which* header governs a cell instead of
                # re-parsing prose brackets (see CatalogRow.column_headers).
                column_headers=tuple(sorted(labels.items())),
                **({"row_classification": "ambiguous",
                    "classification_reason": "subject_row_requires_model_role",
                    "classification_confidence": 0.0,
                    "ambiguities": (*row.ambiguities, "row_role_unconfirmed")}
                   if subject_without_literal_quantity else {}),
            )
        )
    return bound


def quantity_candidates(row: CatalogRow) -> tuple[float, ...]:
    """Source quantity candidates, bound to the governing header column.

    When a quantity column is proven from the header, only that column's value
    is a candidate; an empty / erroneous / unresolved-formula cell yields no
    candidate, so a neighbor number (e.g. a mass) can never be mistaken for the
    quantity and no other value is invented. Only when no quantity column is
    proven does the row fall back to every non-sequence number, keeping an
    ambiguous document ambiguous instead of guessing a value for it.
    """

    header_quantity = getattr(row, "header_quantity", None)
    if header_quantity is not None:
        return (float(header_quantity),)
    if getattr(row, "quantity_column_known", False):
        return ()
    others = tuple(
        number
        for number in row.numbers
        if not (float(number).is_integer() and int(number) == row.excel_row)
    )
    return others or row.numbers


@dataclass(frozen=True)
class AttachmentSourceCatalog:
    original_name: str
    sheets: tuple[tuple[str, int], ...]
    rows: tuple[CatalogRow, ...]
    source_file_sha256: str = ""
    source_file_id: str = ""
    schema: str = CATALOG_SCHEMA
    extraction_complete: bool = True
    extraction_warnings: tuple[str, ...] = ()

    def manifest_text(self) -> str:
        lines = [
            f"Файл: {self.original_name}",
            f"Каталог: {self.schema}",
            f"Идентификатор источника: {self.source_file_id or 'legacy'}",
            "Тип данных: таблица/спецификация. Ниже — полный манифест листов исходника.",
            f"source_extraction_complete: {str(self.extraction_complete).lower()}",
            *[f"Предупреждение извлечения: {warning}" for warning in self.extraction_warnings],
            "",
            "## Манифест листов",
        ]
        lines.extend(f"- {name}: непустых строк {count}" for name, count in self.sheets)
        return "\n".join(lines)

    def interleaved_row_lines(self) -> list[str]:
        if self.rows and all(row.source_start is not None for row in self.rows):
            return [row.line for row in sorted(self.rows, key=lambda row: row.source_start)]
        by_sheet: dict[str, list[CatalogRow]] = {}
        for row in self.rows:
            by_sheet.setdefault(row.sheet, []).append(row)
        order = [name for name, _count in self.sheets if name in by_sheet]
        max_len = max((len(by_sheet[name]) for name in order), default=0)
        lines: list[str] = []
        for index in range(max_len):
            for name in order:
                rows = by_sheet[name]
                if index < len(rows):
                    lines.append(rows[index].line)
        return lines

    def sheet_lines(self, sheet: str) -> list[str]:
        wanted = str(sheet or "").strip()
        return [row.line for row in self.rows if row.sheet == wanted]

    def lookup(self, source_row: Any) -> CatalogRow | None:
        if source_row is None:
            return None
        if isinstance(source_row, bool):
            return None
        if isinstance(source_row, int):
            matches = [row for row in self.rows if row.excel_row == source_row]
            return matches[0] if len(matches) == 1 else None
        label = str(source_row).strip()
        if not label:
            return None
        exact = next((row for row in self.rows if row.locator == label), None)
        if exact is not None:
            return exact
        folded = label.casefold()
        matches = [row for row in self.rows if row.locator.casefold() == folded]
        if matches:
            return matches[0] if len(matches) == 1 else None
        # The manifest exposes the source file identity alongside sheet locators.
        # A model may use that exact identity as a prefix. It identifies a row
        # only when the row number is unique across the original source; it is
        # never permission to strip an arbitrary prefix or choose a sheet.
        file_id, separator, coordinate = label.partition("!")
        if (separator and self.source_file_id and file_id == self.source_file_id
                and re.fullmatch(r"R[1-9][0-9]*", coordinate)):
            # Compare actual catalog row numbers instead of converting an
            # untrusted, potentially arbitrarily long numeral to an integer.
            matches = [row for row in self.rows if str(row.excel_row) == coordinate[1:]]
            return matches[0] if len(matches) == 1 else None
        return None

    def work_locators(self) -> list[str]:
        return [row.locator for row in self.rows if row.is_work_row()]


def catalog_work_sheet_names(catalog: AttachmentSourceCatalog) -> list[str]:
    """Sheet order from the source file, only sheets that have work rows."""

    names: list[str] = []
    for name, _count in catalog.sheets:
        if any(row.sheet == name and row.is_work_row() for row in catalog.rows):
            names.append(name)
    return names


def catalog_sheet_work_locators(
    catalog: AttachmentSourceCatalog,
    sheet: str,
) -> list[str]:
    wanted = str(sheet or "").strip()
    return [
        row.locator
        for row in catalog.rows
        if row.sheet == wanted and row.is_work_row()
    ]


def catalog_sheet_pack_lines(
    catalog: AttachmentSourceCatalog,
    sheet: str,
    remaining_work: Sequence[str] | None = None,
) -> list[str]:
    """Sheet headers/descriptors, current rows and literal notes-only sheet context.

    A numbered description with an empty, proven quantity column can govern
    several subsequent variants. Keep its literal text as source context even
    though its numeric item number makes its classification ambiguous. Neither
    its classification nor the coverage denominator is changed here.
    """

    wanted = str(sheet or "").strip()
    work_sheets = set(catalog_work_sheet_names(catalog))
    pdf_references, pdf_reference_omissions = _pdf_text_reference_rows(catalog, wanted)
    visible_pdf_references = {row.locator for row in pdf_references}
    remaining = (
        {str(item).strip() for item in remaining_work}
        if remaining_work is not None
        else None
    )
    lines: list[str] = []
    ordered_rows = (sorted(catalog.rows, key=lambda row: row.source_start)
                    if catalog.rows and all(row.source_start is not None for row in catalog.rows)
                    else catalog.rows)
    for row in ordered_rows:
        if row.sheet != wanted:
            if row.classification_reason == "application_authored_decision_audit":
                # The exported audit names the exact main-sheet row. Include
                # all fragments for offered rows, without duplicating unrelated
                # decisions in every packet or making audit rows new work.
                requested = remaining if remaining is not None else {
                    item.locator for item in catalog.rows if item.sheet == wanted and item.is_work_row()
                }
                if row.excel_row <= 3 or (row.values and row.values[0] in requested):
                    lines.append(row.line)
                continue
            # A separate conditions/notes sheet otherwise never reaches a work
            # packet. Its literal rows are references, not new requested work.
            if (row.sheet not in work_sheets
                or row.classification_reason == "docx_text_requires_model_role"
                or row.locator in visible_pdf_references):
                lines.append(row.line)
            continue
        descriptor_context = row.quantity_column_known and row.header_quantity is None
        if remaining is not None and row.is_work_row() and row.locator not in remaining and not descriptor_context:
            continue
        lines.append(row.line)
    if pdf_reference_omissions:
        first, last = pdf_reference_omissions[0], pdf_reference_omissions[-1]
        lines.append(
            "model_context_partial: нетабличный PDF-текст не показан в этом пакете целиком. "
            f"Точный каталог: {len(pdf_reference_omissions)} блоков, "
            f"{first.locator} — {last.locator}. Полное содержание сохранено в отдельных "
            "пакетах PDF_TEXT; диапазон адресов не является прочитанным содержанием. "
            "Не считай отсутствующими или просмотренными непоказанные условия исходника."
        )
    return lines


def _pdf_text_reference_rows(
    catalog: AttachmentSourceCatalog, sheet: str,
) -> tuple[list[CatalogRow], list[CatalogRow]]:
    """Show small literal prose context; otherwise expose exact unshown scope."""
    rows = [row for row in catalog.rows if row.sheet != sheet
            and row.classification_reason == "pdf_text_requires_model_role"]
    size = sum(len(row.line) + 1 for row in rows)
    return (rows, []) if size <= _PDF_REFERENCE_CONTEXT_CHARS else ([], rows)


def catalog_sheets_as_lines(
    catalog: AttachmentSourceCatalog,
) -> list[tuple[str, list[str]]]:
    by_sheet: dict[str, list[str]] = {}
    for row in catalog.rows:
        by_sheet.setdefault(row.sheet, []).append(row.line)
    return [(name, by_sheet.get(name, [])) for name, _count in catalog.sheets]


def catalog_sheet_packet_facts(
    catalog: AttachmentSourceCatalog, sheet: str, remaining_work: Sequence[str] | None = None,
) -> dict[str, Any]:
    """Mechanical packet scope and literal header-bound quantity cells."""
    lines = set(catalog_sheet_pack_lines(catalog, sheet, remaining_work))
    visible = [row for row in catalog.rows if row.line in lines]
    _pdf_references, omitted_pdf_references = _pdf_text_reference_rows(catalog, str(sheet or "").strip())
    requested = (set(remaining_work) if remaining_work is not None else {
        row.locator for row in visible if row.sheet == sheet and row.is_work_row()
    })
    cells = []
    for row in visible:
        if row.locator not in requested and not (row.is_work_row() and row.quantity_column_known):
            continue
        item = {"source_row": row.locator, "quantity_column_known": row.quantity_column_known}
        if row.quantity_column_known:
            item.update(quantity=row.header_quantity, unit=header_bound_unit(row))
        cells.append(item)
    facts = {
        "source_extraction_complete": catalog.extraction_complete,
        "source_extraction_warnings": list(catalog.extraction_warnings),
        "requested_source_rows": [row.locator for row in visible if row.locator in requested],
        "reference_source_rows": [row.locator for row in visible if row.locator not in requested],
        "reference_only_source_rows": [row.locator for row in visible if row.sheet != sheet],
        "source_quantity_cells": cells,
        "source_scope_instruction": (
            "requested_source_rows — исходные строки текущего пакета. "
            "reference_source_rows — показанные контекстные строки вне текущей области обработки. "
            "source_quantity_cells — значения исходной колонки количества, не расчётные объёмы. "
            "null означает пустое или не определённое значение, а не количество 0 или 1. "
            "reference_only_source_rows — дословный контекст других листов без заданий: "
            "он не расширяет requested_source_rows и не требует новых операций. "
            "Код не изменял роли, нормы или операции."
        ),
    }
    if any(row.source_start is not None or row.classification_reason == "pdf_text_requires_model_role"
           for row in catalog.rows):
        facts.update(
            source_context_complete=not omitted_pdf_references,
            reference_omitted_source_count=len(omitted_pdf_references),
            reference_omitted_source_locator_range=({
                "sheet": "PDF_TEXT", "first": omitted_pdf_references[0].locator,
                "last": omitted_pdf_references[-1].locator,
                "count": len(omitted_pdf_references), "content_status": "not_shown",
            } if omitted_pdf_references else None),
        )
    if visible and all(row.source_start is not None for row in visible):
        facts["source_order_basis"] = catalog_source_order_basis(visible)
        facts["source_order"] = [
            {"source_row": row.locator, "source_start": row.source_start, "source_end": row.source_end}
            for row in sorted(visible, key=lambda row: row.source_start)
        ]
    return facts


def _docx_source_sheets(
    path: Path, source_sha256: str,
) -> tuple[list[tuple[str, str, bool, list[CatalogRow]]], tuple[str, ...]]:
    """Read literal DOCX text/tables; unread embedded content remains a source gap."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from xml.etree import ElementTree

    document = Document(path)
    sheets: list[tuple[str, str, bool, list[CatalogRow]]] = []
    warnings: list[str] = []
    table_index = 0
    seen_tables: set[Any] = set()

    def empty_notes_scaffolding(blob: bytes) -> bool:
        """Exclude only known-empty Word separator parts, never unread notes."""
        try:
            root = ElementTree.fromstring(blob)
        except (ElementTree.ParseError, ValueError, LookupError):
            return False
        word_ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        note_tag = {
            word_ns + "footnotes": word_ns + "footnote",
            word_ns + "endnotes": word_ns + "endnote",
        }.get(root.tag)
        if note_tag is None:
            return False
        for note in root:
            if note.tag != note_tag or note.get(word_ns + "type") not in {
                "separator", "continuationSeparator",
            }:
                return False
        # A conservative scaffold vocabulary: unsupported nodes remain a gap.
        # In particular no table, image, object, field, textbox or math content
        # is discarded just because its note carries a separator type.
        scaffold_tags = {
            root.tag, note_tag,
            *(word_ns + name for name in (
                "p", "r", "pPr", "rPr", "t", "separator", "continuationSeparator",
                "pStyle", "rStyle", "spacing", "ind", "jc", "contextualSpacing",
                "keepNext", "keepLines", "widowControl", "rFonts", "sz", "szCs",
                "b", "bCs", "i", "iCs", "lang", "color", "vanish", "vertAlign",
                "tabs", "tab", "br",
            )),
        }
        return all(
            node.tag in scaffold_tags
            and not (node.text or "").strip()
            and not (node.tail or "").strip()
            for node in root.iter()
        )

    def rows_for(
        sheet: str, physical_rows: Sequence[tuple[int, Sequence[str]]], *, text_blocks: bool = False,
    ) -> list[CatalogRow]:
        sheet_id = _stable_id(source_sha256, sheet, prefix="sheet")
        rows: list[CatalogRow] = []
        section_context = ""
        for row_number, raw in physical_rows:
            values = tuple(_compact_cell(value) for value in raw)
            if not any(values):
                continue
            numbers = _numbers_in_values(values)
            row_class, reason, confidence, ambiguities = _classify_row(values, numbers, excel_row=row_number)
            if text_blocks:
                # Paragraphs can contain scope, conditions or unquantified work.
                # Keep them in the account-all denominator; only the model may
                # decide whether each is a header, a note or an operation.
                row_class, reason, confidence, ambiguities = (
                    "ambiguous", "docx_text_requires_model_role", 0.0, ("row_role_unconfirmed",),
                )
            if row_class == "section":
                section_context = " | ".join(values)
            locator = f"{sheet}!R{row_number}"
            rows.append(CatalogRow(
                locator=locator, sheet=sheet, excel_row=row_number, values=values,
                line=f"{locator}: " + " | ".join(values), numbers=numbers,
                row_id=_stable_id(source_sha256, sheet_id, row_number, prefix="row"),
                sheet_id=sheet_id,
                cells=tuple(CatalogCell(
                    coordinate=f"{_column_letter(column)}{row_number}", column=column,
                    raw_value=value, display_value=values[column - 1], data_type="s", number_format="General",
                ) for column, value in enumerate(raw, 1)),
                row_classification=row_class, classification_reason=reason,
                classification_confidence=confidence, ambiguities=ambiguities,
                section_context=section_context,
            ))
        return rows

    def add_table(table: Table) -> None:
        nonlocal table_index
        if table._tbl in seen_tables:
            return
        seen_tables.add(table._tbl)
        table_index += 1
        sheet = f"DOCX_T{table_index}"
        rows = rows_for(sheet, [(index, [""] * row.grid_cols_before + [cell.text for cell in row.cells]
                                + [""] * row.grid_cols_after)
                                for index, row in enumerate(table.rows, 1)])
        sheets.append((sheet, _stable_id(source_sha256, sheet, prefix="sheet"), False, rows))
        for row in table.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    add_table(nested)

    def add_container(container: Any, prefix: str) -> None:
        unread = {
            "drawing", "pict", "object", "altChunk", "txbxContent", "sdt", "customXml",
            "ins", "del", "moveFrom", "moveTo", "oMath", "oMathPara",
        }
        for node in container._element.iter():
            local = str(node.tag).rsplit("}", 1)[-1]
            if local in unread:
                warnings.append(f"docx_unread_content:{prefix}:{local}")
        paragraphs: list[tuple[int, Sequence[str]]] = []
        for index, block in enumerate(container.iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                paragraphs.append((index, (block.text,)))
            elif isinstance(block, Table):
                add_table(block)
        if paragraphs:
            sheet = f"{prefix}_TEXT"
            sheets.append((sheet, _stable_id(source_sha256, sheet, prefix="sheet"), False,
                           rows_for(sheet, paragraphs, text_blocks=True)))

    add_container(document, "DOCX")
    seen_parts: set[str] = {str(document.part.partname)}
    for index, section in enumerate(document.sections, 1):
        for variant, label in enumerate(("header", "first_page_header", "even_page_header", "footer", "first_page_footer", "even_page_footer"), 1):
            container = getattr(section, label)
            if container.is_linked_to_previous:
                continue
            part_name = str(container.part.partname)
            if part_name not in seen_parts:
                seen_parts.add(part_name)
                kind = "H" if "header" in label else "F"
                add_container(container, f"DOCX_{kind}{index}_{variant}")
    for part in document.part.package.parts:
        name = str(part.partname)
        if name.startswith("/word/") and name.endswith(".xml") and name not in seen_parts:
            if any(marker in name for marker in ("footnotes", "endnotes")):
                if not empty_notes_scaffolding(part.blob):
                    warnings.append(f"docx_unread_part:{name.rsplit('/', 1)[-1]}")
    return sheets, tuple(dict.fromkeys(warnings))


def _pdf_text_block_spans(markdown: str) -> list[tuple[str, int, int]]:
    """Keep non-table converter text in exact bounded blocks, without role inference."""
    runs: list[tuple[str, int]] = []
    current: list[str] = []
    source_offset = run_start = 0
    for line in str(markdown or "").splitlines(keepends=True):
        stripped = line.strip()
        if stripped.startswith("|") and stripped.endswith("|"):
            if current:
                runs.append(("".join(current), run_start))
                current = []
        else:
            if not current:
                run_start = source_offset
            current.append(line)
        source_offset += len(line)
    if current:
        runs.append(("".join(current), run_start))
    blocks: list[tuple[str, int, int]] = []
    for run, run_start in runs:
        offset = 0
        while offset < len(run):
            end = min(len(run), offset + _PDF_TEXT_BLOCK_CHARS)
            if end < len(run):
                newline = run.rfind("\n", offset, end)
                if newline >= offset:
                    end = newline + 1
            block = run[offset:end]
            start = offset
            offset = end
            if block.strip():
                blocks.append((block, run_start + start, run_start + end))
    return blocks


def _pdf_text_blocks(markdown: str) -> list[str]:
    return [text for text, _start, _end in _pdf_text_block_spans(markdown)]


def catalog_source_order_basis(rows: Sequence[CatalogRow]) -> str:
    """Identify literal-text offset provenance, keeping the legacy default."""
    bases = {row.source_order_basis or "converter_markdown_char_offsets"
             for row in rows if row.source_start is not None}
    if len(bases) > 1:
        raise ValueError("SOURCE_TEXT_OFFSET_BASES_CONFLICT")
    return next(iter(bases), "application_catalog_order")


def _rtf_source_sheets(
    path: Path, source_sha256: str,
) -> tuple[list[tuple[str, str, bool, list[CatalogRow]]], tuple[str, ...], bool]:
    """Expose native literal cells only where boundaries were independently checked."""
    from backend.rtf_parser import RtfParserError, extract_rtf

    result = extract_rtf(path)
    if result.source_file_sha256 != source_sha256:
        raise RtfParserError("RTF_SOURCE_CHANGED")
    verified = result.table_boundaries_verified
    sheet = "RTF_T1" if verified else "RTF_TEXT"
    sheet_id = _stable_id(source_sha256, sheet, prefix="sheet")
    physical = [(row.number, row.cells, row.source_text_start, row.source_text_end)
                for row in result.rows] if verified else []
    if not verified:
        start = 0
        for number, literal in enumerate(result.text.splitlines(keepends=True), 1):
            physical.append((number, (literal,), start, start + len(literal)))
            start += len(literal)
    rows: list[CatalogRow] = []
    section_context = ""
    for number, values, start, end in physical:
        numbers = _numbers_in_values(values) if verified else ()
        if verified:
            row_class, reason, confidence, ambiguities = _classify_row(values, numbers, excel_row=number)
        else:
            row_class, reason, confidence, ambiguities = (
                "ambiguous", "rtf_text_requires_model_role", 0.0, ("row_role_unconfirmed",),
            )
        if row_class == "section":
            section_context = " | ".join(values)
        locator = f"{sheet}!R{number}"
        rows.append(CatalogRow(
            locator=locator, sheet=sheet, excel_row=number, values=values,
            line=f"{locator}: " + " | ".join(values), numbers=numbers,
            row_id=_stable_id(source_sha256, sheet_id, number, prefix="row"), sheet_id=sheet_id,
            cells=tuple(CatalogCell(coordinate=f"{_column_letter(column)}{number}", column=column,
                                   raw_value=value, display_value=value, data_type="s", number_format="General")
                        for column, value in enumerate(values, 1)),
            row_classification=row_class, classification_reason=reason, classification_confidence=confidence,
            ambiguities=ambiguities, section_context=section_context, source_start=start, source_end=end,
            source_order_basis="rtf_native_text_char_offsets",
        ))
    return [(sheet, sheet_id, False, rows)], result.warnings, result.extraction_complete


def _pdf_text_rows(markdown: str, source_sha256: str) -> tuple[list[CatalogRow], tuple[str, ...]]:
    sheet = "PDF_TEXT"
    sheet_id = _stable_id(source_sha256, sheet, prefix="sheet")
    rows: list[CatalogRow] = []
    for index, (text, source_start, source_end) in enumerate(_pdf_text_block_spans(markdown), 1):
        locator = f"{sheet}!R{index}"
        rows.append(CatalogRow(
            locator=locator, sheet=sheet, excel_row=index, values=(text,),
            line=f"{locator}: {text}", numbers=(),
            row_id=_stable_id(source_sha256, sheet_id, index, prefix="row"), sheet_id=sheet_id,
            cells=(CatalogCell(coordinate=f"A{index}", column=1, raw_value=text,
                               display_value=text, data_type="s", number_format="General"),),
            row_classification="ambiguous", classification_reason="pdf_text_requires_model_role",
            classification_confidence=0.0, ambiguities=("row_role_unconfirmed",),
            source_start=source_start, source_end=source_end,
        ))
    warnings: list[str] = []
    if re.search(r"!\[[^\]]*\](?:\(|\[)|<\s*(?:svg|img)\b", markdown, re.IGNORECASE):
        warnings.append("pdf_unread_visual_content")
    if re.search(r"\[WARN\]|\[Ошибка распознавания страницы:", markdown, re.IGNORECASE):
        warnings.append("pdf_converter_incomplete_content")
    return rows, tuple(warnings)


def _is_vor_model_audit(worksheet: Any, workbook: Any) -> bool:
    """Recognize only the complete application-owned reference table shape."""
    if (worksheet.cell(1, 1).value != VOR_MODEL_AUDIT_SCHEMA
            or tuple(worksheet.cell(3, column).value for column in range(1, 4)) != VOR_MODEL_AUDIT_HEADERS
            or "ВОР" not in workbook.sheetnames):
        return False
    groups: dict[str, tuple[int, list[str]]] = {}
    for values in worksheet.iter_rows(min_row=4, values_only=True):
        if not any(value is not None for value in values):
            continue
        if len(values) < 3 or any(value is not None for value in values[3:]):
            return False
        target, part, literal = values[:3]
        match = re.fullmatch(r"ВОР!R([1-9][0-9]{0,6})", str(target or ""))
        part_match = re.fullmatch(r"Фрагмент ([1-9][0-9]{0,6})/([1-9][0-9]{0,6})", str(part or ""))
        if (match is None or not 4 <= int(match.group(1)) <= workbook["ВОР"].max_row
                or part_match is None
                or not isinstance(literal, str) or not literal.startswith("JSON: ")):
            return False
        number, total = (int(value) for value in part_match.groups())
        declared_total, pieces = groups.setdefault(str(target), (total, []))
        if total != declared_total or number != len(pieces) + 1 or number > total:
            return False
        pieces.append(literal[len("JSON: "):])

    def unique_object(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ValueError("duplicate audit JSON field")
            result[key] = value
        return result

    def invalid_constant(_constant):
        raise ValueError("invalid audit JSON constant")

    for total, pieces in groups.values():
        if len(pieces) != total:
            return False
        try:
            decision = json.loads("".join(pieces), object_pairs_hook=unique_object, parse_constant=invalid_constant)
        except (ValueError, RecursionError):
            return False
        if not isinstance(decision, dict):
            return False
    return bool(groups)


def load_tabular_source_catalog(
    path: str | Path,
    *,
    original_name: str = "",
) -> AttachmentSourceCatalog | None:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    name = original_name or file_path.name
    source_sha256 = _sha256_file(file_path)
    source_file_id = _stable_id(source_sha256, prefix="src")
    sheets: list[tuple[str, str, bool, list[CatalogRow]]] = []
    extraction_warnings: tuple[str, ...] = ()
    extraction_complete: bool | None = None

    if suffix in {".xlsx", ".xlsm"}:
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(file_path, data_only=False, read_only=False)
            cached_workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=False)
        except Exception:  # noqa: BLE001 - preview/catalog failure stays visible upstream
            return None
        try:
            for sheet_index, sheet in enumerate(workbook.sheetnames):
                worksheet = workbook[sheet]
                cached_sheet = cached_workbook[sheet]
                sheet_id = _stable_id(source_sha256, sheet_index, sheet, prefix="sheet")
                hidden_sheet = worksheet.sheet_state != "visible"
                authored_audit = _is_vor_model_audit(worksheet, workbook)
                if (not authored_audit and worksheet.cell(1, 1).value == VOR_MODEL_AUDIT_SCHEMA
                        and tuple(worksheet.cell(3, column).value for column in range(1, 4)) == VOR_MODEL_AUDIT_HEADERS):
                    extraction_warnings = (*extraction_warnings, f"vor_model_audit_invalid:{sheet}")
                # These source parts are not represented by worksheet cells.
                # Surface the unread scope instead of claiming full extraction;
                # do not classify their content or invent extra work rows.
                for part_name in ("oddHeader", "evenHeader", "firstHeader", "oddFooter", "evenFooter", "firstFooter"):
                    part = getattr(worksheet, part_name)
                    if any(str(getattr(part, position).text or "").strip()
                           for position in ("left", "center", "right")):
                        extraction_warnings = (*extraction_warnings,
                            f"xlsx_unread_content:{sheet}:{part_name}")
                merged_ranges = tuple(worksheet.merged_cells.ranges)
                rows: list[CatalogRow] = []
                section_context = ""
                has_unread_comments = False
                for excel_row, raw_cells in enumerate(worksheet.iter_rows(), 1):
                    cells: list[CatalogCell] = []
                    values: list[str] = []
                    for cell in raw_cells:
                        # A comment can belong to an otherwise empty/hidden cell.
                        has_unread_comments = has_unread_comments or getattr(cell, "comment", None) is not None
                        raw_value = cell.value
                        if raw_value in (None, ""):
                            continue
                        # JSON fragments are literal application transport:
                        # compacting spaces inside a quoted value or at a
                        # fragment boundary would alter the authored decision.
                        display_value = str(raw_value) if authored_audit else _compact_cell(raw_value)
                        values.append(display_value)
                        merged_range = next(
                            (
                                str(cell_range)
                                for cell_range in merged_ranges
                                if (
                                    cell_range.min_row <= cell.row <= cell_range.max_row
                                    and cell_range.min_col <= cell.column <= cell_range.max_col
                                )
                            ),
                            "",
                        )
                        formula = (
                            str(raw_value)
                            if cell.data_type == "f" or str(raw_value).startswith("=")
                            else ""
                        )
                        cached_value = cached_sheet[cell.coordinate].value if formula else None
                        cells.append(
                            CatalogCell(
                                coordinate=cell.coordinate,
                                column=int(cell.column),
                                raw_value=raw_value,
                                display_value=display_value,
                                data_type=str(cell.data_type or ""),
                                number_format=str(cell.number_format or ""),
                                formula=formula,
                                cached_formula_value=cached_value,
                                merged_range=merged_range,
                                hidden_column=bool(
                                    worksheet.column_dimensions[cell.column_letter].hidden
                                ),
                            )
                        )
                    if not values:
                        continue
                    value_tuple = tuple(values)
                    numbers = _numbers_in_values(value_tuple)
                    row_class, reason, confidence, ambiguities = _classify_row(
                        value_tuple,
                        numbers,
                        excel_row=excel_row,
                    )
                    if authored_audit:
                        row_class, reason, confidence, ambiguities = (
                            "note", "application_authored_decision_audit", 1.0, (),
                        )
                    if row_class == "section":
                        section_context = " | ".join(value_tuple)
                    locator = f"{sheet}!R{excel_row}"
                    rows.append(
                        CatalogRow(
                            locator=locator,
                            sheet=sheet,
                            excel_row=excel_row,
                            values=value_tuple,
                            line=f"{locator}: " + " | ".join(value_tuple),
                            numbers=numbers,
                            row_id=_stable_id(source_sha256, sheet_id, excel_row, prefix="row"),
                            sheet_id=sheet_id,
                            cells=tuple(cells),
                            row_classification=row_class,
                            classification_reason=reason,
                            classification_confidence=confidence,
                            ambiguities=ambiguities,
                            section_context=section_context,
                            hidden_row=bool(worksheet.row_dimensions[excel_row].hidden),
                            hidden_sheet=hidden_sheet,
                        )
                    )
                if has_unread_comments:
                    extraction_warnings = (*extraction_warnings, f"xlsx_unread_content:{sheet}:cell_comments")
                sheets.append((sheet, sheet_id, hidden_sheet, rows))
        finally:
            workbook.close()
            cached_workbook.close()
    elif suffix == ".csv":
        from backend.converter import iter_csv_source_records, read_csv_source_text

        text = read_csv_source_text(file_path)
        sheet = "CSV"
        sheet_id = _stable_id(source_sha256, 0, sheet, prefix="sheet")
        rows: list[CatalogRow] = []
        section_context = ""
        for excel_row, raw in iter_csv_source_records(text):
            # CSV fields have physical column identity, including empty fields.
            # Filtering them per row shifts quantities under adjacent headers.
            values = tuple(_compact_cell(item) for item in raw)
            if not any(values):
                continue
            numbers = _numbers_in_values(values)
            row_class, reason, confidence, ambiguities = _classify_row(
                values, numbers, excel_row=excel_row
            )
            if row_class == "section":
                section_context = " | ".join(values)
            locator = f"{sheet}!R{excel_row}"
            cells = tuple(
                CatalogCell(
                    coordinate=f"{_column_letter(column)}{excel_row}",
                    column=column,
                    raw_value=raw[column - 1],
                    display_value=value,
                    data_type="s",
                    number_format="General",
                )
                for column, value in enumerate(values, 1)
            )
            rows.append(
                CatalogRow(
                    locator=locator,
                    sheet=sheet,
                    excel_row=excel_row,
                    values=values,
                    line=f"{locator}: " + " | ".join(values),
                    numbers=numbers,
                    row_id=_stable_id(source_sha256, sheet_id, excel_row, prefix="row"),
                    sheet_id=sheet_id,
                    cells=cells,
                    row_classification=row_class,
                    classification_reason=reason,
                    classification_confidence=confidence,
                    ambiguities=ambiguities,
                    section_context=section_context,
                )
            )
        sheets.append((sheet, sheet_id, False, rows))
    elif suffix == ".rtf":
        sheets, extraction_warnings, extraction_complete = _rtf_source_sheets(file_path, source_sha256)
    elif suffix == ".docx":
        try:
            sheets, extraction_warnings = _docx_source_sheets(file_path, source_sha256)
        except Exception:  # noqa: BLE001 - unavailable source stays an explicit audit gap
            return None
    elif suffix == ".pdf":
        try:
            from backend.converter import convert_to_markdown

            markdown = convert_to_markdown(file_path, apply_char_limit=False)
        except Exception:  # noqa: BLE001 - catalog failure stays visible upstream
            return None
        sheet = "PDF"
        sheet_id = _stable_id(source_sha256, 0, sheet, prefix="sheet")
        rows = []
        section_context = ""
        for pdf_row, (values, source_start, source_end) in enumerate(_markdown_table_row_spans(markdown), 1):
            numbers = _numbers_in_values(values)
            row_class, reason, confidence, ambiguities = _classify_row(
                values, numbers, excel_row=pdf_row
            )
            if _is_numeric_column_legend(values):
                row_class, reason, confidence, ambiguities = (
                    "header", "numeric_column_legend", 0.99, ()
                )
            if row_class == "section":
                section_context = " | ".join(values)
            locator = f"{sheet}!R{pdf_row}"
            cells = tuple(
                CatalogCell(
                    coordinate=f"{_column_letter(column)}{pdf_row}",
                    column=column,
                    raw_value=value,
                    display_value=value,
                    data_type="s",
                    number_format="General",
                )
                for column, value in enumerate(values, 1)
                if value
            )
            rows.append(
                CatalogRow(
                    locator=locator,
                    sheet=sheet,
                    excel_row=pdf_row,
                    values=values,
                    line=f"{locator}: " + " | ".join(values),
                    numbers=numbers,
                    row_id=_stable_id(source_sha256, sheet_id, pdf_row, prefix="row"),
                    sheet_id=sheet_id,
                    cells=cells,
                    row_classification=row_class,
                    classification_reason=reason,
                    classification_confidence=confidence,
                    ambiguities=ambiguities,
                    section_context=section_context,
                    source_start=source_start,
                    source_end=source_end,
                )
            )
        text_rows, extraction_warnings = _pdf_text_rows(str(markdown or ""), source_sha256)
        if rows:
            sheets.append((sheet, sheet_id, False, rows))
        if text_rows:
            sheets.append(("PDF_TEXT", text_rows[0].sheet_id, False, text_rows))
        if not rows and not text_rows:
            return None
    else:
        return None

    catalog_rows: list[CatalogRow] = []
    sheet_counts: list[tuple[str, int]] = []
    for sheet, _sheet_id, _hidden_sheet, rows in sheets:
        rows = _bind_column_headers(rows)
        sheet_counts.append((sheet, len(rows)))
        catalog_rows.extend(rows)
    if not sheet_counts:
        return None
    return AttachmentSourceCatalog(
        original_name=name,
        sheets=tuple(sheet_counts),
        rows=tuple(catalog_rows),
        source_file_sha256=source_sha256,
        source_file_id=source_file_id,
        extraction_complete=(not extraction_warnings if extraction_complete is None else extraction_complete),
        extraction_warnings=extraction_warnings,
    )


def locators_missing_from_text(
    catalog: AttachmentSourceCatalog,
    text: str,
    *,
    work_rows_only: bool = True,
    locators: Sequence[str] | None = None,
) -> list[str]:
    blob = str(text or "")
    wanted = (
        {str(item).strip() for item in locators if str(item).strip()}
        if locators is not None
        else None
    )
    missing: list[str] = []
    for row in catalog.rows:
        if wanted is not None:
            if row.locator not in wanted:
                continue
        elif work_rows_only and not row.is_work_row():
            continue
        if row.locator not in blob:
            missing.append(row.locator)
    return missing


def sheet_coverage_from_missing(
    catalog: AttachmentSourceCatalog,
    missing_locators: Sequence[str],
) -> dict[str, list[str]]:
    missing = set(missing_locators)
    included: list[str] = []
    omitted: list[str] = []
    for name, count in catalog.sheets:
        sheet_rows = [
            row for row in catalog.rows if row.sheet == name and row.is_work_row()
        ]
        if not sheet_rows:
            included.append(name)
            continue
        missed = [row.locator for row in sheet_rows if row.locator in missing]
        if not missed:
            included.append(name)
        elif len(missed) == len(sheet_rows):
            omitted.append(f"{name} ({count} непустых строк)")
        else:
            included.append(f"{name} (частично)")
            omitted.append(f"{name} (не переданы {len(missed)} строк)")
    return {"included": included, "omitted": omitted}


def header_bound_unit(row: CatalogRow) -> str | None:
    """A unique source unit column, including a known but empty unit cell."""
    columns = []
    for column, label in getattr(row, "column_headers", ()):
        marker = "".join(ch for ch in label.casefold() if ch.isalnum())
        if marker.startswith(("едизм", "единицаизмер", "единицыизмер", "unit")) or marker in {
            "ед", "единица", "единицы",
        }:
            columns.append(column)
    if len(columns) != 1:
        return None
    return next((cell.display_value for cell in row.cells if cell.column == columns[0]), "")


def compare_model_row_to_catalog(
    row: Mapping[str, Any],
    catalog: AttachmentSourceCatalog,
) -> list[str]:
    """Check identity/quantity/unit against the source row without changing the model row."""

    source_row = row.get("source_row")
    catalog_row = catalog.lookup(source_row)
    if catalog_row is None:
        return [f"source_row_not_in_catalog:{source_row}"]
    errors: list[str] = ([] if catalog.extraction_complete else [f"source_extraction_incomplete:{source_row}"])
    quantity = row.get("quantity")
    if isinstance(quantity, bool):
        quantity = None
    candidates = quantity_candidates(catalog_row)
    source_unit = header_bound_unit(catalog_row)
    unit = normalize_unit(row.get("unit"))
    equivalent_measure = False
    comparable_quantity = quantity
    if source_unit is not None and unit:
        source_factor, source_base = norm_measure(source_unit)
        model_factor, model_base = norm_measure(row.get("unit"))
        source_base, model_base = normalize_unit(source_base), normalize_unit(model_base)
        physical_factor = (1.0 if source_base and source_base == model_base else {
            ("км", "м"): 1000.0, ("м", "км"): 0.001,
            ("т", "кг"): 1000.0, ("кг", "т"): 0.001,
        }.get((model_base, source_base)))
        equivalent_measure = bool(
            physical_factor is not None
            and math.isfinite(source_factor) and source_factor > 0
            and math.isfinite(model_factor) and model_factor > 0
        )
        if equivalent_measure and isinstance(quantity, (int, float)):
            comparable_quantity = float(quantity) * model_factor / source_factor * physical_factor
    if isinstance(quantity, (int, float)):
        if candidates:
            if not any(quantities_close(float(comparable_quantity), number) for number in candidates):
                errors.append(f"source_quantity_mismatch:{source_row}")
        elif getattr(catalog_row, "quantity_column_known", False):
            errors.append(f"source_quantity_unconfirmed:{source_row}")
    if unit:
        if source_unit is not None:
            unit_matches = equivalent_measure
        else:
            cell_units = {normalize_unit(item) for item in catalog_row.values}
            joined = normalize_unit(" ".join(catalog_row.values))
            unit_matches = unit in cell_units or unit in joined
        if not unit_matches:
            errors.append(f"source_unit_mismatch:{source_row}")
    return errors


def coverage_errors(
    rows: Sequence[Mapping[str, Any]],
    catalog: AttachmentSourceCatalog,
) -> list[str]:
    cited = {
        str(row.get("source_row") or "").strip()
        for row in rows
        if row.get("source_row") not in (None, "")
    }
    cited.update(
        str(int(row.get("source_row")))
        for row in rows
        if isinstance(row.get("source_row"), int)
        and not isinstance(row.get("source_row"), bool)
    )
    # Compare catalog identities without rewriting the authored source_row.
    # All accepted aliases use the same exact, ambiguity-preserving lookup as
    # source quantity/unit validation.
    cited.update(
        found.locator
        for row in rows
        if (found := catalog.lookup(row.get("source_row"))) is not None
    )
    work = catalog.work_locators()
    uncovered = [locator for locator in work if locator not in cited]
    errors = [] if catalog.extraction_complete else ["source_extraction_incomplete"]
    if uncovered:
        errors.append(f"source_coverage_incomplete:{len(work) - len(uncovered)}/{len(work)}")
    return errors


def load_catalog_from_attachment_id(attachment_id: str) -> AttachmentSourceCatalog | None:
    ident = str(attachment_id or "").strip()
    if not ident:
        return None
    try:
        from proxy.services.chat_attachment_service import resolve_read_attachment

        path, metadata = resolve_read_attachment(ident)
    except Exception:  # noqa: BLE001 - missing attachment is a fallback, not a crash
        return None
    return load_tabular_source_catalog(
        path,
        original_name=str(metadata.get("original_name") or path.name),
    )
