"""Что стало с каждой строкой ведомости в итоговой книге.

ЗАЧЕМ. Итог прогона говорил «12 норма в книге из 70» — и не говорил, где остальные
пятьдесят восемь. Отличить «учтено ресурсом работы» от «забыли» по книге нельзя: обе
строки выглядят в ней одинаково, то есть никак. Сверка нужна построчная и по всем
позициям, а не по тем, что дошли (ревью Codex 11.09, шаг 6).

ЧТО ЗДЕСЬ. Для каждой строки ведомости — цепочка до конца: решение модели, его вид и
полнота, попадание в книгу, связанная работа, ресурсы и стоимость. Источники сверяются
между собой, а не пересказывают друг друга: ведомость читается из вложения, решения — из
хранилища прогресса, книга — из самого XLSX.

ЧЕГО ЗДЕСЬ НЕТ. Суждения о правильности выбранной нормы: инструмент отвечает, доведена
ли строка до денег и чем это подтверждено, а не верно ли выбрана расценка. Строка с
ценой и полным решением попадает в «посчитана», даже если норма выбрана неудачно.

    uv run python -m tools.smeta_book_audit --attachment read_82d9bb0ff0a5 \
        --session sks-mat-20260911-1430 --book run.xlsx --json audit.json
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

from backend.runtime_paths import mutable_path


SCHEMA = "les.smeta.book-audit.v1"

# Куда попала строка ведомости.
PRICED = "посчитана в книге"
PARTIAL_PRICE = "в книге с частичной стоимостью"
IN_BOOK_NO_PRICE = "в книге без стоимости"
FORMULA_UNCACHED = "в книге с формулой стоимости; кеш отсутствует"
RESOURCE_OF_WORK = "учёт ресурсом не подтверждён"
RESOURCE_LINKED = "связан с ресурсом работы; стоимость требует сверки"
OWN_MATERIAL = "стоимость самостоятельного материала не подтверждена"
NEEDS_REVIEW = "требуется уточнение"
DECIDED_NOT_IN_BOOK = "позиция книги не подтверждена"
NO_DECISION = "решения нет"
INCOMPLETE = "решение неполное"

_NORM_IN_CELL = re.compile(r"(?:ГЭСН|ТЕР|ФЕР)[а-яА-Я]{0,2}\d{2}-\d{2}-\d{3}-\d{2}")


class BookEvidence(dict):
    def __init__(self):
        super().__init__()
        self.bindings = {}
        self.operation_bindings = {}


def portions_db() -> Path:
    return Path(mutable_path("storage/smeta_portions.db"))


def attachment_path(attachment_id: str) -> Path:
    return Path(mutable_path(f"storage/chat_attachments/{attachment_id}.xlsx"))


def read_book(path: Path) -> dict[str, dict[str, Any]]:
    """Позиции итоговой книги по шифру нормы: цена и номер строки листа.

    Книга читается как книга, а не как пересказ решений: расхождение между ними и есть
    то, что сверка обязана показать.
    """
    import openpyxl

    found = BookEvidence()
    workbook = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    try:
        # Diagnostic sheets contain rejected codes too. They are not positions.
        for sheet in workbook.worksheets[:1]:
            current = None
            for row in sheet.iter_rows(values_only=False):
                cells = ["" if cell.value is None else str(cell.value).strip() for cell in row]
                match = _NORM_IN_CELL.fullmatch(cells[1]) if len(cells) > 6 else None
                material = not cells[1] and isinstance(row[0].value, int) and bool(cells[3])
                if not match and not material:
                    if current is not None and len(cells) > 11 and cells[2] in {
                        "Всего по позиции", "Итого по позиции", "Известная стоимость позиции"
                    }:
                        total = row[11].value
                        current['total'] = float(total) if isinstance(total, (int, float)) else None
                        current['total_formula'] = formulas[sheet.title].cell(row[0].row, 12).data_type == 'f'
                        current['partial'] = cells[2] == 'Известная стоимость позиции'
                    continue
                code = match.group(0) if match else ''
                # Графа 12 формы ЛСР — стоимость всего в текущем уровне цен.
                total = row[11].value if len(row) > 11 else None
                entry = found.setdefault(
                    code,
                    {"norm_code": code, "rows": [], "positions": []},
                )
                entry["rows"].append(row[0].row)
                provenance = dict(line.split(': ', 1) for line in
                                  (row[1].comment.text.splitlines() if row[1].comment else [])
                                  if ': ' in line)
                entry["positions"].append({
                    "title": cells[2], "unit": cells[3], "quantity": row[6].value,
                    "total": float(total) if isinstance(total, (int, float)) else None,
                    "has_comment": row[5].comment is not None, "taken": False,
                    "sheet": sheet.title, "row": row[0].row,
                    "source_row": provenance.get('source_row', ''),
                    "work_id": provenance.get('work_id', ''),
                })
                current = entry['positions'][-1]
        if 'Проверка' in workbook:
            for values in workbook['Проверка'].iter_rows(min_row=2, values_only=True):
                if len(values) > 8 and values[2] == 'resource_bound' and values[8]:
                    binding = {'status': values[2], 'resource': values[8], 'taken': False}
                    found.bindings.setdefault(str(values[0]), []).append(binding)
                    if values[1]:
                        found.operation_bindings[str(values[1])] = binding
    finally:
        workbook.close()
        formulas.close()
    return found


def read_decisions(session_id: str, *, db_path: Path | None = None, source_revision: str = "") -> dict[str, dict[str, Any]]:
    """Решения последней ревизии ведомости этой сессии, по якорю строки."""
    import sqlite3

    target = Path(db_path or portions_db())
    if not target.is_file():
        return {}
    with sqlite3.connect(f"file:{target}?mode=ro", uri=True) as conn:
        conn.row_factory = sqlite3.Row
        revisions = [
            str(row["source_revision"])
            for row in conn.execute(
                "SELECT source_revision, MAX(recorded_at) AS last FROM smeta_portion_done"
                " WHERE session_id=? GROUP BY source_revision ORDER BY last DESC",
                (str(session_id),),
            )
        ]
        if not revisions:
            return {}
        found: dict[str, dict[str, Any]] = {}
        for row in conn.execute(
            "SELECT anchor, decision_json, row_state, attempts FROM smeta_portion_done"
            " WHERE session_id=? AND source_revision=?",
            (str(session_id), source_revision or revisions[0]),
        ):
            try:
                decision = json.loads(row["decision_json"] or "{}")
            except (TypeError, ValueError):
                decision = {}
            found[str(row["anchor"])] = {
                **(decision or {}),
                "row_state": str(row["row_state"] or ""),
                "attempts": int(row["attempts"] or 0),
            }
    return found


def classify(
    row: dict[str, Any],
    decision: dict[str, Any] | None,
    book: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    """Куда попала одна строка ведомости и чем это подтверждено."""
    from proxy.smeta_core.row_decision import (
        MATERIAL,
        RESOURCE,
        UNCLEAR,
        describe_decision,
    )

    if not decision:
        return {"state": NO_DECISION, "detail": "строка не дошла до решения"}

    if isinstance(decision.get('operations'), list):
        operations = [dict(work_id=operation.get('work_id', ''),
                           **classify(row, operation, book))
                      for operation in decision['operations'] if isinstance(operation, dict)]
        if not operations:
            return {'state': NO_DECISION, 'detail': 'операции не заданы', 'operations': []}
        # Every operation must have evidence. Preserve each outcome instead of
        # letting a priced first operation conceal a missing second one.
        unresolved = next((op for op in operations if op['state'] != PRICED), operations[0])
        return dict(state=unresolved['state'], detail='; '.join(op['detail'] for op in operations),
                    operations=operations)

    described = describe_decision(decision)
    kind = described["kind"]
    code = str(decision.get("norm_code") or "").strip()
    entry = book.get(code) if code else None

    if kind == RESOURCE:
        detail = described["included_in"] or "работа не названа"
        stored = getattr(book, 'operation_bindings', {}).get(str(decision.get('work_id') or ''))
        if not stored and not decision.get('work_id'):
            candidates = getattr(book, 'bindings', {}).get(str(row.get('source_row')), [])
            stored = candidates[0] if len(candidates) == 1 else None
        if stored and not stored.get('taken'):
            stored['taken'] = True
            return {'state': RESOURCE_LINKED, 'detail': str(stored['resource'])}
        state = RESOURCE_OF_WORK if described["complete"] else INCOMPLETE
        return {"state": state, "detail": f"заявлено включение в {detail}; связь с ресурсом и его расчётом не подтверждена", "gaps": described["gaps"]}

    if not described["complete"]:
        return {
            "state": INCOMPLETE,
            "detail": "; ".join(described["gaps"]),
            "gaps": described["gaps"],
        }

    # Материал и уточнение нормой не считаются, и искать их в книге как норму значит
    # выдавать законное отсутствие за пропажу: строка 987 прошлого прогона так и
    # попадала в «решение есть, в книге нет», не имея никакой нормы.
    if kind == MATERIAL:
        entry = book.get('')
        if entry is None:
            return {"state": OWN_MATERIAL, "detail": "модель выбрала самостоятельный материал; позиция и стоимость не подтверждены"}
    if kind == UNCLEAR:
        reason = str(decision.get("analogue") or decision.get("coverage") or "").strip()
        return {"state": NEEDS_REVIEW, "detail": reason[:120] or "причина не названа"}

    if entry is not None:
        from proxy.services.rim_lsr_trace_service import _norm_qty_from_visible_row
        candidates = entry.get("positions", [])
        entry = None
        for candidate in candidates:
            if candidate.get("taken"):
                continue
            work_id = str(decision.get('work_id') or '')
            if work_id and candidate.get('work_id') and candidate['work_id'] != work_id:
                continue
            titles = [str(row.get("title") or "").strip()]
            if str(decision.get("source_row")) == str(row.get("source_row")):
                titles.append(str(decision.get("title") or "").strip())
            exact_source = candidate.get('source_row') == str(row.get('source_row'))
            if not exact_source and not any(title and (candidate["title"].casefold() == title.casefold()
                                  or candidate["title"].casefold().endswith(' / ' + title.casefold()))
                       for title in titles):
                continue
            physical = dict(row)
            if work_id:
                physical.update({key: decision[key] for key in ('quantity', 'unit') if key in decision})
            physical.update({key: decision[key] for key in ('conversion_factor', 'conversion_unit', 'conversion_source') if key in decision})
            expected, _ = _norm_qty_from_visible_row(physical, code, {"unit": candidate["unit"]})
            actual = candidate.get("quantity")
            if expected is None or not isinstance(actual, (int, float)):
                continue
            if abs(actual - expected) > max(1e-6, abs(expected) * 1e-9):
                continue
            candidate["taken"] = True
            entry = candidate
            break
    if entry is None:
        return {
            "state": DECIDED_NOT_IN_BOOK,
            "detail": f"позиция {code or '—'} с этим наименованием и количеством в книге не подтверждена",
        }
    if entry.get("total") is None and entry.get("total_formula"):
        return {"state": FORMULA_UNCACHED,
                "detail": f"{code}: формула есть, сохранённое значение отсутствует",
                "partial": bool(entry.get('partial')),
                "book_cell": f"{entry.get('sheet')}!B{entry.get('row')}"}
    if entry.get("total") is not None:
        return {
            "state": PARTIAL_PRICE if entry.get('partial') else PRICED,
            "detail": f"{code}: {entry['total']:.2f} руб.",
            "coefficient_note": bool(entry.get("has_comment")),
        }
    return {"state": IN_BOOK_NO_PRICE, "detail": f"{code}: стоимость не посчитана",
            "book_cell": f"{entry.get('sheet')}!B{entry.get('row')}"}


def audit(
    *,
    attachment_id: str,
    session_id: str,
    book_path: Path | None,
    db_path: Path | None = None,
    revision: str = "",
) -> dict[str, Any]:
    """Построчная сверка ведомости, решений и книги."""
    from proxy.smeta_core.vor_rows import row_anchor, row_anchors, source_revision

    rows = row_anchors(attachment_path(attachment_id))
    decisions = read_decisions(session_id, db_path=db_path, source_revision=revision or source_revision(rows))
    book = read_book(book_path) if book_path and Path(book_path).is_file() else {}

    report: list[dict[str, Any]] = []
    states: Counter = Counter()
    for row in rows:
        anchor = row_anchor(row)
        verdict = classify(row, decisions.get(anchor), book)
        states[verdict["state"]] += 1
        report.append({
            "anchor": anchor,
            "source_row": str(row.get("source_row") or ""),
            "title": str(row.get("title") or "")[:110],
            "unit": str(row.get("unit") or ""),
            **verdict,
        })
    return {
        "schema": SCHEMA,
        "attachment_id": attachment_id,
        "session_id": session_id,
        "source_revision": revision or source_revision(rows),
        "book": str(book_path or ""),
        "rows_total": len(rows),
        "decisions_stored": len(decisions),
        "norms_in_book": len([key for key in book if key]),
        "work_positions_in_book": sum(len(item["positions"]) for key, item in book.items() if key),
        "material_positions_in_book": len(book.get('', {}).get('positions', [])),
        "states": dict(states),
        "rows": report,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--attachment", required=True, help="идентификатор вложения ведомости")
    parser.add_argument("--session", required=True, help="идентификатор сессии прогона")
    parser.add_argument("--book", type=Path, help="итоговая книга XLSX")
    parser.add_argument("--source-revision", default="", help="явная историческая ревизия из хранилища")
    parser.add_argument("--json", type=Path, help="куда записать полный отчёт")
    parser.add_argument("--show", type=int, default=0, help="сколько строк напечатать")
    args = parser.parse_args()

    result = audit(
        attachment_id=args.attachment,
        session_id=args.session,
        book_path=args.book,
        revision=args.source_revision,
    )
    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(
            json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    print(f"строк ведомости:      {result['rows_total']}")
    print(f"решений сохранено:    {result['decisions_stored']}")
    print(f"норм в книге:         {result['norms_in_book']}")
    print()
    for state, count in sorted(result["states"].items(), key=lambda item: -item[1]):
        print(f"  {state:<26} {count}")
    for row in result["rows"][: max(0, args.show)]:
        print(f"\n  {row['source_row']:>5}  {row['title'][:58]}")
        print(f"         {row['state']} — {row['detail']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
