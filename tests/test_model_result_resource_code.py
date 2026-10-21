"""Keep an authored resource identifier through the visible result transport."""

import copy
import json

import pytest

from proxy.services import chat_evidence_application_service as service


RESOURCE_CODE = "20.4.01.01-0001"


def _visible_answer(format_name: str, role: str, *, with_code: bool = True) -> str:
    fields = {
        "source_row": "Source sheet!R7",
        "row_role": role,
        "title": "Model-authored resource",
        "unit": "m",
        "quantity": "12",
        "norm_code": "ГЭСНм08-03-609-04" if role == "work" else "—",
    }
    if with_code:
        fields["resource_code"] = RESOURCE_CODE
    else:
        # Mentioning a catalog-looking code in prose does not select that code.
        fields["title"] = f"Resource described beside {RESOURCE_CODE}"
    if format_name == "markdown":
        return (
            "| " + " | ".join(fields) + " |\n"
            + "| " + " | ".join("---" for _ in fields) + " |\n"
            + "| " + " | ".join(fields.values()) + " |"
        )
    if format_name == "inline":
        return ", ".join(f"{key}={value}" for key, value in fields.items())
    return "Строка 7\n" + "\n".join(f"{key}: {value}" for key, value in fields.items())


@pytest.mark.parametrize("format_name", ["markdown", "inline", "multiline"])
@pytest.mark.parametrize("role", ["material", "equipment", "work"])
def test_explicit_resource_code_survives_visible_result_formats(format_name, role):
    answer = _visible_answer(format_name, role)

    parsed = service.parse_model_rag_result(answer)

    assert parsed is not None
    visible, rows = parsed
    assert visible == answer
    assert len(rows) == 1
    assert rows[0]["resource_code"] == RESOURCE_CODE
    assert rows[0]["row_role"] == role
    assert rows[0]["quantity"] == 12
    assert rows[0]["norm_code"] == ("ГЭСНм08-03-609-04" if role == "work" else "—")


@pytest.mark.parametrize("role", ["material", "equipment", "work"])
def test_resource_code_survives_compact_decision_replay_without_mutation(role):
    row = {
        "source_row": "Source sheet!R7",
        "operation_id": "op_authored",
        "row_role": role,
        "title": "Model-authored resource",
        "unit": "m",
        "quantity": 12,
        "norm_code": "ГЭСНм08-03-609-04" if role == "work" else "",
        "resource_code": "User-Resource-A",
    }
    original = copy.deepcopy(row)

    compact = service.compact_estimator_draft_rows([row])

    assert compact[0]["resource_code"] == "User-Resource-A"
    assert compact[0]["row_role"] == role
    assert compact[0]["norm_code"] == row["norm_code"]
    assert row == original


def test_missing_resource_code_is_not_inferred_from_title_or_defaulted():
    parsed = service.parse_model_rag_result(_visible_answer("markdown", "material", with_code=False))

    assert parsed is not None
    row = parsed[1][0]
    assert "resource_code" not in row
    assert "resource_code" not in service.compact_estimator_draft_rows([row])[0]


def test_authored_resource_code_reaches_chosen_book_and_rendered_money(tmp_path, monkeypatch):
    from types import SimpleNamespace

    import openpyxl

    from proxy.services import attachment_source_catalog_service as catalog_service
    from proxy.services import lsr_workbook_adapter_service as adapter_service
    from proxy.services import rim_lsr_trace_service as rim_service
    from proxy.services import rim_trace_xlsx_service as xlsx_service

    source_path = tmp_path / "source.xlsx"
    source_book = openpyxl.Workbook()
    source_book.active.title = "Source sheet"
    source_book.active.append(["Наименование", "Ед. изм.", "Количество"])
    source_book.active.append(["Model-authored resource", "m", 12])
    source_book.save(source_path)
    source_book.close()
    catalog = catalog_service.load_tabular_source_catalog(source_path)
    monkeypatch.setattr(catalog_service, "load_catalog_from_attachment_id", lambda _: catalog)

    resource_match = json.dumps({
        "source_title": "Model-authored resource", "resource_code": RESOURCE_CODE,
        "decision": "yes", "reason": "The shown resource names this supplied item.",
        "source_refs": ["Q1.H1"],
    })
    answer = (
        "| source_row | row_role | title | unit | quantity | norm_code | resource_code | evidence_refs | supply_resource_match |\n"
        "| --- | --- | --- | --- | --- | --- | --- | --- | --- |\n"
        f"| Source sheet!R2 | material | Model-authored resource | m | 12 | — | {RESOURCE_CODE} | Q1.H1 | {resource_match} |"
    )
    evidence = SimpleNamespace(
        content=f"Resource: {RESOURCE_CODE}\nName: Model-authored resource\nUnit: m\nReference-only price: 777",
        doc_name="Synthetic resource evidence",
        meta={"model_evidence_ref": "Q1.H1", "global_evidence_id": "selected-resource-card",
              "resource_code": RESOURCE_CODE},
    )

    class ChosenBook:
        def __init__(self):
            self.lookups = []

        def lookup(self, code):
            self.lookups.append(code)
            # This fixture has no name matching or resource selection path.
            return {"price_current": 100, "unit": "m"} if code == RESOURCE_CODE else None

    parsed = service.parse_model_rag_result(answer)
    assert parsed is not None
    decisions = service.compact_estimator_draft_rows(parsed[1])
    decisions = service.bind_row_evidence_ids(decisions, [evidence])
    adapter_service._attach_catalog_facts(decisions, "synthetic-source")
    assert decisions[0]["resource_code"] == RESOURCE_CODE
    assert decisions[0]["evidence_refs"] == ["Q1.H1"]
    assert decisions[0]["evidence_bindings"] == {"Q1.H1": "selected-resource-card"}
    assert decisions[0]["quantity"] == 12 and decisions[0]["unit"] == "m"

    chosen_book = ChosenBook()
    lsr = rim_service.build_lsr_trace_from_visible_rows(decisions, pricebook=chosen_book)
    binding = lsr["row_bindings"][0]
    assert binding["status"] == "bound"
    assert binding["position"]["resources"][0]["code"] == RESOURCE_CODE
    assert binding["readiness"]["volume_verified"] is True
    assert binding["readiness"]["prices_provided"] is True
    assert binding["readiness"]["price_from_norm_card"] is False
    assert chosen_book.lookups == [RESOURCE_CODE]
    assert lsr["summary"]["calculated_positions"] == 1
    assert lsr["summary"]["confirmed_amount"] == 1200
    assert lsr["summary"]["preliminary_amount"] == 0
    resource = next(
        row for row in lsr["sections"][0]["positions"][0]["rows"]
        if row["type"] == "resource_material"
    )
    assert resource["source"] == "fgis_current"
    assert resource["columns"]["2"] == RESOURCE_CODE
    assert resource["columns"]["7"] == 12
    assert resource["columns"]["10"] == 100
    assert resource["columns"]["12"] == 1200

    output = xlsx_service.render_lsr_xlsx(lsr, tmp_path / "priced-resource.xlsx")
    rendered = openpyxl.load_workbook(output, data_only=False)
    try:
        sheet = rendered["ЛСР РИМ"]
        resource_row = next(row[1].row for row in sheet.iter_rows() if row[1].value == RESOURCE_CODE)
        assert sheet.cell(resource_row, 7).value == 12
        assert sheet.cell(resource_row, 10).value == 100
        assert sheet.cell(resource_row, 12).value == f"=ROUND(G{resource_row}*J{resource_row}*K{resource_row},2)"
        review_row = next(
            row for row in rendered["Проверка"].iter_rows()
            if row[0].value == "Source sheet!R2"
        )
        assert review_row[4].value == "bound"
        assert "позиция полностью рассчитана=да" in review_row[5].value
        assert "цены=да" in review_row[5].value
    finally:
        rendered.close()
