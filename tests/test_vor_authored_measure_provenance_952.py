"""VOR measure labels describe authored cells without asserting source origin."""

from copy import deepcopy
import hashlib
import json

import openpyxl
import pytest

from proxy.services.attachment_source_catalog_service import (
    header_bound_unit, load_tabular_source_catalog, quantity_candidates,
)
from proxy.services.spec_to_bor_service import _map_headers
from proxy.services.workbook_tool_service import _default_vor_adapter


@pytest.mark.parametrize("quantity", [None, 0, 12.5])
def test_neutral_model_measure_labels_keep_the_existing_reader_binding(tmp_path, quantity):
    path = tmp_path / "ВОР 'Ω'.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "ВОР"
    sheet.append(["Ведомость объёмов работ"])
    sheet.append(["Авторское решение; сверка с исходником отражена отдельно"])
    headers = ["№", "Раздел", "Наименование", "Код/Шифр", "Ед. изм. в решении", "Кол-во в решении",
               "Источник", "Предполагаемая работа", "Статус решения", "Тип строки", "Связанная работа",
               "Замечания / проверка исходника", "ID операции", "Учтено работой", "Полное решение модели",
               "Расчётная единица", "Результат расчёта", "Коэффициент перевода", "Основание расчёта"]
    sheet.append(headers)
    sheet.append([1, "", "Авторская операция", "", "м", quantity, "Источник!R14", "", "", "work",
                  "", "", "op-14", "", "", "м²", 777, 2, "Авторская формула и источник"])
    book.save(path)
    book.close()

    catalog = load_tabular_source_catalog(path)
    row = catalog.lookup("ВОР!R4")
    assert row.quantity_column_known is True
    assert quantity_candidates(row) == (() if quantity is None else (quantity,))
    assert header_bound_unit(row) == "м"
    assert _map_headers(headers)["qty"] == 5
    assert _map_headers(headers)["unit"] == 4


@pytest.mark.asyncio
async def test_derived_layout_keeps_disputed_authored_zero_visible_without_calling_it_source_quantity(tmp_path):
    source = tmp_path / "Источник 'Ω'.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Лист Ω"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Исходная операция", "м", 12])
    sheet.append(["Изделие без исходного количества", "м", None])
    workbook.save(source)
    workbook.close()
    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    decisions = [
        {"source_row": "Лист Ω!R2", "row_role": "work", "title": "Авторская операция",
         "quantity": 12, "unit": "м", "operation_id": "work-1",
         "calc_quantity": 24, "calc_unit": "м²", "conversion_factor": 2,
         "conversion_source": "12 м × 2 м = 24 м²; Лист Ω!R2"},
        {"source_row": "Лист Ω!R3", "row_role": "material", "title": "Авторская поставка",
         "quantity": 0, "unit": "м", "operation_id": "supply-1", "work_row": "work-1",
         "open_gap": "Количество в исходнике не указано"},
    ]
    before = deepcopy(decisions)
    output = tmp_path / "Результат 'Ω'.xlsx"
    result = await _default_vor_adapter(source, {"decisions": decisions}, output, lambda *_: None)
    assert decisions == before
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_sha256
    assert "source_quantity_unconfirmed:Лист Ω!R3" in result["vor_source_errors"]
    source_catalog = load_tabular_source_catalog(source)
    assert source_catalog.lookup("Лист Ω!R3").header_quantity is None

    workbook = openpyxl.load_workbook(output, data_only=True)
    sheet = workbook["ВОР"]
    assert [sheet["E3"].value, sheet["F3"].value] == ["Ед. изм. в решении", "Кол-во в решении"]
    assert [sheet["E4"].value, sheet["F4"].value] == ["м", 12]
    assert [sheet["E5"].value, sheet["F5"].value] == ["м", 0]
    assert [sheet["P4"].value, sheet["Q4"].value] == ["м²", 24]
    assert sheet["I5"].value == "ТРЕБУЕТ РЕШЕНИЯ"
    assert "Количество не подтверждено исходной строкой" in sheet["L5"].value
    assert "ЧЕРНОВИК" in sheet["A2"].value
    assert json.loads(workbook["Решения модели"]["C5"].value.removeprefix("JSON: ")) == decisions[1]
    workbook.close()

    exported_catalog = load_tabular_source_catalog(output)
    assert quantity_candidates(exported_catalog.lookup("ВОР!R4")) == (12,)
    assert quantity_candidates(exported_catalog.lookup("ВОР!R5")) == (0,)
    assert header_bound_unit(exported_catalog.lookup("ВОР!R4")) == "м"
