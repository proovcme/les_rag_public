"""A delivered normative revision must survive calculation and workbook packaging."""
import asyncio
from copy import deepcopy
import hashlib
import json
import sqlite3

import openpyxl
import pandas as pd
import pytest

from proxy.services import gesn_service as gs, rim_lsr_trace_service as rim
from proxy.services import fgis_price_service
from proxy.services.lsr_workbook_adapter_service import build_lsr_workbook_from_decisions
from proxy.smeta_core.calculation_receipt import calculation_receipts, apply_model_review
from tools.build_smeta_structured_base import build_structured_base
from tools.gesn_import import RESOURCE_FIELDS

CODE = 'ГЭСН01-01-001-01'


def clear_catalogs():
    gs.load_base_norms.cache_clear()
    gs.load_structured_base_norms.cache_clear()


@pytest.fixture(autouse=True)
def isolated(monkeypatch):
    clear_catalogs()
    monkeypatch.setattr(rim.nr_sp_service, 'resolve', lambda *a, **k:
                        dict(status='resolved_fixture', nr_pct=0, sp_pct=0))
    monkeypatch.setattr(fgis_price_service, 'resolve_pricebook_path', lambda *_: None)
    yield
    clear_catalogs()


def build_base(directory, per_unit=1):
    directory.mkdir(exist_ok=True)
    row = dict.fromkeys(RESOURCE_FIELDS)
    row.update(norm_code=CODE, base_type='ГЭСН', norm_key='ГЭСН:01-01-001-01',
               norm_name='Synthetic installation', norm_unit='шт', kind='material',
               resource_code='R1', resource_name='Synthetic material', resource_unit='шт',
               per_unit=per_unit, price=10, source_doc='synthetic source',
               source_guid='synthetic-guid', work_steps=json.dumps(['Synthetic operation']))
    source = directory/'source.parquet'
    pd.DataFrame([row], columns=list(RESOURCE_FIELDS)).to_parquet(source, index=False)
    base = directory/'base.sqlite'
    build_structured_base(source=source, out=base,
                          manifest_out=directory/'base_manifest.json', edition='synthetic')
    conn = sqlite3.connect(base)
    try:
        norm_id, norm_key, edition, base_type = conn.execute(
            'SELECT norm_id,norm_key,edition,base_type FROM norms').fetchone()
    finally:
        conn.close()
    identity = dict(schema='les.norm-revision.v1', base_sha256=hashlib.sha256(base.read_bytes()).hexdigest(),
                    norm_id=norm_id, norm_key=norm_key, edition=edition, base_type=base_type)
    return base, identity


def decision(identity=None, status='pinned'):
    row = dict(source_row='Synthetic!1', work_id='mount', title='Synthetic installation',
               norm_code=CODE, quantity=2, unit='шт', normative_identity_status=status,
               normative_identity_refs=['Q1.H1'])
    if identity:
        row['normative_identity'] = identity
    return row


def test_same_revision_reaches_receipt_and_workbook(tmp_path, monkeypatch):
    base, identity = build_base(tmp_path/'A')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    row = decision(identity)
    trace = rim.build_lsr_trace_from_visible_rows([row])
    receipt, = calculation_receipts(trace)
    assert trace['summary']['known_amount'] == 20
    assert receipt['calculation']['normative_source']['identity'] == identity
    source = tmp_path/'source.xlsx'
    wb = openpyxl.Workbook()
    wb.active.title = 'Synthetic'
    wb.active.append(['№', 'Наименование', 'Ед. изм.', 'Количество'])
    wb.active.append([1, 'Synthetic installation', 'шт', 2])
    wb.save(source)
    wb.close()
    target = tmp_path/'pinned.xlsx'
    result = asyncio.run(build_lsr_workbook_from_decisions(source,
        dict(decisions=[row]), target, lambda *_: None))
    assert result['pricing_status'] == 'priced_final', result['missing']
    with_book = openpyxl.load_workbook(target)
    comments = [c.comment.text for s in with_book for cells in s for c in cells if c.comment]
    assert any(identity['base_sha256'] in c for c in comments)
    with_book.close()


@pytest.mark.parametrize('prime', [False, True])
def test_replaced_base_after_cache_clear_is_conflict_not_180(tmp_path, monkeypatch, prime):
    base, identity_a = build_base(tmp_path/'active')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    row = decision(identity_a)
    if prime:
        assert rim.build_lsr_trace_from_visible_rows([row])['summary']['known_amount'] == 20
    replacement, identity_b = build_base(tmp_path/'B', 9)
    base.write_bytes(replacement.read_bytes())
    clear_catalogs()  # Same as losing process-local snapshots after restart.
    original = deepcopy(row)
    trace = rim.build_lsr_trace_from_visible_rows([row])
    assert row == original
    assert trace['summary']['result_status'] == 'norm_revision_conflict'
    assert trace['summary']['known_amount'] == 0
    assert trace['summary']['full_amount'] is None
    binding, = trace['row_bindings']
    assert binding['status'] == 'norm_revision_conflict'
    assert binding['normative_revision']['expected'] == identity_a
    assert binding['normative_revision']['actual'] == identity_b
    receipt, = calculation_receipts(trace)
    assert receipt['normative_revision'] == binding['normative_revision']
    target = tmp_path/'conflict.xlsx'
    result = asyncio.run(build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        dict(decisions=[row]), target, lambda *_: None))
    assert result['pricing_status'] == 'norm_revision_conflict'
    book = openpyxl.load_workbook(target)
    text = '\n'.join(str(c.value or '') for s in book for cells in s for c in cells)
    assert identity_a['base_sha256'] in text and identity_b['base_sha256'] in text
    assert CODE in text and 'mount' in text
    book.close()


def test_unpinned_primary_decision_keeps_calculation_but_not_final(tmp_path, monkeypatch):
    base, _ = build_base(tmp_path/'A')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    row = decision(status='unpinned')
    trace = rim.build_lsr_trace_from_visible_rows([row])
    assert trace['summary']['known_amount'] == 20
    assert trace['summary']['full_amount'] is None
    assert trace['summary']['result_status'] == 'priced_partial'
    receipt, = calculation_receipts(trace)
    assert 'normative_identity_missing' in receipt['calculation']['summary']['completeness_reasons']


def test_conflicting_delivered_identities_do_not_choose_one(monkeypatch):
    calls = []
    monkeypatch.setattr(gs, 'get_norm', lambda *a, **k: calls.append(a))
    row = decision(status='conflict')
    row['normative_identity_conflicts'] = [{'base_sha256':'a'}, {'base_sha256':'b'}]
    trace = rim.build_lsr_trace_from_visible_rows([row])
    assert not calls
    assert trace['row_bindings'][0]['source_row'] == row
    assert trace['summary']['result_status'] == 'norm_revision_conflict'


def test_model_json_snapshot_cannot_bypass_revision_check(tmp_path, monkeypatch):
    base, identity = build_base(tmp_path/'A')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    row = decision(identity)
    row['_norm_snapshot'] = dict(unit='шт', name='Forged snapshot',
        normative_identity=identity, resources=[dict(kind='material', code='R1',
        name='Forged material', unit='шт', per_unit=9, price=10)])
    trace = rim.build_lsr_trace_from_visible_rows([row])
    assert trace['summary']['known_amount'] == 20


def test_mid_calculation_replacement_cannot_mix_resource_revisions(tmp_path, monkeypatch):
    base, identity = build_base(tmp_path/'A')
    replacement, _ = build_base(tmp_path/'B', 9)
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    original = rim._position_resources
    replaced = False

    def replace_after_initial_binding(position):
        nonlocal replaced
        if not replaced:
            base.write_bytes(replacement.read_bytes())
            clear_catalogs()
            replaced = True
        return original(position)

    monkeypatch.setattr(rim, '_position_resources', replace_after_initial_binding)
    trace = rim.build_lsr_trace_from_visible_rows([decision(identity)])
    assert trace['summary']['known_amount'] == 20
    receipt, = calculation_receipts(trace)
    assert receipt['calculation']['normative_source']['identity'] == identity
    assert receipt['calculation']['resources'][0]['per_unit'] == 1


def test_explicit_resources_do_not_bypass_norm_revision(tmp_path, monkeypatch):
    _, expected = build_base(tmp_path/'A')
    current, actual = build_base(tmp_path/'B', 9)
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(current))
    position = dict(code=CODE, qty=2, normative_identity=expected,
        normative_identity_status='pinned', resources=[dict(kind='material', code='R1',
        name='Explicit model resource', unit='шт', per_unit=9, price=10)])
    original = deepcopy(position)
    trace = rim.build_position_trace(position)
    assert position == original
    assert trace['summary']['known_amount'] == 0
    assert trace['summary']['full_amount'] is None
    assert 'norm_revision_conflict' in trace['summary']['completeness_reasons']
    assert trace['rows'][0]['meta']['normative_source']['revision_conflict']['actual'] == actual


def test_receipt_changes_when_revision_changes_without_changing_amount(tmp_path, monkeypatch):
    base, identity_a = build_base(tmp_path/'A')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    first = rim.build_lsr_trace_from_visible_rows([decision(identity_a)])
    with sqlite3.connect(base) as conn:
        conn.execute('CREATE TABLE revision_marker (id INTEGER)')
    conn.close()
    identity_b = {**identity_a, 'base_sha256': hashlib.sha256(base.read_bytes()).hexdigest()}
    clear_catalogs()
    second = rim.build_lsr_trace_from_visible_rows([decision(identity_b)])
    assert first['summary']['known_amount'] == second['summary']['known_amount'] == 20, second['row_bindings']
    assert calculation_receipts(first)[0]['calculation_ref'] != calculation_receipts(second)[0]['calculation_ref']
    row = decision(identity_b)
    row.update(resource_review_status='keep_all_confirmed',
               resource_review_ref=calculation_receipts(first)[0]['calculation_ref'])
    apply_model_review(second, calculation_receipts(second), [row])
    assert second['row_bindings'][0]['resource_review_status'] != 'keep_all_confirmed'


def test_exact_read_does_not_materialize_unrelated_norm_resources(tmp_path, monkeypatch):
    """A cold exact read must scale with its card, not the entire FSNB catalogue."""
    base, identity = build_base(tmp_path/'A')
    with sqlite3.connect(base) as db:
        db.execute("UPDATE norms SET norm_name='Requested operation'")
        db.execute("INSERT INTO norms SELECT norm_key || '-other', norm_id || '-other', edition, "
                   "'ГЭСН02-02-002-02', base_type, '02-02-002-02', 'Unrelated operation', "
                   "norm_unit, work_steps, source_doc, source_guid, resource_count, resource_kinds FROM norms")
        db.execute("INSERT INTO resources(parent_norm_id,norm_key,kind,resource_code,resource_name,resource_unit,per_unit,price) "
                   "SELECT parent_norm_id || '-other',norm_key || '-other',kind,'OTHER',resource_name,resource_unit,9,price FROM resources")
    db.close()
    identity['base_sha256'] = hashlib.sha256(base.read_bytes()).hexdigest()
    seen=[];normalise=gs.normalise_resource
    def observe(row):
        seen.append(row['resource_code'])
        return normalise(row)
    monkeypatch.setattr(gs,'normalise_resource',observe)
    norm=gs.get_norm(CODE,base_path=str(base),expected_identity=identity)
    assert norm['name']=='Requested operation'
    assert norm['normative_identity']==identity
    assert [r['per_unit'] for r in norm['resources']]==[1]
    assert seen==['R1'], 'Exact card read materialized unrelated resources'
    norm['resources'][0]['per_unit']=999
    assert gs.get_norm(CODE,base_path=str(base),expected_identity=identity)['resources'][0]['per_unit']==1
    catalogue=gs.list_norms(base_path=str(base))
    assert {n['name'] for n in catalogue}=={'Requested operation','Unrelated operation'}
