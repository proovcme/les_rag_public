"""The official adapter must account for source rows outside the model's portion."""
import copy

import openpyxl
import pytest

from proxy.services import lsr_workbook_adapter_service as adapter
from tests.test_smeta_resource_receipt import norm, work


def source_book(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'S'
    ws.append(['№', 'Наименование', 'Ед. изм.', 'Количество'])
    ws.append([1, 'Fixture', 'шт', 2])
    ws.append([2, 'Pending cable', 'м', 12.5])
    ws = wb.create_sheet('Other')
    ws.append(['№', 'Наименование', 'Ед. изм.', 'Количество'])
    ws.append([1, 'Other pending item', 'шт', 3])
    wb.save(path)
    wb.close()


@pytest.mark.asyncio
async def test_partial_decisions_keep_every_source_in_download(norm, monkeypatch, tmp_path):
    monkeypatch.setattr(adapter.fgis_price_service, 'resolve_pricebook_path', lambda _: None)
    source = tmp_path/'source.xlsx'
    source_book(source)
    decisions = [work(), work(work_id='second_operation', quantity=1)]
    original = copy.deepcopy(decisions)
    target = tmp_path/'book.xlsx'
    result = await adapter.build_lsr_workbook_from_decisions(
        source, dict(decisions=decisions), target, lambda *a: None)
    assert result['source_rows'] == 3
    assert result['bound_rows'] == 1  # Two operations still belong to one source row.
    assert result['pricing_status'] == 'priced_partial'
    assert decisions == original
    assert len(result['row_calculation']) == 2
    baseline = await adapter.build_lsr_workbook_from_decisions(
        tmp_path/'absent.xlsx', dict(decisions=decisions), tmp_path/'baseline.xlsx', lambda *a: None)
    assert result['row_calculation'] == baseline['row_calculation']
    assert 'source:S!2:not_processed' in result['missing']
    assert 'source:Other!1:not_processed' in result['missing']
    wb = openpyxl.load_workbook(target, data_only=False)
    records = list(wb['Исходные строки'].values)
    assert [r[0] for r in records[1:]] == ['S!1', 'S!2', 'Other!1']
    assert records[1][6] == 'bound'
    assert records[1][7].split('\n') == ['mount', 'second_operation']
    assert records[2][3:7] == ('Pending cable', 'м', 12.5, 'not_processed')
    assert records[3][3:7] == ('Other pending item', 'шт', 3, 'not_processed')
    main_text = '\n'.join(str(c.value or '') for row in wb['ЛСР РИМ'] for c in row)
    assert 'привязано 1 из 3 строк' in main_text
    assert 'РАСЧЁТ НЕ ЗАВЕРШЁН' in main_text
    assert len(list(wb['Проверка'].values)) > 2
    wb.close()


@pytest.mark.asyncio
async def test_unreadable_source_is_explicit_instead_of_full_estimate(norm, monkeypatch, tmp_path):
    monkeypatch.setattr(adapter.fgis_price_service, 'resolve_pricebook_path', lambda _: None)
    result = await adapter.build_lsr_workbook_from_decisions(
        tmp_path/'absent.xlsx', dict(decisions=[work()]), tmp_path/'book.xlsx', lambda *a: None)
    assert 'SOURCE_REGISTRY_UNAVAILABLE' in result['missing']
    assert result['pricing_status'] == 'priced_partial'
    wb = openpyxl.load_workbook(tmp_path/'book.xlsx')
    assert 'не прочитан' in wb['Исходные строки']['A2'].value
    wb.close()


def numbered_source_book(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in ('S', 'Other'):
        ws = wb.create_sheet(sheet)
        ws.append(['№ п/п', 'Наименование работ', 'Ед.изм.', 'Кол-во', 'Стоимость'])
        ws.append(['1', 'Раздел 1. Монтаж'])
        ws.append(['1.1', 'Оборудование'])
        ws.append(['1.2', 'First fixture', 'шт', 3])
        ws.append(['1.10', 'Second fixture', 'шт', '2,50'])
        ws.append(['1.10.1', 'Cable', 'м', 12.75])
        ws.append(['Итого', None, None, None, 200])
    wb.save(path)
    wb.close()


@pytest.mark.parametrize('reader', ['attachment', 'direct'])
def test_hierarchical_position_labels_preserve_identity_and_sections(tmp_path, reader):
    from proxy.services.tabular_attachment_service import format_tabular_attachment_context
    from proxy.smeta_core.vor_rows import row_anchors, rows_from_attachment_text, decision_anchor
    source = tmp_path/'numbered.xlsx'
    numbered_source_book(source)
    if reader == 'attachment':
        context, truncated = format_tabular_attachment_context(
            source, source.name, max_chars=None, max_cell_chars=None)
        assert not truncated
        rows = rows_from_attachment_text(context)
    else:
        rows = [r for sheet in ('S', 'Other') for r in row_anchors(source, sheet=sheet)]
    anchors = [r['anchor'] for r in rows]
    assert anchors == ['S!1.1', 'S!1.2', 'S!1.10', 'S!1.10.1', 'Other!1.1', 'Other!1.2', 'Other!1.10', 'Other!1.10.1']
    assert [r['quantity'] for r in rows] == [None, 3, 2.5, 12.75] * 2
    assert [r['sheet_row'] for r in rows] == [3, 4, 5, 6] * 2
    assert {r['section'] for r in rows} == {'Раздел 1. Монтаж'}
    # A numbered title without unit/quantity has no proven structural role.
    assert rows[0]['title'] == 'Оборудование' and rows[0]['unit'] == ''
    assert all(r['source_row'] == r['cells'][0] for r in rows)
    assert decision_anchor({'source_row': '1.10'}, anchors) == ''
    assert decision_anchor({'source_row': 'S!1.10'}, anchors) == 'S!1.10'


@pytest.mark.asyncio
@pytest.mark.parametrize('label', ['S!1.10', '[Строка ВОР: S!1.10]', '[Строка ВОР: source_row=S!1.10]'])
async def test_hierarchical_positions_reach_official_workbook_registry(norm, monkeypatch, tmp_path, label):
    monkeypatch.setattr(adapter.fgis_price_service, 'resolve_pricebook_path', lambda _: None)
    source = tmp_path/'numbered.xlsx'
    numbered_source_book(source)
    decisions = [work(source_row=label, quantity=2.5)]
    original = copy.deepcopy(decisions)
    from proxy.services.chat_evidence_application_service import portion_structure_errors
    assert portion_structure_errors(decisions, {'source_anchors': ['S!1.10', 'Other!1.10'],
                                               'shown': ['S!1.10']}) == []
    target = tmp_path/'book.xlsx'
    result = await adapter.build_lsr_workbook_from_decisions(
        source, dict(decisions=decisions), target, lambda *a: None)
    assert result['source_rows'] == 8
    assert result['bound_rows'] == 1
    assert decisions == original
    wb = openpyxl.load_workbook(target)
    records = list(wb['Исходные строки'].values)[1:]
    assert [r[0] for r in records] == ['S!1.1', 'S!1.2', 'S!1.10', 'S!1.10.1', 'Other!1.1', 'Other!1.2', 'Other!1.10', 'Other!1.10.1']
    assert records[0][3:7] == ('Оборудование', None, None, 'not_processed')
    assert records[2][4:7] == ('шт', 2.5, 'bound')
    assert records[6][4:7] == ('шт', 2.5, 'not_processed')
    assert 'source:S!1.1:not_processed' in result['missing']
    assert 'source:Other!1.10:not_processed' in result['missing']
    wb.close()


@pytest.mark.parametrize('label', [
    '[Строка ВОР: 1.10]',  # Two sheets: no guessed identity.
    'prefix [Строка ВОР: S!1.10]',
    '[Строка ВОР: S!1.99]',
    '[Строка ВОР: S!1.10] trailing',
    '[Строка ВОР: S!1.10|Other!1.10]',
])
def test_source_marker_does_not_invent_an_unambiguous_anchor(label):
    from proxy.smeta_core.vor_rows import decision_anchor
    row = {'source_row': label, 'work_id': 'authored-id', 'quantity': 2.5}
    original = copy.deepcopy(row)
    assert decision_anchor(row, ['S!1.10', 'Other!1.10']) == ''
    assert row == original
