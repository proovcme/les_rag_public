"""Offline sheet-pass integration: model-owned reads reach the real XLSX executor."""

import asyncio
import hashlib
import json
import re
import time
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

from proxy.routers import chat
from proxy.services import chat_evidence_application_service as service
from proxy.services import fgis_price_service
from proxy.services.canonical_route_service import CanonicalRouteDecision, CanonicalRouteMode
from proxy.services.chat_attachment_service import preserve_read_attachment
from proxy.services.model_connection_contracts import ConnectionLocality, ConnectionRole
from proxy.services.model_execution_preset_service import ModelExecutionPreset
from proxy.services.openai_compatible_transport_service import InferenceResponse


NORM_CODE = "ГЭСН99-01-001-01"
RESOURCE_A = "20.4.01.01-0001"
RESOURCE_B = "20.4.01.01-0002"
BOOK = "sankt-peterburg_3kv2026"


def _native(tool, **arguments):
    identity = hashlib.sha256(json.dumps(arguments, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest()[:16]
    return {"id": tool + ":" + identity, "type": "function",
            "function": {"name": tool, "arguments": json.dumps(arguments, ensure_ascii=False)}}


def _answer(*, refs=("Q1.H1", "Q2.H1")):
    first, second = refs
    return (
        "artifact_kind: lsr\n"
        "Model-authored final rows\n\n"
        "| source_row | row_role | title | unit | quantity | norm_code | resource_code | evidence_refs |\n"
        "| --- | --- | --- | --- | --- | --- | --- | --- |\n"
        f"| S!R2 | material | Cable A | m | 12 | — | {RESOURCE_A} | {first} |\n"
        f"| S!R3 | material | Cable B | m | 3 | — | {RESOURCE_B} | {second} |"
    )


async def _run_sheet(tmp_path, monkeypatch, *, reads, variant="B1_CANDIDATE", unknown_tool=False,
                     reasoning_enabled=False, query_reasoning_enabled=None, query_repair=False,
                     empty_final=False, empty_query=False,
                     query_count=2, group_char_budget=None, partial_final=False, packet_count=1,
                     input_token_limit=100_000, candidate_padding=0, native_read_batch=False,
                     malformed_first_packet=False, malformed_all_first_packet=False,
                     malformed_second_packet=False, duplicate_json_first_packet=False,
                     native_discovery_queries=(), native_discovery_padding=0, json_final=False,
                     exact_resource_padding=0, query_declaration=None):
    # All stores, attachment bytes, norm cards and prices belong to this test.
    from proxy.services import nr_sp_service
    real_load_catalog = nr_sp_service._load
    catalog_path = str(Path(__file__).resolve().parents[1] / "config/domain/nr_sp_catalog.json")
    # Changing the fixture cwd must not cache an absent catalog for later tests.
    monkeypatch.setattr(nr_sp_service, "_load", lambda path=None: real_load_catalog(path or catalog_path))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(tmp_path / "attachments"))
    source = tmp_path / "source.xlsx"
    source_book = openpyxl.Workbook()
    source_book.active.title = "S"
    source_book.active.append(["Наименование", "Ед. изм.", "Количество"])
    for _ in range(packet_count):
        source_book.active.append(["Cable A", "m", 12])
        source_book.active.append(["Cable B", "m", 3])
    source_book.save(source)
    source_book.close()
    attachment = preserve_read_attachment(source, attachment_id="read_123456abcdef",
                                         original_name="source.xlsx", root=tmp_path / "attachments")
    lookups = []
    book = fgis_price_service.PriceBook([
        {"code": RESOURCE_A, "name": "Cable A", "unit": "m", "price_current": 100},
        {"code": RESOURCE_B, "name": "Cable B", "unit": "m", "price_current": 200},
    ], region="Санкт-Петербург", quarter="3 квартал 2026 г.")
    real_lookup = book.lookup

    def lookup(code):
        lookups.append(code)
        return real_lookup(code)

    book.lookup = lookup
    monkeypatch.setattr(fgis_price_service, "available_pricebooks", lambda *a, **k: [BOOK + ".json"])

    def resolve_book(selected=None, **kwargs):
        assert kwargs.get("explicit_only") is True
        assert selected == BOOK
        return BOOK + ".json"

    monkeypatch.setattr(fgis_price_service, "resolve_pricebook_path", resolve_book)
    monkeypatch.setattr(fgis_price_service, "get_pricebook", lambda path: book)
    from proxy.services import model_pricebook_evidence_service as price_evidence
    source_bytes = json.dumps([real_lookup(RESOURCE_A), real_lookup(RESOURCE_B)], ensure_ascii=False).encode("utf-8")
    book.source_sha256 = hashlib.sha256(source_bytes).hexdigest()
    book.source_name = BOOK + ".parquet"
    book.source_byte_size = len(source_bytes)
    monkeypatch.setattr(price_evidence, "_provider", lambda name: price_evidence.PriceBookEvidenceSource(
        book, book.source_sha256, book.source_name, book.source_byte_size))
    search_queries = []

    def norm_search(query, *, limit):
        search_queries.append(query)
        assert limit == 6
        if query in native_discovery_queries:
            return {"cards": [{"norm_key": f"new-discovery-{index}", "norm_code": f"DISCOVERY-{index}",
                              "title": f"Literal discovered candidate {index}", "measure_unit": "m",
                              "source_ref": f"fixture#new-discovery-{index}",
                              "work_steps": [f"withheld_literal_body_{index}_" + query + "x" * native_discovery_padding],
                              "resource_count": 0, "resource_preview": []} for index in range(6)],
                    "retrieval_trace": {"rag": {"status": "ok", "fusion": "native_rrf"}}}
        return {"cards": [{"norm_key": "immutable-norm", "norm_code": NORM_CODE,
                           "title": "Synthetic exact norm", "measure_unit": "m",
                           "source_ref": "fixture#immutable-norm",
                           "work_steps": ["Install" + (("x" * (candidate_padding or 4000)) + query
                                                       if group_char_budget or candidate_padding else "")],
                           "resource_count": 24,
                           "resource_preview": [{"code": RESOURCE_A, "name": "Cable A", "unit": "m"},
                                                {"code": RESOURCE_B, "name": "Cable B", "unit": "m"}]}],
                "retrieval_trace": {"rag": {"status": "ok", "fusion": "native_rrf"}}}

    monkeypatch.setattr(service, "retrieve_smeta_norm_cards", norm_search)
    from proxy.smeta_core import norm_browser
    exact_reads = []
    full_resources = [{"kind": "material", "code": f"fixture-resource-{i}",
                       "name": f"Resource {i}", "unit": "m", "per_unit": i + 0.25}
                      for i in range(24)]
    full_resources[0]["name"] += "x" * exact_resource_padding
    full_resources[-2] = {"kind": "material", "code": RESOURCE_A, "name": "Cable A", "unit": "m", "per_unit": 12}
    full_resources[-1] = {"kind": "material", "code": RESOURCE_B, "name": "Cable B", "unit": "m", "per_unit": 3}

    def exact_card(**pointer):
        exact_reads.append(pointer)
        assert pointer == {"norm_key": "immutable-norm", "norm_code": NORM_CODE}
        return {"status": "ok", "schema": "smeta_exact_norm_card.v1", "missing": [], "warnings": [],
                "card": {"norm_key": "immutable-norm", "norm_code": NORM_CODE,
                         "title": "Synthetic exact norm", "measure_unit": "m", "work_steps": ["Install"],
                         "source_ref": "fixture#immutable-norm", "resource_count": 24,
                         "resources": full_resources, "resources_returned": 24,
                         "resources_complete": True, "resource_consumption_known": True}}

    monkeypatch.setattr(norm_browser, "read_norm_card_exact", exact_card)
    preset = ModelExecutionPreset(preset_id="fixture-sheet", model_family="fixture",
                                  input_token_limit=input_token_limit, generation_reserve_tokens=5000,
                                  safety_reserve_tokens=100, normal_tool_count=6, max_tools=6,
                                  max_batch_items=6, parallel_read_limit=1, reasoning_enabled=False,
                                  reasoning_supported=True,
                                  source_chain=("test",))
    connection = SimpleNamespace(connection_id="conn:sheet", revision_id="conn:sheet:r1",
                                 display_name="Offline sheet model", model_id="fixture-model",
                                 locality=ConnectionLocality.LOOPBACK, base_url="http://127.0.0.1:1919/v1",
                                 secret_ref=None, effective_preset=preset)

    class Resolver:
        def resolve(self, role, **kwargs):
            assert role is ConnectionRole.ANSWER
            return connection

        def resolve_fallback(self, *args, **kwargs):
            pytest.fail("offline model must not use fallback")

    packet_queries = [
        ([f"S!R{2 + offset * 2} — Cable A", f"S!R{3 + offset * 2} — Cable B"] if query_count == 2 else
         [f"S!R{2 + offset * 2 + index % 2} — Cable {'A' if index % 2 == 0 else 'B'} query {index}"
          for index in range(query_count)]) for offset in range(packet_count)]
    declared_queries = [query for queries in packet_queries for query in queries]
    if packet_count > 1:
        original_packet = service.adaptive_row_packet
        monkeypatch.setattr(service, "adaptive_row_packet", lambda *a, **kw: original_packet(*a, **kw)[:2])
    if group_char_budget is not None:
        monkeypatch.setattr(service, "model_rag_group_char_budget", lambda **kwargs: group_char_budget)
    responses = [InferenceResponse(text="\n".join(packet_queries[0]), tool_calls=(),
                                   finish_reason="stop", usage={})]
    if query_declaration is not None:
        responses[0] = InferenceResponse(text=query_declaration, tool_calls=(), finish_reason="stop", usage={})
    if query_repair:
        responses.insert(0, InferenceResponse(text='{"calls":[]}', tool_calls=(), finish_reason="stop", usage={}))
    if empty_query:
        responses[0] = InferenceResponse(text="", tool_calls=(), finish_reason="length",
                                         usage={"completion_tokens": 5000})
    if unknown_tool:
        responses.append(InferenceResponse(text="", tool_calls=(_native("not_a_reader", code=RESOURCE_A),),
                                          finish_reason="tool_calls", usage={}))
    elif reads:
        if native_discovery_queries:
            responses.append(InferenceResponse(text="", tool_calls=tuple(
                _native("search_sources", q=query, target="norms") for query in native_discovery_queries),
                finish_reason="tool_calls", usage={}))
        read_responses = [
            InferenceResponse(text="", tool_calls=(_native("read_norm_card", norm_key="immutable-norm"),),
                              finish_reason="tool_calls", usage={}),
            InferenceResponse(text="", tool_calls=(_native("read_resource_price", code=RESOURCE_A),
                                                    _native("read_resource_price", code=RESOURCE_B)),
                              finish_reason="tool_calls", usage={}),
        ]
        if native_read_batch:
            read_responses = [InferenceResponse(text="", tool_calls=tuple(
                call for response in read_responses for call in response.tool_calls),
                finish_reason="tool_calls", usage={})]
        responses.extend(read_responses)
    def final_response(packet_index=0):
        # The model fixture uses the source addresses actually assigned to the
        # two book reads, including any candidate groups shown before them.
        refs = (f"Q{len(search_queries) + 2}.H1", f"Q{len(search_queries) + 3}.H1") if reads else ("Q1.H1", "Q2.H1")
        if reads and group_char_budget is not None:
            shown_prices = {}
            for message in requests[-1].messages:
                for line in str(message["content"]).splitlines():
                    try:
                        registry = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    if isinstance(registry, dict):
                        for item in registry.get("previously_shown_evidence_registry", []):
                            if item.get("doc_name") == "fgis_pricebook_resources.v1":
                                shown_prices[item["resource_code"]] = item["evidence_ref"]
            refs = shown_prices[RESOURCE_A], shown_prices[RESOURCE_B]
        answer = _answer(refs=refs)
        if json_final:
            answer = json.dumps({"artifact_kind": "lsr", "rows": [
                {"source_row": "S!R2", "row_role": "material", "title": "Cable A", "unit": "m",
                 "quantity": 12, "norm_code": "", "resource_code": RESOURCE_A, "evidence_refs": [refs[0]]},
                {"source_row": "S!R3", "row_role": "material", "title": "Cable B", "unit": "m",
                 "quantity": 3, "norm_code": "", "resource_code": RESOURCE_B, "evidence_refs": [refs[1]]}]})
        if packet_index:
            answer = answer.replace("S!R2", f"S!R{2 + packet_index * 2}").replace("S!R3", f"S!R{3 + packet_index * 2}")
        if partial_final and packet_index == 0:
            answer = "\n".join(answer.splitlines()[:-1])
        if ((malformed_first_packet and packet_index == 0)
                or (malformed_second_packet and packet_index == 1)):
            # Keep the authored line literal but omit one declared cell. The
            # harness may preserve this proposal; it must not repair its fields.
            answer = answer.replace(f"| {RESOURCE_A} | {refs[0]} |", f"| {RESOURCE_A} |")
            if malformed_all_first_packet:
                answer = answer.replace(f"| {RESOURCE_B} | {refs[1]} |", f"| {RESOURCE_B} |")
        if duplicate_json_first_packet and packet_index == 0:
            conflict = ('{"source_row":"S!R2","row_role":"material","title":"Cable A",'
                        '"unit":"m","quantity":12,"norm_code":"authored-first","norm_code":"",'
                        '"resource_code":' + json.dumps(RESOURCE_A) + ',"evidence_refs":' + json.dumps([refs[0]])
                        + ',"evidence_refs":[]}')
            unique = {"source_row": "S!R3", "row_role": "material", "title": "Cable B", "unit": "m",
                      "quantity": 3, "norm_code": "", "resource_code": RESOURCE_B, "evidence_refs": [refs[1]]}
            answer = '{"artifact_kind":"lsr","rows":[' + conflict + ',' + json.dumps(unique) + ']}'
        return InferenceResponse(text="" if empty_final else answer, tool_calls=(),
                                 finish_reason="length" if empty_final else "stop",
                                 usage={"completion_tokens": 120})

    if not unknown_tool:
        responses.append(final_response)
    for packet_index in range(1, packet_count):
        responses.extend([
            InferenceResponse(text="\n".join(packet_queries[packet_index]), tool_calls=(), finish_reason="stop", usage={}),
            InferenceResponse(text="", tool_calls=(_native("read_norm_card", norm_key="immutable-norm"),), finish_reason="tool_calls", usage={}),
            InferenceResponse(text="", tool_calls=(_native("read_resource_price", code=RESOURCE_A),
                                                   _native("read_resource_price", code=RESOURCE_B)), finish_reason="tool_calls", usage={}),
            lambda index=packet_index: final_response(index),
        ])
    requests = []

    class Transport:
        async def complete(self, selected, request):
            assert selected.revision_id == connection.revision_id
            requests.append(request)
            assert responses, "harness forced another model cycle after final rows"
            response = responses.pop(0)
            return response() if callable(response) else response

    monkeypatch.setattr(service, "resolve_canonical_route", lambda **kw: CanonicalRouteDecision(
        requested=CanonicalRouteMode.ACTIVE, effective=CanonicalRouteMode.ACTIVE,
        source="test", reason="offline fixture", restart_required=False))
    monkeypatch.setattr(service, "maybe_answer_table_query", lambda *a, **k: None)
    monkeypatch.setattr(service, "dataset_memory_prompt_excerpt", lambda *a, **k: "")
    monkeypatch.setattr(service, "rules_pre_verdict", lambda *a, **k: pytest.fail("no semantic validator"))
    state = SimpleNamespace(reranker_available=False, reranker_cls=None, llm_semaphore=asyncio.Semaphore(1),
                            metrics_cache={}, crag_stats={"no_data": 0, "hallucination": 0, "verified": 0},
                            chat_metrics={"latency_search": [], "latency_gen": [], "tokens": [], "crag_fail": 0})
    config = SimpleNamespace(provider="mlx", model="fixture-model", base_url="http://fixture.invalid",
                             chat_url="http://fixture.invalid/v1/chat/completions", api_key="",
                             supports_validation=False)
    executions = []

    async def execute_workbook(call, context, progress):
        result = await chat._execute_chat_workbook_tool(call, context, progress)
        executions.append((call, result))
        return result

    def no_shared_retrieval(**kwargs):
        pytest.fail("norm queries must not use the document retriever")

    runtime = service.EvidenceRuntimeDeps(
        state=state, rag_backend=SimpleNamespace(collection_name="fixture"), cache=SimpleNamespace(),
        cache_embedding=None, cache_marker="miss", cache_scope="",
        assistant_text=lambda message: str(message.get("content") or ""),
        augment_model_tool_args=lambda call, **kw: call,
        chat_model_final_answer=lambda answer, status: (answer, status, {}),
        cloud_body_for_model=lambda body, *a: body,
        compact_tool_result_for_prompt=lambda item, **kw: item,
        dataset_ids_from_chunks=lambda chunks: [], dataset_sensitivities=lambda ids: [],
        env_bool=lambda key, default=False: default, env_float=lambda key, default=0: default,
        env_int=lambda key, default=0: 3 if native_read_batch and key == "LES_CHAT_TOOL_MAX_CALLS" else default,
        expand_context_windows=lambda chunks, *a, **kw: SimpleNamespace(chunks=chunks, payload=lambda: {}),
        format_tool_results_for_model=lambda rows: json.dumps(rows, ensure_ascii=False),
        generation_token_budget=lambda **kw: 5000, llm_runtime=lambda: config,
        local_context_budget=lambda **kw: {"focus_max_chunks": 0, "context_max_chunks": 0,
                                          "context_chars_limit": 300_000, "context_window_chars": 1000},
        mlx_runtime=lambda: config, names_for_dataset_ids=lambda *a: [],
        notebook_study_validation_status=lambda status, **kw: status,
        ollama_native_complete=lambda *a, **kw: pytest.fail("unbound model call"),
        parse_model_tool_calls=chat._parse_model_tool_calls,
        prepare_notebook_reader_memory=lambda *a, **kw: None, record_cloud_cost=lambda *a, **kw: None,
        retrieve_chat_chunks=no_shared_retrieval, source_excerpts=lambda *a, **kw: [],
        table_query_response=lambda **kw: None, cloud_fallback_models=lambda runtime: [],
        cloud_model_timeout=lambda: 1, model_connection_resolver=lambda: (Resolver(), object()),
        model_connection_transport=lambda client, secret_store: Transport(),
        workbook_tool_executor=execute_workbook,
    )
    request = service.EvidenceRequestContext(
        req=SimpleNamespace(question="Собери ЛСР; ВОР не создавай. Санкт-Петербург, 3 квартал 2026",
                            mode="estimator", response_length="short", output_directive="", session_id="sheet-session",
                            dataset_filter=None, project_id=0, attachment_id=attachment["attachment_id"],
                            attachment_context="Cable A — 12 m; Cable B — 3 m", candidate_acceptance=False,
                            reranker_enabled=None, comparative_acceptance_variant=variant),
        dataset_ids=["selected"], effective_dataset_filter="", resolved_dataset_names=[], dataset_name_by_id={},
        query_route_payload={"channel": "rag"}, target_doc_filter=[], target_file_ref=None, topic_doc_filter=[],
        topic_retrieval_plan=None, inventory_requested=False, study_requested=False, memory_block="", session_block="",
        class_suggestions=[], use_semantic_cache=False, use_validation=False, validation_skip_reason="test",
        route=SimpleNamespace(intent="rag"), table_result=None, request_started_at=time.time(),
        profile_snapshot={"revision_id": "fixture-profile:1", "mode": "estimator",
                          "model_policy": {"reasoning_enabled": reasoning_enabled,
                                           **({"query_reasoning_enabled": query_reasoning_enabled}
                                              if query_reasoning_enabled is not None else {})},
                          "tools": ["search_sources", "build_vor_workbook", "build_lsr_workbook"],
                          "rag_policy": {"iterative": True, "system_datasets": ["smeta"],
                                         "model_authored_initial_query": True},
                          "prompt_text": "Answer only from shown evidence."},
    )
    histories = []
    try:
        result = await service.run_chat_evidence_application(request, runtime, service.ResponseBoundary(
            save_chat_history=lambda **row: histories.append(row) or "history-sheet", token_sink=None,
            version_stamp=lambda: {}))
    except service.HTTPException as error:
        if not empty_query:
            raise
        result = {"error": error.detail, "status": error.status_code}
    async def resume_same_request():
        return await service.run_chat_evidence_application(request, runtime, service.ResponseBoundary(
            save_chat_history=lambda **row: histories.append(row) or "history-resume", token_sink=None,
            version_stamp=lambda: {}))

    async def corrected_source_row(locator):
        resource, title, quantity = ((RESOURCE_A, "Cable A", 12) if int(locator.rsplit("R", 1)[1]) % 2 == 0
                                     else (RESOURCE_B, "Cable B", 3))
        responses.extend([
            InferenceResponse(text=f"{locator} — {title} correction", tool_calls=(), finish_reason="stop", usage={}),
            InferenceResponse(text="", tool_calls=(_native("read_norm_card", norm_key="immutable-norm"),),
                              finish_reason="tool_calls", usage={}),
            InferenceResponse(text="", tool_calls=(_native("read_resource_price", code=RESOURCE_A),
                                                    _native("read_resource_price", code=RESOURCE_B)),
                              finish_reason="tool_calls", usage={}),
        ])
        def corrected_response():
            shown_ref = None
            for message in requests[-1].messages:
                for line in str(message["content"]).splitlines():
                    try:
                        payload = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    if isinstance(payload, dict):
                        for item in payload.get("previously_shown_evidence_registry", []):
                            if item.get("doc_name") == "fgis_pricebook_resources.v1" and item.get("resource_code") == resource:
                                shown_ref = item["evidence_ref"]
            assert shown_ref
            rows = [{"source_row": locator, "row_role": "material", "title": title, "unit": "m",
                     "quantity": quantity, "norm_code": "", "resource_code": resource, "evidence_refs": shown_ref}]
            return InferenceResponse(text=json.dumps({"artifact_kind": "lsr", "rows": rows}), tool_calls=(),
                                     finish_reason="stop", usage={"completion_tokens": 80})
        responses.append(corrected_response)
        return await resume_same_request()

    return SimpleNamespace(result=result, executions=executions, requests=requests, responses=responses,
                           exact_reads=exact_reads, lookups=lookups, queries=search_queries, histories=histories,
                           declared_queries=declared_queries, resume_same_request=resume_same_request,
                           corrected_source_row=corrected_source_row, connection=connection)


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_native_sheet_reads_keep_full_resources_distinct_price_arguments_and_real_workbook(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant=variant)
    assert len(run.requests) == 4 and not run.responses
    for request in run.requests:
        frame = "\n".join(str(message["content"]) for message in request.messages)
        assert "ВОР не создавай" in frame and "Санкт-Петербург, 3 квартал 2026" in frame
        assert "source.xlsx" in frame
        assert "S!R2" in frame and "Cable A" in frame
        assert "S!R3" in frame and "Cable B" in frame
    for name in ("read_norm_card", "search_pricebook_resources", "read_resource_price", "read_source"):
        assert name in str(run.requests[1].tools)
    assert len(run.exact_reads) == 1
    registries = []
    for message in run.requests[2].messages:
        for line in message["content"].splitlines():
            try:
                item = json.loads(line)
            except json.JSONDecodeError:
                continue
            if isinstance(item, dict) and "previously_shown_evidence_registry" in item:
                registries.extend(item["previously_shown_evidence_registry"])
    full_card = next(item["content"] for item in registries if "fixture-resource-21" in item["content"])
    assert '"per_unit": 21.25' in full_card and "resources_complete=True" in full_card
    assert sum("fixture-resource-21" in item["content"] for item in registries) == 1
    assert RESOURCE_A in str(run.requests[3].messages) and RESOURCE_B in str(run.requests[3].messages)
    # Same tool name with different literal codes must execute both reads.
    assert run.lookups.count(RESOURCE_A) == 2
    assert run.lookups.count(RESOURCE_B) == 2
    assert len(run.executions) == 1
    call, receipt = run.executions[0]
    trace = run.result["retrieval_trace"]
    assert all(read["model_delivery_status"] == "delivered" for read in trace["tool_loop"]["evidence_reads"])
    assert not trace["evidence_manifest"].get("prepared_not_delivered")
    assert call["tool"] == "build_lsr_workbook" and receipt["status"] == "complete"
    rows = call["args"]["decisions"]
    assert [(row["source_row"], row["resource_code"], row["quantity"]) for row in rows] == [
        ("S!R2", RESOURCE_A, 12), ("S!R3", RESOURCE_B, 3)]
    assert all(row["evidence_bindings"] for row in rows)
    assert all("confirmed_price" not in row for row in rows)
    assert receipt["calculated_positions"] == 2
    assert receipt["confirmed_amount"] == 1800 and receipt["preliminary_amount"] == 0
    from proxy.services.artifact_revision_service import ArtifactRevisionStore
    artifacts = ArtifactRevisionStore(tmp_path / "storage/artifacts/meta.db", tmp_path / "storage/artifacts/files")
    output = artifacts.resolve_path(receipt["artifact"]["revision_id"])
    rendered = openpyxl.load_workbook(output, data_only=False)
    try:
        sheet = rendered["ЛСР РИМ"]
        for code, quantity, price in ((RESOURCE_A, 12, 100), (RESOURCE_B, 3, 200)):
            row = next(cells[1].row for cells in sheet.iter_rows() if cells[1].value == code)
            assert sheet.cell(row, 7).value == quantity
            assert sheet.cell(row, 10).value == price
            assert sheet.cell(row, 12).value == f"=ROUND(G{row}*J{row}*K{row},2)"
    finally:
        rendered.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_native_discovery_inventory_keeps_exact_reads_and_real_json_workbook_within_capacity(tmp_path, monkeypatch, variant):
    queries = tuple(f"literal discovery query {index}" for index in range(4))
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant=variant,
                           input_token_limit=30_000, native_discovery_queries=queries,
                           native_discovery_padding=10_000, json_final=True)
    assert len(run.requests) == 5 and not run.responses
    assert run.queries[:len(run.declared_queries)] == run.declared_queries
    assert run.queries[-len(queries):] == list(queries)
    assert len(run.exact_reads) == 1 and len(run.executions) == 1
    trace = run.result["retrieval_trace"]
    assert not trace["tool_loop"].get("evidence_read_stop")
    reads = trace["tool_loop"]["evidence_reads"]
    discoveries = [read for read in reads if read["call"]["tool"] == "search_sources"]
    assert len(discoveries) == 4
    assert all(read["model_inventory_delivery_status"] == "delivered" for read in discoveries)
    assert all(read["model_delivery_status"] == "prepared" for read in discoveries)
    assert all(len(read["deferred_body_evidence_refs"]) == 6 for read in discoveries)
    frames = [call for call in trace["context_governor"]["calls"]
              if call["purpose"] in {"sheet_pass_evidence_read", "direct_evidence_read"}]
    assert len(frames) == 3
    assert all(frame["included_tokens"] <= frame["input_budget_tokens"] for frame in frames)
    first = "\n".join(message["content"] for message in run.requests[2].messages)
    assert "new-discovery-0" in first and "withheld_literal_body_0_" not in first
    last = "\n".join(message["content"] for message in run.requests[-1].messages)
    registry = next(json.loads(line) for line in last.splitlines()
                    if line.startswith('{"body_delivery":'))
    full_card = next(row["content"] for row in registry["previously_shown_evidence_registry"]
                     if "fixture-resource-21" in row["content"])
    assert "resources_complete=True" in full_card and '"per_unit": 21.25' in full_card
    assert sum("fixture-resource-21" in message["content"] for message in run.requests[-1].messages) == 1
    inventory_refs = {ref for read in discoveries for ref in read["evidence_refs"]}
    indexed = {item["id"]: item for item in trace["evidence_manifest"]["model_visible"]}
    assert inventory_refs <= indexed.keys()
    assert all(indexed[ref]["body_visible"] is False for ref in inventory_refs)
    assert all("withheld_literal_body_" not in str(item) for item in trace["source_map"])
    rows = run.executions[0][0]["args"]["decisions"]
    assert [(row["source_row"], row["quantity"], row["resource_code"]) for row in rows] == [
        ("S!R2", 12, RESOURCE_A), ("S!R3", 3, RESOURCE_B)]
    receipt = run.executions[0][1]
    assert receipt["status"] == "complete" and receipt["confirmed_amount"] == 1800
    from proxy.services.artifact_revision_service import ArtifactRevisionStore
    artifact = ArtifactRevisionStore(tmp_path / "storage/artifacts/meta.db", tmp_path / "storage/artifacts/files")
    output = artifact.resolve_path(receipt["artifact"]["revision_id"])
    assert output.is_file() and output.suffix == ".xlsx"
    if variant == "B1_CANDIDATE":
        from proxy.services.workbook_research_job_service import WorkbookResearchJobService
        job = WorkbookResearchJobService().get(trace["workbook_research_job"]["job_id"])
        stored = [item for item in job["evidence"] if item["meta"].get("model_read_tool") == "search_sources"]
        assert len(stored) == 24
        assert all("withheld_literal_body_" in item["content"] for item in stored)
        assert all(item["meta"]["model_delivery_status"] == "prepared" for item in stored)
        assert all(item["meta"]["model_inventory_delivery_status"] == "delivered" for item in stored)


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_plain_final_sheet_rows_do_not_force_an_evidence_read_cycle(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=False, variant=variant)
    assert len(run.requests) == 2 and not run.responses
    assert run.exact_reads == []
    for request in run.requests:
        frame = "\n".join(str(message["content"]) for message in request.messages)
        assert "ВОР не создавай" in frame
        assert "S!R2" in frame and "S!R3" in frame
    assert len(run.executions) == 1
    assert run.executions[0][1]["status"] == "complete"
    assert [row["resource_code"] for row in run.executions[0][0]["args"]["decisions"]] == [RESOURCE_A, RESOURCE_B]


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_bounded_inventory_reserves_native_growth_and_finishes_after_third_read(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant=variant,
                           query_count=35, input_token_limit=30_000,
                           candidate_padding=1200, native_read_batch=True)
    assert len(run.requests) == 3 and not run.responses
    assert len(run.exact_reads) == 1 and len(run.executions) == 1
    trace = run.result["retrieval_trace"]
    grouped = (trace["grouped_model_evidence"] if variant == "B0_DIRECT"
               else trace["grouped_model_evidence_pages"][0])
    budget = grouped["native_read_growth_budget"]
    assert 0 < budget["reserved_tokens"] <= budget["requested_tokens"]
    assert budget["remaining_group_tokens"] >= budget["first_whole_group_tokens"]
    packets = trace["context_governor"]["calls"]
    answer_packet = next(packet for packet in packets if packet["purpose"] in {"sheet_pass_answer", "answer"})
    read_packet = next(packet for packet in packets if packet["purpose"] in {"sheet_pass_evidence_read", "direct_evidence_read"})
    assert answer_packet["included_tokens"] + budget["reserved_tokens"] <= answer_packet["input_budget_tokens"]
    assert read_packet["included_tokens"] <= read_packet["input_budget_tokens"]
    messages = run.requests[-1].messages
    assert [message["role"] for message in messages] == ["system", "user", "assistant", "tool", "tool", "tool"]
    assert len(messages[2]["tool_calls"]) == 3
    assert [message["tool_call_id"] for message in messages[3:]] == [call["id"] for call in messages[2]["tool_calls"]]
    frame = messages[1]["content"]
    assert "source.xlsx" in frame and "S!R2" in frame and "S!R3" in frame
    registry = next(json.loads(line) for line in frame.splitlines()
                    if line.startswith('{"body_delivery":'))
    full_card = next(row["content"] for row in registry["previously_shown_evidence_registry"]
                     if "fixture-resource-21" in row["content"])
    assert '"per_unit": 21.25' in full_card and "resources_complete=True" in full_card
    assert sum("fixture-resource-21" in str(message["content"]) for message in messages) == 1
    assert not trace["tool_loop"].get("evidence_read_stop")
    assert [(row["source_row"], row["resource_code"], row["quantity"])
            for row in run.executions[0][0]["args"]["decisions"]] == [
                ("S!R2", RESOURCE_A, 12), ("S!R3", RESOURCE_B, 3)]


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_actual_native_results_exceeding_bound_pause_without_false_final(tmp_path, monkeypatch, variant):
    # The exact card itself must exceed the bound, independently of how much
    # lossless sharing saves on repeated unread inventory fields.
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant=variant,
                           query_count=35, input_token_limit=20_000,
                           candidate_padding=1200, native_read_batch=True, exact_resource_padding=40_000)
    assert len(run.requests) == 2 and len(run.responses) == 1
    assert len(run.exact_reads) == 1 and not run.executions
    trace = run.result["retrieval_trace"]
    capacity = trace["context_governor"]["error"]
    assert capacity["required_tokens"] > capacity["available_tokens"]
    assert capacity["overflow_tokens"] > 0
    assert trace["tool_loop"]["evidence_read_stop"] == "evidence_read_capacity"
    assert len(trace["tool_loop"]["evidence_reads"]) == 3
    assert all(read["status"] == "ok" for read in trace["tool_loop"]["evidence_reads"])
    assert all(read["model_delivery_status"] == "prepared" for read in trace["tool_loop"]["evidence_reads"])
    prepared_refs = {ref for read in trace["tool_loop"]["evidence_reads"] for ref in read["evidence_refs"]}
    visible_refs = {item["id"] for item in trace["evidence_manifest"]["model_visible"]}
    assert not prepared_refs.intersection(visible_refs)
    assert prepared_refs == {item["id"] for item in trace["evidence_manifest"]["prepared_not_delivered"]}
    assert not any(call["purpose"] in {"sheet_pass_evidence_read", "direct_evidence_read"}
                   for call in trace["actual_model_calls"]["calls"])


@pytest.mark.asyncio
async def test_resumed_zero_decision_job_can_cite_saved_native_reads_and_preserve_delivered_status(tmp_path, monkeypatch):
    from dataclasses import replace
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService

    run = await _run_sheet(tmp_path, monkeypatch, reads=True, query_count=35, input_token_limit=20_000,
                           candidate_padding=1200, native_read_batch=True)
    store = WorkbookResearchJobService()
    job_id = run.result["retrieval_trace"]["workbook_research_job"]["job_id"]
    paused = store.get(job_id)
    assert paused["accepted_rows"] == []
    assert run.result["retrieval_trace"]["tool_loop"]["evidence_read_stop"] == "evidence_read_capacity"
    saved = [record for record in paused["evidence"] if record["meta"].get("model_read_packet_id")]
    assert len(saved) == 3 and all(record["meta"]["model_delivery_status"] == "prepared" for record in saved)
    prices = {record["meta"]["resource_code"]: record["meta"]["model_evidence_ref"] for record in saved
              if record["doc_name"] == "fgis_pricebook_resources.v1"}
    # A new declared runtime bound is explicit; source/model/revision/task are identical.
    run.connection.effective_preset = replace(run.connection.effective_preset, input_token_limit=100_000)
    run.responses[:] = [InferenceResponse(text="\n".join(run.declared_queries), tool_calls=(), finish_reason="stop", usage={}),
                        InferenceResponse(text=_answer(refs=(prices[RESOURCE_A], prices[RESOURCE_B])),
                                          tool_calls=(), finish_reason="stop", usage={})]
    requests_before = len(run.requests)
    reads_before = len(run.exact_reads)
    resumed = await run.resume_same_request()
    assert len(run.requests) == requests_before + 2 and len(run.exact_reads) == reads_before
    assert len(run.executions) == 1 and not run.responses
    rows = run.executions[0][0]["args"]["decisions"]
    assert all(row["evidence_bindings"] and row["_resource_evidence_check"]["code_in_refs"] for row in rows)
    frame = "\n".join(str(m["content"]) for m in run.requests[-1].messages)
    assert "fixture-resource-21" in frame and all(ref in frame for ref in prices.values())
    visible_refs = [item["id"] for item in resumed["retrieval_trace"]["evidence_manifest"]["model_visible"]]
    mapped_refs = [item["evidence_ref"] for item in resumed["retrieval_trace"]["source_map"]]
    for ref in prices.values():
        assert visible_refs.count(ref) == 1 and mapped_refs.count(ref) == 1
    recovered = WorkbookResearchJobService(store.db_path).get(job_id)
    by_id = {r["global_evidence_id"]: r for r in recovered["evidence"]}
    for previous in saved:
        record = by_id[previous["global_evidence_id"]]
        assert record["meta"]["model_delivery_status"] == "delivered"
        assert record["content"] == previous["content"]
    assert resumed["retrieval_trace"]["model_execution_profile"]["context_input_tokens"]["effective"] == 100_000


@pytest.mark.asyncio
@pytest.mark.parametrize("partial_final", [False, True])
async def test_terminal_sheet_result_packages_before_35_candidate_groups_are_drained(tmp_path, monkeypatch, partial_final):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, query_count=35,
                           group_char_budget=7000, partial_final=partial_final)
    assert len(run.requests) == 4 and not run.responses
    assert len(run.exact_reads) == 1 and len(run.executions) == 1
    rows = run.executions[0][0]["args"]["decisions"]
    assert [(row["source_row"], row["resource_code"], row["quantity"]) for row in rows] == (
        [("S!R2", RESOURCE_A, 12)] if partial_final else
        [("S!R2", RESOURCE_A, 12), ("S!R3", RESOURCE_B, 3)])
    trace = run.result["retrieval_trace"]
    pages = trace["grouped_model_evidence_pages"]
    assert len(pages) == 1 and pages[0]["full_body_hit_count"] == 1
    assert pages[0]["deferred_body_refs"]
    from proxy.services.workbook_research_job_service import WorkbookResearchIdentity, WorkbookResearchJobService
    store = WorkbookResearchJobService()
    job = store.get(trace["workbook_research_job"]["job_id"])
    assert job["status"] == ("paused" if partial_final else "complete")
    assert job["remaining"] == (["S!R3"] if partial_final else [])
    assert run.executions[0][1]["status"] == ("partial" if partial_final else "complete")
    assert run.executions[0][1]["checkpoint"]["status"] == "complete"  # Mechanical packaging succeeded.
    assert trace["attachment_integrity"]["omitted_source_locators"] == (["S!R3"] if partial_final else [])
    assert job["held_packet"]["queries"]
    assert len(job["held_packet"]["terminal_packets"]) == 1
    assert set(run.declared_queries) <= set(job["query_cache"])
    for query in job["held_packet"]["queries"]:
        hits = service.restore_query_hits_from_cache([query], job["query_cache"])
        assert hits and query in hits[0][1][0].content
    reopened = WorkbookResearchJobService(store.db_path).begin_or_resume(WorkbookResearchIdentity(**job["identity"]))
    assert reopened["held_packet"] == job["held_packet"]
    assert len(reopened["accepted_rows"]) == len(rows)
    assert "source.xlsx" in str(run.requests[-1].messages)
    assert "S!R2" in str(run.requests[-1].messages) and "S!R3" in str(run.requests[-1].messages)


@pytest.mark.asyncio
@pytest.mark.parametrize("partial_final", [False, True])
async def test_terminal_packets_continue_only_unseen_rows_and_keep_independent_evidence_cursors(tmp_path, monkeypatch, partial_final):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, query_count=35,
                           group_char_budget=7000, partial_final=partial_final, packet_count=2)
    assert len(run.requests) == 8 and not run.responses
    assert len(run.exact_reads) == 2 and len(run.executions) == 1
    rows = run.executions[0][0]["args"]["decisions"]
    assert [row["source_row"] for row in rows] == (
        ["S!R2", "S!R4", "S!R5"] if partial_final else ["S!R2", "S!R3", "S!R4", "S!R5"])
    assert len({row["operation_id"] for row in rows}) == len(rows)
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    store = WorkbookResearchJobService()
    job = store.get(run.result["retrieval_trace"]["workbook_research_job"]["job_id"])
    cursors = job["held_packet"]["terminal_packets"]
    assert [cursor["locators"] for cursor in cursors] == [["S!R2", "S!R3"], ["S!R4", "S!R5"]]
    assert all(cursor["queries"] and cursor["terminal_model_result_saved"] for cursor in cursors)
    assert set(run.declared_queries) <= set(job["query_cache"])
    assert job["remaining"] == (["S!R3"] if partial_final else [])
    assert job["status"] == ("paused" if partial_final else "complete")
    assert run.executions[0][1]["status"] == ("partial" if partial_final else "complete")
    assert run.executions[0][1]["checkpoint"]["status"] == "complete"
    assert run.result["retrieval_trace"]["attachment_integrity"]["omitted_source_locators"] == (["S!R3"] if partial_final else [])
    assert job["status_dimensions"]["completeness"] == ("partial" if partial_final else "full")
    reopened = WorkbookResearchJobService(store.db_path).get(job["job_id"])
    assert reopened["held_packet"] == job["held_packet"] and reopened["evidence"] == job["evidence"]
    request_count = len(run.requests)
    await run.resume_same_request()
    assert len(run.requests) == request_count
    assert store.get(job["job_id"])["accepted_rows"] == job["accepted_rows"]


@pytest.mark.asyncio
@pytest.mark.parametrize("packet_count,all_malformed", [(1, False), (2, False), (1, True), (2, True)])
async def test_malformed_packet_keeps_raw_pending_and_advances_only_unseen_source_rows(tmp_path, monkeypatch, packet_count, all_malformed):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, query_count=35,
                           group_char_budget=7000, packet_count=packet_count,
                           malformed_first_packet=True, malformed_all_first_packet=all_malformed)
    assert len(run.requests) == packet_count * 4 and not run.responses
    assert [query for query in run.queries if query in run.declared_queries] == run.declared_queries
    has_draft = not all_malformed or packet_count > 1
    assert len(run.executions) == int(has_draft)
    if has_draft:
        assert run.executions[-1][1]["artifact"]["draft"] is True
        assert run.executions[-1][1]["pricing_status"] == "priced_partial"
        assert run.executions[-1][1]["document_issues"]
    trace = run.result["retrieval_trace"]
    assert [entry["packet_locators"] for entry in trace["tool_loop"]["passes"]] == (
        [["S!R2", "S!R3"]] if packet_count == 1 else
        [["S!R2", "S!R3"], ["S!R4", "S!R5"]])
    blocked_status = "unstructured" if all_malformed and packet_count == 1 else "structural_error"
    assert trace["model_result"]["status"] == blocked_status
    assert trace["model_result"]["packaged"] is has_draft
    diagnostic = trace["model_result_transport_diagnostics"][0]
    assert diagnostic["code"] == "MODEL_RESULT_TABLE_ROW_WIDTH_MISMATCH"
    assert diagnostic["expected_cells"] == 8 and diagnostic["actual_cells"] == 7
    assert diagnostic["decision"]["source_row"] == "S!R2"
    assert f"| {RESOURCE_A} |" in diagnostic["decision"]["raw_line"]
    from proxy.services.workbook_research_job_service import WorkbookResearchIdentity, WorkbookResearchJobService
    store = WorkbookResearchJobService()
    job = store.get(trace["workbook_research_job"]["job_id"])
    assert job["status"] == "paused" and job["status_dimensions"]["completeness"] == "partial"
    assert job["remaining"] == (["S!R2", "S!R3"] if all_malformed else ["S!R2"]), {
        "held_flags": {key: job["held_packet"].get(key) for key in
                       ("model_result_transport_pending", "model_read_pending", "model_final_output_missing")},
        "projection": sorted(service.model_rag_pending_result_projection(
            service.model_rag_pending_result_packets(job["held_packet"]))[0]),
        "accepted_locators": [row["source_row"] for row in job["accepted_rows"]],
    }
    assert [row["source_row"] for row in job["accepted_rows"]] == (
        ([] if all_malformed else ["S!R3"]) + ([] if packet_count == 1 else ["S!R4", "S!R5"]))
    if has_draft:
        assert run.executions[-1][0]["args"]["decisions"] == job["accepted_rows"]
    assert diagnostic["decision"] in [proposal["decision"] for proposal in job["unapplied_proposals"]]
    pending = job["held_packet"]["pending_model_result_packets"]
    assert len(pending) == 1 and pending[0]["model_result_transport_pending"] is True
    assert pending[0]["locators"] == ["S!R2", "S!R3"] and pending[0]["queries"]
    assert pending[0]["transport_diagnostics"] == trace["model_result_transport_diagnostics"]
    for query in pending[0]["queries"]:
        assert service.restore_query_hits_from_cache([query], job["query_cache"])
    restarted = WorkbookResearchJobService(store.db_path).begin_or_resume(WorkbookResearchIdentity(**job["identity"]))
    assert restarted["held_packet"] == job["held_packet"]
    assert restarted["unapplied_proposals"] == job["unapplied_proposals"]
    request_count = len(run.requests)
    resumed = await run.resume_same_request()
    assert len(run.requests) == request_count
    assert len(run.executions) == 2 * int(has_draft)
    assert resumed["retrieval_trace"]["model_result"]["status"] == blocked_status
    assert resumed["retrieval_trace"]["model_result"]["packaged"] is has_draft
    assert store.get(job["job_id"])["accepted_rows"] == job["accepted_rows"]
    assert store.get(job["job_id"])["held_packet"]["pending_model_result_packets"] == pending


@pytest.mark.asyncio
@pytest.mark.parametrize("same_packet", [False, True])
async def test_explicit_reopen_api_resolves_only_the_corrected_pending_source_packet(tmp_path, monkeypatch, same_packet):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, query_count=35,
                           group_char_budget=7000, packet_count=1 if same_packet else 2,
                           malformed_first_packet=True, malformed_second_packet=not same_packet,
                           malformed_all_first_packet=same_packet)
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    store = WorkbookResearchJobService()
    job = store.get(run.result["retrieval_trace"]["workbook_research_job"]["job_id"])
    original_proposals = job["unapplied_proposals"]
    unchanged_rows = job["accepted_rows"]
    assert len(job["held_packet"]["pending_model_result_packets"]) == (1 if same_packet else 2)
    for index, locator in enumerate(("S!R2", "S!R3" if same_packet else "S!R4")):
        patched = await chat.patch_lsr_rows(chat.LsrPatchRequest(
            session_id="sheet-session", job_id=job["job_id"], reopen_locators=[locator], repack=False,
        ), _user=SimpleNamespace())
        assert patched["model_called"] is False
        corrected = await run.corrected_source_row(locator)
        assert len(run.requests) == (8 if same_packet else 12) + index * 4 and not run.responses
        job = store.get(job["job_id"])
        assert job["unapplied_proposals"] == original_proposals
        assert all(row in job["accepted_rows"] for row in unchanged_rows)
        pending = job["held_packet"].get("pending_model_result_packets", [])
        if index == 0:
            assert len(pending) == 1
            assert pending[0]["locators"] == (["S!R2", "S!R3"] if same_packet else ["S!R4", "S!R5"])
            if same_packet:
                assert pending[0]["resolved_source_rows"] == ["S!R2"]
                assert job["remaining"] == ["S!R3"]
            assert job["status"] == "paused"
            assert run.executions[-1][1]["artifact"]["draft"] is True
            assert run.executions[-1][0]["args"]["decisions"] == job["accepted_rows"]
            assert [item["decision"]["source_row"] for item in corrected["retrieval_trace"]["model_result_transport_diagnostics"]] == ["S!R3" if same_packet else "S!R4"]
        else:
            assert pending == [] and job["status"] == "complete" and job["remaining"] == []
            assert len(run.executions) == (2 if same_packet else 3)
            assert run.executions[-1][1]["status"] == "complete"
            assert not run.executions[-1][1]["document_issues"]


@pytest.mark.asyncio
async def test_joined_sheet_refreshes_owner_transport_pending_before_empty_query_guard(tmp_path, monkeypatch):
    async def owner_saved_malformed_packet(job_id):
        from proxy.services.workbook_research_job_service import WorkbookResearchJobService
        store = WorkbookResearchJobService()
        diagnostic = {"code": "MODEL_RESULT_TABLE_ROW_WIDTH_MISMATCH", "reason": "ambiguous_markdown_row_width",
                      "expected_cells": 8, "actual_cells": 7, "packet_locators": ["S!R2", "S!R3"],
                      "decision": {"source_row": "S!R2", "raw_line": "| S!R2 | material | Cable A | m | 12 | | |"}}
        pending = {"sheet": "S", "locators": ["S!R2", "S!R3"], "queries": ["S!R2 — Cable A"],
                   "model_result_transport_pending": True, "transport_diagnostics": [diagnostic]}
        store.record_packet(job_id, sheet="S", packet_locators=["S!R2", "S!R3"], accepted_rows=[],
                            evidence=[], remaining=["S!R2", "S!R3"], unapplied_proposals=[diagnostic],
                            held_packet=service.model_rag_with_pending_results(pending, [pending]))
        store.pause(job_id, remaining=["S!R2", "S!R3"])
        return False
    monkeypatch.setattr(service, "claim_or_join_sheet_pass", owner_saved_malformed_packet)
    run = await _run_sheet(tmp_path, monkeypatch, reads=False)
    assert run.requests == [] and run.executions == []
    trace = run.result["retrieval_trace"]
    assert trace["workbook_research_job"]["status"] == "paused"
    assert trace["workbook_research_job"]["remaining"] == 2
    assert trace["model_result"]["status"] == "unstructured" and trace["model_result"]["packaged"] is False
    assert trace["model_result_transport_diagnostics"][0]["decision"]["source_row"] == "S!R2"


@pytest.mark.asyncio
async def test_duplicate_json_row_keeps_other_decisions_and_only_conflicting_locator_pending_until_explicit_correction(tmp_path, monkeypatch):
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, packet_count=2, duplicate_json_first_packet=True)
    store = WorkbookResearchJobService()
    job = store.latest(session_id="sheet-session")
    assert len(run.requests) == 8 and not run.responses
    assert run.result["retrieval_trace"]["tool_loop"]["passes"][0]["parsed_rows"] == 1
    accepted = {row["source_row"]: row for row in job["accepted_rows"]}
    assert set(accepted) == {"S!R3", "S!R4", "S!R5"}
    assert job["status"] == "paused" and job["remaining"] == ["S!R2"]
    assert accepted["S!R3"]["quantity"] == 3 and accepted["S!R3"]["resource_code"] == RESOURCE_B
    pending = service.model_rag_pending_result_packets(job["held_packet"])
    assert service.model_rag_pending_result_projection(pending)[0] == {"S!R2"}
    assert len(job["unapplied_proposals"]) == 1
    raw_proposal = job["unapplied_proposals"][0]
    assert raw_proposal["code"] == "MODEL_RESULT_JSON_DUPLICATE_FIELD"
    assert raw_proposal["decision"]["source_row"] == "S!R2"
    assert raw_proposal["decision"]["raw_json"].count('"norm_code":') == 2
    assert raw_proposal["decision"]["raw_json"].count('"evidence_refs":') == 2
    assert "norm_code" not in raw_proposal["decision"]
    assert run.result["retrieval_trace"]["model_result"]["packaged"] is True
    assert run.executions[-1][1]["artifact"]["draft"] is True
    request_count = len(run.requests)
    await run.resume_same_request()
    assert len(run.requests) == request_count  # plain continuation cannot repair it
    reopened = await chat.patch_lsr_rows(chat.LsrPatchRequest(
        session_id="sheet-session", job_id=job["job_id"], reopen_locators=["S!R2"], repack=False), _user=SimpleNamespace())
    assert reopened["job"]["remaining"] == ["S!R2"]
    result = await run.corrected_source_row("S!R2")
    resolved = store.get(job["job_id"])
    assert resolved["status"] == "complete" and resolved["remaining"] == []
    assert service.model_rag_pending_result_packets(resolved["held_packet"]) == []
    assert resolved["unapplied_proposals"] == [raw_proposal]
    for row in resolved["accepted_rows"]:
        if row["source_row"] in accepted:
            assert row == accepted[row["source_row"]]
    assert result["retrieval_trace"]["model_result_transport_diagnostics"] == []
    assert result["retrieval_trace"]["model_result"]["packaged"] is True
    assert len(run.executions) == 3
    assert not run.executions[-1][1]["document_issues"]


@pytest.mark.asyncio
async def test_joined_sheet_clears_diagnostics_resolved_by_owner(tmp_path, monkeypatch):
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    run = await _run_sheet(tmp_path, monkeypatch, reads=True,
                           malformed_first_packet=True, malformed_all_first_packet=True)
    store = WorkbookResearchJobService()
    job = store.latest(session_id="sheet-session")
    assert run.result["retrieval_trace"]["model_result_transport_diagnostics"]
    old_proposals = job["unapplied_proposals"]
    request_count = len(run.requests)

    async def owner_resolved_packet(job_id):
        store.record_packet(job_id, sheet="S", packet_locators=["S!R2", "S!R3"],
            accepted_rows=[{"source_row": locator, "row_role": "header", "title": "Authored context"}
                           for locator in ["S!R2", "S!R3"]], evidence=[], remaining=[],
            held_packet={"terminal_model_result_saved": True, "sheet": "S", "locators": ["S!R2", "S!R3"]})
        store.finish(job_id, status="complete")
        return False
    monkeypatch.setattr(service, "claim_or_join_sheet_pass", owner_resolved_packet)
    resumed = await run.resume_same_request()
    assert len(run.requests) == request_count
    trace = resumed["retrieval_trace"]
    assert trace["model_result_transport_diagnostics"] == []
    assert trace["workbook_research_job"]["remaining"] == 0
    assert store.get(job["job_id"])["unapplied_proposals"] == old_proposals


@pytest.mark.asyncio
async def test_joined_sheet_refreshes_owner_terminal_locators_before_empty_query_guard(tmp_path, monkeypatch):
    async def owner_completed_empty_packet(job_id):
        from proxy.services.workbook_research_job_service import WorkbookResearchJobService
        store = WorkbookResearchJobService()
        store.record_packet(job_id, sheet="S", packet_locators=["S!R2", "S!R3"], accepted_rows=[],
            evidence=[], remaining=["S!R2", "S!R3"],
            held_packet={"terminal_model_result_saved": True, "sheet": "S", "locators": ["S!R2", "S!R3"]})
        store.pause(job_id, remaining=["S!R2", "S!R3"])
        return False
    monkeypatch.setattr(service, "claim_or_join_sheet_pass", owner_completed_empty_packet)
    run = await _run_sheet(tmp_path, monkeypatch, reads=False)
    assert run.requests == [] and run.executions == []
    trace = run.result["retrieval_trace"]
    assert trace["workbook_research_job"]["remaining"] == 2
    assert trace["model_result"]["status"] == "unstructured" and trace["model_result"]["packaged"] is False


@pytest.mark.asyncio
async def test_completed_sheet_admits_new_catalog_context_for_unchanged_attachment_bytes(tmp_path, monkeypatch):
    from dataclasses import replace
    from proxy.services.attachment_source_catalog_service import CatalogRow
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    run = await _run_sheet(tmp_path, monkeypatch, reads=True)
    store = WorkbookResearchJobService()
    before = store.latest(session_id="sheet-session")
    assert before["status"] == "complete" and before["document"]["state"] == "ready"
    old_decisions = before["accepted_rows"]
    old_revision = before["document"]["final_revision_id"]
    original_bind = service.bind_tabular_estimator_profile
    literal = "Literal source context now exposed by the server catalog"
    def expanded_catalog(*args, **kwargs):
        profile, catalog, bound = original_bind(*args, **kwargs)
        extra = CatalogRow(locator="PDF_TEXT!R1", sheet="PDF_TEXT", excel_row=1, values=(literal,),
            line="PDF_TEXT!R1: " + literal, numbers=(), row_classification="ambiguous",
            classification_reason="pdf_text_requires_model_role")
        return profile, replace(catalog, rows=(*catalog.rows, extra), sheets=(*catalog.sheets, ("PDF_TEXT", 1))), bound
    monkeypatch.setattr(service, "bind_tabular_estimator_profile", expanded_catalog)
    run.responses.extend([
        InferenceResponse(text="PDF_TEXT!R1 — literal source context", tool_calls=(), finish_reason="stop", usage={}),
        InferenceResponse(text=json.dumps({"rows": [{"source_row": "PDF_TEXT!R1", "row_role": "header", "title": literal}]}),
                          tool_calls=(), finish_reason="stop", usage={"completion_tokens": 30}),
    ])
    request_count = len(run.requests)
    resumed = await run.resume_same_request()
    after = store.get(before["job_id"])
    assert len(run.requests) == request_count + 2 and not run.responses
    assert after["identity"]["source_sha256"] == before["identity"]["source_sha256"]
    assert after["accepted_rows"][:2] == old_decisions
    assert after["accepted_rows"][2]["source_row"] == "PDF_TEXT!R1"
    assert after["accepted_rows"][2]["row_role"] == "header"
    assert after["remaining"] == [] and after["status"] == "complete"
    assert after["document"]["state"] == "ready" and after["document"]["final_revision_id"] != old_revision
    expansion = [event for event in store.events_after(before["job_id"]) if event["event"] == "source_scope_expanded"]
    assert expansion[0]["data"]["added_locators"] == ["PDF_TEXT!R1"]
    assert resumed["retrieval_trace"]["model_result"]["packaged"] is True
    assert run.executions[-1][0]["args"]["decisions"][:2] == run.executions[0][0]["args"]["decisions"]
    assert run.executions[-1][0]["args"]["decisions"][2]["source_row"] == "PDF_TEXT!R1"


@pytest.mark.asyncio
async def test_direct_baseline_keeps_35_literal_queries_and_single_terminal_packaging(tmp_path, monkeypatch):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant="B0_DIRECT",
                           query_count=35, group_char_budget=7000)
    assert len(run.requests) == 4 and len(run.executions) == 1 and not run.responses
    assert run.queries == run.declared_queries
    assert [row["resource_code"] for row in run.executions[0][0]["args"]["decisions"]] == [RESOURCE_A, RESOURCE_B]


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_unknown_native_reader_never_becomes_a_terminal_workbook(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=False, variant=variant, unknown_tool=True)
    assert len(run.requests) == 2 and not run.responses
    assert run.executions == []
    trace = run.result["retrieval_trace"]
    assert run.result["model_connection"]["pending_tool_calls"] == 1
    if variant == "B0_DIRECT":
        assert trace["tool_loop"]["evidence_read_stop"] == "evidence_read_unknown_tool"
    else:
        assert "evidence_read_unknown_tool" in str(trace)


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
@pytest.mark.parametrize("reasoning_enabled", [False, True])
async def test_bound_sheet_profile_reasoning_reaches_every_existing_model_turn(tmp_path, monkeypatch, variant, reasoning_enabled):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, variant=variant,
                           reasoning_enabled=reasoning_enabled)
    assert len(run.requests) == 4 and not run.responses
    assert all(request.reasoning_enabled is reasoning_enabled for request in run.requests)
    assert len(run.executions) == 1
    final_prompt = str(run.requests[-1].messages)
    assert ("/no_think" in final_prompt) is (not reasoning_enabled)
    reasoning = run.result["retrieval_trace"]["model_execution_profile"]["reasoning"]
    assert reasoning == {"requested": reasoning_enabled, "effective": reasoning_enabled,
                         "source": "profile_model_policy", "restart_required": False}


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_reasoning_without_final_length_stops_without_workbook_or_new_model_turn(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=False, variant=variant,
                           reasoning_enabled=True, empty_final=True)
    assert len(run.requests) == 2 and not run.responses
    assert run.executions == []
    assert run.result["retrieval_trace"]["model_generation"] == {
        "status": "technical_pending", "reason": "model_final_output_missing", "finish_reason": "length"}


@pytest.mark.asyncio
@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
async def test_empty_length_query_never_becomes_repair_search_or_empty_workbook(tmp_path, monkeypatch, variant):
    run = await _run_sheet(tmp_path, monkeypatch, reads=False, variant=variant,
                           reasoning_enabled=True, empty_query=True)
    assert len(run.requests) == 1 and run.requests[0].max_output_tokens == 5000
    assert run.executions == [] and run.queries == []
    assert run.result == {"error": {"error": "MODEL_QUERY_OUTPUT_INCOMPLETE"}, "status": 503}
    from proxy.services.public_error_service import public_error_payload
    assert public_error_payload(status_code=503, detail=run.result["error"])["code"] == "MODEL_QUERY_OUTPUT_INCOMPLETE"
