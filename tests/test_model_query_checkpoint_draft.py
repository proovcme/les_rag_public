"""Completed durable packets remain downloadable after a typed query failure."""
from __future__ import annotations

import copy
from dataclasses import replace
import json
from types import SimpleNamespace

import openpyxl
import pytest

from proxy.routers import chat
from proxy.services import chat_evidence_application_service as service
from proxy.services.artifact_revision_service import ArtifactRevisionStore
from proxy.services.openai_compatible_transport_service import InferenceResponse
from proxy.services.workbook_research_job_service import WorkbookResearchJobService
from tests.test_model_rag_sheet_evidence_reads import BOOK, RESOURCE_A, RESOURCE_B, _run_sheet
from tests import test_model_rag_sheet_evidence_reads as sheet_fixture


def install_later_query_failure(monkeypatch, *, fail_call=5, failure_kind='primary-empty-length', artifact_kind='lsr'):
    original_run = service.run_chat_evidence_application
    snapshots = {'before_failure':None,'requests':[], 'boundaries':[], 'fail_enabled':True}
    original_begin = WorkbookResearchJobService.begin_or_resume
    def begin(store, *args, **kwargs):
        job=original_begin(store,*args,**kwargs)
        snapshots['job_id']=job['job_id']
        snapshots.setdefault('initial_remaining',copy.deepcopy(job['remaining']))
        return job
    monkeypatch.setattr(WorkbookResearchJobService,'begin_or_resume',begin)

    async def run(request, runtime, boundary):
        factory = runtime.model_connection_transport
        def transport_factory(*args, **kwargs):
            transport = factory(*args, **kwargs)
            complete = transport.complete
            async def interrupted(selected, inference):
                snapshots['requests'].append(inference)
                index=len(snapshots['requests'])
                if index==fail_call:
                    snapshots['before_failure']=copy.deepcopy(WorkbookResearchJobService().get(snapshots['job_id']))
                if index>=fail_call and snapshots['fail_enabled']:
                    if failure_kind=='repair-empty-length' and index==fail_call:
                        return InferenceResponse(text='| output | result |\n| --- | --- |\n| repeated | table |',
                                                 tool_calls=(),finish_reason='stop',usage={'completion_tokens':30})
                    return InferenceResponse(text='',tool_calls=(),finish_reason='length',usage={'completion_tokens':5000})
                response=await complete(selected,inference)
                if artifact_kind=='vor' and index==1:
                    return replace(response,text=json.dumps({'artifact_kind':'vor','queries':response.text.splitlines()}))
                if artifact_kind=='vor' and index==4:
                    payload=json.loads(response.text)
                    payload['artifact_kind']='vor'
                    for row in payload['rows']:
                        row['row_role']='work'
                        row['title']='Install '+row['title']
                        row['work_row']=row['source_row']
                    return replace(response,text=json.dumps(payload))
                return response
            transport.complete=interrupted
            return transport
        snapshots['boundaries'].append(boundary)
        return await original_run(request,replace(runtime,model_connection_transport=transport_factory),boundary)
    monkeypatch.setattr(service,'run_chat_evidence_application',run)
    return snapshots


@pytest.mark.asyncio
@pytest.mark.parametrize('failure_kind',['primary-empty-length','repair-empty-length'])
@pytest.mark.parametrize('artifact_kind',['lsr','vor'])
async def test_completed_packet_query_failure_delivers_physical_draft_without_new_model_call(tmp_path,monkeypatch,failure_kind,artifact_kind):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch,failure_kind=failure_kind,artifact_kind=artifact_kind)
    run=await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    assert len(snapshots['requests'])==(5 if failure_kind=='primary-empty-length' else 6)
    assert len(run.executions)==1
    trace=run.result['retrieval_trace']
    job_id=trace['workbook_research_job']['job_id']
    job=WorkbookResearchJobService().get(job_id)
    before=snapshots['before_failure']
    assert job['status']=='paused' and job['resumable'] is True
    assert job['phase'] in {'model_query_output_incomplete','model_query_format_pending'}
    assert job['remaining']==before['remaining']==['S!R4','S!R5']
    for field in ('accepted_rows','query_cache','evidence','identity'):
        assert job[field]==before[field],field
    assert job['task_context']['original_question']==before['task_context']['original_question']
    assert job['held_packet']['locators']==['S!R4','S!R5']
    assert job['held_packet']['model_query_output_missing'] is True
    expected_code='MODEL_QUERY_OUTPUT_INCOMPLETE'
    issue=next(i for i in run.result['document_issues'] if i['code']==expected_code)
    assert issue['origin']=='application' and issue['technical_failure'] is True
    assert run.result['retrieval_trace']['query_checkpoint_failure']['code']==expected_code
    assert run.result['retrieval_trace']['query_checkpoint_failure']['model_decisions_rewritten'] is False
    assert job_id not in service._LIVE_SHEET_PASS_GATES
    call,result=run.executions[0]
    assert call['tool']=='build_'+artifact_kind+'_workbook'
    assert call['args']['decisions']==before['accepted_rows']
    assert call['args']['draft'] is True
    assert result['status']=='partial'
    if artifact_kind=='lsr':
        assert result['pricing_status']=='priced_partial'
        assert result['confirmed_amount']==0 and result['price_basis']['actual_book']==BOOK
    assert result['artifact']['download_url']
    assert any(i['code']==expected_code for i in result['document_issues'])
    assert 'NO_DATA' not in run.result['answer']
    store=ArtifactRevisionStore(tmp_path/'storage/artifacts/meta.db',tmp_path/'storage/artifacts/files')
    path=store.resolve_path(result['artifact']['revision_id'])
    book=openpyxl.load_workbook(path,data_only=True)
    try:
        if artifact_kind=='lsr':
            sheet=book['ЛСР РИМ']
            assert 'НЕ ЗАВЕРШ' in str(sheet['A15'].value)
            for code,qty in ((RESOURCE_A,12),(RESOURCE_B,3)):
                row=next(cells for cells in sheet if cells[1].value==code)
                assert row[6].value==qty
        else:
            sheet=book['ВОР']
            assert 'НЕ ЗАВЕРШ' in str(sheet['A2'].value)
            assert [sheet.cell(i,6).value for i in (4,5)]==[12,3]
        checks='\n'.join(str(c.value) for row in book['Проверка'] for c in row if c.value is not None)
        assert expected_code in checks
    finally: book.close()
    requests_before=len(snapshots['requests'])
    repack=await chat.repack_lsr_without_model(chat.LsrRepackRequest(session_id='sheet-session',job_id=job_id),_user=SimpleNamespace())
    assert repack['model_called'] is False and len(snapshots['requests'])==requests_before
    assert any(i['code']==expected_code for i in repack['result']['document_issues'])
    assert repack['result']['status']=='partial'
    if artifact_kind=='lsr': assert repack['result']['price_basis']['actual_book']==BOOK
    repacked_id=repack['result']['artifact']['revision_id']
    prior_bytes=store.read_bytes(result['artifact']['revision_id'])
    repack_bytes=store.read_bytes(repacked_id)
    replay=await chat.repack_lsr_without_model(chat.LsrRepackRequest(session_id='sheet-session',job_id=job_id),_user=SimpleNamespace())
    assert replay['result']['artifact']['revision_id']==repacked_id
    assert store.read_bytes(repacked_id)==repack_bytes
    assert store.read_bytes(result['artifact']['revision_id'])==prior_bytes
    after=WorkbookResearchJobService().get(job_id)
    for field in ('accepted_rows','remaining','identity','query_cache','evidence'):
        assert after[field]==before[field],field


@pytest.mark.asyncio
async def test_no_accepted_packet_keeps_typed_query_error_and_never_makes_empty_file(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'empty-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch,fail_call=1)
    with pytest.raises(service.ModelQueryFailure) as caught:
        await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2)
    assert caught.value.detail=={'error':'MODEL_QUERY_OUTPUT_INCOMPLETE'}
    assert len(snapshots['requests'])==1
    job=WorkbookResearchJobService().get(snapshots['before_failure']['job_id'])
    assert job['accepted_rows']==[] and job['status']=='paused'
    assert not (tmp_path/'storage/artifacts/files').exists()


@pytest.mark.asyncio
async def test_failed_packet_preserves_never_reached_sheet_global_remainder(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'global-jobs.sqlite'))
    preserve=sheet_fixture.preserve_read_attachment
    def multiple_sheets(source,*args,**kwargs):
        book=openpyxl.load_workbook(source)
        sheet=book.create_sheet('T')
        sheet.append(['Наименование','Ед. изм.','Количество'])
        sheet.append(['Other source operation','m',19])
        book.save(source); book.close()
        return preserve(source,*args,**kwargs)
    monkeypatch.setattr(sheet_fixture,'preserve_read_attachment',multiple_sheets)
    snapshots=install_later_query_failure(monkeypatch)
    run=await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    before=snapshots['before_failure']
    assert snapshots['initial_remaining']==['S!R2','S!R3','S!R4','S!R5','T!R2']
    job=WorkbookResearchJobService().get(run.result['retrieval_trace']['workbook_research_job']['job_id'])
    assert job['remaining']==['S!R4','S!R5','T!R2']
    assert job['accepted_rows']==before['accepted_rows']
    assert len(snapshots['requests'])==5 and len(run.executions)==1
    assert any('T!R2' in issue.get('remaining_locators',[]) for issue in run.result['document_issues'])


@pytest.mark.asyncio
async def test_sse_terminal_result_contains_draft_and_technical_resume_instead_of_only_red_error(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'sse-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch)
    completed={}
    async def actual_producer(request,token_sink=None):
        completed['run']=await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
        return completed['run'].result
    monkeypatch.setattr(chat,'_run_chat',actual_producer)
    response=await chat.chat_stream(chat.ChatRequest(question='same task',session_id='sheet-session'),_user=None)
    body=''.join([piece if isinstance(piece,str) else piece.decode('utf-8') async for piece in response.body_iterator])
    assert 'event: error\n' not in body
    payload=json.loads(body.split('event: final\ndata: ',1)[1].split('\n\n',1)[0])
    assert payload['artifact']['download_url']
    assert payload['artifact']['resume']['job_id']==completed['run'].result['retrieval_trace']['workbook_research_job']['job_id']
    assert payload['artifact']['resume']['remaining_locators']==['S!R4','S!R5','S!R2','S!R3']
    # The original fixture's unparented material rows remain quality gaps;
    # the durable unprocessed scope stays separate from these known gaps.
    assert payload['retrieval_trace']['workbook_research_job']['remaining_locators']==['S!R4','S!R5']
    assert payload['artifact']['result_readiness']['estimate_complete'] is False
    assert payload['pricing_status']=='priced_partial'
    assert any(i['code']=='MODEL_QUERY_OUTPUT_INCOMPLETE' for i in payload['document_issues'])
    assert len(snapshots['requests'])==5


@pytest.mark.asyncio
async def test_explicit_resume_reads_only_remaining_packet_and_chains_from_retained_draft(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'resume-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch)
    run=await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    first=run.executions[0][1]['artifact']
    job_id=run.result['retrieval_trace']['workbook_research_job']['job_id']
    store=ArtifactRevisionStore(tmp_path/'storage/artifacts/meta.db',tmp_path/'storage/artifacts/files')
    first_bytes=store.read_bytes(first['revision_id'])
    accepted_before=copy.deepcopy(WorkbookResearchJobService().get(job_id)['accepted_rows'])
    snapshots['fail_enabled']=False
    resumed=await run.resume_same_request()
    assert len(snapshots['requests'])==9, 'one explicitly requested unfinished-packet pass, no hidden retry'
    job=WorkbookResearchJobService().get(job_id)
    assert job['status']=='complete' and job['remaining']==[]
    assert job['accepted_rows'][:2]==accepted_before
    assert len(job['accepted_rows'])==4
    assert 'query_checkpoint_failure' not in job['held_packet']
    assert not any(i['code']=='MODEL_QUERY_OUTPUT_INCOMPLETE' for i in resumed['document_issues'])
    assert len(run.executions)==2
    second=run.executions[1][1]['artifact']
    assert second['revision_id']!=first['revision_id']
    assert second['parent_revision_id']==first['revision_id']
    assert store.read_bytes(first['revision_id'])==first_bytes
    assert resumed['price_basis']['actual_book']==BOOK


@pytest.mark.asyncio
async def test_checkpoint_draft_never_promotes_private_failure_diagnostics_to_public_json(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'private-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch)
    original=service.model_query_failure
    markers=('PRIVATE_REASONING_SENTINEL','PRIVATE_TOKEN_SENTINEL','PRIVATE_SOURCE_DUMP_SENTINEL')
    def poisoned(*args,**kwargs):
        failure=original(*args,**kwargs)
        failure.private_query_trace.update(reasoning=markers[0],authorization=markers[1],
                                           source_debug={'raw_request':markers[2]},final_text=markers[0])
        return failure
    monkeypatch.setattr(service,'model_query_failure',poisoned)
    run=await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    summary=run.result['retrieval_trace']['query_checkpoint_failure']
    assert set(summary)=={'schema','code','reason','origin','technical_failure','model_decisions_rewritten'}
    assert summary['code']=='MODEL_QUERY_OUTPUT_INCOMPLETE'
    async def completed(request,token_sink=None): return run.result
    monkeypatch.setattr(chat,'_run_chat',completed)
    response=await chat.chat_stream(chat.ChatRequest(question='same task',session_id='sheet-session'),_user=None)
    body=''.join([piece if isinstance(piece,str) else piece.decode('utf-8') async for piece in response.body_iterator])
    public=json.dumps(run.result,ensure_ascii=False)+body
    assert all(marker not in public for marker in markers)
    assert 'event: final\n' in body and 'event: error\n' not in body
    assert len(snapshots['requests'])==5


@pytest.mark.asyncio
@pytest.mark.parametrize('scope_field',['job_id','identity','task_context',None,'missing_lookup'])
async def test_changed_fresh_checkpoint_scope_keeps_original_error_without_draft(tmp_path,monkeypatch,scope_field):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'scope-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch)
    original=WorkbookResearchJobService.get
    original_failure=service.model_query_failure
    fresh_fetch={'armed':False}
    def failure(*args,**kwargs):
        error=original_failure(*args,**kwargs)
        fresh_fetch['armed']=True
        return error
    monkeypatch.setattr(service,'model_query_failure',failure)
    def fetched(store,job_id):
        job=original(store,job_id)
        # The existing typed error has already paused its original job;
        # only the new fallback fetch presents this changed checkpoint.
        if fresh_fetch['armed']:
            fresh_fetch['armed']=False
            if scope_field=='missing_lookup':
                raise service.WorkbookResearchJobError('workbook research job not found')
            if scope_field is None:
                return None
            job=copy.deepcopy(job)
            job[scope_field]='other-checkpoint-scope'
        return job
    monkeypatch.setattr(WorkbookResearchJobService,'get',fetched)
    with pytest.raises(service.ModelQueryFailure) as caught:
        await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    assert caught.value.detail=={'error':'MODEL_QUERY_OUTPUT_INCOMPLETE'}
    assert len(snapshots['requests'])==5
    assert not (tmp_path/'storage/artifacts/files').exists()


@pytest.mark.asyncio
async def test_unrecognized_typed_query_failure_is_not_swallowed_after_completed_packet(tmp_path,monkeypatch):
    monkeypatch.setenv('RAG_META_DB_PATH',str(tmp_path/'unknown-jobs.sqlite'))
    snapshots=install_later_query_failure(monkeypatch)
    original=service.model_query_failure
    def unexpected(*args,**kwargs):
        error=original(*args,**kwargs)
        error.detail={'error':'UNRECOGNIZED_QUERY_FAILURE'}
        return error
    monkeypatch.setattr(service,'model_query_failure',unexpected)
    with pytest.raises(service.ModelQueryFailure) as caught:
        await _run_sheet(tmp_path,monkeypatch,reads=True,packet_count=2,json_final=True)
    assert caught.value.detail=={'error':'UNRECOGNIZED_QUERY_FAILURE'}
    assert len(snapshots['requests'])==5
    assert not (tmp_path/'storage/artifacts/files').exists()
