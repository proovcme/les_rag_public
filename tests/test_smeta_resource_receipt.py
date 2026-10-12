"""Resource links and operation identities across the tool/artifact audit boundary."""
import pytest

from proxy.services import chat_evidence_application_service as chat
from proxy.services import rim_lsr_trace_service as rim
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
from tools import smeta_book_audit as audit


@pytest.fixture
def norm(monkeypatch):
    monkeypatch.setattr(rim.gesn_service,'get_norm',lambda _:dict(name='Fixture',unit='шт',
        norm_id='fixture-norm',edition='fixture',source_ref='fixture source',resources=[
            dict(kind='material',code='R',name='Fixture material',unit='шт',per_unit=3,price=2)]))
    monkeypatch.setattr(rim.nr_sp_service,'resolve',lambda *a,**k:dict(
        status='resolved_fixture',nr_pct=0,sp_pct=0))


def work(**changes):
    return dict(dict(source_row='S!1',work_id='mount',title='Fixture',
        norm_code='ГЭСН01-01-001-01',quantity=2,unit='шт'),**changes)


@pytest.mark.parametrize('boundary', ['receipt', 'workbook'])
def test_unit_conflict_explains_both_units_without_changing_decision(norm, tmp_path, boundary):
    from proxy.smeta_core.calculation_receipt import calculation_receipts
    from openpyxl import load_workbook

    trace = rim.build_lsr_trace_from_visible_rows([work(unit='м')], pricebook=None)
    binding = trace['row_bindings'][0]
    assert binding['status'] == 'unit_conflict'
    quantity = binding['quantity_trace']
    assert (quantity['source_quantity'], quantity['source_unit'], quantity['norm_unit']) == (2, 'м', 'шт')
    expected = quantity['message']
    if boundary == 'receipt':
        receipt = calculation_receipts(trace)[0]
        assert receipt['calculation'] is None
        assert receipt['reason'] == expected
    else:
        path = tmp_path / 'unit-conflict.xlsx'
        render_lsr_xlsx(trace, path)
        book = load_workbook(path, read_only=True)
        assert book['Проверка']['F2'].value == expected
        book.close()


def resource(**changes):
    return dict(dict(source_row='S!2',work_id='allocate',title='Fixture material',
        decision_kind='ресурс',norm_code='',included_in='mount',resource_action='included',
        target_resource_code='R',quantity=3,unit='шт'),**changes)


def test_work_ids_survive_real_tool_projection_and_harvest():
    calculations=[dict(source_row='S!1',work_id=w,status='bound') for w in ('mount','terminate')]
    result=dict(schema='les.workbook_tool_result.v1',status='complete',bound_rows=1,
        row_calculation=calculations,missing=[],blockers=[],
        source=dict(attachment_id='fixture'),checkpoint=dict(checkpoint_id='cp1',status='complete'),
        artifact=dict(revision_id='rev1',filename='fixture.xlsx',download_url='/api/artifacts/rev1',
                      decision_checkpoint_id='cp1',source_scope=['attachment:fixture']))
    safe=chat.safe_workbook_history_projection(result)
    assert safe['row_calculation']==calculations
    harvested=chat.harvest_workbook_tool_result(safe)
    assert chat.workbook_row_statuses(harvested)=={'S!1':'bound'}


def test_receipt_guidance_keeps_unresolved_work_instead_of_deleting_it():
    from proxy.smeta_core.calculation_receipt import receipt_context
    text=receipt_context([{'work_id':'w','status':'norm_selection_required'}])
    assert 'decision_kind=уточнение' in text
    assert 'либо убери саму операцию' not in text


def test_explicit_norm_withdrawal_reaches_workbook_without_reviving_old_choice(norm,tmp_path):
    old=work(norm_code='ГЭСН11-01-038-05')
    answer='''| source_row | work_id | title | unit | quantity | norm_code | evidence_refs | decision_kind | coverage |
|---|---|---|---|---|---|---|---|---|
| S!1 | mount | Work awaiting a suitable norm | шт | 2 | | | уточнение | Model rejected previous norm |
'''
    rows=chat.parse_model_rag_result(answer)[1]
    rows=chat.merged_pass_decisions(rows,closed=[],portion_state={
        'shown':['S!1'],'source_anchors':['S!1'],'prior_decisions':{'S!1':old}})
    trace=rim.build_lsr_trace_from_visible_rows([*rows,work(source_row='S!2',work_id='neighbor')])
    assert trace['row_bindings'][0]['status']=='norm_selection_required'
    assert trace['summary']['bound_rows']==1
    out=tmp_path/'withdrawal.xlsx'
    render_lsr_xlsx(trace,out)
    book=audit.read_book(out)
    assert old['norm_code'] not in book
    assert len(book['ГЭСН01-01-001-01']['positions'])==1
    assert old['norm_code']=='ГЭСН11-01-038-05'


def test_audit_reads_each_provenance_field_and_all_operations(norm,tmp_path):
    rows=[work(),work(work_id='terminate',quantity=4)]
    out=tmp_path/'fixture.xlsx'
    render_lsr_xlsx(rim.build_lsr_trace_from_visible_rows(rows),out)
    book=audit.read_book(out)
    positions=book[rows[0]['norm_code']]['positions']
    assert [p['source_row'] for p in positions]==['S!1','S!1']
    assert [p['work_id'] for p in positions]==['mount','terminate']
    verdict=audit.classify(dict(source_row='S!1',title='Fixture',quantity=2,unit='шт'),
        {'source_row':'S!1','operations':rows},book)
    assert len(verdict['operations'])==2
    assert all(op['state']==audit.FORMULA_UNCACHED for op in verdict['operations']), verdict


def test_resource_cannot_confirm_an_unresolved_parent(norm):
    result=rim.positions_from_visible_lsr_rows([work(coefficient=1.5),resource()])
    assert result['row_bindings'][1]['status']=='resource_parent_unresolved'


def test_norm_code_alone_is_not_an_operation_reference(norm):
    result=rim.positions_from_visible_lsr_rows([work(),resource(included_in=work()['norm_code'])])
    assert result['row_bindings'][1]['status']=='resource_parent_missing'


def test_unresolved_source_prevents_full_estimate_amount(norm):
    result=rim.build_lsr_trace_from_visible_rows([work(),resource(included_in='missing')])
    assert result['summary']['full_amount'] is None
    assert result['summary']['amount_status']=='partial'


@pytest.mark.asyncio
async def test_missing_norm_resource_price_reaches_tool_result(norm,monkeypatch,tmp_path):
    from proxy.services import lsr_workbook_adapter_service as adapter
    monkeypatch.setattr(adapter.fgis_price_service,'resolve_pricebook_path',lambda _:None)
    monkeypatch.setattr(rim.gesn_service,'get_norm',lambda _:dict(name='Fixture',unit='шт',
        resources=[dict(kind='material',code='UNPRICED',name='Unknown',unit='шт',per_unit=1)]))
    result=await adapter.build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        dict(decisions=[work()]),tmp_path/'book.xlsx',lambda *a:None)
    assert result['missing']
    assert any('UNPRICED' in item for item in result['missing'])


def test_calculation_receipt_keeps_actual_resources_and_price_identity(norm):
    from copy import deepcopy
    from proxy.smeta_core.calculation_receipt import calculation_receipts, receipt_context
    from proxy.services.fgis_price_service import PriceBook
    prices=PriceBook([dict(code='R',name='Different pricebook description',unit='шт',
                           price_current=5)],region='Fixture region',quarter='Fixture quarter')
    trace=rim.build_lsr_trace_from_visible_rows([work(),resource()],pricebook=prices)
    before=deepcopy(trace)
    receipts=calculation_receipts(trace,pricebook=prices,pricebook_name='fixture')
    assert trace==before
    material=receipts[0]['calculation']['resources'][0]
    assert material['quantity']==6
    assert material['price_card']['name']=='Different pricebook description'
    assert material['name']=='Fixture material'
    assert receipts[0]['pricebook']['region']=='Fixture region'
    allocation=receipts[1]['resource_bindings'][0]
    assert allocation['parent_resource_quantity']==6
    assert allocation['allocated_known_cost']==material['known_cost']/2
    assert allocation['added_cost']==0
    assert receipts[1]['parent_normative_source']['norm_id']=='fixture-norm'
    assert calculation_receipts(trace,pricebook=prices,pricebook_name='fixture')==receipts
    assert 'Different pricebook description' in receipt_context(receipts)


@pytest.mark.parametrize('action',['replace','add'])
def test_unsupported_resource_action_never_changes_original_norm(norm,action):
    result=rim.positions_from_visible_lsr_rows([work(),resource(resource_action=action)])
    assert result['row_bindings'][1]['status']=='resource_accounting_unverified'
    assert len(result['positions'])==1
    _,resources=rim._position_resources(result['positions'][0])
    assert [r['code'] for r in resources]==['R']


def test_book_resource_proof_is_per_operation_and_never_taken_twice(norm,tmp_path):
    rows=[work(),resource(quantity=3,coverage='Model supplied purpose'),
          resource(source_row='S!2',work_id='allocate2',quantity=3,coverage='Second allocation')]
    out=tmp_path/'allocation.xlsx'
    render_lsr_xlsx(rim.build_lsr_trace_from_visible_rows(rows),out)
    book=audit.read_book(out)
    result=audit.classify(rows[1],{'operations':rows[1:]},book)
    assert all(op['state']==audit.RESOURCE_LINKED for op in result['operations'])
    assert audit.classify(rows[1],rows[1],book)['state']!=audit.RESOURCE_LINKED


@pytest.mark.asyncio
async def test_chat_review_requires_current_calculation_and_preserves_original(norm,monkeypatch,tmp_path):
    from proxy.services import lsr_workbook_adapter_service as adapter
    monkeypatch.setattr(adapter.fgis_price_service,'resolve_pricebook_path',lambda _:None)
    async def build(rows,filename):
        return await adapter.build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
            dict(decisions=rows,review_calculation=True),tmp_path/filename,lambda *a:None)
    original=work()
    first=await build([original],'first.xlsx')
    assert first['row_calculation'][0]['status']=='resource_review_required'
    assert first['bound_rows']==0
    ref=first['row_calculation'][0]['calculation_ref']
    confirmed=work(resource_review_ref=ref,resource_review_status='keep_all_confirmed')
    second=await build([confirmed],'second.xlsx')
    assert second['row_calculation'][0]['status']=='bound'
    assert second['bound_rows']==1
    assert first['row_calculation'][0]['calculation_ref']==second['row_calculation'][0]['calculation_ref']
    changed=await build([dict(confirmed,quantity=3)],'changed.xlsx')
    assert changed['row_calculation'][0]['status']=='resource_review_required'
    assert 'resource_review_ref' not in original
    assert (tmp_path/'first.xlsx').is_file()


def test_norm_resource_price_unit_mismatch_is_not_a_priced_result():
    from proxy.services.fgis_price_service import PriceBook
    trace=rim.build_position_trace(dict(code='',name='Fixture',unit='шт',qty=3,nr_pct=0,sp_pct=0,
        resources=[dict(kind='material',code='R',name='Connector',unit='шт',per_unit=1)]),
        pricebook=PriceBook([dict(code='R',name='Tape',unit='кг',price_current=10)]))
    assert trace['summary']['full_amount'] is None
    assert trace['summary']['mat']==0
    material=next(r for r in trace['rows'] if r['type']=='resource_material')
    assert material['meta']['price_action']=='price_unit_conflict'


def test_review_fields_survive_model_parser_and_revision():
    from proxy.services.portion_cursor_service import canonical_decision,decisions_revision
    raw='| source_row | work_id | title | unit | quantity | norm_code | resource_review_ref | resource_review_status |\n'
    raw+='|---|---|---|---|---|---|---|---|\n'
    raw+='| S!1 | mount | Fixture | шт | 2 | ГЭСН01-01-001-01 | abc | keep_all_confirmed |\n'
    _,rows=chat.parse_model_rag_result(raw)
    assert rows[0]['resource_review_ref']=='abc'
    assert canonical_decision(rows[0])['resource_review_status']=='keep_all_confirmed'
    assert decisions_revision(rows)!=decisions_revision([dict(rows[0],resource_review_ref='changed')])


def test_calculated_facts_survive_cursor_restart_into_next_model_context(tmp_path):
    from proxy.services.portion_cursor_service import PortionCursorService
    text='Файл: fixture.xlsx\n\n## Лист: S\nS!R1: № П.П | НАИМЕНОВАНИЕ РАБОТ И ЗАТРАТ | ЕД. ИЗМ | ОБЪЕМ\nS!R2: 1 | Fixture | шт | 2\n'
    cursor=PortionCursorService(tmp_path/'cursor.sqlite')
    _,state=chat.smeta_portion_context(text,session_id='s',limit=1,cursors=cursor)
    fact=dict(source_row='S!1',work_id='mount',status='resource_review_required',
              calculation_ref='actual-receipt-ref',calculation=dict(resources=[
                  dict(code='R',name='Fixture material',quantity=6,unit='шт',known_cost=12)]))
    chat.record_portion_pass(cursor,session_id='s',portion_state=state,rows=[work()],
        closed=['S!1'],artifact_revision_id='rev',row_statuses={'S!1':'resource_review_required'},
        calculation_receipts=[fact],pass_id='turn1')
    restarted=PortionCursorService(tmp_path/'cursor.sqlite')
    context,next_state=chat.smeta_portion_context(text,session_id='s',limit=1,cursors=restarted)
    assert next_state['completed']==0
    assert 'actual-receipt-ref' in context
    assert '"quantity":6' in context
    assert 'keep_all_confirmed' in context


def test_registry_preserves_all_authored_operations(norm,tmp_path):
    from proxy.services.portion_cursor_service import PortionCursorService
    from tools.smeta_row_registry import registry,CALCULATED
    text='Файл: fixture.xlsx\nS!R1: № | Наименование | Ед. изм | Количество\nS!R2: 1 | Fixture | шт | 2\n'
    cursor=PortionCursorService(tmp_path/'cursor.sqlite')
    _,state=chat.smeta_portion_context(text,session_id='s',limit=1,cursors=cursor)
    rows=[work(),work(work_id='terminate',quantity=4)]
    chat.record_portion_pass(cursor,session_id='s',portion_state=state,rows=rows,closed=['S!1'],
        artifact_revision_id=None,row_statuses={'S!1':'bound'})
    out=tmp_path/'operations.xlsx'
    render_lsr_xlsx(rim.build_lsr_trace_from_visible_rows(rows),out)
    result=registry(session='s',attachment_text=text,book=out,db=tmp_path/'cursor.sqlite')
    assert result[0]['accounted_by']==CALCULATED
    assert len(result[0]['operation_evidence']['operations'])==2
    assert all(op['state']==audit.FORMULA_UNCACHED for op in result[0]['operation_evidence']['operations'])


# --- Проход 8 на 0.30.144: подтверждение не сходилось из-за переписанного названия ---------
# Модель подтверждала чек прошлого прохода, а расчёт давал новый хэш: в него входило
# resource_name, то есть её собственная формулировка. Переименование материала — не
# изменение расчёта; личность ресурса задаёт код, и пересматривать по ней нечего.

def _supply_receipt(name, code='21.1.04.01-1046'):
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {'row_bindings': [{'row': 'СКС!5.6', 'work_id': 'supply', 'status': 'material_bound',
                               'resource_bindings': [{'action': 'standalone', 'resource_code': code,
                                                      'resource_name': name, 'quantity': 8500.0,
                                                      'unit': 'м', 'reason': 'material_bound'}]}],
             'sections': []}
    return calculation_receipts(trace)[0]


def test_renaming_a_priced_material_does_not_reopen_the_confirmed_calculation():
    first = _supply_receipt('Кабель витая пара кат.5Е 4 пары серая')
    second = _supply_receipt('Кабель витая пара кат.5Е 4 пары серая, HF UTP-4P-Cat.5e-SOLID-LSZH-GY Cabeus')
    assert first['calculation_ref'] == second['calculation_ref']
    assert second['resource_bindings'][0]['resource_name'].endswith('Cabeus'), 'название остаётся в чеке'


def test_another_resource_code_is_a_different_calculation():
    assert _supply_receipt('Кабель', '21.1.04.01-1046')['calculation_ref'] != \
        _supply_receipt('Кабель', '21.1.04.01-1024')['calculation_ref']


def test_without_a_code_the_name_identifies_the_material():
    assert _supply_receipt('Плитка тактильная', '')['calculation_ref'] != \
        _supply_receipt('Ковровая плитка Atlas 05', '')['calculation_ref']


def test_instruction_requires_confirmation_for_every_operation_with_a_receipt():
    """Проход 10: модель подтверждала только работу, а поставка с чеком оставалась открытой."""
    from proxy.smeta_core.calculation_receipt import receipt_context
    from proxy.services.chat_evidence_application_service import model_answer_request

    for text in (receipt_context([{'work_id': 'supply', 'calculation_ref': 'b' * 64}]),
                 model_answer_request('Смета', model_driven=True)):
        assert 'каждой операции' in text and 'самостоятельный материал' in text


# --- Проход 11: ссылка на чек и есть подтверждение ----------------------------------------
# Таблица ответа разрослась до 27 колонок, и на строке поставки модель сдвинула ячейки: код
# ресурса попал в колонку ссылки. Отдельная колонка статуса рядом со ссылкой ничего не
# добавляет: точную ссылку нельзя получить, не прочитав именно этот расчёт. Явный отказ
# (rejected) и пометки материала по-прежнему сильнее.

def _review(row_extra, ref='r' * 64):
    from proxy.smeta_core.calculation_receipt import apply_model_review

    trace = {'row_bindings': [{'row': 'СКС!5.6', 'work_id': 'supply', 'status': 'material_bound'}],
             'summary': {}, 'sections': []}
    receipts = [{'work_id': 'supply', 'calculation_ref': ref, 'status': 'material_bound'}]
    apply_model_review(trace, receipts, [dict({'work_id': 'supply'}, **row_extra)])
    return trace['row_bindings'][0]['status']


def test_matching_reference_alone_confirms_the_calculation():
    assert _review({'resource_review_ref': 'r' * 64}) == 'material_bound'
    assert _review({'resource_review_ref': 'r' * 64, 'resource_review_status': 'keep_all_confirmed'}) == 'material_bound'


def test_wrong_or_missing_reference_keeps_the_row_open():
    assert _review({'resource_review_ref': 'x' * 64}) == 'resource_review_required'
    assert _review({}) == 'resource_review_required'


def test_explicit_rejection_outweighs_the_reference():
    assert _review({'resource_review_ref': 'r' * 64, 'resource_review_status': 'rejected'}) == 'resource_review_required'
    assert _review({'resource_review_ref': 'r' * 64, 'not_covered': 'плитка не та'}) == 'resource_review_required'


# --- Проход 13: кодом изделия назван ресурс самой нормы ------------------------------------
# Модель поставила в поставку кабеля код 01.7.06.05-0041 — изоляционную ленту из состава
# нормы этой же работы, и расчёт посчитал «изделие» по цене ленты. Совпадение кода с
# ресурсом нормы строки — проверяемый факт, и чек называет его. Решение остаётся за моделью.

def test_receipt_names_a_product_code_that_is_the_norms_own_resource():
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': 'ГЭСНм08-02-398-01'}, 'meta': {'work_id': 'lay', 'source_row': 'СКС!5.6'}},
            {'type': 'resource_material', 'columns': {'2': '01.7.06.05-0041', '3': 'Лента изоляционная', '4': 'м'},
             'meta': {'kind': 'material'}},
        ]}]}],
        'row_bindings': [
            {'row': 'СКС!5.6', 'work_id': 'lay', 'status': 'bound', 'resource_bindings': []},
            {'row': 'СКС!5.6', 'work_id': 'supply', 'status': 'material_bound',
             'resource_bindings': [{'action': 'standalone', 'resource_code': '01.7.06.05-0041',
                                    'resource_name': 'Витая пара кат.5Е', 'quantity': 8500.0, 'unit': 'м'}]},
        ],
    }
    supply = next(item for item in calculation_receipts(trace) if item['work_id'] == 'supply')
    binding = supply['resource_bindings'][0]
    assert binding['matches_norm_resource'] == {'work_id': 'lay', 'norm_code': 'ГЭСНм08-02-398-01',
                                                'name': 'Лента изоляционная'}
    other = trace['row_bindings'][1]['resource_bindings'][0]
    assert 'matches_norm_resource' not in other, 'исходная трасса не меняется'


def test_a_product_code_outside_the_norm_is_not_flagged():
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': 'ГЭСНм08-02-398-01'}, 'meta': {'work_id': 'lay', 'source_row': 'СКС!5.6'}},
            {'type': 'resource_material', 'columns': {'2': '01.7.06.05-0041', '3': 'Лента изоляционная', '4': 'м'},
             'meta': {'kind': 'material'}},
        ]}]}],
        'row_bindings': [{'row': 'СКС!5.6', 'work_id': 'supply', 'status': 'material_bound',
                          'resource_bindings': [{'action': 'standalone', 'resource_code': '21.1.04.01-1046',
                                                 'resource_name': 'Кабель витая пара', 'quantity': 8500.0, 'unit': 'м'}]}],
    }
    supply = calculation_receipts(trace)[0]
    assert 'matches_norm_resource' not in supply['resource_bindings'][0]


# --- Проход 15: подсказку модель прочитала и проигнорировала -------------------------------
# Кодом изделия остался ресурс нормы (изоляционная лента), и строка закрылась по цене ленты.
# Код по-прежнему не выбирает код за модель, но и не даёт закрыть такую строку молча:
# нужно её собственное объяснение, почему ресурс нормы назван изделием ведомости.

def _norm_resource_review(row_extra):
    from proxy.smeta_core.calculation_receipt import apply_model_review

    ref = 'n' * 64
    trace = {'row_bindings': [{'row': 'СКС!5.6', 'work_id': 'supply', 'status': 'material_bound'}],
             'summary': {}, 'sections': []}
    receipts = [{'work_id': 'supply', 'calculation_ref': ref, 'status': 'material_bound',
                 'resource_bindings': [{'action': 'standalone', 'resource_code': '01.7.06.05-0041',
                                        'matches_norm_resource': {'work_id': 'lay', 'norm_code': 'ГЭСНм08-02-398-01',
                                                                  'name': 'Лента изоляционная'}}]}]
    apply_model_review(trace, receipts, [dict({'work_id': 'supply', 'resource_review_ref': ref}, **row_extra)])
    return trace['row_bindings'][0]


def test_norm_resource_as_a_product_needs_the_authors_explanation():
    binding = _norm_resource_review({})
    assert binding['status'] == 'resource_review_required'
    assert 'ресурс нормы' in binding['resource_review_reason']


def test_explained_choice_is_accepted_as_the_authors_decision():
    binding = _norm_resource_review({'material_review_status': 'confirmed',
                                     'material_review_reason': 'в ведомости именно эта лента, а не кабель'})
    assert binding['status'] == 'material_bound'


def test_a_standalone_material_is_not_its_own_norm_resource():
    """Проход 15: подсказка сработала на самой строке материала (её же work_id, шифр нормы пуст)."""
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': ''}, 'meta': {'work_id': '1.31-2', 'source_row': 'АР!1.31'}},
            {'type': 'resource_material', 'columns': {'2': '01.6.03.03-0011', '3': 'Покрытие ковровое', '4': 'м2'},
             'meta': {'kind': 'material'}},
        ]}]}],
        'row_bindings': [{'row': 'АР!1.31', 'work_id': '1.31-2', 'status': 'material_bound',
                          'resource_bindings': [{'action': 'standalone', 'resource_code': '01.6.03.03-0011',
                                                 'resource_name': 'Ковровая плитка Atlas 05', 'quantity': 66.7}]}],
    }
    receipt = calculation_receipts(trace)[0]
    assert 'matches_norm_resource' not in receipt['resource_bindings'][0]


# --- Проход 22: к демонтажу добавлено изделие ----------------------------------------------
# У строки СОУЭ!7.2 работа — демонтаж оповещателя, а модель добавила сам оповещатель по цене
# нового (4 016,04 ₽). Вид работы — проверяемый факт наименования нормы, и чек его называет;
# решение, нужно ли изделие, остаётся за моделью.

def _removal_receipt(work_name, material_name='Оповещатель'):
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': 'ГЭСНм10-04-101-07', '3': work_name},
             'meta': {'work_id': 'w', 'source_row': 'СОУЭ!7.2'}}]}]}],
        'row_bindings': [{'row': 'СОУЭ!7.2', 'work_id': 'supply', 'status': 'material_bound',
                          'resource_bindings': [{'action': 'standalone', 'resource_code': '61.2.04.05-0004',
                                                 'resource_name': material_name, 'quantity': 7}]}],
    }
    return calculation_receipts(trace)[0]['resource_bindings'][0]


def test_receipt_names_that_the_row_work_is_a_removal():
    hint = _removal_receipt('Демонтаж громкоговорителя в помещении')['row_work_is_removal']
    assert hint == {'work_id': 'w', 'norm_code': 'ГЭСНм10-04-101-07', 'name': 'Демонтаж громкоговорителя в помещении'}
    assert 'row_work_is_removal' not in _removal_receipt('Громкоговоритель или звуковая колонка: в помещении')


def test_removal_hint_names_the_supply_operation_not_the_work():
    """Проход 23: модель отвергла саму работу-демонтаж, а изделие оставила.

    Прежний текст «убери эту операцию» не называл, какую именно. Чек говорит про
    операцию поставки и прямо снимает требование отвергать работу строки.
    """
    from proxy.smeta_core.calculation_receipt import receipt_context

    text = receipt_context([{'work_id': 'x'}])
    assert 'row_work_is_removal' in text
    assert 'операцию поставки' in text
    assert 'работу строки при этом отвергать не нужно' in text.casefold()


# --- Проход 24: подтверждена операция без шифра нормы -------------------------------------
# Модель убрала изделие при демонтаже, но оставила вторую работу «Снятие оповещателя» без
# norm_code и подтвердила её по чеку. Расчёт всё равно отверг строку: norm_selection_required.
# Подтверждать нечего, пока нормы нет, и чек называет это прямо.

def _receipt_for(norm_code, resource_bindings):
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': norm_code, '3': 'Снятие оповещателя Глагол-Н1-1'},
             'meta': {'work_id': 'w2', 'source_row': 'СОУЭ!7.2'}}]}]}],
        'row_bindings': [{'row': 'СОУЭ!7.2', 'work_id': 'w2', 'status': 'resource_review_required',
                          'resource_bindings': resource_bindings}],
    }
    return calculation_receipts(trace)[0]


def test_receipt_names_an_operation_left_without_a_norm_code():
    assert _receipt_for('', [])['norm_code_missing'] is True
    assert 'norm_code_missing' not in _receipt_for('ГЭСНм10-04-101-07', [])
    # Самостоятельный материал шифра нормы не имеет по своей природе — это не пробел.
    assert 'norm_code_missing' not in _receipt_for('', [{'action': 'standalone', 'resource_code': '61.2.04.05-0004'}])


def test_missing_norm_code_is_explained_in_the_receipt_instruction():
    from proxy.smeta_core.calculation_receipt import receipt_context

    text = receipt_context([{'work_id': 'x'}])
    assert 'norm_code_missing' in text
    assert 'подтверждение такой операции строку не закроет' in text.casefold()


def test_a_supply_operation_is_not_scolded_for_having_no_norm_code():
    """Проход 28: подсказка про пробел нормы прилипла к поставке и сбила решение.

    В 27-м проходе модель по этой строке ответила «Демонтаж; изделие не поставляется»,
    а в 28-м вернула пустую поставку без решения: её операция получила norm_code_missing
    с советом «назови norm_code или убери операцию», хотя у поставки нормы и не бывает.
    """
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    trace = {
        'sections': [{'positions': [{'summary': {}, 'rows': [
            {'type': 'work', 'columns': {'2': '', '3': 'Поставка оповещателя настенного'},
             'meta': {'work_id': 'w2', 'source_row': 'СОУЭ!7.2'}}]}]}],
        'row_bindings': [{'row': 'СОУЭ!7.2', 'work_id': 'w2', 'status': 'material_price_missing',
                          'resource_bindings': []}],
    }
    assert 'norm_code_missing' not in calculation_receipts(trace)[0]


@pytest.mark.parametrize('confirmed', [False, True])
@pytest.mark.parametrize('reason', ['LSZH/SOLID не подтверждены: модель выбрала аналог.', '=LSZH/SOLID не подтверждены'])
def test_material_caveat_reaches_xlsx_without_granting_confirmation(tmp_path, confirmed, reason):
    from copy import deepcopy
    import openpyxl
    from proxy.services.fgis_price_service import PriceBook
    from proxy.smeta_core.calculation_receipt import calculation_receipts, apply_model_review

    choice = dict(source_row='fixture!1', work_id='supply', title='Fixture cable LSZH',
                  unit='м', quantity=2, norm_code='', decision_kind='материал',
                  resource_code='21.1.04.01-1046', material_review_status='confirmed' if confirmed else '',
                  material_review_reason=reason, resource_review_status='keep_all_confirmed')
    book = PriceBook([dict(code=choice['resource_code'], name='Fixture cable', unit='м', price_current=100)])
    trace = rim.build_lsr_trace_from_visible_rows([choice], pricebook=book)
    receipts = calculation_receipts(trace, pricebook=book)
    choice['resource_review_ref'] = receipts[0]['calculation_ref'] if confirmed else 'wrong-receipt'
    original = deepcopy(choice)
    apply_model_review(trace, receipts, [choice])
    assert trace['row_bindings'][0]['status'] == ('material_bound' if confirmed else 'resource_review_required')
    assert receipts[0]['calculation']['summary']['known_amount'] == 200
    path = render_lsr_xlsx(trace, tmp_path/'caveat.xlsx')
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        sheet = wb['Проверка']
        row = next(cells for cells in sheet.iter_rows(min_row=2) if cells[1].value == 'supply')
        assert reason in str(row[11].value)
        assert row[11].data_type == 's'
        assert row[2].value == trace['row_bindings'][0]['status']
    finally:
        wb.close()
    assert choice == original


@pytest.mark.parametrize('catalog_name', ['Керамические плитки из справочника', '=Каталожное название'])
def test_supply_xlsx_discloses_catalog_name_without_changing_model_choice(tmp_path, catalog_name):
    from copy import deepcopy
    import openpyxl
    from proxy.services.fgis_price_service import PriceBook
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    choice = dict(source_row='S!1', work_id='supply', title='Ковровая плитка проекта',
                  unit='м2', quantity=2, norm_code='', decision_kind='материал',
                  resource_code='06.2.02.01-0002')
    original = deepcopy(choice)
    book = PriceBook([dict(code=choice['resource_code'], name=catalog_name, unit='м2', price_current=100)])
    trace = rim.build_lsr_trace_from_visible_rows([choice], pricebook=book)
    receipts = calculation_receipts(trace, pricebook=book)
    without_display_metadata = deepcopy(trace)
    for section in without_display_metadata['sections']:
        for position in section['positions']:
            for row in position['rows']:
                (row.get('meta') or {}).pop('price_card_name', None)
    assert calculation_receipts(without_display_metadata, pricebook=book) == receipts
    assert receipts[0]['calculation']['summary']['known_amount'] == 200
    path = render_lsr_xlsx(trace, tmp_path/'price-source.xlsx')
    wb = openpyxl.load_workbook(path, data_only=False)
    try:
        rows = list(wb['ЛСР РИМ'].iter_rows())
        row = next(cells for cells in rows if cells[1].value == choice['resource_code'])
        assert catalog_name in str(row[2].value)
        assert choice['title'] in str(row[2].value)
        assert row[2].data_type == 's'
    finally:
        wb.close()
    assert choice == original
