"""Real renderer/trace/export regression; no model, book or normative DB is read."""

import copy
import csv
from pathlib import Path

import openpyxl
import pytest

from proxy.services import gesn_service, rim_lsr_trace_service, smeta_artifact_service
from proxy.smeta_core import application


@pytest.fixture
def literal_trace_boundary(monkeypatch):
    calls = []

    def calculate(rows, **_kwargs):
        # The compatibility application wrapper persists historical revisions.
        # Exercise its actual downstream trace builder without that persistence.
        before = copy.deepcopy(rows)
        result = rim_lsr_trace_service.build_lsr_trace_from_visible_rows(rows)
        assert rows == before
        calls.append(copy.deepcopy(result))
        return result

    monkeypatch.setattr(application, "calculate_visible_rows_revision", calculate)
    monkeypatch.setattr(smeta_artifact_service, "_select_pricebook_for_question", lambda _question: (None, ""))
    monkeypatch.setattr(gesn_service, "get_norm", lambda *_args, **_kwargs: pytest.fail("conflict renderer must not read or replace a norm"))
    return calls


def authored_rows(code, quantity):
    return [{"source_row": 1, "row_role": "material", "title": "Authored separate supply",
             "unit": "шт.", "quantity": 1, "confirmed_price": 5,
             "price_kind": "unit_price", "price_source": "literal source"},
            {"source_row": 2, "row_role": "work", "title": "Authored operation",
             "norm_code": code, "norm_status": "ANALOGUE_CANDIDATE", "unit": "шт.",
             "quantity": quantity, "_evidence_check": {"normative_snapshot_conflict": True}}]


@pytest.mark.parametrize("code,status", [
    ("ГЭСН01-01-001-01", "normative_generation_conflict"),
    ("01-01-001-01", "norm_identity_conflict"),
])
@pytest.mark.parametrize("quantity", [0, ""])
def test_conflict_form_and_physical_exports_keep_authored_cipher_and_explicit_zero(
    literal_trace_boundary, tmp_path, code, status, quantity,
):
    rows = authored_rows(code, quantity)
    before = copy.deepcopy(rows)
    form = smeta_artifact_service._build_rim_trace_form({"title": "Literal rows", "rows": rows})
    assert form is not None
    assert form["rows"][1]["quantity"] == quantity
    assert form["rows"][1]["basis"] == code
    assert form["rows"][1]["status"] == status
    assert form["rows"][1]["amount"] == 0
    assert form["trace"]["row_bindings"][1]["source_row"]["norm_code"] == code
    assert form["amount_total"] == literal_trace_boundary[0]["summary"]["total"]
    assert rows == before

    artifact = smeta_artifact_service.build_smeta_artifact_from_rim_form(form)
    expected_quantity = str(quantity)
    assert artifact["tables"][0]["data"][1][1] == code
    assert artifact["tables"][0]["data"][1][6] == expected_quantity
    rendered_row = next(line for line in artifact["content"].splitlines() if line.startswith(f"| 2 | {code} |"))
    assert rendered_row.split("|")[5].strip() == expected_quantity
    assert rendered_row.split("|")[7].strip() == expected_quantity

    exported = smeta_artifact_service.persist_smeta_artifact_exports(artifact, output_dir=tmp_path)
    workbook = openpyxl.load_workbook(tmp_path / exported["downloads"]["xlsx"].split("path=")[1], data_only=True)
    try:
        sheet = workbook["ЛСР РИМ"]
        assert sheet.cell(18, 2).value == code
        assert sheet.cell(18, 5).value == (None if quantity == "" else 0)
        assert sheet.cell(18, 7).value == (None if quantity == "" else 0)
    finally:
        workbook.close()
    csv_path = Path(tmp_path / exported["downloads"]["csv"].split("path=")[1])
    with csv_path.open(encoding="utf-8-sig", newline="") as source:
        csv_rows = list(csv.reader(source, delimiter=";"))
    assert csv_rows[4][1] == code
    assert csv_rows[4][4] == expected_quantity
    assert csv_rows[4][6] == expected_quantity


def test_no_authored_cipher_keeps_existing_unselected_placeholder(literal_trace_boundary):
    rows = authored_rows("", 0)
    rows[1].pop("_evidence_check")
    form = smeta_artifact_service._build_rim_trace_form({"rows": rows})
    assert form["rows"][1]["basis"] == "нужен подбор нормы"
    assert form["rows"][1]["status"] == "norm_selection_required"
