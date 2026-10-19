import json

import openpyxl
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from backend import converter
from proxy.services import attachment_source_catalog_service as catalog_service
from proxy.services import workbook_tool_service as workbook_service
from proxy.services.artifact_revision_service import ArtifactRevisionStore
from proxy.services.chat_attachment_service import preserve_read_attachment
from proxy.services.workflow_checkpoint_service import WorkflowCheckpointService


def _table(document, name="Исходный ресурс", quantity="0"):
    table = document.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Наименование", "Ед. изм.", "Кол-во")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, (name, "м", quantity)):
        cell.text = text
    return table


def _context(tmp_path):
    return workbook_service.WorkbookExecutionContext(
        session_id="transport", idempotency_key="transport-key", model_decision_revision="decision",
        profile_revision_id="profile", model_identity="fixture", model_preset="fixture",
        attachment_root=tmp_path / "attachments", work_dir=tmp_path / "work",
        checkpoints=WorkflowCheckpointService(tmp_path / "checkpoints.db"),
        artifacts=ArtifactRevisionStore(tmp_path / "artifacts.db", tmp_path / "artifacts"),
    )


def _assert_printed_vor(workbook, expected_work_values, decisions):
    assert workbook.active.title == "Ведомость работ" and workbook.active.max_column == 5
    printed = [values for values in workbook.active.iter_rows(min_row=9, values_only=True)
               if isinstance(values[0], int) and not isinstance(values[0], bool)]
    assert [(values[1], values[2], values[3]) for values in printed] == expected_work_values
    assert workbook.active["A6"].value == workbook["ВОР"]["A2"].value
    assert [json.loads(row[2].removeprefix("JSON: "))
            for row in workbook["Решения модели"].iter_rows(min_row=4, values_only=True)] == decisions


def test_pdf_catalog_bypasses_preview_cap_without_losing_later_rows(tmp_path, monkeypatch):
    path = tmp_path / "Источник.pdf"
    path.write_bytes(b"%PDF-fixture")
    long_name = "Длинное исходное условие " * 80 + "END-CELL"
    markdown = (
        "| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n"
        f"| {long_name} | м | 0 |\n| Последний ресурс | м | 17 |\n"
    )
    monkeypatch.setattr(converter, "_parse_pdf", lambda _path, route=None: markdown)
    monkeypatch.setenv("RAG_PDF_MAX_FILE_CHARS", "100")

    assert len(converter.convert_to_markdown(path)) == 100
    assert converter.convert_to_markdown(path, apply_char_limit=False) == markdown
    catalog = catalog_service.load_tabular_source_catalog(path)

    assert len(catalog.rows) == 3
    assert catalog.lookup("PDF!R2").values[0] == long_name
    assert "END-CELL" in catalog.lookup("PDF!R2").line
    assert catalog.lookup("PDF!R2").header_quantity == 0
    assert catalog.lookup("PDF!R3").header_quantity == 17


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_exact_catalog_keeps_zero_and_full_cells_after_480_chars(tmp_path, suffix):
    path = tmp_path / f"Источник{suffix}"
    long_name = "x" * 700 + "END-CELL"
    if suffix == ".csv":
        path.write_text(f"Наименование;Ед. изм.;Кол-во\n{long_name};м;0\n", encoding="utf-8")
        locator = "CSV!R2"
    else:
        workbook = openpyxl.Workbook()
        workbook.active.append(["Наименование", "Ед. изм.", "Кол-во"])
        workbook.active.append([long_name, "м", 0])
        workbook.save(path)
        workbook.close()
        locator = "Sheet!R2"
    catalog = catalog_service.load_tabular_source_catalog(path)
    row = catalog.lookup(locator)

    assert row.values == (long_name, "м", "0")
    assert "0 [Кол-во]" in row.line and "END-CELL" in row.line
    assert row.header_quantity == 0
    assert catalog_service.compare_model_row_to_catalog(
        {"source_row": locator, "quantity": 0, "unit": "м"}, catalog,
    ) == []


def test_docx_catalog_keeps_all_tables_text_parts_and_stable_identity(tmp_path):
    path = tmp_path / "Спецификация.docx"
    document = Document()
    document.add_paragraph("Исходные условия объекта перед таблицей")
    _table(document, "Первый ресурс", "0")
    document.add_paragraph("Условия следующей таблицы")
    second = document.add_table(rows=2, cols=3)
    for cell, text in zip(second.rows[0].cells, ("Кол-во", "Наименование", "Ед. изм.")):
        cell.text = text
    for cell, text in zip(second.rows[1].cells, ("12", "Второй ресурс", "шт.")):
        cell.text = text
    document.sections[0].header.paragraphs[0].text = "Исходное условие в колонтитуле"
    document.save(path)

    catalog = catalog_service.load_tabular_source_catalog(path)
    again = catalog_service.load_tabular_source_catalog(path)

    assert catalog.extraction_complete is True
    assert catalog.lookup("DOCX_T1!R2").header_quantity == 0
    assert catalog.lookup("DOCX_T2!R2").header_quantity == 12
    assert catalog.lookup("DOCX_TEXT!R1").values == ("Исходные условия объекта перед таблицей",)
    assert "DOCX_TEXT!R1" in catalog.work_locators()  # role stays model-owned
    assert any("колонтитуле" in row.line for row in catalog.rows)
    lines = catalog_service.catalog_sheet_pack_lines(catalog, "DOCX_T1", ["DOCX_T1!R2"])
    facts = catalog_service.catalog_sheet_packet_facts(catalog, "DOCX_T1", ["DOCX_T1!R2"])
    assert "DOCX_TEXT!R1: Исходные условия объекта перед таблицей" in lines
    assert "DOCX_TEXT!R1" in facts["reference_only_source_rows"]
    assert facts["requested_source_rows"] == ["DOCX_T1!R2"]
    assert [row.row_id for row in catalog.rows] == [row.row_id for row in again.rows]
    assert len({row.locator for row in catalog.rows}) == len(catalog.rows)


def test_docx_nested_table_is_not_silently_ignored(tmp_path):
    path = tmp_path / "Вложенная таблица.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "Контейнер таблицы"
    nested = outer.cell(0, 0).add_table(rows=2, cols=3)
    for cell, text in zip(nested.rows[0].cells, ("Наименование", "Ед. изм.", "Кол-во")):
        cell.text = text
    for cell, text in zip(nested.rows[1].cells, ("Вложенный ресурс", "м", "5")):
        cell.text = text
    document.save(path)

    catalog = catalog_service.load_tabular_source_catalog(path)

    assert catalog.lookup("DOCX_T2!R2").header_quantity == 5
    assert "Вложенный ресурс" in catalog.lookup("DOCX_T2!R2").line


def test_docx_omitted_leading_grid_cell_does_not_shift_the_quantity_column(tmp_path):
    path = tmp_path / "Неполная строка.docx"
    document = Document()
    table = _table(document, quantity="5")
    row_xml = table.rows[1]._tr
    row_xml.remove(row_xml.tc_lst[0])
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), "1")
    row_xml.get_or_add_trPr().append(grid_before)
    document.save(path)

    catalog = catalog_service.load_tabular_source_catalog(path)
    row = catalog.lookup("DOCX_T1!R2")

    assert row.values == ("", "м", "5")
    assert row.header_quantity == 5
    assert catalog_service.header_bound_unit(row) == "м"


def test_catalog_preserves_later_hidden_sheet_after_a_nonwork_first_sheet(tmp_path):
    path = tmp_path / "Все листы.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Описание"
    workbook.active.append(["Условия объекта"])
    later = workbook.create_sheet("Позиции")
    later.sheet_state = "hidden"
    later.append(["Наименование", "Ед. изм.", "Кол-во"])
    later.append(["Поздний ресурс", "м", 7])
    workbook.save(path)
    workbook.close()

    catalog = catalog_service.load_tabular_source_catalog(path)

    assert catalog.sheets == (("Описание", 1), ("Позиции", 2))
    assert catalog.work_locators() == ["Позиции!R2"]
    assert catalog.lookup("Позиции!R2").hidden_sheet is True
    assert catalog.lookup("Позиции!R2").header_quantity == 7


def test_separate_conditions_sheet_is_reference_context_without_confirming_blank_quantity(tmp_path):
    path = tmp_path / "Условия и позиции.xlsx"
    workbook = openpyxl.Workbook()
    conditions = workbook.active
    conditions.title = "Условия"
    conditions.append(["Дословные исходные условия объекта"])
    positions = workbook.create_sheet("Позиции")
    positions.append(["№", "Наименование", "Ед. изм.", "Кол-во"])
    positions.append(["1.2", "Исходный ресурс", "шт.", None])
    workbook.save(path)
    workbook.close()
    catalog = catalog_service.load_tabular_source_catalog(path)

    lines = catalog_service.catalog_sheet_pack_lines(catalog, "Позиции", ["Позиции!R2"])
    facts = catalog_service.catalog_sheet_packet_facts(catalog, "Позиции", ["Позиции!R2"])

    assert "Условия!R1: Дословные исходные условия объекта" in lines
    assert facts["requested_source_rows"] == ["Позиции!R2"]
    assert facts["reference_only_source_rows"] == ["Условия!R1"]
    assert "Условия!R1" in facts["reference_source_rows"]
    assert facts["source_quantity_cells"] == [{
        "source_row": "Позиции!R2", "quantity_column_known": True, "quantity": None, "unit": "шт.",
    }]
    assert "не расширяет requested_source_rows" in facts["source_scope_instruction"]


@pytest.mark.parametrize("tag", ["drawing", "txbxContent", "altChunk", "ins"])
def test_docx_unread_embedded_content_is_independent_of_the_model_decision(tmp_path, tag):
    path = tmp_path / "Смешанный документ.docx"
    document = Document()
    paragraph = document.add_paragraph("Прочитанное условие")
    paragraph.add_run()._r.append(OxmlElement(f"w:{tag}"))
    _table(document, quantity="8")
    document.save(path)

    catalog = catalog_service.load_tabular_source_catalog(path)
    decision = {"source_row": "DOCX_T1!R2", "quantity": 8, "unit": "м"}

    assert catalog.extraction_complete is False
    assert any(tag in warning for warning in catalog.extraction_warnings)
    assert "source_extraction_complete: false" in catalog.manifest_text()
    assert catalog_service.compare_model_row_to_catalog(decision, catalog) == [
        "source_extraction_incomplete:DOCX_T1!R2",
    ]
    assert decision == {"source_row": "DOCX_T1!R2", "quantity": 8, "unit": "м"}
    assert "source_extraction_incomplete" in catalog_service.coverage_errors(
        [{"source_row": row.locator} for row in catalog.rows], catalog,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix,locator", [(".docx", "DOCX_T1!R2"), (".csv", "CSV!R2")])
async def test_authored_docx_csv_vor_reaches_export_without_code_choosing_work(tmp_path, monkeypatch, suffix, locator):
    source = tmp_path / f"Исходник{suffix}"
    if suffix == ".docx":
        document = Document()
        _table(document)
        document.save(source)
    else:
        source.write_text("Наименование;Ед. изм.;Кол-во\nИсходный ресурс;м;0\n", encoding="utf-8")
    meta = preserve_read_attachment(source, attachment_id="read_123456abcdef", original_name=source.name,
                                    root=tmp_path / "attachments")
    context = _context(tmp_path)
    decision = {"source_row": locator, "row_role": "work", "title": "Работа, выбранная моделью",
                "quantity": 0, "unit": "м"}
    monkeypatch.setattr("proxy.services.spec_to_bor_service.derive_work_for_item",
                        lambda *_a, **_k: pytest.fail("no professional choice in authored export"))

    result = await workbook_service.build_vor_workbook(
        {"attachment_id": meta["attachment_id"], "decisions": [decision]}, context,
    )

    assert result["status"] == "complete"
    assert result["vor_source_errors"] == []
    assert result["source"]["sha256"] == meta["sha256"]
    workbook = openpyxl.load_workbook(context.artifacts.resolve_path(result["artifact"]["revision_id"]), data_only=True)
    assert workbook["ВОР"]["G4"].value == f"{source.name}#{locator}"
    assert workbook["ВОР"]["C4"].value == decision["title"]
    assert workbook["ВОР"]["F4"].value == 0
    _assert_printed_vor(workbook, [(decision["title"], decision["unit"], 0)], [decision])
    workbook.close()


@pytest.mark.asyncio
async def test_header_only_draft_cannot_hide_unread_docx_content(tmp_path):
    source = tmp_path / "Заголовок.docx"
    document = Document()
    document.add_paragraph("Заголовок").add_run()._r.append(OxmlElement("w:drawing"))
    document.save(source)
    target = tmp_path / "result.xlsx"

    result = await workbook_service._default_vor_adapter(
        source, {"decisions": [{"source_row": "DOCX_TEXT!R1", "row_role": "header", "title": "Заголовок"}]},
        target, lambda *_: None,
    )

    assert result["vor_source_errors"] == ["source_extraction_incomplete"]
    workbook = openpyxl.load_workbook(target, data_only=True)
    assert "ЧЕРНОВИК" in workbook["ВОР"]["A2"].value
    assert workbook["ВОР"]["C4"].value == "Заголовок"
    assert workbook["ВОР"]["F4"].value is None
    _assert_printed_vor(workbook, [], [{"source_row": "DOCX_TEXT!R1", "row_role": "header", "title": "Заголовок"}])
    workbook.close()


@pytest.mark.asyncio
async def test_empty_source_does_not_publish_an_empty_vor_artifact(tmp_path):
    source = tmp_path / "Пустой.xlsx"
    workbook = openpyxl.Workbook()
    workbook.save(source)
    workbook.close()
    meta = preserve_read_attachment(source, attachment_id="read_123456abcdef", original_name=source.name,
                                    root=tmp_path / "attachments")
    context = _context(tmp_path)

    result = await workbook_service.build_vor_workbook({"attachment_id": meta["attachment_id"]}, context)

    assert result["status"] == "rejected"
    assert result["code"] == "MODEL_WORK_ROWS_REQUIRED"
    assert "checkpoint" not in result
    assert "artifact" not in result
    assert not list(context.work_dir.rglob("*.xlsx"))
    from proxy.services.spec_to_bor_service import rows_from_spec_xlsx
    assert rows_from_spec_xlsx(source) == []


@pytest.mark.asyncio
async def test_authored_header_from_empty_source_keeps_row_and_exposes_source_gap(tmp_path):
    source = tmp_path / "Пустой исходник.xlsx"
    workbook = openpyxl.Workbook()
    workbook.save(source)
    workbook.close()
    target = tmp_path / "Авторская ВОР.xlsx"
    decision = {"source_row": "Sheet!R1", "row_role": "header", "title": "Авторский заголовок"}

    result = await workbook_service._default_vor_adapter(
        source, {"decisions": [decision]}, target, lambda *_: None,
    )

    assert result["vor_source_errors"] == ["source_catalog_empty"]
    workbook = openpyxl.load_workbook(target, data_only=True)
    assert workbook["ВОР"]["C4"].value == decision["title"]
    assert "ЧЕРНОВИК" in workbook["ВОР"]["A2"].value
    assert workbook["ВОР"]["F4"].value is None
    _assert_printed_vor(workbook, [], [decision])
    workbook.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix,status,code", [
    (".rtf", "failed", "WORKBOOK_GENERATION_FAILED"),
    (".zip", "rejected", "UNSUPPORTED_ATTACHMENT_TYPE"),
])
async def test_malformed_rtf_and_unsupported_zip_fail_closed_even_with_decisions(tmp_path, suffix, status, code):
    source = tmp_path / f"Источник{suffix}"
    source.write_bytes(b"unsupported fixture")
    meta = preserve_read_attachment(source, attachment_id="read_123456abcdef", original_name=source.name,
                                    root=tmp_path / "attachments")

    result = await workbook_service.build_vor_workbook(
        {"attachment_id": meta["attachment_id"], "decisions": [{"source_row": "X!R1", "title": "Авторская строка"}]},
        _context(tmp_path),
    )

    assert result["status"] == status
    assert result["code"] == code
    if suffix == ".rtf":
        assert result["error_type"] == "RtfParserError"
        assert result["message"] == "RTF_INVALID_HEADER"
