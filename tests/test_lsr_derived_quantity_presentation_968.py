"""A cited calculation must not masquerade as the physical source volume."""
from copy import deepcopy

from openpyxl import load_workbook

from proxy.services.model_lsr_result_contract import OPTIONAL_FIELDS, model_lsr_result_field_instruction
from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
from tests.test_rim_lsr_trace_service import _active_resource_norm


def test_lsr_contract_explains_repeated_norm_without_a_second_coefficient_protocol():
    text = model_lsr_result_field_instruction()
    assert "Кратность: calc_quantity/calc_unit" in text
    assert "формула: conversion_source" in text
    assert "calc_quantity" in text and "calc_unit" in text and "conversion_source" in text
    assert "Исходные unit и quantity сохраняй" in text
    assert "covered_by_work_id" in text
    assert "Добавка — новая операция с norm_code/evidence_refs, без covered_by_work_id" in text
    assert "norm_components — аудит без расчёта" in text
    assert "можно добавить norm_components" not in text
    assert all(field in text for field in OPTIONAL_FIELDS)


def test_printed_lsr_explains_source_and_calculated_volume(monkeypatch, tmp_path):
    row = _active_resource_norm(monkeypatch)
    row.update(calc_quantity=3500, calc_unit="м", conversion_source="S!R2: 1000 м × 3,5 повторения по Q1.H1")
    original = deepcopy(row)
    trace = build_lsr_trace_from_visible_rows([row])
    output = render_lsr_xlsx(trace, tmp_path / "расчёт с основанием.xlsx")
    wb = load_workbook(output, data_only=True)
    body = "\n".join(str(cell.value or "") for line in wb.worksheets[0] for cell in line)
    assert "Исходный объём: 1000 м" in body
    assert "расчётный эквивалент" in body
    assert "3500 м" in body
    assert row["conversion_source"] in body
    note_row = next(line[0].row for line in wb.worksheets[0] if any("Исходный объём:" in str(cell.value) for cell in line))
    assert wb.worksheets[0].row_dimensions[note_row].height >= 30
    assert row == original


def test_unconverted_work_does_not_invent_a_derived_volume(monkeypatch, tmp_path):
    row = _active_resource_norm(monkeypatch)
    trace = build_lsr_trace_from_visible_rows([row])
    output = render_lsr_xlsx(trace, tmp_path / "обычный объём.xlsx")
    wb = load_workbook(output, data_only=True)
    body = "\n".join(str(cell.value or "") for line in wb.worksheets[0] for cell in line)
    assert "расчётный эквивалент" not in body
