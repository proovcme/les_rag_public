"""A printable work statement is a checked view of the unchanged decision register."""
from __future__ import annotations

from copy import deepcopy
import json

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
import pytest

from proxy.services.attachment_source_catalog_service import catalog_sheet_pack_lines, load_tabular_source_catalog
from proxy.services.bor_service import source_rows_to_vor_xlsx
from proxy.routers import chat
from proxy.services.openai_compatible_transport_service import InferenceResponse
from proxy.services.workbook_research_job_service import WorkbookResearchJobService
from tests.test_model_rag_sheet_evidence_reads import _run_sheet


def _render(tmp_path, decisions, **kwargs):
    rows = [{"section": item.get("section", ""), "name": item["title"],
             "unit": item.get("unit", ""), "qty": item.get("quantity"),
             "source_file": item.get("source_file", "Раздел 'Ω'.xlsx"),
             "pos": item["source_row"], "row_role": item.get("row_role", ""),
             "work_row": item.get("work_row", ""), "open_gap": item.get("open_gap", ""),
             "_authored_decision": deepcopy(item),
             "_derived_quantity": {key: item.get(key) for key in (
                 "calc_unit", "calc_quantity", "conversion_factor", "conversion_source")}
                 if any(item.get(key) is not None for key in ("calc_unit", "calc_quantity")) else {}}
            for item in decisions]
    path = tmp_path / "Ведомость 'Ω'.xlsx"
    assert source_rows_to_vor_xlsx(rows, path, **kwargs) == len(decisions)
    return path


def _works(sheet):
    return [row for row in sheet.iter_rows() if row[1].hyperlink is not None]


def test_five_column_print_view_preserves_zero_unknowns_and_literal_register(tmp_path):
    decisions = [
        {"source_row": "A!R4", "row_role": "work", "title": "Работа с нулём", "unit": "шт.",
         "quantity": 0, "operation_id": "op-0", "section": "Раздел 1. Монтаж"},
        {"source_row": "B!R8", "row_role": "material", "title": "Изделие", "unit": "шт.",
         "quantity": 0, "work_row": "op-0"},
        {"source_row": "C!R20", "row_role": "work", "title": "Другая работа", "unit": "",
         "quantity": None, "open_gap": "Уточнить геометрию", "section": "Раздел 2. Отделка"},
    ]
    before = deepcopy(decisions)
    path = _render(tmp_path, decisions, draft=True)
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        assert book.sheetnames[0] == "Ведомость работ"
        view = book.active
        assert view.max_column == 5
        assert [cell.value for cell in view[7]] == ["№ пп", "Наименование", "Ед. изм.", "Кол-во", "Примечание"]
        works = _works(view)
        assert [(row[1].value, row[2].value, row[3].value) for row in works] == [
            ("Работа с нулём", "шт.", 0), ("Другая работа", None, None)]
        assert "Уточнить геометрию" in works[1][4].value
        assert "Объём не указан" in works[1][4].value
        assert "Единица измерения не указана" in works[1][4].value
        assert "НЕ ЗАВЕРШЁН" in view["A6"].value
        assert view.page_setup.paperSize == int(view.PAPERSIZE_A4)
        assert view.page_setup.orientation == "portrait"
        assert view.page_setup.fitToWidth == 1 and view.page_setup.fitToHeight == 0
        assert view.print_title_rows == "$7:$8"
        assert [row[2].value for row in book["ВОР"].iter_rows(min_row=4)] == [item["title"] for item in decisions]
        audit = [json.loads(row[2].value.removeprefix("JSON: "))
                 for row in book["Решения модели"].iter_rows(min_row=4)]
        assert audit == decisions == before
    finally:
        book.close()


def test_print_view_uses_only_complete_authored_conversion_without_recalculating(tmp_path):
    decisions = [
        {"source_row": "A!R1", "row_role": "work", "title": "Полный перевод", "unit": "м",
         "quantity": 4, "calc_unit": "м²", "calc_quantity": 12, "conversion_factor": 3,
         "conversion_source": "4 м × 3 м = 12 м²; проект A!R1"},
        {"source_row": "A!R2", "row_role": "work", "title": "Неполный перевод", "unit": "кг",
         "quantity": 0, "calc_unit": "т", "calc_quantity": None},
    ]
    book = openpyxl.load_workbook(_render(tmp_path, decisions), data_only=True)
    try:
        works = _works(book["Ведомость работ"])
        assert [(row[2].value, row[3].value) for row in works] == [("м²", 12), ("кг", 0)]
        assert decisions[0]["conversion_source"] in works[0][4].value
        assert [book["ВОР"].cell(row, 6).value for row in (4, 5)] == [4, 0]
    finally:
        book.close()


@pytest.mark.parametrize("count", [1, 37, 318])
def test_any_work_count_and_source_files_roundtrip_without_duplicate_positions(tmp_path, count):
    decisions = [{"source_row": f"member-{index % 3}::Лист!R{index + 2}", "row_role": "work",
                  "title": f"Операция {index}", "quantity": index / 2, "unit": "м²",
                  "section": f"Раздел {index // 11}", "operation_id": f"op-{index}",
                  "source_file": f"Часть {index % 3}.xlsx"}
                 for index in range(count)]
    path = _render(tmp_path, decisions)
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete and catalog.extraction_warnings == ()
    assert catalog.work_locators() == [f"ВОР!R{index + 4}" for index in range(count)]
    view_rows = [row for row in catalog.rows if row.sheet == "Ведомость работ"]
    assert view_rows and all(row.classification_reason in {
        "application_vor_print_view", "xlsx_literal_metadata"} for row in view_rows)
    assert all(not row.is_work_row() for row in view_rows)
    packet = catalog_sheet_pack_lines(catalog, "ВОР", ["ВОР!R4"])
    assert not any(line.startswith("Ведомость работ!R") for line in packet)
    assert any(line.startswith("ВОР!R4:") for line in packet)
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        assert len(_works(book["Ведомость работ"])) == count
        assert [book["ВОР"].cell(index + 4, 7).value for index in range(count)] == [
            f"Часть {index % 3}.xlsx#member-{index % 3}::Лист!R{index + 2}" for index in range(count)]
    finally:
        book.close()


def test_empty_draft_view_does_not_invent_a_work(tmp_path):
    path = _render(tmp_path, [], draft=True)
    book = openpyxl.load_workbook(path)
    try:
        assert book.active.title == "Ведомость работ"
        assert not _works(book.active)
        assert any("Работы не сформированы" in str(cell.value) for row in book.active for cell in row)
    finally:
        book.close()
    catalog = load_tabular_source_catalog(path)
    assert not any(row.sheet == "Ведомость работ" and row.is_work_row() for row in catalog.rows)


@pytest.mark.parametrize("fault", ["quantity", "extra_row", "source_link", "formula", "missing_work", "marker_version"])
def test_modified_print_view_is_never_silently_admitted_as_a_duplicate(tmp_path, fault):
    path = _render(tmp_path, [{"source_row": "A!R4", "row_role": "work", "title": "Операция",
                              "unit": "м", "quantity": 12}])
    book = openpyxl.load_workbook(path)
    view = book["Ведомость работ"]
    work = _works(view)[0]
    if fault == "quantity":
        work[3].value = 999
    elif fault == "extra_row":
        view.append([1, "Новая операция", "м", 77])
    elif fault == "source_link":
        work[1].hyperlink = "#'ВОР'!C999"
    elif fault == "formula":
        work[3].value = "=12"
    elif fault == "missing_work":
        view.delete_rows(work[0].row)
    else:
        del book.defined_names["_LES_VOR_PRINT_VIEW_V1"]
        book.defined_names.add(DefinedName("_LES_VOR_PRINT_VIEW_V99", attr_text="'Ведомость работ'!$A$1"))
    book.save(path)
    book.close()
    catalog = load_tabular_source_catalog(path)
    assert not catalog.extraction_complete
    assert "vor_print_view_invalid:Ведомость работ" in catalog.extraction_warnings
    assert all(row.classification_reason != "application_vor_print_view"
               for row in catalog.rows if row.sheet == "Ведомость работ")
    assert any(row.sheet == "Ведомость работ" and row.is_work_row() for row in catalog.rows)


def test_marker_on_arbitrary_sheet_does_not_hide_original_work(tmp_path):
    path = tmp_path / "Чужая ведомость.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Ведомость работ"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Настоящая исходная работа", "м", 27])
    book.defined_names.add(DefinedName("_LES_VOR_PRINT_VIEW_V1", attr_text="'Ведомость работ'!$A$1"))
    book.save(path)
    book.close()
    catalog = load_tabular_source_catalog(path)
    assert "Ведомость работ!R2" in catalog.work_locators()
    assert not catalog.extraction_complete
    assert "vor_print_view_invalid:Ведомость работ" in catalog.extraction_warnings


@pytest.mark.parametrize("role", ["работа", "works"])
def test_existing_work_role_aliases_are_presented_without_rewriting_raw_decisions(tmp_path, role):
    decisions = [{"source_row": "Лист!R7", "row_role": role, "title": "Авторская операция",
                  "unit": "м", "quantity": 2.5}]
    book = openpyxl.load_workbook(_render(tmp_path, decisions), data_only=True)
    try:
        assert [(row[1].value, row[3].value) for row in _works(book["Ведомость работ"])] == [
            ("Авторская операция", 2.5)]
        assert book["ВОР"]["J4"].value == role
        assert json.loads(book["Решения модели"]["C4"].value.removeprefix("JSON: ")) == decisions[0]
    finally:
        book.close()


@pytest.mark.asyncio
@pytest.mark.parametrize(("latest", "missing_original"), [
    ("Продолжи", False),
    ("Продолжи", True),
    ("Перепроверь связи материалов с работами по замечанию инженера", False),
], ids=["plain-original-title", "legacy-no-original-fallback", "nonplain-message-unchanged"])
async def test_actual_vor_resume_print_title_does_not_replace_latest_audit_message_977(
        tmp_path, monkeypatch, latest, missing_original):
    original = "Собери ВОР для объекта 'Ω'"
    rows = [
        {"source_row": source, "row_role": "work", "title": title,
         "unit": "m", "quantity": quantity}
        for source, title, quantity in (("S!R2", "Cable A", 12), ("S!R3", "Cable B", 3))
    ]
    run = await _run_sheet(
        tmp_path, monkeypatch, reads=False, question=original,
        query_declaration=json.dumps({"artifact_kind": "vor", "queries": []}),
        final_answer=json.dumps({"artifact_kind": "vor", "rows": rows}),
    )
    store = WorkbookResearchJobService()
    job_id = run.result["retrieval_trace"]["workbook_research_job"]["job_id"]
    before = store.get(job_id)
    assert before["task_context"]["original_question"] == original
    old_rows = deepcopy(before["accepted_rows"])
    calls_before = len(run.requests)
    executions_before = len(run.executions)
    contexts = []
    executor = chat._execute_chat_workbook_tool

    async def capture_context(call, context, progress):
        contexts.append(deepcopy(context))
        return await executor(call, context, progress)

    monkeypatch.setattr(chat, "_execute_chat_workbook_tool", capture_context)
    missing_observed = []
    if missing_original:
        getter = WorkbookResearchJobService.get

        def legacy_snapshot(self, current_id, **kwargs):
            snapshot = getter(self, current_id, **kwargs)
            snapshot = deepcopy(snapshot)
            snapshot["task_context"].pop("original_question", None)
            missing_observed.append(current_id)
            return snapshot

        monkeypatch.setattr(WorkbookResearchJobService, "get", legacy_snapshot)
    if latest != "Продолжи":
        # This is the existing natural review path, not a new header workflow.
        reviewed = [{**row, "expected_revision": row["revision"]} for row in old_rows]
        run.responses.extend([
            InferenceResponse(text='{"artifact_kind":"vor","queries":[]}',
                              tool_calls=(), finish_reason="stop", usage={}),
            InferenceResponse(text=json.dumps({"artifact_kind": "vor", "rows": reviewed}),
                              tool_calls=(), finish_reason="stop", usage={}),
        ])
    result = await run.resume_same_request(question=latest)
    assert len(run.requests) - calls_before == (0 if latest == "Продолжи" else 2)
    assert not run.responses
    resumed = run.executions[executions_before:]
    assert len(resumed) == 1
    call, payload = resumed[0]
    expected = original if latest == "Продолжи" and not missing_original else latest
    assert call["tool"] == "build_vor_workbook"
    assert call["args"]["question"] == expected
    assert contexts[0]["question"] == latest
    assert result["artifact"]["resume"]["question"] == latest
    if missing_original:
        assert missing_observed, "the fallback must exercise an absent original field"
    fields = ("source_row", "row_role", "title", "unit", "quantity", "operation_id")
    assert [{key: row.get(key) for key in fields} for row in call["args"]["decisions"]] == [
        {key: row.get(key) for key in fields} for row in old_rows]
    # Working copies share a minute-based filename. Inspect the exact delivered
    # immutable revision, never an older checkpoint's identically named copy.
    from hashlib import sha256
    output = next(tmp_path.rglob(payload["artifact"]["revision_id"] + ".xlsx"))
    assert sha256(output.read_bytes()).hexdigest() == payload["artifact"]["sha256"]
    book = openpyxl.load_workbook(output, data_only=False)
    try:
        assert book["Ведомость работ"]["A5"].value == expected
        assert book["ВОР"]["A1"].value == expected
    finally:
        book.close()
