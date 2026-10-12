"""Presentation must distinguish calculated previews from model-confirmed cost."""
import copy

import openpyxl
import pytest

from proxy.services import lsr_workbook_adapter_service as adapter
from tests.test_smeta_resource_receipt import norm, work


async def build(tmp_path, monkeypatch, rows, name, **args):
    monkeypatch.setattr(adapter.fgis_price_service, 'resolve_pricebook_path', lambda _: None)
    path = tmp_path/name
    result = await adapter.build_lsr_workbook_from_decisions(
        tmp_path/'source.xlsx', dict(decisions=rows, review_calculation=True, **args),
        path, lambda *args: None,
    )
    return result, openpyxl.load_workbook(path, data_only=False)


@pytest.mark.asyncio
async def test_question_is_not_workbook_title(norm, monkeypatch, tmp_path):
    instruction = 'Продолжай. INTERNAL_ACCEPTANCE_FEEDBACK\nПроверь evidence_refs и resource_review_ref.'
    rows = [work()]; before = copy.deepcopy(rows)
    result, wb = await build(tmp_path, monkeypatch, rows, 'title.xlsx',
                             question=instruction, _source_name='SKS-70.xlsx')
    assert 'SKS-70.xlsx' in wb['ЛСР РИМ']['A8'].value
    text = '\n'.join(str(c.value or '') for s in wb for row in s for c in row)
    assert 'INTERNAL_ACCEPTANCE_FEEDBACK' not in text
    assert rows == before and result['bound_rows'] == 0
    wb.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('confirmed_count', [0, 1])
async def test_unreviewed_amount_is_labelled_as_preview(norm, monkeypatch, tmp_path, confirmed_count):
    rows = [work(), work(source_row='S!2', work_id='second')]
    first, wb = await build(tmp_path, monkeypatch, rows, 'initial.xlsx')
    wb.close()
    if confirmed_count:
        rows[0].update(resource_review_ref=first['row_calculation'][0]['calculation_ref'],
                       resource_review_status='keep_all_confirmed')
    result, wb = await build(tmp_path, monkeypatch, rows, 'review.xlsx')
    assert result['bound_rows'] == confirmed_count
    ws = wb['ЛСР РИМ']
    labels = [str(ws.cell(r,1).value or '') for r in range(1,24)]
    assert not any('Стоимость только привязанной части' in s for s in labels)
    assert not any('Сумма в шапке — только по привязанным' in s for s in labels)
    label = 'Предварительная стоимость рассчитанной части'
    assert label in labels
    assert any(f'Требуют проверки модели: {2-confirmed_count} операций' in s for s in labels)
    assert ws.cell(labels.index(label)+1,4).data_type == 'f'
    assert len(result['row_calculation']) == 2
    # Review changes state/presentation, never the previously calculated amounts.
    assert [r['calculation']['summary'] for r in result['row_calculation']] == [
        r['calculation']['summary'] for r in first['row_calculation']]
    wb.close()
