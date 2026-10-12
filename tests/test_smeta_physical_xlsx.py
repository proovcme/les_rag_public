"""Physical conversion and coefficient scope through the production XLSX adapter."""
import asyncio

import openpyxl
import pytest

from proxy.services import rim_lsr_trace_service as rim
from proxy.services import lsr_workbook_adapter_service as adapter


@pytest.fixture
def norm(monkeypatch):
    data = {'name':'Synthetic operation', 'unit':'шт', 'nr_pct':0, 'sp_pct':0,
        'resources':[
            {'kind':'labor','code':'L','name':'Worker','unit':'чел.-ч','per_unit':2,'price':10},
            {'kind':'machine','code':'M','name':'Machine','unit':'маш.-ч','per_unit':1,'price':20},
            {'kind':'machinist','code':'D','name':'Driver','unit':'чел.-ч','per_unit':0.25,'price':5},
            {'kind':'material','code':'R','name':'Material','unit':'шт','per_unit':3,'price':2}]}
    monkeypatch.setattr(rim.gesn_service,'get_norm',lambda _:data)
    monkeypatch.setattr(rim.nr_sp_service,'resolve',lambda *args,**kwargs:{
        'status':'resolved_fixture','nr_pct':0,'sp_pct':0,'basis':'synthetic test rule'})
    monkeypatch.setattr(adapter.fgis_price_service,'resolve_pricebook_path',lambda _:None)
    return data


def row(**changes):
    return dict(dict(source_row='fixture',work_id='op',title='Synthetic only',
        norm_code='ГЭСН01-01-001-01',quantity=2,unit='шт'),**changes)


@pytest.mark.parametrize('quantity,unit,factor,target,norm_unit,expected', [
    (2,'упак',100,'шт','шт',200),
    (120,'бухт',305,'м','100 м',366),
    (4,'бухт',500,'м','100 м',20),
    (6,'м',3.03,'кг','т',0.01818),
    (2,'шт',None,None,'100 шт',0.02),
])
def test_physical_quantity_reaches_xlsx_without_second_division(norm,tmp_path,quantity,unit,factor,target,norm_unit,expected):
    norm['unit'] = norm_unit
    choice = row(quantity=quantity,unit=unit)
    if factor is not None:
        choice.update(conversion_factor=factor,conversion_unit=target,conversion_source='synthetic manufacturer fixture')
    out = tmp_path/'physical.xlsx'
    meta = asyncio.run(adapter.build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        {'decisions':[choice]},out,lambda *args:None))
    assert meta['bound_rows'] == 1
    wb = openpyxl.load_workbook(out)
    sheet = wb['ЛСР РИМ']
    work = next(c.row for cells in sheet for c in cells if c.comment and 'work_id: op' in c.comment.text)
    assert sheet.cell(work,5).value == pytest.approx(expected)
    assert sheet.cell(work,7).value == pytest.approx(expected)
    if expected == 0.01818:
        precision = sheet.cell(work,7).number_format.split('.')[-1]
        assert sum(precision.count(c) for c in ('0','#')) >= 5
    wb.close()


@pytest.mark.parametrize('factor',[0,-1,'NaN','Infinity'])
def test_invalid_conversion_cannot_create_position(norm,factor):
    bound = rim.positions_from_visible_lsr_rows([row(unit='упак',conversion_factor=factor,
        conversion_unit='шт',conversion_source='fixture')])
    assert not bound['positions']
    assert bound['row_bindings'][0]['status'] == 'conversion_basis_required'


def test_pack_size_is_not_guessed_from_product_title(norm):
    bound = rim.positions_from_visible_lsr_rows([row(quantity=12,title='Connector (100 шт)',unit='шт')])
    assert bound['positions'][0]['qty'] == 12


@pytest.mark.parametrize('quantity',['NaN','Infinity'])
def test_nonfinite_quantity_cannot_reach_excel(norm,quantity):
    bound = rim.positions_from_visible_lsr_rows([row(quantity=quantity)])
    assert not bound['positions']
    assert bound['row_bindings'][0]['status'] in {'invalid_quantity','missing_quantity'}


@pytest.mark.parametrize('target,expected', [
    ('labor',{'labor':6,'machine':2,'machinist':0.5,'material':6}),
    ('machines',{'labor':4,'machine':3,'machinist':0.75,'material':6}),
    ('materials',{'labor':4,'machine':2,'machinist':0.5,'material':9}),
])
def test_authored_coefficient_changes_only_declared_components(norm,tmp_path,target,expected):
    from proxy.services.chat_evidence_application_service import parse_model_rag_result
    choice=row(coefficient=1.5,coefficient_target=target,
               coefficient_basis='synthetic normative clause 4, model selected')
    raw='\n'.join(['|'+'|'.join(choice)+'|','|'+'|'.join('---' for _ in choice)+'|',
                   '|'+'|'.join(str(v) for v in choice.values())+'|'])
    original,parsed=parse_model_rag_result(raw)
    assert original == raw and parsed[0]['coefficient_target'] == target
    trace = rim.build_lsr_trace_from_visible_rows(parsed)
    assert trace['row_bindings'][0]['coefficient_status'] == 'applied'
    pos = trace['sections'][0]['positions'][0]
    values = {r['meta']['kind']:float(r['columns']['7']) for r in pos['rows']
              if r['type'].startswith('resource_')}
    assert values == expected
    assert pos['summary']['total'] == sum(expected[k]*price for k,price in
        {'labor':10,'machine':20,'machinist':5,'material':2}.items())
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
    out=tmp_path/'scope.xlsx'
    render_lsr_xlsx(trace,out)
    wb=openpyxl.load_workbook(out)
    sheet=wb['ЛСР РИМ']
    for kind,code in {'labor':'L','machine':'M','machinist':'D','material':'R'}.items():
        number=next(c.row for cells in sheet for c in cells if c.column==2 and c.value==code)
        assert sheet.cell(number,7).value == expected[kind]
    wb.close()


def test_unapplied_coefficient_cannot_produce_final_total(norm,tmp_path):
    out = tmp_path/'needs-basis.xlsx'
    meta = asyncio.run(adapter.build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        {'decisions':[row(coefficient=1.5)]},out,lambda *args:None))
    trace = rim.build_lsr_trace_from_visible_rows([row(coefficient=1.5)])
    assert trace['summary']['full_amount'] is None
    assert 'coefficient_basis_required' in trace['sections'][0]['positions'][0]['summary']['completeness_reasons']
    assert meta['bound_rows'] == 0
    assert meta['row_calculation'][0]['status'] == 'coefficient_basis_required'
    wb = openpyxl.load_workbook(out)
    assert any(c.comment and '1.5' in c.comment.text for s in wb for cells in s for c in cells)
    wb.close()


@pytest.mark.parametrize('basis,target',[('нет','labor'),('clause 4',''),('clause 4','all')])
def test_coefficient_needs_explicit_source_and_component_scope(norm,basis,target):
    trace=rim.build_lsr_trace_from_visible_rows([row(coefficient=1.5,
        coefficient_basis=basis,coefficient_target=target)])
    assert trace['row_bindings'][0]['coefficient_status']=='needs_basis'
    assert trace['summary']['bound_rows']==0


def test_legacy_rows_without_ids_keep_distinct_fallback_coordinates(norm):
    rows=[row(),row(norm_code='')]
    for choice in rows:
        choice.pop('source_row')
        choice.pop('work_id')
    trace=rim.build_lsr_trace_from_visible_rows(rows)
    assert trace['summary']['source_rows']==2
    assert trace['summary']['bound_rows']==1
    assert trace['summary']['unbound_rows']==1


def test_material_allocation_uses_authored_coefficient_once(norm):
    parent=row(coefficient=1.5,coefficient_target='materials',coefficient_basis='synthetic clause 4')
    resource=row(source_row='material',work_id='allocation',title='Material',norm_code='',quantity=9,
                 decision_kind='ресурс',included_in='op',resource_action='included',target_resource_code='R')
    result=rim.positions_from_visible_lsr_rows([parent,resource])
    assert result['row_bindings'][1]['status']=='resource_bound'
    assert len(result['positions'])==1
