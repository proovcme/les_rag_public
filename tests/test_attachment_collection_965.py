"""Collections preserve original files and source identities across the active path."""
from __future__ import annotations

import hashlib
import json
import zipfile
from dataclasses import replace
from pathlib import Path

import openpyxl
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from proxy.services.chat_attachment_collection_service import (
    AttachmentCollectionError, COLLECTION_METADATA_SCHEMA,
    create_read_attachment_collection,
)
from proxy.services.chat_attachment_service import preserve_read_attachment, resolve_read_attachment
from proxy.services.attachment_source_catalog_service import (
    catalog_sheet_pack_lines, catalog_sheet_packet_facts,
    load_catalog_from_attachment_id, load_tabular_source_catalog,
)
from proxy.services.model_source_frame_service import build_model_source_frame, render_model_source_frame


def _child(tmp_path, index, *, name='Лист "А".xlsx', values=None):
    source = tmp_path / f"source-{index}{Path(name).suffix}"
    if source.suffix == ".xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Раздел"
        sheet.append(["№", "Наименование", "Ед. изм.", "Количество"])
        for number, row in enumerate(values or [["Работа", "м", 0], ["Неизвестный объём", "шт", None]], 1):
            sheet.append([number, *row])
        book.save(source)
        book.close()
    else:
        source.write_text("literal unsupported source", encoding="utf-8")
    return preserve_read_attachment(source, attachment_id=f"read_{index:012x}",
        original_name=name, root=tmp_path / "attachments")


def _parent(tmp_path, children):
    metadata = create_read_attachment_collection([item["attachment_id"] for item in children],
        max_bytes=10_000_000, root=tmp_path / "attachments")
    path, loaded = resolve_read_attachment(metadata["attachment_id"], root=tmp_path / "attachments")
    catalog = load_tabular_source_catalog(path, trusted_collection_metadata=loaded)
    return path, metadata, catalog


def _rewrite_container(path, metadata, mutate):
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    manifest = json.loads(entries["manifest.json"])
    mutate(entries, manifest)
    entries["manifest.json"] = json.dumps(manifest, ensure_ascii=False).encode()
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_STORED) as archive:
        for name, body in entries.items():
            archive.writestr(name, body)
    return {**metadata, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "size": path.stat().st_size,
        "manifest_sha256": hashlib.sha256(entries["manifest.json"]).hexdigest()}


def test_collection_preserves_identical_files_as_distinct_sources_and_original_cells(tmp_path):
    first = _child(tmp_path, 1)
    second = preserve_read_attachment(tmp_path / "source-1.xlsx", attachment_id="read_000000000002",
        original_name='Другой "раздел".xlsx', root=tmp_path / "attachments")
    _, _, catalog = _parent(tmp_path, [second, first])
    assert catalog.extraction_complete and not catalog.extraction_warnings
    assert len(catalog.rows) == 6 and len(catalog.work_locators()) == 4
    assert catalog.sheets == (("read_000000000002::Раздел", 3), ("read_000000000001::Раздел", 3))
    original = load_tabular_source_catalog(tmp_path / "source-1.xlsx")
    for child_index, member in enumerate([second, first]):
        for original_row, row in zip(original.rows, catalog.rows[child_index * 3:(child_index + 1) * 3], strict=True):
            assert row.values == original_row.values and row.cells == original_row.cells
            assert row.header_quantity == original_row.header_quantity
            assert row.column_headers == original_row.column_headers
            assert row.locator == member["attachment_id"] + "::" + original_row.locator
            origin = catalog.row_origin(row.locator)
            assert origin["original_name"] == member["original_name"]
            assert origin["source_file_sha256"] == member["sha256"]
            assert origin["original_locator"] == original_row.locator
    assert catalog.lookup("Раздел!R2") is None and catalog.lookup(2) is None
    assert len({row.row_id for row in catalog.rows}) == 6
    assert catalog.rows[1].header_quantity == 0
    assert catalog.rows[2].header_quantity is None


def test_parent_is_self_contained_after_original_uploads_expire(tmp_path, monkeypatch):
    children = [_child(tmp_path, 1), _child(tmp_path, 2)]
    path, metadata, before = _parent(tmp_path, children)
    from proxy.services.chat_attachment_service import cleanup_expired
    monkeypatch.setattr("proxy.services.chat_attachment_service.time.time", lambda: metadata["created_at_epoch"] + 1)
    for child in children:
        meta_path = tmp_path / "attachments" / f"{child['attachment_id']}.json"
        data = json.loads(meta_path.read_text(encoding="utf-8"))
        data["created_at_epoch"] = 1
        meta_path.write_text(json.dumps(data), encoding="utf-8")
    assert cleanup_expired(root=tmp_path / "attachments", max_age_sec=60) == 2
    parent_path, trusted = resolve_read_attachment(metadata["attachment_id"], root=tmp_path / "attachments")
    after = load_tabular_source_catalog(parent_path, trusted_collection_metadata=trusted)
    assert after == before and path.is_file()
    assert not list(path.parent.glob("collection-read-*"))


@pytest.mark.parametrize("ids", [[], ["read_000000000001", "read_000000000001"],
                                  ["read_000000000001", "READ_000000000001"]])
def test_repeated_or_empty_attachment_ids_do_not_commit_parent(tmp_path, ids):
    _child(tmp_path, 1)
    with pytest.raises(AttachmentCollectionError, match="DISTINCT_ATTACHMENT_IDS"):
        create_read_attachment_collection(ids, max_bytes=100000, root=tmp_path / "attachments")
    assert not list((tmp_path / "attachments").glob("*.lescollection"))


def test_aggregate_size_error_uses_actual_raw_member_bytes(tmp_path):
    children = [_child(tmp_path, 1), _child(tmp_path, 2)]
    actual = sum(item["size"] for item in children)
    with pytest.raises(AttachmentCollectionError) as caught:
        create_read_attachment_collection([item["attachment_id"] for item in children],
            max_bytes=actual - 1, root=tmp_path / "attachments")
    assert (caught.value.code, caught.value.actual_bytes, caught.value.max_bytes) == ("UPLOAD_TOO_LARGE", actual, actual - 1)


@pytest.mark.parametrize("include_readable", [True, False])
def test_unknown_member_is_retained_and_blocks_extraction_completeness(tmp_path, include_readable):
    children = [_child(tmp_path, 1, name="Неизвестный.txt")]
    if include_readable:
        children.append(_child(tmp_path, 2))
    _, _, catalog = _parent(tmp_path, children)
    assert len(catalog.source_members) == len(children)
    assert catalog.source_members[0]["status"] == "unread"
    assert not catalog.extraction_complete
    assert "collection_member_unread:read_000000000001" in catalog.extraction_warnings[0]
    assert len(catalog.work_locators()) == (2 if include_readable else 0)
    assert "Неизвестный.txt" in catalog.manifest_text()


def test_user_container_suffix_without_authoritative_metadata_is_not_admitted(tmp_path):
    path, metadata, _ = _parent(tmp_path, [_child(tmp_path, 1)])
    assert load_tabular_source_catalog(path) is None
    with pytest.raises(AttachmentCollectionError, match="SERVER_METADATA_MISMATCH"):
        load_tabular_source_catalog(path, trusted_collection_metadata={**metadata, "schema": "chat_read_attachment_v1"})
    with pytest.raises(AttachmentCollectionError, match="SERVER_METADATA_MISMATCH"):
        load_tabular_source_catalog(path, trusted_collection_metadata={**metadata, "sha256": "0" * 64})


@pytest.mark.parametrize("fault", ["missing", "extra", "traversal", "hash", "size", "nested"])
def test_invalid_container_members_fail_closed(tmp_path, fault):
    path, metadata, _ = _parent(tmp_path, [_child(tmp_path, 1)])
    def mutate(entries, manifest):
        member = manifest["members"][0]
        if fault == "missing":
            entries.pop(member["archive_path"])
        elif fault == "extra":
            entries["unlisted.txt"] = b"unlisted"
        elif fault == "traversal":
            member["archive_path"] = "../../escaped.xlsx"
        elif fault == "hash":
            member["sha256"] = "0" * 64
        elif fault == "size":
            member["size"] += 1
        else:
            member["original_name"] = "nested.lescollection"
            member["suffix"] = ".lescollection"
    trusted = _rewrite_container(path, metadata, mutate)
    with pytest.raises(AttachmentCollectionError):
        load_tabular_source_catalog(path, trusted_collection_metadata=trusted)
    assert not (tmp_path / "escaped.xlsx").exists()


def test_duplicate_manifest_fields_are_not_silently_selected(tmp_path):
    path, metadata, _ = _parent(tmp_path, [_child(tmp_path, 1)])
    with zipfile.ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    entries["manifest.json"] = entries["manifest.json"].replace(b'{', b'{"schema":"ambiguous",', 1)
    with zipfile.ZipFile(path, "w") as archive:
        for name, body in entries.items():
            archive.writestr(name, body)
    trusted = {**metadata, "size": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "manifest_sha256": hashlib.sha256(entries["manifest.json"]).hexdigest()}
    with pytest.raises(AttachmentCollectionError, match="DUPLICATE_MANIFEST_KEY"):
        load_tabular_source_catalog(path, trusted_collection_metadata=trusted)


def test_collection_offsets_are_member_local_and_frame_preserves_manifest_order(tmp_path):
    _, _, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2)])
    # Distinct source formats may both start at offset zero with different bases.
    rows = tuple(replace(row, source_start=index % 3, source_end=index % 3 + 1,
                         source_order_basis=f"member-{index // 3}") for index, row in enumerate(catalog.rows))
    catalog = replace(catalog, rows=rows)
    frame = build_model_source_frame(catalog, [rows[4].locator, rows[1].locator])
    payload = frame.to_payload()
    assert frame.requested_locators == (rows[1].locator, rows[4].locator)
    assert payload["source_order_basis"] == "collection_manifest_member_order"
    assert len(payload["source_manifest"]["source_members"]) == 2
    assert len(payload["source_manifest"]["source_row_origins"]) == 6
    rendered = render_model_source_frame(frame).to_payload()
    assert rendered["source_manifest"]["source_members"] == payload["source_manifest"]["source_members"]
    assert catalog.interleaved_row_lines() == [row.line for row in rows]
    current_sheet = rows[4].sheet
    assert all("read_000000000001::" not in line for line in catalog_sheet_pack_lines(catalog, current_sheet))
    facts = catalog_sheet_packet_facts(catalog, current_sheet)
    assert facts["source_order_basis"] == "collection_manifest_member_order"
    assert all(item["member_key"] == "read_000000000002" for item in facts["source_row_origins"])


@pytest.mark.asyncio
async def test_native_read_source_resolves_exact_member_and_original_provenance(tmp_path):
    from proxy.services.model_research_tool_service import ModelResearchToolService
    _, _, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2, values=[["Другая работа", "м2", 13]])])
    async def forbidden(**kwargs):
        pytest.fail("literal source read must not call RAG or a model")
    service = ModelResearchToolService(retrieve=forbidden, frozen_dataset_ids=(), retrieval_kwargs={},
        fallback=lambda *args: pytest.fail("unexpected fallback"), attachment_catalog=catalog,
        expected_attachment_sha256=catalog.source_file_sha256)
    result = await service.execute({"tool": "read_source", "args": {"source_row": "read_000000000002::Раздел!R2"}})
    assert result.payload["status"] == "ok"
    origin = result.payload["result"]["source_member"]
    assert origin["original_locator"] == "Раздел!R2"
    assert origin["attachment_id"] == "read_000000000002"
    assert "Другая работа" in result.chunks[0].content
    assert result.payload["result"]["source_sha256"] == catalog.source_file_sha256
    missing = await service.execute({"tool": "read_source", "args": {"source_row": "Раздел!R2"}})
    assert missing.payload["status"] == "missing"


def test_collection_api_is_existing_read_attachment_shape_and_enforces_total_size(tmp_path, monkeypatch):
    from proxy.routers import datasets
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(tmp_path / "attachments"))
    children = [_child(tmp_path, 1), _child(tmp_path, 2)]
    app = FastAPI()
    app.include_router(datasets.search_router)
    app.dependency_overrides[datasets.require_user] = lambda: {"id": "test"}
    with TestClient(app) as client:
        response = client.post("/api/chat/attachment-collections", json={"attachment_ids": [item["attachment_id"] for item in children]})
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["mode"] == "read" and body["attachment_id"].startswith("read_")
        assert body["name"] == "Комплект документов (2)"
        assert body["source_extraction_complete"] is True
        assert body["collection"]["member_count"] == 2
        assert body["chars"] >= len(body["text"]) and body["source_sha256"]
        assert load_catalog_from_attachment_id(body["attachment_id"]).work_locators()
        monkeypatch.setattr(datasets, "max_upload_bytes", lambda: 1)
        oversized = client.post("/api/chat/attachment-collections", json={"attachment_ids": [item["attachment_id"] for item in children]})
        assert oversized.status_code == 413
        assert oversized.json()["detail"] == {"code": "UPLOAD_TOO_LARGE", "max_bytes": 1,
            "actual_bytes": sum(item["size"] for item in children)}
        assert client.post("/api/chat/attachment-collections", json={"attachment_ids": []}).status_code == 422
        assert client.post("/api/chat/attachment-collections", json={"attachment_ids": ["read_ffffffffffff"]}).status_code == 422


def test_collection_identity_remaining_and_evidence_survive_store_restart(tmp_path):
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    _, metadata, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2)])
    db = tmp_path / "jobs.sqlite"
    store = WorkbookResearchJobService(db)
    def identity(service, question):
        return service.identity(session_id="collection-session", attachment_id=metadata["attachment_id"],
            source_sha256=metadata["sha256"], dataset_ids=["dataset"], profile_revision_id="profile",
            model_identity="model", question=question)
    work = catalog.work_locators()
    first = store.begin_or_resume(identity(store, "Собери ВОР"), source_locators=work)
    row = {"source_row": work[0], "operation_id": "permanent", "row_role": "work",
        "title": "Буквальная работа", "quantity": 0, "evidence_refs": "Q1.H1"}
    store.record_packet(first["job_id"], sheet=catalog.lookup(work[0]).sheet,
        packet_locators=[work[0]], accepted_rows=[row], evidence=[{"source_ref": "unchanged"}], remaining=work[1:])
    saved = store.get(first["job_id"])
    reopened = WorkbookResearchJobService(db)
    resumed = reopened.begin_or_resume(identity(reopened, "Продолжи"), source_locators=work)
    assert resumed["continues_job_id"] == first["job_id"]
    assert resumed["accepted_rows"] == saved["accepted_rows"]
    assert resumed["evidence"] == saved["evidence"]
    assert resumed["remaining"] == work[1:]


def test_collection_module_is_in_runtime_alignment_contract():
    from proxy.services.version_service import _CRITICAL_FILES
    assert "proxy/services/chat_attachment_collection_service.py" in _CRITICAL_FILES
    assert "proxy/services/chat_attachment_service.py" in _CRITICAL_FILES


def _workbook_context(tmp_path, *, lsr=False):
    from proxy.services.artifact_revision_service import ArtifactRevisionStore
    from proxy.services.workflow_checkpoint_service import WorkflowCheckpointService
    from proxy.services.workbook_tool_service import WorkbookExecutionContext
    from proxy.services.lsr_workbook_adapter_service import build_lsr_workbook_from_decisions
    return WorkbookExecutionContext(session_id="collection-session", idempotency_key="lsr" if lsr else "vor",
        model_decision_revision="literal-1", profile_revision_id="profile", model_identity="model", model_preset="preset",
        attachment_root=tmp_path / "attachments", work_dir=tmp_path / "work",
        checkpoints=WorkflowCheckpointService(tmp_path / "checkpoints.sqlite"),
        artifacts=ArtifactRevisionStore(tmp_path / "artifacts.sqlite", tmp_path / "artifacts"),
        lsr_adapter=build_lsr_workbook_from_decisions if lsr else None)


def _decisions(catalog):
    return [{"source_row": row.locator, "row_role": "work", "work_row": "",
             "title": row.values[1], "quantity": row.header_quantity, "unit": row.values[2],
             "open_gap": "Объём не определён в источнике" if row.header_quantity is None else "",
             "literal_unknown_field": {"text": 'Кавычки "и"\nстрока'}}
            for row in catalog.rows if row.is_work_row()]


@pytest.mark.asyncio
async def test_real_vor_exporter_uses_union_and_keeps_authored_zero_null_and_original_names(tmp_path):
    from proxy.services.workbook_tool_service import build_vor_workbook
    _, metadata, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2, name="Второй файл.xlsx")])
    decisions = _decisions(catalog)
    before = json.dumps(decisions, ensure_ascii=False)
    ctx = _workbook_context(tmp_path)
    result = await build_vor_workbook({"attachment_id": metadata["attachment_id"], "decisions": decisions}, ctx)
    assert "artifact" in result, result
    assert result["source"]["sha256"] == metadata["sha256"]
    assert len(result["source"]["collection_members"]) == 2
    assert not result.get("vor_source_errors")
    assert any(":quantity" in item for item in result["missing"])
    path = ctx.artifacts.resolve_path(result["artifact"]["revision_id"])
    book = openpyxl.load_workbook(path, data_only=False)
    try:
        audit = book["Решения модели"]
        authored = [json.loads(row[2].removeprefix("JSON: ")) for row in audit.iter_rows(min_row=4, values_only=True)]
        assert authored == decisions
        values = [value for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None]
        assert any('Второй файл.xlsx#' in str(value) for value in values)
        assert any('Лист "А".xlsx#' in str(value) for value in values)
        assert all(any(item["source_row"] in str(value) for value in values) for item in decisions)
        assert [item["quantity"] for item in authored] == [0, None, 0, None]
    finally:
        book.close()
    assert json.dumps(decisions, ensure_ascii=False) == before


@pytest.mark.asyncio
async def test_real_lsr_adapter_receives_same_union_source_cells_and_explicit_price_basis(tmp_path, monkeypatch):
    from proxy.services.workbook_tool_service import build_lsr_workbook
    from proxy.services import lsr_workbook_adapter_service as adapter
    from proxy.services import fgis_price_service
    _, metadata, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2)])
    decisions = _decisions(catalog)
    before = json.dumps(decisions, ensure_ascii=False)
    captured = []
    actual_calculate = adapter.build_lsr_trace_from_visible_rows
    def calculate(rows, **kwargs):
        captured.append((json.loads(json.dumps(rows)), kwargs["source_catalog"]))
        return actual_calculate(rows, **kwargs)
    monkeypatch.setattr(adapter, "build_lsr_trace_from_visible_rows", calculate)
    monkeypatch.setattr(fgis_price_service, "resolve_pricebook_path", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(fgis_price_service, "available_pricebooks", lambda *_args, **_kwargs: [])
    ctx = _workbook_context(tmp_path, lsr=True)
    result = await build_lsr_workbook({"attachment_id": metadata["attachment_id"], "decisions": decisions,
        "question": "Собери ЛСР, Санкт-Петербург, 3 квартал 2026 года"}, ctx)
    assert "artifact" in result, result
    rows, supplied = captured[0]
    assert supplied == catalog
    assert [item["source_row"] for item in rows] == catalog.work_locators()
    assert rows[0]["_catalog_source_numbers"] == [0]
    assert rows[1]["_catalog_source_numbers"] == []
    assert rows[1]["quantity"] is None
    assert result["calculated_positions"] == 0 and result["confirmed_amount"] == 0
    assert result["price_basis"]["requested_region"]
    assert result["price_basis"]["match"] == "no_book"
    assert result["source"]["sha256"] == metadata["sha256"]
    assert len(result["source"]["collection_members"]) == 2
    output = ctx.artifacts.resolve_path(result["artifact"]["revision_id"])
    assert output.is_file()
    book = openpyxl.load_workbook(output, data_only=False)
    try:
        audit = book["Проверка"]
        headers = [cell.value for cell in audit[1]]
        name_column, original_column = headers.index("Файл источника"), headers.index("Исходная строка")
        actual = list(audit.iter_rows(min_row=2, max_row=1 + len(decisions), values_only=True))
        assert [row[0] for row in actual] == catalog.work_locators()
        assert [row[name_column] for row in actual] == [catalog.row_origin(item["source_row"])["original_name"] for item in decisions]
        assert [row[original_column] for row in actual] == [catalog.row_origin(item["source_row"])["original_locator"] for item in decisions]
    finally:
        book.close()
    assert json.dumps(decisions, ensure_ascii=False) == before


@pytest.mark.asyncio
async def test_canonical_packer_rejects_suffix_only_user_collection_before_checkpoint(tmp_path):
    from proxy.services.workbook_tool_service import build_vor_workbook
    path, _, catalog = _parent(tmp_path, [_child(tmp_path, 1)])
    forged = preserve_read_attachment(path, attachment_id="read_ffffffffffff",
        original_name="user.lescollection", root=tmp_path / "attachments")
    result = await build_vor_workbook({"attachment_id": forged["attachment_id"], "decisions": _decisions(catalog)},
                                     _workbook_context(tmp_path))
    assert result["code"] == "ATTACHMENT_INVALID"
    assert "checkpoint" not in result and "artifact" not in result


def test_model_render_has_only_shown_origins_but_full_audit_and_manifest(tmp_path):
    children = [_child(tmp_path, number, values=[[f"Буквальная работа {index}", "м", index]
                for index in range(1, 31)]) for number in range(1, 4)]
    _, _, catalog = _parent(tmp_path, children)
    frame = build_model_source_frame(catalog, catalog.work_locators()[:12])
    full = frame.to_payload()["source_manifest"]
    rendered = render_model_source_frame(frame).to_payload()["source_manifest"]
    assert len(full["source_row_origins"]) == len(catalog.rows) == 93
    assert len(rendered["source_row_origins"]) == len(frame.shown_locators)
    assert len(rendered["source_row_origins"]) < len(full["source_row_origins"])
    assert rendered["source_row_origins_not_shown_count"] == 93 - len(frame.shown_locators)
    assert len(rendered["records"]) == 93 and len(rendered["source_members"]) == 3
    assert frame.to_payload()["source_manifest"] == full


@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["vor", "lsr"])
async def test_unread_member_stays_a_document_blocker_in_delivered_workbook(tmp_path, monkeypatch, kind):
    from proxy.services.workbook_tool_service import build_vor_workbook, build_lsr_workbook
    from proxy.services import fgis_price_service
    children = [_child(tmp_path, 1, values=[["Работа", "м", 0]]),
                _child(tmp_path, 2, name="Непрочитанный.txt")]
    _, metadata, catalog = _parent(tmp_path, children)
    monkeypatch.setattr(fgis_price_service, "resolve_pricebook_path", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(fgis_price_service, "available_pricebooks", lambda *_args, **_kwargs: [])
    handler = build_vor_workbook if kind == "vor" else build_lsr_workbook
    result = await handler({"attachment_id": metadata["attachment_id"], "decisions": _decisions(catalog)},
                           _workbook_context(tmp_path, lsr=kind == "lsr"))
    assert result["status"] == "partial" and "artifact" in result
    assert any(issue["code"] == "SOURCE_EXTRACTION_INCOMPLETE" for issue in result["document_issues"])
    assert result["source"]["collection_members"][1]["status"] == "unread"


@pytest.mark.asyncio
async def test_source_outside_small_frame_keeps_original_member_access(tmp_path):
    from proxy.services.model_research_tool_service import ModelResearchToolService
    _, _, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2, name="Другой.xlsx")])
    frame = build_model_source_frame(catalog, catalog.work_locators()[:1])
    outside = catalog.work_locators()[-1]
    rendered = render_model_source_frame(frame).to_payload()["source_manifest"]
    assert outside not in {item["source_row"] for item in rendered["source_row_origins"]}
    assert [outside] in rendered["records"]
    service = ModelResearchToolService(retrieve=lambda **kw: pytest.fail("unexpected retrieval"),
        frozen_dataset_ids=(), retrieval_kwargs={}, fallback=lambda *a: pytest.fail("unexpected fallback"),
        attachment_catalog=catalog, expected_attachment_sha256=catalog.source_file_sha256)
    result = await service.execute({"tool": "read_source", "args": {"source_row": outside}})
    assert result.payload["result"]["source_member"] == catalog.row_origin(outside)
    assert result.chunks[0].doc_name == "Другой.xlsx"


def test_model_cannot_forge_collection_origin_in_single_file_or_missing_source(tmp_path, monkeypatch):
    from proxy.services.lsr_workbook_adapter_service import _attach_catalog_facts
    child = _child(tmp_path, 1)
    catalog = load_tabular_source_catalog(tmp_path / "source-1.xlsx")
    rows = [{"source_row": "Раздел!R2", "quantity": 0,
             "_catalog_source_origin": {"original_name": "forged.xlsx"}}]
    assert _attach_catalog_facts(rows, child["attachment_id"], catalog_override=catalog) == catalog
    assert "_catalog_source_origin" not in rows[0]
    missing = [{"source_row": "absent", "_catalog_source_origin": {"original_name": "forged.xlsx"}}]
    _attach_catalog_facts(missing, "")
    assert "_catalog_source_origin" not in missing[0]


def test_failed_collection_metadata_commit_removes_only_own_temporary_files(tmp_path, monkeypatch):
    child = _child(tmp_path, 1)
    replace_path = Path.replace
    def fail_metadata(path, target):
        if path.name.endswith(".json.tmp"):
            raise OSError("synthetic metadata commit failure")
        return replace_path(path, target)
    monkeypatch.setattr(Path, "replace", fail_metadata)
    with pytest.raises(OSError, match="synthetic metadata commit failure"):
        create_read_attachment_collection([child["attachment_id"]], max_bytes=100000,
                                          root=tmp_path / "attachments")
    assert not list((tmp_path / "attachments").glob("*.tmp"))
    assert not list((tmp_path / "attachments").glob("*.lescollection"))
    assert resolve_read_attachment(child["attachment_id"], root=tmp_path / "attachments")[0].is_file()


@pytest.mark.parametrize("literal_link", [True, False])
def test_operation_links_keep_member_namespace_and_never_choose_equal_original_coordinate(tmp_path, literal_link):
    from proxy.services.vor_result_integrity_service import audit_vor_result_integrity
    _, _, catalog = _parent(tmp_path, [_child(tmp_path, 1), _child(tmp_path, 2)])
    first = "read_000000000001::Раздел!R2"
    second = "read_000000000002::Раздел!R2"
    rows = [{"source_row": first, "operation_id": "authored-work", "row_role": "work", "title": "Работа",
             "unit": "м", "quantity": 0},
            {"source_row": second, "row_role": "material", "title": "Материал", "unit": "м", "quantity": 0,
             "work_row": first if literal_link else "Раздел!R2"}]
    before = json.dumps(rows, ensure_ascii=False)
    issues = audit_vor_result_integrity(rows, source_catalog=catalog)
    link_issues = [item for item in issues if item["code"] == "invalid_work_relationship"]
    assert bool(link_issues) is not literal_link
    assert json.dumps(rows, ensure_ascii=False) == before


@pytest.mark.parametrize("name", ["=1+1.csv", "'Кириллица \"А\".xlsx"])
def test_new_lsr_origin_columns_are_literal_even_when_filename_looks_like_formula(tmp_path, name):
    from proxy.services.rim_trace_xlsx_service import _review_sheet
    book = openpyxl.Workbook()
    try:
        _review_sheet(book, {"row_bindings": [{"row": "read_000000000001::=Раздел!R2", "source_row": {
            "_catalog_source_origin": {"original_name": name, "original_locator": "=Раздел!R2"}}}]})
        output = tmp_path / "literal.xlsx"
        book.save(output)
    finally:
        book.close()
    reread = openpyxl.load_workbook(output, data_only=False)
    try:
        sheet = reread["Проверка"]
        headers = [cell.value for cell in sheet[1]]
        name_cell = sheet.cell(2, headers.index("Файл источника") + 1)
        source_cell = sheet.cell(2, headers.index("Исходная строка") + 1)
        assert name_cell.value == name and name_cell.data_type == "s"
        assert source_cell.value == "=Раздел!R2" and source_cell.data_type == "s"
    finally:
        reread.close()


@pytest.mark.parametrize("collection", [False, True])
def test_used_typed_source_survives_age_cleanup_consume_and_fresh_service_resume(tmp_path, monkeypatch, collection):
    from proxy.services.chat_attachment_service import cleanup_expired, consume_read_attachment
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(tmp_path / "attachments"))
    children = [_child(tmp_path, 1), _child(tmp_path, 2)] if collection else [_child(tmp_path, 1)]
    if collection:
        path, metadata, _ = _parent(tmp_path, children)
    else:
        metadata = children[0]
        path, _ = resolve_read_attachment(metadata["attachment_id"])
    before = path.read_bytes()
    catalog = load_catalog_from_attachment_id(metadata["attachment_id"])
    retained = json.loads((path.parent / f"{metadata['attachment_id']}.json").read_text(encoding="utf-8"))
    assert retained["retained_for_workbook"] is True
    assert retained["sha256"] == metadata["sha256"] == catalog.source_file_sha256
    store_path = tmp_path / "durable.sqlite"
    store = WorkbookResearchJobService(store_path)
    def identity(service, question):
        return service.identity(session_id="aged-source-session", attachment_id=metadata["attachment_id"],
            source_sha256=metadata["sha256"], dataset_ids=[], profile_revision_id="profile",
            model_identity="model", question=question)
    locators = catalog.work_locators()
    job = store.begin_or_resume(identity(store, "Собери ВОР"), source_locators=locators)
    store.record_packet(job["job_id"], sheet=catalog.lookup(locators[0]).sheet,
        packet_locators=[locators[0]], accepted_rows=[{"source_row": locators[0], "operation_id": "kept",
        "row_role": "work", "title": "Работа", "unit": "м", "quantity": 0}], evidence=[], remaining=locators[1:])
    saved = store.get(job["job_id"])
    retained["created_at_epoch"] = 1
    (path.parent / f"{metadata['attachment_id']}.json").write_text(json.dumps(retained), encoding="utf-8")
    raw = _child(tmp_path, 99, name="unused.txt")
    raw_meta_path = path.parent / f"{raw['attachment_id']}.json"
    raw_metadata = json.loads(raw_meta_path.read_text(encoding="utf-8"))
    raw_metadata["created_at_epoch"] = 1
    raw_meta_path.write_text(json.dumps(raw_metadata), encoding="utf-8")
    assert cleanup_expired(root=path.parent) == 1
    assert not raw_meta_path.exists()
    consume_read_attachment(metadata["attachment_id"])
    reloaded = load_catalog_from_attachment_id(metadata["attachment_id"])
    reopened = WorkbookResearchJobService(store_path)
    resumed = reopened.begin_or_resume(identity(reopened, "Продолжи"), source_locators=reloaded.work_locators())
    assert resumed["continues_job_id"] == job["job_id"]
    assert resumed["accepted_rows"] == saved["accepted_rows"] and resumed["remaining"] == saved["remaining"]
    assert reloaded == catalog and path.read_bytes() == before
    assert hashlib.sha256(path.read_bytes()).hexdigest() == metadata["sha256"]


def test_wrong_hash_cannot_pin_and_unread_or_unused_sources_keep_original_ttl(tmp_path, monkeypatch):
    from proxy.services.chat_attachment_service import retain_read_attachment_for_workbook
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(tmp_path / "attachments"))
    child = _child(tmp_path, 1)
    raw = _child(tmp_path, 2, name="unread.txt")
    with pytest.raises(ValueError, match="hash mismatch"):
        retain_read_attachment_for_workbook(child["attachment_id"], expected_sha256="0" * 64)
    path, metadata = resolve_read_attachment(child["attachment_id"])
    assert "retained_for_workbook" not in metadata
    assert load_catalog_from_attachment_id(raw["attachment_id"]) is None
    for item in (child, raw):
        meta_path = path.parent / f"{item['attachment_id']}.json"
        data = json.loads(meta_path.read_text(encoding="utf-8"))
        assert "retained_for_workbook" not in data
        data["created_at_epoch"] = 1
        meta_path.write_text(json.dumps(data), encoding="utf-8")
        with pytest.raises(FileNotFoundError, match="expired"):
            resolve_read_attachment(item["attachment_id"])
        assert not meta_path.exists()


def test_pin_does_not_delete_source_when_ttl_crosses_during_successful_catalog_read(tmp_path, monkeypatch):
    from proxy.services import attachment_source_catalog_service as catalogs
    from proxy.services.chat_attachment_service import DEFAULT_MAX_AGE_SEC
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(tmp_path / "attachments"))
    child = _child(tmp_path, 1)
    raw = _child(tmp_path, 2, name="unused.txt")
    path, _ = resolve_read_attachment(child["attachment_id"])
    original_bytes = path.read_bytes()
    for metadata in (child, raw):
        meta_path = path.parent / f"{metadata['attachment_id']}.json"
        data = json.loads(meta_path.read_text(encoding="utf-8"))
        data["created_at_epoch"] = 100
        meta_path.write_text(json.dumps(data), encoding="utf-8")
    clock = [100 + DEFAULT_MAX_AGE_SEC - 1]
    monkeypatch.setattr("proxy.services.chat_attachment_service.time.time", lambda: clock[0])
    actual_load = catalogs.load_tabular_source_catalog
    def crossing_read(*args, **kwargs):
        catalog = actual_load(*args, **kwargs)
        clock[0] += 2
        return catalog
    monkeypatch.setattr(catalogs, "load_tabular_source_catalog", crossing_read)
    catalog = catalogs.load_catalog_from_attachment_id(child["attachment_id"])
    stored, metadata = resolve_read_attachment(child["attachment_id"])
    assert metadata["retained_for_workbook"] is True
    assert stored.read_bytes() == original_bytes
    assert metadata["sha256"] == child["sha256"] == catalog.source_file_sha256
    with pytest.raises(FileNotFoundError, match="expired"):
        resolve_read_attachment(raw["attachment_id"])


def test_concurrent_first_readers_pin_with_independent_atomic_writers(tmp_path, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    from proxy.services import chat_attachment_service as attachments
    child = _child(tmp_path, 1)
    source, _ = resolve_read_attachment(child["attachment_id"], root=tmp_path / "attachments")
    original_bytes = source.read_bytes()
    both_read, both_written = Barrier(2), Barrier(2)
    original_resolve, original_write = attachments.resolve_read_attachment, Path.write_text
    def read_before_either_pin(*args, **kwargs):
        result = original_resolve(*args, **kwargs)
        assert result[1].get("retained_for_workbook") is not True
        both_read.wait(timeout=5)
        return result
    def write_before_either_replace(path, *args, **kwargs):
        result = original_write(path, *args, **kwargs)
        if path.name.endswith(".retain.tmp"):
            both_written.wait(timeout=5)
        return result
    monkeypatch.setattr(attachments, "resolve_read_attachment", read_before_either_pin)
    monkeypatch.setattr(Path, "write_text", write_before_either_replace)
    with ThreadPoolExecutor(max_workers=2) as workers:
        futures = [workers.submit(attachments.retain_read_attachment_for_workbook,
            child["attachment_id"], expected_sha256=child["sha256"], root=tmp_path / "attachments") for _ in range(2)]
        results = [future.result(timeout=10) for future in futures]
    assert all(item["retained_for_workbook"] is True for item in results)
    stored, metadata = original_resolve(child["attachment_id"], root=tmp_path / "attachments")
    assert metadata["retained_for_workbook"] is True and metadata["sha256"] == child["sha256"]
    assert stored.read_bytes() == original_bytes
    assert not list(stored.parent.glob("*.retain.tmp"))
