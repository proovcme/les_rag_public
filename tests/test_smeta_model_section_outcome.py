"""Explicit model section outcomes, never inferred from titles or blank quantities."""

import pytest
from openpyxl import Workbook, load_workbook

from proxy.services import chat_evidence_application_service as chat
from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
from proxy.services.smeta_source_registry_service import attach_source_registry
from proxy.services.tabular_attachment_service import format_tabular_attachment_context
from proxy.services.portion_cursor_service import PortionCursorService
from proxy.smeta_core.calculation_receipt import calculation_receipts, apply_model_review


def visible(kind="раздел", reason="Это название раздела исходной ведомости; отдельной работы здесь нет", **extra):
    values = dict(source_row="S!1", work_id="heading", title="Электроснабжение", unit="", quantity="",
                  norm_code="", decision_kind=kind, analogue=reason, coverage="", evidence_refs="")
    values.update(extra)
    headers = list(values)
    text = "| " + " | ".join(headers) + " |\n| " + " | ".join("---" for _ in headers) + " |\n"
    text += "| " + " | ".join(str(values[key]) for key in headers) + " |"
    return chat.parse_model_rag_result(text)[1]


@pytest.mark.parametrize("kind", ["раздел", "section"])
def test_explicit_model_heading_survives_registry_receipt_and_reopened_continuation(tmp_path, kind):
    rows = visible(kind)
    assert rows[0]["decision_kind"] == kind, "Preserve the literal model decision before binding"
    trace = build_lsr_trace_from_visible_rows(rows)
    assert trace["row_bindings"][0]["status"] == "section_recognized"
    assert trace["row_bindings"][0]["decision_kind"] in {kind, "раздел"}
    assert trace["summary"]["bound_operations"] == 0
    assert trace["summary"]["work_positions"] == 0
    assert not any(section.get("positions") for section in trace["sections"])
    receipts = calculation_receipts(trace)
    assert not any(receipt.get("calculation") or receipt.get("status") == "resource_review_required" for receipt in receipts)
    apply_model_review(trace, receipts, rows)
    assert trace["row_bindings"][0]["status"] == "section_recognized"

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "S"
    workbook.active.append(["№", "Наименование", "Ед. изм", "Количество"])
    workbook.active.append([1, "Электроснабжение", None, None])
    workbook.active.append([2, "Работа с неизвестным количеством", "шт", None])
    workbook.save(source)
    workbook.close()
    attach_source_registry(trace, source)
    inventory = trace["source_registry"]["rows"]
    assert inventory[0]["status"] == "section_recognized"
    assert inventory[0]["title"] == "Электроснабжение"
    assert trace["summary"]["bound_rows"] == 0, "Recognized headings are not calculated works"
    assert trace["summary"]["non_work_rows"] == 1
    assert trace["summary"]["resolved_source_rows"] == 1
    assert trace["summary"]["open_rows"] == 1
    target = tmp_path / "result.xlsx"
    render_lsr_xlsx(trace, target)
    rendered = load_workbook(target, read_only=True, data_only=True)
    try:
        cells = list(rendered["Исходные строки"].values)
        assert any("section_recognized" in row and "Электроснабжение" in row for row in cells)
        labels = '\n'.join(str(row[0] or '') for row in rendered['ЛСР РИМ'].values)
        assert 'привязано 0 из 1 строк' in labels
        assert 'Строк-разделов по решению модели: 1' in labels
    finally:
        rendered.close()

    original = format_tabular_attachment_context(source, source.name, max_chars=None, max_cell_chars=None)[0]
    store = PortionCursorService(tmp_path / "cursor.sqlite")
    _, state = chat.smeta_portion_context(original, session_id="s", limit=2, cursors=store)
    chat.record_portion_pass(store, session_id="s", portion_state=state, rows=rows, closed=[],
        artifact_revision_id=None, row_statuses={"S!1": "section_recognized"}, calculation_receipts=receipts,
        pass_id="heading-turn-1")
    store.record_calculation(session_id="s", source_revision=state['source_revision'],
                             statuses={"S!1": "section_recognized"}, receipts=receipts)
    restarted = PortionCursorService(tmp_path / "cursor.sqlite")
    context, next_state = chat.smeta_portion_context(original, session_id="s", limit=2, cursors=restarted)
    assert "S!1" in next_state["prior_decisions"]
    assert next_state["prior_decisions"]["S!1"]["decision_kind"] == kind
    assert next_state["shown"] == ["S!2"]
    assert "Работа с неизвестным количеством" in context
    # remaining counts rows after this shown portion; pending includes it.
    assert next_state["pending"] == 1
    assert next_state["completed"] == 1
    assert next_state["remaining"] == 0


def test_heading_without_model_reason_stays_unresolved():
    rows = visible(reason="")
    binding = build_lsr_trace_from_visible_rows(rows)["row_bindings"][0]
    assert binding["status"] != "section_recognized"
    assert binding.get("message") or binding.get("missing"), "The missing reason must be observable"


def test_explicit_heading_conflicting_with_work_or_resource_is_observable():
    for extra in ({"norm_code": "ГЭСН01-01-001-01"},
                  {"resource_action": "included", "included_in": "mount"},
                  {"material_review_status": "confirmed"}):
        rows = visible(**extra)
        binding = build_lsr_trace_from_visible_rows(rows)["row_bindings"][0]
        assert "conflict" in binding["status"]
        assert binding.get("message"), "Do not silently discard a contradictory professional instruction"
        for key, value in extra.items():
            assert rows[0][key] == value


def test_mixed_heading_work_operations_conflict_and_blank_quantity_never_infers_heading():
    heading = visible()[0]
    work = visible("работа", reason="", work_id="mount", norm_code="", unit="шт")[0]
    work_trace = build_lsr_trace_from_visible_rows([work])
    assert work_trace["row_bindings"][0]["status"] != "section_recognized"
    for mixed_rows in ([heading, work], [work, heading]):
        mixed = build_lsr_trace_from_visible_rows(mixed_rows)
        assert all("conflict" in binding["status"] for binding in mixed["row_bindings"])
        assert mixed["summary"]["bound_operations"] == 0


def test_heading_does_not_hide_calculation_fields_or_placeholder_reason():
    for extra in ({'resource_quantity': 4}, {'resource_unit': 'kg'},
                  {'price_source_ref': 'price:1'}, {'coefficient': 2},
                  {'coefficient_target': 'labor'}, {'coefficient_basis': 'Explicit basis'},
                  {'conversion_factor': 3}):
        rows = visible()
        rows[0].update(extra)
        binding = build_lsr_trace_from_visible_rows(rows)['row_bindings'][0]
        assert binding['status'] == 'source_role_conflict', extra
        assert all(binding['source_row'][key] == value for key, value in extra.items())
    assert build_lsr_trace_from_visible_rows(visible(reason='—'))['row_bindings'][0]['status'] == 'section_reason_required'
    assert build_lsr_trace_from_visible_rows(visible(coefficient=1))['row_bindings'][0]['status'] == 'section_recognized'


def test_heading_does_not_bypass_duplicate_decision_identity():
    rows = [visible()[0], visible(source_row='S!2')[0]]
    result = build_lsr_trace_from_visible_rows(rows)
    assert [b['status'] for b in result['row_bindings']] == ['duplicate_work_id', 'duplicate_work_id']
    assert result['summary']['resolved_source_rows'] == 0


def test_heading_reason_in_coverage_is_not_hidden_by_empty_marker_in_analogue():
    rows = visible(reason='—', coverage='Explicit explanation of the source heading')
    result = build_lsr_trace_from_visible_rows(rows)
    assert result['row_bindings'][0]['status'] == 'section_recognized'
    assert result['row_bindings'][0]['message'] == rows[0]['coverage']
    assert result['row_bindings'][0]['decision']['complete'] is True
