"""Open applicability remains preliminary without rewriting the model's rows.

No runtime/model/API access. Uses the existing synthetic normative fixture.
"""
from copy import deepcopy

import pytest
from openpyxl import load_workbook

from proxy.services.lsr_norm_quality_gate_service import (
    applicability_gap_kind,
    norm_applicability_substantiated,
)
from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness, classify_norm_gap
from proxy.services.rim_lsr_trace_service import _row_readiness, build_lsr_trace_from_visible_rows
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
from tests.lsr_citation_support import cited_snapshot
from tests.test_rim_lsr_trace_service import _active_resource_norm, _ChosenBook


MATCH_YES = {key: "yes" for key in (
    "action_match", "object_match", "condition_match", "unit_compatible", "composition_covers"
)}


@pytest.mark.parametrize("status,gap", [
    ("ANALOGUE_CANDIDATE", ""),
    ("ANALOGUE_CANDIDATE", "Applicability is unresolved; another norm may be needed."),
    ("ANALOGUE", ""),
])
def test_analogue_with_an_open_match_fact_is_not_confirmed(status, gap):
    row = {"source_row": "S!R7", "row_role": "work", "norm_code": "ГЭСНм08-02-409-09",
           "norm_status": status, "open_gap": gap, **{**MATCH_YES, "condition_match": "unknown"},
           "_evidence_check": {"card_body_shown": True,
                               "cited_ciphers": ["ГЭСНм08-02-409-09"]}}
    before = deepcopy(row)
    assert applicability_gap_kind(row) == "analogue_needs_confirmation"
    assert norm_applicability_substantiated(row) is False
    assert classify_norm_gap(row) == "analogue_needs_confirmation"
    binding = {"code": row["norm_code"], "norm_status": row["norm_status"],
               "row_role": "work", "status": "bound", "position": {"op_key": "model-operation"}}
    assert _row_readiness(binding, row)["norm_selected"] is True
    assert _row_readiness(binding, row)["applicability_substantiated"] is False
    ready = build_lsr_result_readiness(packaged=True, rows=[row], pricing_status="priced_final",
        calculated_positions=1, catalog_work_count=1, catalog_locators=["S!R7"])
    assert ready["counts"]["norms_selected"] == 1
    assert ready["counts"]["norms_applicability_substantiated"] == 0
    assert ready["gap_states"]["norm_unsubstantiated"]["count"] == 1
    assert ready["dimensions"]["norms"]["complete"] is False
    assert row == before


@pytest.mark.parametrize("status", ["analog", "аналог"])
def test_existing_informal_analogue_aliases_remain_canonical_candidates(status):
    row = {"norm_code": "ГЭСНм08-02-409-09", "norm_status": status,
           **{**MATCH_YES, "condition_match": "unknown"}}
    before = deepcopy(row)
    assert applicability_gap_kind(row) == "analogue_needs_confirmation"
    assert norm_applicability_substantiated(row) is False
    assert row == before

def test_price_only_gap_does_not_reject_verified_analogue_status():
    row = {"norm_code": "ГЭСНм08-02-409-09", "norm_status": "ANALOGUE_VERIFIED", **MATCH_YES,
           "open_gap": "The selected price book does not provide the material price.",
           "_evidence_check": {"card_body_shown": True,
                               "cited_ciphers": ["ГЭСНм08-02-409-09"]}}
    before = deepcopy(row)
    assert applicability_gap_kind(row) == ""
    assert norm_applicability_substantiated(row) is True
    assert row == before


def test_active_allowed_analogue_token_can_substantiate_with_five_match_facts():
    # The active contract explicitly forbids *_VERIFIED as a first-pass label;
    # do not make every supported analogue permanently unconfirmable.
    row = {"norm_code": "ГЭСНм08-02-409-09", "norm_status": "ANALOGUE_CANDIDATE", **MATCH_YES,
           "_evidence_check": {"card_body_shown": True,
                               "cited_ciphers": ["ГЭСНм08-02-409-09"]}}
    before = deepcopy(row)
    assert applicability_gap_kind(row) == ""
    assert norm_applicability_substantiated(row) is True
    assert row == before


def test_candidate_keeps_its_norm_and_calculation_but_cannot_be_confirmed_money(monkeypatch, tmp_path):
    row = cited_snapshot(_active_resource_norm(monkeypatch), synthetic=True)
    row.update(norm_status="ANALOGUE_CANDIDATE", **{**MATCH_YES, "condition_match": "unknown"},
               open_gap="Applicability is unresolved; another norm may be needed.")
    before = deepcopy(row)
    book = _ChosenBook({"A-1": 100.0, "A-10": 100.0}, units={"A-1": "m", "A-10": "m"})
    trace = build_lsr_trace_from_visible_rows([row], pricebook=book)
    binding = trace["row_bindings"][0]
    assert binding["status"] == "bound"
    assert binding["code"] == row["norm_code"]
    assert binding["norm_status"] == "ANALOGUE_CANDIDATE"
    assert binding["source_row"] == before
    assert binding["readiness"]["norm_selected"] is True
    assert binding["readiness"]["applicability_substantiated"] is False
    assert binding["readiness"]["norm_applicability_ok"] is False
    assert binding["readiness"]["price_ok"] is True
    assert trace["summary"]["positions_built"] == 1
    # 1000 m / (100 m/norm unit) * (2 + 3) m * 100 per m, with no labour/overhead.
    assert trace["summary"]["total"] == 5000.0
    assert trace["summary"]["confirmed_amount"] == 0.0
    assert trace["summary"]["preliminary_amount"] == 5000.0
    assert trace["summary"]["result_status"] == "priced_partial"
    path = render_lsr_xlsx(trace, tmp_path / "candidate.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        review = next(values for values in workbook["Проверка"].iter_rows(values_only=True)
                      if values[0] == row["source_row"])
        assert "применимость обоснована=нет" in review[5]
        assert "применимость обоснована=да" not in review[5]
        assert review[6] == row["norm_code"]
        assert review[7] == "ANALOGUE_CANDIDATE"
    finally:
        workbook.close()
    assert row == before
