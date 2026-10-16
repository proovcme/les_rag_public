"""Source supplies and authored calculated work measures stay distinct."""
from copy import deepcopy
import json

import openpyxl
import pytest

from proxy.services.vor_result_integrity_service import audit_vor_result_integrity
from proxy.services.workbook_tool_service import _default_vor_adapter


def _work(quantity=12, unit="м²", **derivation):
    return {"source_row": "Объёмы Ω!R2", "row_role": "work", "title": "Авторская операция",
            "quantity": quantity, "unit": unit, "operation_id": "work-1", **derivation}


async def _export(tmp_path, row):
    source = tmp_path / "Вход 'Ω'.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Объёмы Ω"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Исходная позиция", row["unit"], row["quantity"]])
    workbook.save(source)
    workbook.close()
    rows = [row]
    before = deepcopy(rows)
    output = tmp_path / "ВОР 'Ω'.xlsx"
    result = await _default_vor_adapter(source, {"decisions": rows}, output, lambda *_args: None)
    assert rows == before, "the original model decision must never be replaced"
    return output, result


@pytest.mark.asyncio
@pytest.mark.parametrize("source_quantity,source_unit,calculated,calc_unit,factor,basis", [
    (12, "м²", 6, "м²", 0.5, "12 м² / 2 слоя = 6 м²; Объёмы Ω!R2"),
    (1250, "кг", 1.25, "т", 0.001, "1250 кг / 1000 = 1.25 т; Объёмы Ω!R2"),
    (8, "шт", 192, "м", 24, "8 изделий по 24 м = 192 м; Объёмы Ω!R2"),
    (0, "м²", 0, "м²", 0, "Исходный нулевой объём 0 м²; Объёмы Ω!R2"),
])
async def test_both_source_and_authored_derived_work_quantities_are_visible(
    tmp_path, source_quantity, source_unit, calculated, calc_unit, factor, basis,
):
    row = _work(source_quantity, source_unit, calc_quantity=calculated, calc_unit=calc_unit,
                conversion_factor=factor, conversion_source=basis)
    output, result = await _export(tmp_path, row)
    assert not result.get("document_issues") and not result["vor_source_errors"]
    workbook = openpyxl.load_workbook(output, data_only=True)
    sheet = workbook["ВОР"]
    assert [sheet["E3"].value, sheet["F3"].value] == ["Ед. изм. источника", "Кол-во источника"]
    assert [sheet["E4"].value, sheet["F4"].value] == [source_unit, source_quantity]
    assert [sheet.cell(4, c).value for c in range(16, 20)] == [calc_unit, calculated, factor, basis]
    assert sheet["O4"].hyperlink.target == "#'Решения модели'!C4"
    assert json.loads(workbook["Решения модели"]["C4"].value.removeprefix("JSON: ")) == row
    assert sheet["A2"].value is None
    workbook.close()
    from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog, quantity_candidates, header_bound_unit

    catalog = load_tabular_source_catalog(output)
    exported_row = catalog.lookup("ВОР!R4")
    assert exported_row.quantity_column_known is True
    assert quantity_candidates(exported_row) == (source_quantity,)
    assert header_bound_unit(exported_row) == source_unit
    assert catalog.extraction_complete is True


@pytest.mark.asyncio
@pytest.mark.parametrize("calculated", [True, "1/2", "NaN", "Infinity", {"expression": "12/2"}, [6]])
async def test_invalid_calculation_is_visible_and_never_made_into_a_valid_numeric_total(tmp_path, calculated):
    row = _work(calc_quantity=calculated, calc_unit="м²", conversion_source="Авторское основание")
    output, result = await _export(tmp_path, row)
    assert "invalid_calc_quantity" in {i["code"] for i in result["document_issues"]}
    assert result["requires_decision_count"] == 1
    workbook = openpyxl.load_workbook(output, data_only=False)
    sheet = workbook["ВОР"]
    assert sheet["F4"].value == 12
    assert isinstance(sheet["Q4"].value, str) and sheet["Q4"].data_type != "f"
    assert sheet["I4"].value == "ТРЕБУЕТ РЕШЕНИЯ"
    assert "не является конечным числом" in sheet["L4"].value
    assert "НЕ ЗАВЕРШЁН" in sheet["A2"].value
    assert json.loads(workbook["Решения модели"]["C4"].value.removeprefix("JSON: ")) == row
    workbook.close()


@pytest.mark.parametrize("fields,code", [
    ({"calc_quantity": 6, "conversion_source": "12 / 2; исходная строка"}, "missing_calc_unit"),
    ({"calc_quantity": 6, "calc_unit": "м²"}, "missing_calc_basis"),
    ({"calc_unit": "м²", "conversion_factor": 0.5, "conversion_source": "Два слоя"}, "missing_calc_quantity"),
    ({"calc_quantity": 6, "calc_unit": "м²", "conversion_source": "12 / 2", "conversion_factor": "1/2"}, "invalid_conversion_factor"),
])
def test_incomplete_derivation_remains_open_even_with_a_known_source_quantity(fields, code):
    from proxy.services.chat_evidence_application_service import build_vor_result_readiness

    row = _work(**fields)
    assert code in {i["code"] for i in audit_vor_result_integrity([row])}
    readiness = build_vor_result_readiness(rows=[row], has_file=True, remaining_locators=[],
        catalog_work_count=1, model_errors=[])
    assert readiness["vor_complete"] is False
    assert code in {i["code"] for i in readiness["row_integrity_issues"]}


def test_a_complete_derived_measure_is_possible_with_unknown_source_quantity_and_unit():
    row = _work(None, "", calc_quantity=6, calc_unit="м²", conversion_source="Размеры 2 × 3 м; чертёж А!R7")
    assert audit_vor_result_integrity([row]) == []


@pytest.mark.asyncio
async def test_known_derived_measure_does_not_require_invented_source_quantity_or_unit(tmp_path):
    row = _work(None, "", calc_quantity=6, calc_unit="м²", conversion_source="Размеры 2 × 3 м; чертёж А!R7")
    output, result = await _export(tmp_path, row)
    assert not result["missing"] and not result.get("document_issues")
    workbook = openpyxl.load_workbook(output, data_only=True)
    sheet = workbook["ВОР"]
    assert sheet["E4"].value is None and sheet["F4"].value is None
    assert [sheet["P4"].value, sheet["Q4"].value] == ["м²", 6]
    workbook.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("source_quantity,source_unit,model_quantity,model_unit,expected_codes", [
    (8, "шт", None, "", {"missing_source_quantity", "missing_source_unit"}),
    (0, "шт", None, "шт", {"missing_source_quantity"}),
    (8, "шт", 8, "", {"missing_source_unit"}),
    (None, "", None, "", set()),
    (None, "шт", None, "шт", set()),
])
async def test_calculated_measure_cannot_hide_omitted_known_source_values(
    tmp_path, source_quantity, source_unit, model_quantity, model_unit, expected_codes,
):
    from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog
    from proxy.services.chat_evidence_application_service import build_vor_result_readiness

    source = tmp_path / "Разные величины.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Объёмы Ω"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Источник с собственной величиной", source_unit, source_quantity])
    workbook.save(source)
    workbook.close()
    catalog = load_tabular_source_catalog(source)
    row = _work(model_quantity, model_unit, calc_quantity=192, calc_unit="м",
                conversion_source="8 изделий по 24 м; Объёмы Ω!R2")
    before = deepcopy(row)
    issues = audit_vor_result_integrity([row], source_catalog=catalog)
    assert {issue["code"] for issue in issues} == expected_codes
    readiness = build_vor_result_readiness(rows=[row], has_file=True, remaining_locators=[],
        catalog_work_count=1, model_errors=[], source_catalog=catalog)
    assert readiness["vor_complete"] is not bool(expected_codes)
    output = tmp_path / "Результат.xlsx"
    result = await _default_vor_adapter(source, {"decisions": [row]}, output, lambda *_: None)
    assert {i["code"] for i in result.get("document_issues", [])} == expected_codes
    assert row == before, "missing original values are reported, never filled from the source"
    workbook = openpyxl.load_workbook(output, data_only=True)
    sheet = workbook["ВОР"]
    assert sheet["F4"].value == model_quantity
    assert (sheet["E4"].value or "") == model_unit
    assert [sheet["P4"].value, sheet["Q4"].value] == ["м", 192]
    workbook.close()


@pytest.mark.parametrize("quantity,factor,calculated,expected_codes", [
    (12, 0.5, 999, {"conversion_result_mismatch"}),
    (3, 0.1, 0.3, set()),
    (0, 0.5, 0, set()),
    (None, 24, 192, {"missing_conversion_source_quantity"}),
])
def test_only_explicit_scalar_arithmetic_is_checked_without_replacing_authored_values(
    quantity, factor, calculated, expected_codes,
):
    row = _work(quantity, calc_quantity=calculated, calc_unit="м²", conversion_factor=factor,
                conversion_source="Модель указала формулу и исходную строку")
    before = deepcopy(row)
    codes = {i["code"] for i in audit_vor_result_integrity([row])}
    assert codes == expected_codes
    assert row == before


@pytest.mark.asyncio
async def test_an_inconsistent_explicit_scalar_result_stays_visible_as_an_unfinished_document(tmp_path):
    row = _work(calc_quantity=999, calc_unit="м²", conversion_factor=0.5,
                conversion_source="12 / 2 = 6; исходная строка Объёмы Ω!R2")
    output, result = await _export(tmp_path, row)
    assert "conversion_result_mismatch" in {i["code"] for i in result["document_issues"]}
    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook["ВОР"]["F4"].value == 12 and workbook["ВОР"]["Q4"].value == 999
    assert workbook["ВОР"]["I4"].value == "ТРЕБУЕТ РЕШЕНИЯ"
    workbook.close()


@pytest.mark.asyncio
async def test_factor_alone_does_not_invent_the_model_calculation_result(tmp_path):
    row = _work(conversion_factor=0.5, calc_unit="м²", conversion_source="12 / 2; Объёмы Ω!R2")
    output, result = await _export(tmp_path, row)
    assert "missing_calc_quantity" in {i["code"] for i in result["document_issues"]}
    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook["ВОР"]["F4"].value == 12
    assert workbook["ВОР"]["Q4"].value is None
    assert workbook["ВОР"]["R4"].value == 0.5
    workbook.close()


@pytest.mark.asyncio
async def test_an_authored_expression_is_text_and_never_an_excel_formula(tmp_path):
    output, _ = await _export(tmp_path, _work(calc_quantity="=12/2", calc_unit="м²", conversion_source="=12/2"))
    workbook = openpyxl.load_workbook(output, data_only=False)
    assert workbook["ВОР"]["Q4"].data_type != "f" and workbook["ВОР"]["S4"].data_type != "f"
    workbook.close()


@pytest.mark.asyncio
async def test_plain_work_does_not_gain_a_derived_measure(tmp_path):
    output, result = await _export(tmp_path, _work())
    assert not result.get("document_issues")
    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook["ВОР"].max_column == 15 and workbook["ВОР"]["F4"].value == 12
    workbook.close()


@pytest.mark.parametrize("format_name", ["json", "markdown", "inline"])
def test_calc_unit_survives_result_transports(format_name):
    from proxy.services.chat_evidence_application_service import parse_model_rag_result

    row = _work(calc_quantity=6, calc_unit="м²", conversion_factor=0.5, conversion_source="12 / 2; source row")
    if format_name == "json":
        text = json.dumps({"artifact_kind": "vor", "rows": [row]}, ensure_ascii=False)
    elif format_name == "markdown":
        keys = tuple(row)
        text = "artifact_kind: vor\n\n" + "\n".join([
            "| " + " | ".join(keys) + " |", "| " + " | ".join("---" for _ in keys) + " |",
            "| " + " | ".join(str(row[k]) for k in keys) + " |",
        ])
    else:
        text = "artifact_kind: vor\n\n" + "; ".join(f"{k}: {v}" for k, v in row.items())
    parsed = parse_model_rag_result(text)
    assert parsed is not None
    actual = parsed[1][0]
    assert (actual["calc_unit"], actual["calc_quantity"], actual["quantity"], actual["unit"]) == ("м²", 6, 12, "м²")
