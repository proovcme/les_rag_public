"""РИМ-трасса ЛСР: графы 2-12, цены ФГИС ЦС и итог позиции."""

from __future__ import annotations

import pytest

from proxy.services.rim_lsr_trace_service import (
    build_lsr_trace,
    build_lsr_trace_from_visible_rows,
    build_position_trace,
    positions_from_visible_lsr_rows,
)


CODE = "ГЭСН12-01-034-02"


def test_explicit_covered_work_is_preserved_and_charged_once(monkeypatch, tmp_path):
    from copy import deepcopy
    from openpyxl import load_workbook
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
    parent = _active_resource_norm(monkeypatch)
    parent["norm_status"] = "NEEDS_RESEARCH"
    child = dict(parent, source_row="S!R3", operation_id="child", title="Included operation",
                 covered_by_work_id="R2", norm_status="EXACT", coverage="The cited combined norm includes this operation")
    original = deepcopy([child, parent])
    trace = build_lsr_trace_from_visible_rows(original)
    base = build_lsr_trace_from_visible_rows([parent])
    assert trace["summary"]["positions_built"] == 1
    assert trace["summary"]["covered_rows"] == 1
    assert trace["summary"]["preliminary_amount"] == base["summary"]["preliminary_amount"]
    covered, billed = trace["row_bindings"]
    assert covered["status"] == "covered"
    assert covered["covered_by_work_id"] == "S!R2"
    assert covered["covered_by_operation_id"] == "op:op_model"
    assert covered["source_row"] == child
    assert original == [child, parent]
    assert covered["readiness"]["prices_provided"] == billed["readiness"]["prices_provided"]
    assert covered["readiness"]["applicability_substantiated"] is False
    assert trace["row_projection"][0]["operation_id"] == "op:child"
    output = render_lsr_xlsx(trace, tmp_path / "covered.xlsx")
    wb = load_workbook(output, data_only=True)
    body = " ".join(str(cell.value or "") for row in wb.worksheets[0] for cell in row)
    assert "Included operation" in body and "S!R2" in body and "повторно не начисляется" in body
    check = " ".join(str(cell.value or "") for row in wb["Проверка"] for cell in row)
    assert "S!R3" in check and "S!R2" in check


@pytest.mark.parametrize("changes, reason", [
    ({"covered_by_work_id": "R99"}, "coverage_target_missing"),
    ({"covered_by_work_id": "R3"}, "coverage_cycle"),
    ({"coverage": ""}, "coverage_basis_missing"),
    ({"evidence_refs": []}, "coverage_evidence_missing"),
    ({"evidence_bindings": {}}, "coverage_evidence_missing"),
    ({"row_role": "material"}, "coverage_role_conflict"),
    ({"norm_status": "CONFLICT"}, "coverage_basis_unresolved"),
    ({"condition_match": "unknown"}, "coverage_basis_unresolved"),
    ({"quantity": None}, "coverage_quantity_conflict"),
    ({"_evidence_check": {"card_body_shown": False}}, "coverage_basis_unresolved"),
])
def test_invalid_work_coverage_is_visible_and_never_silently_billed(monkeypatch, changes, reason):
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", covered_by_work_id="R2",
                 coverage="Included by the cited norm", **{})
    child.update(changes)
    result = build_lsr_trace_from_visible_rows([parent, child])
    assert result["summary"]["positions_built"] == 1
    binding = result["row_bindings"][1]
    assert binding["status"] == "coverage_conflict"
    assert binding["coverage_error"] == reason
    assert result["summary"]["result_status"] == "priced_partial"
    assert any(item["code"] == "S!R3" for item in result["blockers"])
    assert binding["source_row"] == child


def test_identical_codes_and_prose_do_not_infer_work_coverage(monkeypatch):
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", coverage="Already included, charge once")
    result = build_lsr_trace_from_visible_rows([parent, child])
    assert result["summary"]["positions_built"] == 2
    assert result["summary"]["covered_rows"] == 0


def test_coverage_chains_cycles_and_ambiguous_operations(monkeypatch):
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", covered_by_work_id="R2", coverage="Cited composition")
    nested = dict(child, source_row="S!R4", operation_id="nested", covered_by_work_id="child")
    result = positions_from_visible_lsr_rows([nested, child, parent])
    assert len(result["positions"]) == 1
    assert [b.get("covered_by_operation_id") for b in result["row_bindings"][:2]] == ["op:op_model"] * 2
    sibling = dict(parent, operation_id="sibling")
    result = positions_from_visible_lsr_rows([parent, sibling, child])
    assert result["row_bindings"][2]["coverage_error"] == "coverage_target_ambiguous"
    child["covered_by_work_id"] = "op_model"
    assert positions_from_visible_lsr_rows([parent, sibling, child])["row_bindings"][2]["status"] == "covered"
    child["covered_by_work_id"] = "nested"
    result = positions_from_visible_lsr_rows([parent, child, nested])
    assert [b["coverage_error"] for b in result["row_bindings"][1:]] == ["coverage_cycle"] * 2


def test_covered_work_cannot_close_an_unbound_parent(monkeypatch):
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", covered_by_work_id="R2", coverage="Cited composition")
    parent.update(norm_code="", norm_status="MISSING")
    result = build_lsr_trace_from_visible_rows([child, parent])
    assert result["summary"]["positions_built"] == 0
    assert result["row_bindings"][0]["coverage_error"] == "coverage_target_unbound"


def test_covered_chain_cannot_hide_intermediate_missing_evidence(monkeypatch):
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", covered_by_work_id="R2", coverage="Cited composition", evidence_refs=[])
    nested = dict(child, source_row="S!R4", operation_id="nested", covered_by_work_id="child", evidence_refs=["Q1.H1"])
    result = positions_from_visible_lsr_rows([nested, parent, child])
    assert result["row_bindings"][0]["status"] == "coverage_conflict"
    assert result["row_bindings"][2]["coverage_error"] == "coverage_evidence_missing"


def test_readiness_asks_about_billed_work_without_asking_covered_row_for_another_norm(monkeypatch):
    from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness
    from proxy.services.lsr_workbook_adapter_service import collect_lsr_workbook_gaps
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", covered_by_work_id="R2", coverage="Cited composition", norm_code="")
    parent["norm_status"] = "NEEDS_RESEARCH"
    trace = build_lsr_trace_from_visible_rows([parent, child])
    gaps = collect_lsr_workbook_gaps(trace)
    assert not any(item.startswith("row:S!R3:") for item in gaps)
    readiness = build_lsr_result_readiness(packaged=True, rows=[parent, child], missing=gaps,
        row_projection=trace["row_projection"], pricing_status=trace["summary"]["result_status"])
    assert not any(gap["locator"] == "S!R3" for gap in readiness["dimensions"]["norms"]["gaps"])
    assert readiness["dimensions"]["norms"]["complete"] is False


def _active_resource_norm(monkeypatch):
    from proxy.services import gesn_service
    from tests.lsr_citation_support import issued_norm_metadata
    norm = {"code": "ГЭСНм08-02-409-09", "name": "Authored work", "unit": "100 м",
            "resources": [
                {"kind": "material", "code": "A-1", "name": "Target", "unit": "m", "per_unit": 2, "price": 777},
                {"kind": "material", "code": "A-10", "name": "Other", "unit": "m", "per_unit": 3, "price": 888},
            ]}
    issued = issued_norm_metadata(norm, synthetic=True)
    norm["_source_sha256"] = issued["normative_base_sha256"]

    def get_norm(*_args, **kwargs):
        if kwargs.get("expected_base_sha256"):
            assert kwargs["expected_base_sha256"] == issued["normative_base_sha256"]
        return norm

    monkeypatch.setattr(gesn_service, "get_norm", get_norm)
    return {"source_row": "S!R2", "operation_id": "op_model", "row_role": "work", "title": "Authored work",
            "quantity": 1000, "unit": "м", "norm_code": norm["code"], "norm_status": "EXACT",
            "nr_sp_rule_id": "nrsp-49.3", "evidence_refs": ["Q1.H1"]}


@pytest.mark.parametrize("resource_first", [False, True])
def test_explicit_actions_on_linked_resource_rows_reach_the_chosen_work(monkeypatch, resource_first):
    from copy import deepcopy
    parent = _active_resource_norm(monkeypatch)
    child = dict(parent, source_row="S!R3", operation_id="child", norm_code="",
                 covered_by_work_id="S!R2", coverage="The cited composition includes reinstallation")
    resource = dict(parent, row_role="resource", operation_id="resource-action", work_row="S!R2",
                    resource_bindings=[{"action": "reuse", "target_resource_code": "A-1",
                                        "reason": "Authored reuse", "source_refs": ["S!R2"]}])
    rows = [resource, child, parent] if resource_first else [parent, child, resource]
    before = deepcopy(rows)
    trace = build_lsr_trace_from_visible_rows(rows)
    work = next(binding for binding in trace["row_bindings"] if binding["status"] == "bound")
    reused = next(item for item in work["resource_decisions"] if item["code"] == "A-1")
    assert reused["price"] == 0 and reused["cost"] == 0
    assert reused["per_unit"] == 2 and reused["qty"] == 20
    assert sum(binding["status"] == "covered" for binding in trace["row_bindings"]) == 1
    assert len(trace["sections"][0]["positions"]) == 1
    assert work["resource_binding_audit"][0]["source_row"] == "S!R2"
    assert rows == before


@pytest.mark.parametrize("bad_link", ["ambiguous", "missing", "different_norm"])
def test_linked_resource_action_does_not_choose_a_parent_or_reinterpret_its_norm(monkeypatch, bad_link):
    from copy import deepcopy
    parent = _active_resource_norm(monkeypatch)
    resource = dict(parent, row_role="resource", operation_id="resource-action", work_row="S!R2",
                    resource_bindings=[{"action": "reuse", "target_resource_code": "A-1",
                                        "reason": "Authored reuse", "source_refs": ["S!R2"]}])
    rows = [parent, resource]
    if bad_link == "ambiguous":
        rows.append(dict(parent, operation_id="another-work"))
    elif bad_link == "missing":
        resource["work_row"] = "S!R999"
    else:
        resource["norm_code"] = "ГЭСНм08-02-409-01"
    before = deepcopy(rows)
    trace = build_lsr_trace_from_visible_rows(rows)
    works = [binding for binding in trace["row_bindings"] if binding["status"] == "bound"]
    assert all(next(item for item in work["resource_decisions"] if item["code"] == "A-1")["price"] == 777 for work in works)
    if bad_link == "different_norm":
        assert works[0]["resource_binding_conflicts"][0]["reason"] == "linked resource norm disagrees with its work"
    else:
        assert any(binding["status"] == "resource_link_conflict" for binding in trace["row_bindings"])
    assert trace["summary"]["result_status"] == "priced_partial" and rows == before


def test_active_model_resource_reuse_preserves_consumption_and_zero_outranks_book(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": "reuse", "target_resource_code": "A-1",
                                 "reason": "Authored reuse", "basis_ref": "S!R2", "source_refs": ["Q1.H1"]}]

    class Book:
        def __init__(self):
            self.lookups = []

        def lookup(self, code):
            self.lookups.append(code)
            return {"price_current": 100, "unit": "m"}

    book = Book()
    trace = build_lsr_trace_from_visible_rows([row], pricebook=book)
    binding = trace["row_bindings"][0]
    resources = binding["resource_decisions"]
    reused = next(item for item in resources if item["code"] == "A-1")
    assert (reused["unit"], reused["per_unit"], reused["qty"], reused["price"], reused["cost"]) == ("m", 2, 20, 0, 0)
    assert reused["price_source"] == "reuse"
    assert reused["price_basis"] == "S!R2"
    assert book.lookups == ["A-10"]
    assert binding["resource_binding_audit"][0]["target"]["per_unit"] == 2
    assert binding["source_row"]["resource_bindings"] == row["resource_bindings"]
    assert binding["position"]["source_quantity"] == 1000
    assert binding["position"]["qty"] == 10


def test_active_model_resource_exclusion_can_produce_empty_resources_without_rehydration(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": "exclude", "target_resource_code": code,
                                 "reason": "Authored exclusion", "source_refs": ["Q1.H1"]}
                                for code in ("A-1", "A-10")]
    position = positions_from_visible_lsr_rows([row])["positions"][0]
    assert position["resources"] == []
    result = build_position_trace(position)
    assert result["summary"]["resource_decisions"] == []
    assert [item["status"] for item in result["summary"]["resource_binding_audit"]] == ["applied", "applied"]


def test_active_model_resource_invalid_replacement_preserves_model_and_target_and_is_visible(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    action = {"action": "replace", "target_resource_code": "A-1", "resource_name": "Replacement",
              "unit": "m", "quantity": -1, "reason": "Authored replacement", "source_refs": ["Q1.H1"]}
    row["resource_bindings"] = [action]
    trace = build_lsr_trace_from_visible_rows([row])
    binding = trace["row_bindings"][0]
    assert binding["code"] == row["norm_code"]
    assert binding["norm_status"] == "EXACT"
    assert [item["code"] for item in binding["resource_decisions"]] == ["A-1", "A-10"]
    assert binding["resource_bindings"] == [action]
    assert binding["resource_binding_conflicts"][0]["proposal"] == action
    assert binding["readiness"]["conflicts_ok"] is False
    assert "resource_binding_conflict" in binding["readiness"]["open_reasons"]
    assert any("Конфликт действия над ресурсом" in flag for flag in trace["summary"]["flags"])
    assert trace["sections"][0]["positions"][0]["summary"]["full_amount"] is None


def test_active_model_resource_source_work_quantity_is_not_norm_quantity(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": "add", "resource_name": "Authored material", "unit": "m",
                                 "quantity_basis": "source_work", "reason": "Source quantity", "basis_ref": "S!R2"}]
    position = positions_from_visible_lsr_rows([row])["positions"][0]
    added = position["resources"][-1]
    assert position["qty"] == 10
    assert position["source_quantity"] == 1000
    assert added["per_unit"] == 100
    assert added["resource_binding"]["physical_quantity"] == 1000
    trace = build_position_trace(position)
    added_decision = next(item for item in trace["summary"]["resource_decisions"] if item["name"] == "Authored material")
    assert added_decision["qty"] == 1000


def test_active_model_resource_optional_review_does_not_invent_component_gaps(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    row["resource_review"] = {"status": "keep_all_confirmed", "reason": "Authored whole-resource review"}
    position = positions_from_visible_lsr_rows([row])["positions"][0]
    assert "resources" not in position
    summary = build_position_trace(position)["summary"]
    assert summary["resource_review_status"] == "keep_all_confirmed"
    assert not {"labor_review_unresolved", "machine_review_unresolved", "material_review_unresolved"} & set(summary["completeness_reasons"])


def test_active_model_resource_supply_preserves_materials_and_exposes_unresolved_overlap(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": "reuse", "target_resource_code": "A-1",
                                 "reason": "Authored reuse", "basis_ref": "S!R2"}]
    supply = {"source_row": "S!R3", "row_role": "material", "work_row": "S!R2", "supply_kind": "supply",
              "title": "Separate supply", "unit": "m", "quantity": 1000}
    trace = build_lsr_trace_from_visible_rows([row, supply])
    binding = trace["row_bindings"][0]
    assert binding["resource_binding_audit"][0]["status"] == "applied"
    overlap = next(item for item in binding["resource_binding_conflicts"] if item["code"] == "separate_supply_possible_overlap")
    assert {item["code"] for item in overlap["norm_material_resources"]} == {"A-1", "A-10"}
    assert {item["code"] for item in binding["resource_decisions"]} == {"A-1", "A-10"}
    assert next(item for item in binding["resource_decisions"] if item["code"] == "A-1")["price"] == 0
    assert binding["readiness"]["conflicts_ok"] is False
    assert trace["sections"][0]["positions"][0]["summary"]["full_amount"] is None


@pytest.mark.parametrize("ambiguous", [False, True])
def test_active_model_resource_linked_supply_exact_exclusion_keeps_consumables_and_parent_ambiguity(monkeypatch, ambiguous):
    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": "exclude", "target_resource_code": "A-1",
                                 "reason": "Supplied separately", "basis_ref": "S!R3", "source_refs": ["Q1.H1"]}]
    row["_resource_binding_evidence_checks"] = [{"index": 0, "target_code_in_refs": True}]
    supply = {"source_row": "S!R3", "row_role": "material", "work_row": "S!R2", "resource_code": "A-1",
              "title": "Separate supply", "unit": "m", "quantity": 20}
    rows = [row, supply]
    if ambiguous:
        rows.append({**row, "operation_id": "second_work", "title": "Other authored work"})
    trace = build_lsr_trace_from_visible_rows(rows)
    binding = trace["row_bindings"][0]
    assert [item["code"] for item in binding["resource_decisions"]] == ["A-10"]
    assert binding["resource_decisions"][0]["qty"] == 30
    assert bool(binding["resource_binding_conflicts"]) is ambiguous
    if ambiguous:
        assert binding["resource_binding_conflicts"][-1]["parent_ambiguous"] is True
        assert binding["readiness"]["conflicts_ok"] is False
    else:
        assert binding["readiness"]["conflicts_ok"] is True
    assert trace["row_bindings"][1]["row_role"] == "material"


def test_active_model_resource_legacy_supply_marker_preserves_exact_card(monkeypatch):
    row = _active_resource_norm(monkeypatch)
    position = positions_from_visible_lsr_rows([row])["positions"][0]
    position["exclude_gesn_materials"] = True
    trace = build_position_trace(position)
    assert [item["code"] for item in trace["summary"]["resource_decisions"]] == ["A-1", "A-10"]
    assert trace["summary"]["resource_binding_conflicts"][0]["code"] == "separate_supply_possible_overlap"
    assert trace["summary"]["full_amount"] is None


@pytest.mark.asyncio
@pytest.mark.parametrize("action,linked_supply", [("reuse", False), ("exclude", False), ("exclude", True)])
async def test_active_model_resource_public_result_reaches_actual_adapter_and_xlsx(tmp_path, monkeypatch, action, linked_supply):
    import copy
    import json
    from types import SimpleNamespace

    import openpyxl

    from proxy.services import attachment_source_catalog_service as catalog_service
    from proxy.services import chat_evidence_application_service as chat_service
    from proxy.services import lsr_workbook_adapter_service as adapter_service
    from proxy.services import gesn_service
    from tests.lsr_citation_support import issued_norm_metadata

    row = _active_resource_norm(monkeypatch)
    row["resource_bindings"] = [{"action": action, "target_resource_code": "A-1",
                                 "reason": "Source says it is supplied", "basis_ref": "S!R2", "source_refs": ["Q2.H1"]}]
    if action == "reuse":
        row["resource_bindings"][0]["explicit_price"] = 0
    authored_rows = [row]
    if linked_supply:
        row["resource_bindings"][0]["basis_ref"] = "S!R3"
        authored_rows.append({"source_row": "S!R3", "row_role": "material", "work_row": "S!R2",
                              "title": "Separate supply", "unit": "м", "quantity": 20,
                              "norm_code": "", "resource_code": "A-1", "evidence_refs": ["Q2.H1"]})
    answer = json.dumps({"rows": authored_rows}, ensure_ascii=False)
    chunks = [
        SimpleNamespace(content=f"Norm {row['norm_code']}\nA-1 Target 2 m\nA-10 Other 3 m", doc_name="smeta_norm_cards.v1",
                        meta={"model_evidence_ref": "Q1.H1", "global_evidence_id": "norm-card",
                              **issued_norm_metadata(gesn_service.get_norm(row["norm_code"]), synthetic=True)}),
        SimpleNamespace(content=f"Original source {'S!R3' if linked_supply else 'S!R2'}: Target A-1 is supplied", doc_name="Original source",
                        meta={"model_evidence_ref": "Q2.H1", "global_evidence_id": "source-card"}),
    ]
    parsed = chat_service.parse_model_rag_result(answer)
    assert parsed is not None and parsed[0] == answer
    decisions = chat_service.bind_row_evidence_ids(parsed[1], chunks)
    assert decisions[0]["evidence_refs"] == ["Q1.H1"]
    assert decisions[0]["_resource_binding_evidence_checks"][0]["evidence_bindings"] == {"Q2.H1": "source-card"}

    source_path = tmp_path / "source.xlsx"
    source_book = openpyxl.Workbook()
    source_book.active.title = "S"
    source_book.active.append(["Наименование", "Ед. изм.", "Количество"])
    source_book.active.append(["Authored work", "м", 1000])
    if linked_supply:
        source_book.active.append(["Separate supply", "м", 20])
    source_book.save(source_path)
    source_book.close()
    catalog = catalog_service.load_tabular_source_catalog(source_path)
    monkeypatch.setattr(catalog_service, "load_catalog_from_attachment_id", lambda _: catalog)

    class Book:
        region = "Synthetic region"
        source_sha256 = "a" * 64

        def __init__(self):
            self.lookups = []

        def lookup(self, code):
            self.lookups.append(code)
            return {"price_current": 100, "unit": "m"}

    book = Book()
    monkeypatch.setattr(adapter_service, "visible_pricebook_stems", lambda: ["synthetic_2kv2026"])
    monkeypatch.setattr(adapter_service.fgis_price_service, "resolve_pricebook_path", lambda *_args, **_kwargs: tmp_path / "synthetic_2kv2026.parquet")
    monkeypatch.setattr(adapter_service.fgis_price_service, "get_pricebook", lambda _: book)
    captured = {}
    calculate = adapter_service.build_lsr_trace_from_visible_rows

    def capture_calculation(*args, **kwargs):
        trace = calculate(*args, **kwargs)
        captured["trace"] = copy.deepcopy(trace)
        return trace

    monkeypatch.setattr(adapter_service, "build_lsr_trace_from_visible_rows", capture_calculation)
    output = tmp_path / "authored-reuse.xlsx"
    result = await adapter_service.build_lsr_workbook_from_decisions(
        source_path, {"attachment_id": "synthetic-source", "decisions": decisions,
                      "pricebook": "synthetic_2kv2026", "_source_text": "Authored source"},
        output, lambda *_args: None,
    )
    assert result["file_path"] == output and output.is_file()
    binding = captured["trace"]["row_bindings"][0]
    assert binding["resource_binding_conflicts"] == []
    assert binding["resource_bindings"] == row["resource_bindings"]
    assert binding["resource_binding_audit"][0]["provenance"]["evidence_bindings"] == {"Q2.H1": "source-card"}
    if action == "reuse":
        reused = next(item for item in binding["resource_decisions"] if item["code"] == "A-1")
        assert (reused["qty"], reused["price"], reused["cost"]) == (20, 0, 0)
        assert reused["resource_binding"]["reason"] == "Source says it is supplied"
    else:
        assert all(item["code"] != "A-1" for item in binding["resource_decisions"])
    assert book.lookups == (["A-10", "A-1"] if linked_supply else ["A-10"])
    workbook = openpyxl.load_workbook(output, data_only=False)
    try:
        sheet = workbook["ЛСР РИМ"]
        if action == "reuse":
            resource_row = next(cells[1].row for cells in sheet.iter_rows() if cells[1].value == "A-1")
            assert sheet.cell(resource_row, 4).value == "m"
            assert sheet.cell(resource_row, 5).value == 2
            assert sheet.cell(resource_row, 7).value == 20
            assert sheet.cell(resource_row, 10).value == 0
            assert sheet.cell(resource_row, 12).value == f"=ROUND(G{resource_row}*J{resource_row}*K{resource_row},2)"
        elif linked_supply:
            supply_resource_rows = [cells for cells in sheet.iter_rows() if cells[1].value == "A-1"]
            assert len(supply_resource_rows) == 1
            assert supply_resource_rows[0][6].value == 20
            other = next(cells for cells in sheet.iter_rows() if cells[1].value == "A-10")
            assert other[6].value == 30
        else:
            assert all(cells[1].value != "A-1" for cells in sheet.iter_rows())
        check_text = "\n".join(str(cell.value or "") for cells in workbook["Проверка"].iter_rows() for cell in cells)
        assert "Source says it is supplied" in check_text
        assert "Q2.H1" in check_text and "S!R2" in check_text
        if linked_supply:
            assert "S!R3" in check_text
        assert ("повторно использовать A-1" if action == "reuse" else "исключить A-1") in check_text
    finally:
        workbook.close()


def test_mounting_norm_never_prices_an_explicit_supply_or_header(monkeypatch):
    from proxy.services import gesn_service
    calls = []
    monkeypatch.setattr(gesn_service, "get_norm", lambda *args: calls.append(args))
    for role in ("material", "equipment", "header"):
        row = {"source_row": "S!R7", "row_role": role, "norm_code": CODE, "norm_status": "EXACT",
               "title": "Authored item", "unit": "м²", "quantity": 61}
        bound = positions_from_visible_lsr_rows([row])
        assert not bound["positions"]
        binding = bound["row_bindings"][0]
        assert binding["status"] == "role_norm_conflict"
        assert binding["code"] == CODE and binding["row_role"] == role
        assert binding["source_row"] == row
    assert not calls


def test_model_accrual_rule_reaches_calculation_without_invented_default(monkeypatch):
    from proxy.services import gesn_service
    norm = {"code": "ГЭСНм08-02-409-09", "name": "Прокладка труб", "unit": "100 м", "resources": []}
    monkeypatch.setattr(gesn_service, "get_norm", lambda *_args, **_kwargs: norm)
    source = {"source_row": "СКС!R69", "title": "Прокладка труб", "unit": "м", "quantity": 1200,
              "norm_code": norm["code"], "row_role": "work", "nr_sp_rule_id": "nrsp-49.3"}
    position = positions_from_visible_lsr_rows([source])["positions"][0]
    assert position["nr_sp_rule_id"] == "nrsp-49.3"
    assert position["code"] == source["norm_code"] and position["qty"] == 12
    summary = build_position_trace(position)["summary"]
    assert summary["nr_sp_ownership"]["rule_id"] == "nrsp-49.3"
    assert (summary["nr_pct"], summary["sp_pct"]) == (97, 51)
    assert "nr_sp_unresolved" not in summary["completeness_reasons"]
    for rule_id, expected_status in [("", "ambiguous"), ("nrsp-1.1", "rejected")]:
        row = {**source, "nr_sp_rule_id": rule_id}
        other = positions_from_visible_lsr_rows([row])["positions"][0]
        gap = build_position_trace(other)["summary"]
        assert gap["nr_sp_ownership"]["catalog_status"] == expected_status
        assert "nr_sp_unresolved" in gap["completeness_reasons"]
        assert gap["nr_pct"] is None and gap["sp_pct"] is None


# ── эталон позиции: карточка основной базы + детерминированная ценовая книга ──
#
# Семя `config/domain/gesn_seed.yaml` больше не участвует в production-расчёте
# (`proxy/services/gesn_service.py::_merged_norms` merged только canonical-базу
# и оверлей ФГИС), а база цен не несёт вовсе.  Карточка `ГЭСН12-01-034-02`
# (source `structured_sqlite`, единица «100 м2») даёт нормативные расходы:
#   labor     1-100-25 «Средний разряд работы 2.5»  12,94 чел.-ч
#   machine   91.05.01-017 кран башенный 8 т         0,97 маш.-ч
#   machine   91.05.05-015 кран автомобильный 16 т   0,01 маш.-ч
#   machine   91.14.02-001 бортовой автомобиль 5 т   0,03 маш.-ч
#   machinist «Затраты труда машинистов», без шифра  1,01 маш.-ч (ФСЭМ → 4-100-*)
#   material  01.7.15.06-0111 гвозди строительные    0,0015 т
#   material  11.1.03.01-0076 бруски обрезные         0,4 м3
# Цены ниже — строки установленной книги `sankt-peterburg_2kv2026` (замер:
# `stage_b/live_830/gold_rebaseline_830.json`), зафиксированные здесь, чтобы
# тесты не зависели от того, какая книга лежит у оператора.
GOLD_QTY = 0.61
GOLD_PRICES = {
    "1-100-25": (507.56, "чел.-ч"),
    "91.05.01-017": (1048.03, "маш.-ч"),
    "91.05.05-015": (1644.26, "маш.-ч"),
    "91.14.02-001": (570.54, "маш.-ч"),
    # водители машинистов, на которые ФСЭМ детализирует строку «Затраты труда
    # машинистов» 1,01 маш.-ч при наличии выбранной книги
    "4-100-040": (596.61, "чел.-ч"),
    "4-100-060": (801.41, "чел.-ч"),
    "01.7.15.06-0111": (85761.36, "т"),
    # 11.1.03.01-0076 «Бруски обрезные…» в книге отсутствует намеренно: этой
    # строки нет в установленной книге, поэтому ресурс остаётся названной дырой.
    #
    # Две строки той же книги для второй эталонной позиции — карточка
    # «ГЭСНм08-03-575-01» (электромонтаж, «шт»), которую проверяет
    # tests/test_lsr_assembly_service.py: замер `data/price_base/sankt-peterburg_2kv2026.parquet`
    # (у болтов `price_current` пустой, а `price_current_eff` = 174,93 × 1,08 = 188,92 —
    # ровно то, что отдаёт движок сборки по полю `price_current_eff`).
    "1-100-42": (614.42, "чел.-ч"),
    "01.7.15.03-0042": (188.92, "кг"),
}

# Каждый эталонный иток — сумма строк карточки × цены книги, проверенная вручную:
GOLD_OZP = 4006.37        # 12,94 × 0,61 = 7,8934 чел.-ч × 507,56
GOLD_MACHINE = 640.59     # 0,5917×1048,03 + 0,0061×1644,26 + 0,0183×570,54
GOLD_ZPM = 490.00         # 0,0183×596,61 + 0,0061×801,41 + 0,5917×801,41
GOLD_EM = 1130.59         # ЭМ = машинам 640,59 + ЗПМ 490,00
GOLD_MAT = 78.47          # 0,0015 × 0,61 = 0,000915 т × 85 761,36
GOLD_DIRECT = 5215.43     # 4006,37 + 1130,59 + 78,47
GOLD_FOT = 4496.37        # ОЗП 4006,37 + ЗПМ 490,00
GOLD_NR = 4901.04         # ФОТ × 109 % (правило каталога nrsp-12, «Приказ 812/пр п.12»)
GOLD_SP = 2562.93         # ФОТ × 57 %  (правило каталога nrsp-12, «Приказ 774/пр п.12»)
GOLD_TOTAL = 12679.40     # прямые 5215,43 + НР 4901,04 + СП 2562,93
GOLD_LABOR_QTY = 7.8934
GOLD_MACHINIST_QTY = 0.6161
GOLD_HOLE = "price_missing:11.1.03.01-0076"

# Без выбранной книги позиция обязана сойтись в нуль и назвать каждую дыру:
# семь строк карточки, семь вопросов оператору (замер `no_book` в
# `stage_b/live_830/gold_rebaseline_830.json`).
NO_BASIS_REASONS = [
    "price_missing:91.14.02-001",
    "price_missing:91.05.05-015",
    "price_missing:91.05.01-017",
    "price_missing:Затраты труда машинистов",
    "price_missing:1-100-25",
    "price_missing:01.7.15.06-0111",
    "price_missing:11.1.03.01-0076",
]


class _Book:
    """Детерминированная ценовая книга: та же форма, что в `test_lsr_price_ownership`.

    `lookup(code)` отвечает строкой книги целиком, с «unit» — `_resolve_price_trace`
    сверяет единицу строки с единицей ресурса (`book_price_unit_conflict`), и книга,
    отвечающая только ценой, этого check не изображает.

    Поле `price_current_eff` дублирует `price_current`, как это делает настоящая
    конвертация `fgis_price_service` (текущая цена, иначе база × индекс): рутинный
    движок сборки ЛСР читает именно его, а РИМ-трасса — `price_current`.  Один
    двойник с обоими полями заведомо даёт двум движкам одну и ту же сумму на одной
    и той же книге (см. `tests/test_lsr_assembly_service.py`).
    """

    def __init__(self) -> None:
        self._prices = dict(GOLD_PRICES)

    def lookup(self, code: str):
        hit = self._prices.get(str(code).strip())
        if hit is None:
            return None
        return {
            "price_current": hit[0],
            "price_current_eff": hit[0],
            "unit": hit[1],
        }


def test_invalid_conversion_does_not_fall_back_to_pricing_otherwise_valid_quantity():
    from proxy.services.rim_lsr_trace_service import _norm_qty_from_visible_row

    for field in ("conversion_factor", "calc_quantity"):
        row = {"quantity": 120, "unit": "м", "calc_quantity": 36600,
               "conversion_factor": 305, "conversion_source": "СКС!R17"}
        row[field] = "305 м * 120 шт"
        quantity, trace = _norm_qty_from_visible_row(row, CODE, {"unit": "100 м"})
        assert quantity is None
        assert trace["status"] == "needs_input"
        assert trace["invalid_field"] == field
        assert trace["authored_value"] == row[field]


def test_equivalent_source_measures_produce_identical_money_without_double_scaling():
    """61 m² and 0.61 hundreds of m² are the same physical work volume."""
    for quantity, unit in [(61, "м2"), (0.61, "100 м2"), (0.061, "1000 м²"), (6.1, "10 м²")]:
        authored = {"source_row": "R1", "row_role": "work", "title": "Same physical volume",
                    "norm_code": CODE, "quantity": quantity, "unit": unit}
        original = dict(authored)
        result = build_lsr_trace_from_visible_rows([authored], pricebook=_Book())
        assert result["summary"]["total"] == GOLD_TOTAL
        assert result["row_bindings"][0]["quantity_trace"]["norm_quantity"] == GOLD_QTY
        assert authored == original


def test_source_measure_multiplier_survives_si_conversion(monkeypatch):
    from proxy.services import gesn_service
    norm = {"code": CODE, "name": "Authored work", "unit": "т", "resources": []}
    monkeypatch.setattr(gesn_service, "get_norm", lambda *_args, **_kwargs: norm)
    authored = {"source_row": "R1", "row_role": "work", "title": "Authored work",
                "norm_code": CODE, "quantity": 1, "unit": "1000 кг"}
    position = positions_from_visible_lsr_rows([authored])["positions"][0]
    assert position["qty"] == 1
    assert position["quantity_trace"]["source_quantity"] == 1
    assert position["quantity_trace"]["source_unit"] == "1000 кг"


def test_source_measure_overflow_is_an_explicit_quantity_gap():
    from proxy.services.rim_lsr_trace_service import _norm_qty_from_visible_row
    quantity, trace = _norm_qty_from_visible_row(
        {"quantity": 2, "unit": str(10 ** 308) + " м"}, CODE, {"unit": "м"}
    )
    assert quantity is None
    assert trace["status"] == "needs_input"


def test_bad_numeric_row_preserves_good_row_calculation_and_gap():
    valid = {"source_row": "R1", "row_role": "work", "title": "First",
             "norm_code": CODE, "quantity": 61, "unit": "м2"}
    invalid = {**valid, "source_row": "R2", "conversion_factor": "305 м * 120 шт"}
    good = build_lsr_trace_from_visible_rows([valid], pricebook=_Book())
    combined = build_lsr_trace_from_visible_rows([valid, invalid], pricebook=_Book())
    assert combined["summary"]["total"] == good["summary"]["total"] > 0
    # The broken row is the only difference between the two revisions, so the one
    # surviving position is the gold one (base card rows × `_Book` prices).
    assert combined["summary"]["total"] == GOLD_TOTAL
    assert len(combined["row_bindings"]) == 2
    gap = combined["row_bindings"][1]
    assert gap["status"] == "needs_input"
    assert gap["source_row"] == invalid
    assert gap["quantity_trace"]["invalid_field"] == "conversion_factor"


def _rows_by_type(trace):
    return {row["type"]: row for row in trace["rows"]}


def test_rim_trace_from_base_norm_card_reproduces_gold_position():
    trace = build_position_trace({"code": CODE, "qty": GOLD_QTY}, pricebook=_Book())
    summary = trace["summary"]

    assert summary["flags"]
    # ADR-16 § 7 still bites, just on the other side of the card: the canonical base
    # carries no prices at all (0 of 63 673 labour rows), so nothing is counted from a
    # card and the only open flag is the resource the chosen book does not price.
    assert all(
        not str(flag).startswith("цена взята из карточки нормы") for flag in summary["flags"]
    )
    assert all("11.1.03.01-0076" in str(flag) for flag in summary["flags"])
    assert summary["card_priced_resources"] == []
    assert summary["amount_status"] == "partial"
    assert summary["full_amount"] is None
    assert summary["completeness_reasons"] == [GOLD_HOLE]
    assert summary["ozp"] == GOLD_OZP
    assert summary["em"] == GOLD_EM
    assert summary["mat"] == GOLD_MAT
    assert summary["direct"] == GOLD_DIRECT
    assert summary["fot"] == GOLD_FOT
    assert summary["nr"] == GOLD_NR
    assert summary["sp"] == GOLD_SP
    assert summary["total"] == GOLD_TOTAL
    # нормативные чел.-ч (в шапку многопозиционной формы) — те же, что в групп-строках
    assert summary["labor_qty"] == GOLD_LABOR_QTY
    assert summary["machinist_qty"] == GOLD_MACHINIST_QTY

    rows = _rows_by_type(trace)
    assert rows["work"]["columns"]["2"] == CODE
    assert rows["work"]["columns"]["4"] == "100 м2"
    assert rows["work"]["columns"]["7"] == 0.61
    assert rows["group_labor"]["columns"]["7"] == GOLD_LABOR_QTY
    assert rows["group_labor"]["columns"]["12"] == GOLD_OZP
    assert rows["group_machine"]["columns"]["12"] == GOLD_EM
    assert rows["group_machinist"]["columns"]["12"] == GOLD_ZPM
    assert rows["group_material"]["columns"]["12"] == GOLD_MAT
    # итог позиции — сумма напечатанных колонок, а не отдельное число калькулятора:
    # прямые = ОЗП+ЭМ+М, ФОТ = ОЗП+ЗПМ, всего = прямые+НР+СП
    assert rows["direct_total"]["columns"]["12"] == round(GOLD_OZP + GOLD_EM + GOLD_MAT, 2)
    assert rows["fot"]["columns"]["12"] == round(GOLD_OZP + GOLD_ZPM, 2)
    assert rows["position_total"]["columns"]["12"] == round(GOLD_DIRECT + GOLD_NR + GOLD_SP, 2)
    assert rows["nr"]["source"] == "Пр/812"
    assert rows["sp"]["source"] == "Пр/774"


def test_base_norm_card_without_a_chosen_book_totals_zero_and_names_every_hole():
    """The intended honest outcome of the new contract: 0, and seven named questions.

    The canonical base carries no prices (labour, machine and material rows alike), so
    a position calculated from a base card and no ценовая книга must not invent a sum:
    it lands on 0.0, stays ``partial``, and lists one ``price_missing:<шифр>`` reason per
    card row with the operator question that closes it.
    """
    summary = build_position_trace({"code": CODE, "qty": GOLD_QTY})["summary"]

    assert summary["total"] == 0.0
    assert summary["direct"] == 0.0
    assert summary["ozp"] == 0.0
    assert summary["em"] == 0.0
    assert summary["mat"] == 0.0
    assert summary["nr"] == 0.0
    assert summary["sp"] == 0.0
    assert summary["amount_status"] == "partial"
    assert summary["full_amount"] is None
    assert summary["card_priced_resources"] == []
    assert summary["completeness_reasons"] == NO_BASIS_REASONS
    # нормативные количества при этом не теряются: их даёт карточка, а не книга
    assert summary["labor_qty"] == GOLD_LABOR_QTY
    assert summary["machinist_qty"] == GOLD_MACHINIST_QTY
    questions = [str(flag) for flag in summary["flags"]]
    assert any(flag.startswith("нужна ставка ОЗП: Средний разряд работы 2.5 (1-100-25)")
               for flag in questions)
    assert any(flag.startswith("нужен КАЦ: Гвозди строительные (01.7.15.06-0111)")
               for flag in questions)
    assert any(flag.startswith("нужна цена эксплуатации машины: Краны башенные")
               for flag in questions)
    assert any(flag.startswith("нужна ставка ЗПМ: Затраты труда машинистов (—)")
               for flag in questions)
    assert len(questions) == len(NO_BASIS_REASONS)


def test_visible_rows_accept_colon_prefixed_norm_codes(monkeypatch):
    from proxy.services import gesn_service

    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda *_args, **_kwargs: {
            "code": "ГЭСНм08-03-575-01",
            "name": "Монтаж прибора или аппарата",
            "unit": "шт",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    )
    trace = build_lsr_trace_from_visible_rows([
        {
            "basis": "ГЭСНм:08-03-575-01",
            "title": "Монтаж прибора или аппарата",
            "unit": "шт",
            "quantity": 1,
        }
    ])

    summary = trace["summary"]
    assert summary["bound_rows"] == 1
    assert trace["row_bindings"][0]["status"] == "bound"
    assert trace["sections"][0]["positions"][0]["code"] == "ГЭСНм08-03-575-01"


def test_visible_rows_accept_engineering_count_unit_aliases(monkeypatch):
    from proxy.services import gesn_service

    catalog = {
        "ГЭСНм08-03-575-01": {
            "code": "ГЭСНм08-03-575-01",
            "name": "Установка электротехнического аппарата",
            "unit": "шт",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
        "ГЭСН17-01-010-01": {
            "code": "ГЭСН17-01-010-01",
            "name": "Установка ревизионных люков",
            "unit": "100 шт",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    }
    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda code, *args, **kwargs: catalog.get(str(code or "")),
    )
    trace = build_lsr_trace_from_visible_rows([
        {
            "basis": "ГЭСНм:08-03-575-01",
            "title": "Установка электротехнического аппарата",
            "unit": "компл",
            "quantity": 1,
        },
        {
            "basis": "ГЭСН:17-01-010-01",
            "title": "Установка ревизионных люков",
            "unit": "шт",
            "quantity": 10,
        },
    ])

    summary = trace["summary"]
    assert summary["bound_rows"] == 2
    assert trace["row_bindings"][0]["quantity_trace"]["status"] == "direct_from_row"
    assert trace["row_bindings"][1]["quantity_trace"]["status"] == "unit_conversion"
    assert trace["sections"][0]["positions"][0]["code"] == "ГЭСНм08-03-575-01"
    assert trace["sections"][0]["positions"][1]["code"] == "ГЭСН17-01-010-01"


def test_visible_rows_convert_piece_dimensions_to_area_norm_qty(monkeypatch):
    from proxy.services import gesn_service

    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda *_args, **_kwargs: {
            "code": "ГЭСН15-01-052-01",
            "name": "Разработка проема",
            "unit": "100 м2",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    )
    guessed = build_lsr_trace_from_visible_rows([
        {
            "source_row": "PDF!R2",
            "basis": "ГЭСН:15-01-052-01",
            "title": "Разработка проема 400х400 мм",
            "unit": "шт",
            "quantity": 10,
            "row_role": "work",
        },
    ])
    binding = guessed["row_bindings"][0]
    assert binding["status"] == "needs_input"
    assert binding["code"] == "ГЭСН15-01-052-01"
    assert binding["quantity_trace"]["source_quantity"] == 10.0

    cited = build_lsr_trace_from_visible_rows([
        {
            "source_row": "PDF!R2",
            "basis": "ГЭСН:15-01-052-01",
            "title": "Разработка проема 400х400 мм",
            "unit": "шт",
            "quantity": 10,
            "conversion_factor": 0.16,
            "conversion_source": "400×400 мм = 0,16 м2",
            "row_role": "work",
        },
    ])
    assert cited["row_bindings"][0]["status"] == "bound"
    assert cited["row_bindings"][0]["position"]["qty"] == 0.016
    assert cited["row_bindings"][0]["quantity_trace"]["status"] == "cited_conversion_factor"


def test_rim_trace_applies_resource_coefficients_by_kind():
    trace = build_position_trace(
        {"code": CODE, "qty": GOLD_QTY}, pricebook=_Book(), k_ozp=1.15, k_em=1.15
    )
    summary = trace["summary"]
    rows = _rows_by_type(trace)

    assert rows["coefficient"]["meta"] == {"k_ozp": 1.15, "k_em": 1.15}
    labor = next(row for row in trace["rows"] if row["type"] == "resource_labor")
    machine = next(row for row in trace["rows"] if row["type"] == "resource_machine")
    machinist = next(row for row in trace["rows"] if row["type"] == "resource_machinist")
    material = next(row for row in trace["rows"] if row["type"] == "resource_material")
    assert labor["columns"]["6"] == 1.15
    assert machine["columns"]["6"] == 1.15
    assert machinist["columns"]["6"] == 1.15
    assert material["columns"]["6"] == 1.0
    # each raised column is the uncoefficiented gold sum × 1,15 (7,8934×1,15 чел.-ч ×
    # 507,56 = 4607,33; 640,59×1,15 = 736,68; 490,00×1,15 = 563,50)
    assert summary["ozp"] == 4607.33
    assert summary["machine"] == 736.68
    assert summary["zpm"] == 563.50
    assert summary["ozp"] > GOLD_OZP
    assert summary["machine"] > GOLD_MACHINE
    assert summary["zpm"] > GOLD_ZPM
    # material takes no ОЗП/ЭМ coefficient, so it stays the gold 0,000915 т × 85 761,36
    assert summary["mat"] == GOLD_MAT
    assert summary["total"] == 14569.55
    assert summary["total"] > GOLD_TOTAL


def test_rim_trace_preserves_fgis_current_vs_base_index_columns():
    class FakePriceBook:
        def lookup(self, code):
            return {
                "91.05.05-015": {
                    "price_current": 1663.18,
                    "price_base": 1167.7,
                    "index": None,
                    "price_current_eff": 1663.18,
                },
                "01.7.15.06-0111": {
                    "price_current": None,
                    "price_base": 70296.2,
                    "index": 1.3,
                    "price_current_eff": 91385.06,
                },
            }.get(code)

    trace = build_position_trace(
        {
            "code": "ГЭСН00-00-000-00",
            "name": "Тестовая работа",
            "unit": "шт",
            "qty": 2,
            "nr_pct": 0,
            "sp_pct": 0,
            "resources": [
                {"kind": "machine", "code": "91.05.05-015", "name": "Кран", "unit": "маш.-ч", "per_unit": 1},
                {"kind": "material", "code": "01.7.15.06-0111", "name": "Гвозди", "unit": "т", "per_unit": 0.001},
            ],
        },
        pricebook=FakePriceBook(),
    )

    machine = next(row for row in trace["rows"] if row["type"] == "resource_machine")
    material = next(row for row in trace["rows"] if row["type"] == "resource_material")

    assert machine["source"] == "fgis_current"
    assert machine["columns"]["10"] == 1663.18
    assert "8" not in machine["columns"]
    assert "9" not in machine["columns"]

    assert material["source"] == "fgis_base_index"
    assert material["columns"]["8"] == 70296.2
    assert material["columns"]["9"] == 1.3
    assert material["columns"]["10"] == 91385.06
    assert material["columns"]["12"] == 182.77
    assert trace["summary"]["direct"] == 3509.13


def test_missing_price_and_unresolved_resource_review_separate_known_and_full_amount():
    trace = build_position_trace({
        "code": "ГЭСН00-00-000-00",
        "name": "Тестовая аналоговая работа",
        "unit": "шт",
        "qty": 2,
        "nr_pct": 0,
        "sp_pct": 0,
        "resource_review_status": "unresolved",
        "resource_review_reason": "не проверен чужой материал",
        "resources": [
            {"kind": "material", "code": "01.1", "name": "Материал без цены", "unit": "шт", "per_unit": 1},
        ],
    })

    summary = trace["summary"]
    assert summary["known_amount"] == 0.0
    assert summary["full_amount"] is None
    assert summary["amount_status"] == "partial"
    assert "resource_review_unresolved" in summary["completeness_reasons"]
    assert any(str(reason).startswith("price_missing:") for reason in summary["completeness_reasons"])
    assert _rows_by_type(trace)["position_total"]["columns"]["3"] == "Известная стоимость позиции"


def test_explicit_empty_resource_revision_is_not_rehydrated_from_norm(monkeypatch):
    """An empty model revision is data, not a request to restore raw norm resources."""
    from proxy.services import gesn_service

    norm = {
        "code": "ГЭСНм08-02-409-09",
        "name": "Прокладка труб гофрированных",
        "unit": "100 м",
        "resources": [
            {"kind": "labor", "code": "1-100-36", "name": "Средний разряд", "unit": "чел.-ч", "per_unit": 15.2, "price": 500},
        ],
        "_source_kind": "seed_yaml",
    }
    monkeypatch.setattr(gesn_service, "get_norm", lambda *_args, **_kwargs: norm)

    trace = build_position_trace({
        "code": norm["code"], "qty": 1.6, "resources": [], "nr_pct": 0, "sp_pct": 0,
        "resource_review_status": "unresolved",
        "resource_review_reason": "модель не подтвердила компоненты",
    })

    assert not [row for row in trace["rows"] if str(row.get("type") or "").startswith("resource_")]
    assert trace["summary"]["labor_qty"] == 0
    assert trace["summary"]["known_amount"] == 0


def test_hydrated_norm_resources_are_labor_normalized_at_render_boundary(monkeypatch):
    """The final renderer must not add average, aggregate and grade detail together."""
    from proxy.services import gesn_service

    norm = {
        "code": "ГЭСНм08-02-409-09",
        "name": "Прокладка труб гофрированных",
        "unit": "100 м",
        "resources": [
            {"kind": "labor", "code": "1-100-36", "name": "Средний разряд работы 3.6", "unit": "чел.-ч", "per_unit": 15.2, "price": 500},
            {"kind": "labor", "code": "", "name": "ЗАТРАТЫ ТРУДА РАБОЧИХ, ВСЕГО: В ТОМ ЧИСЛЕ:", "unit": "чел.-ч", "per_unit": 15.6, "price": 500},
            {"kind": "labor", "code": "2-100-02", "name": "Рабочий 2 разряда", "unit": "чел.-ч", "per_unit": 0.02, "price": 500},
            {"kind": "labor", "code": "2-100-03", "name": "Рабочий 3 разряда", "unit": "чел.-ч", "per_unit": 10.75, "price": 500},
            {"kind": "labor", "code": "2-100-04", "name": "Рабочий 4 разряда", "unit": "чел.-ч", "per_unit": 4.83, "price": 500},
        ],
        "_source_kind": "seed_yaml",
    }
    monkeypatch.setattr(gesn_service, "get_norm", lambda *_args, **_kwargs: norm)

    trace = build_position_trace({"code": norm["code"], "qty": 1, "nr_pct": 0, "sp_pct": 0})

    labor_rows = [row for row in trace["rows"] if row.get("type") == "resource_labor"]
    assert [row["columns"]["2"] for row in labor_rows] == ["2-100-02", "2-100-03", "2-100-04"]
    assert trace["summary"]["labor_qty"] == 15.6
    assert trace["summary"]["ozp"] == 7800
    # ADR-16 § 7 on card-carried rates: an explicit resource line with a `price` but
    # no chosen book still travels through the `card_price` / `card_price_without_basis`
    # labelling.  The rule (3) re-baseline keeps that label check alive on data which
    # still triggers it, so the route is not silently retired when the demo seed
    # becomes test-only.  Every one of the three surviving labour rows was priced
    # from the norm card, so the position's amount is preliminary and each row is
    # named by its code both in ``card_priced_resources`` and in the operator flags.
    # Because the card is marked ``_source_kind == "seed_yaml"``, the engine also
    # raises the demo-seed integrity flag («норма взята из демо-семени…») alongside
    # them — a second, orthogonal ADR-16 § 7 check on the same resource line.
    summary = trace["summary"]
    assert {item["code"] for item in summary["card_priced_resources"]} == {
        "2-100-02", "2-100-03", "2-100-04",
    }
    assert summary["amount_status"] == "partial"
    assert summary["full_amount"] is None
    card_flags = [str(flag) for flag in summary["flags"]
                  if str(flag).startswith("цена взята из карточки нормы")]
    assert {flag.rsplit("(", 1)[1].rstrip(")") for flag in card_flags} == {
        "2-100-02", "2-100-03", "2-100-04",
    }
    assert "норма взята из демо-семени (LES_GESN_SEED_ENABLED/path), не из базы" in [
        str(flag) for flag in summary["flags"]
    ]
    assert "normative_source_demo_seed" in summary["completeness_reasons"]
    card_reasons = {
        str(reason).split(":", 1)[1]
        for reason in summary["completeness_reasons"]
        if str(reason).startswith("card_price_without_basis:")
    }
    assert card_reasons == {"2-100-02", "2-100-03", "2-100-04"}


def test_build_lsr_trace_multi_position_sections_and_grand_total():
    # две gold-позиции в двух разделах → общий итог = 2× эталон; итоги разделов и свод — Σ позиций
    positions = [
        {"code": CODE, "qty": 0.61, "section": "Раздел 1. Кровля"},
        {"code": CODE, "qty": 0.61, "section": "Раздел 2. Прочее"},
    ]
    lsr = build_lsr_trace(positions, pricebook=_Book(), name="Кровля и прочее")

    assert lsr["name"] == "Кровля и прочее"
    assert lsr["summary"]["positions"] == 2
    # Grouping into sections introduces no new problem: the only flags are the one
    # resource the chosen book does not price (11.1.03.01-0076), one per position.
    assert lsr["summary"]["flags"]
    assert all("11.1.03.01-0076" in str(flag) for flag in lsr["summary"]["flags"])
    assert lsr["summary"]["total"] == 25358.80  # 2 × 12 679,40 — gold позиции сохранён
    # свод суммирует позиции (прямые/ФОТ/НР/СП) — числа считает код, не LLM
    assert lsr["summary"]["direct"] == round(2 * GOLD_DIRECT, 2)
    assert lsr["summary"]["fot"] == round(2 * GOLD_FOT, 2)
    assert lsr["summary"]["nr"] == round(2 * GOLD_NR, 2)
    assert lsr["summary"]["sp"] == round(2 * GOLD_SP, 2)
    assert lsr["summary"]["labor_qty"] == round(2 * GOLD_LABOR_QTY, 6)

    secs = lsr["sections"]
    assert [s["section"] for s in secs] == ["Раздел 1. Кровля", "Раздел 2. Прочее"]
    assert secs[0]["total"] == GOLD_TOTAL and secs[1]["total"] == GOLD_TOTAL
    # каждый раздел несёт полные трассы позиций (для рендера формы)
    assert len(secs[0]["positions"]) == 1
    assert secs[0]["positions"][0]["summary"]["total"] == GOLD_TOTAL


def test_build_lsr_trace_groups_same_section_contiguously():
    # позиции одного раздела, вперемешку с другим, группируются → 2 раздела, не 3
    positions = [
        {"code": CODE, "qty": 0.61, "section": "A"},
        {"code": CODE, "qty": 0.61, "section": "B"},
        {"code": CODE, "qty": 0.61, "section": "A"},
    ]
    lsr = build_lsr_trace(positions, pricebook=_Book())

    assert [s["section"] for s in lsr["sections"]] == ["A", "B"]
    assert len(lsr["sections"][0]["positions"]) == 2
    assert lsr["sections"][0]["total"] == round(2 * GOLD_TOTAL, 2)
    assert lsr["summary"]["total"] == round(3 * GOLD_TOTAL, 2)


def test_build_lsr_trace_default_section_when_unspecified():
    # без поля section → один раздел «Без раздела», свод = одна позиция
    lsr = build_lsr_trace([{"code": CODE, "qty": 0.61}], pricebook=_Book())
    assert [s["section"] for s in lsr["sections"]] == ["Без раздела"]
    assert lsr["summary"]["total"] == GOLD_TOTAL


def test_positions_from_visible_lsr_rows_converts_physical_unit_to_norm_qty():
    bound = positions_from_visible_lsr_rows([
        {"basis": "ГЭСН 12-01-034-02", "title": "Устройство обрешетки", "quantity": "61", "unit": "м2"}
    ])

    assert bound["row_bindings"][0]["status"] == "bound"
    assert bound["positions"][0]["code"] == CODE
    assert bound["positions"][0]["qty"] == 0.61
    assert bound["row_bindings"][0]["quantity_trace"]["formula"] == "61.0 м2 × 1 / 100 = 0.61"


def test_positions_from_visible_lsr_rows_preserves_model_source_row():
    bound = positions_from_visible_lsr_rows([
        {
            "source_row": 42,
            "basis": "ГЭСН 12-01-034-02",
            "title": "Устройство обрешетки",
            "quantity": "61",
            "unit": "м2",
        }
    ])

    assert bound["positions"][0]["source_row"] == 42
    assert bound["row_bindings"][0]["row"] == 42


def test_lsr_trace_from_visible_rows_builds_priced_trace_from_selected_norms():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": "61",
                "unit": "м2",
                "section": "Кровля",
            }
        ],
        pricebook=_Book(),
        name="Видимые строки",
    )

    assert lsr["summary"]["result_status"] == "priced_partial"
    assert lsr["summary"]["input_rows"] == 1
    assert lsr["summary"]["bound_rows"] == 1
    assert lsr["summary"]["total"] == GOLD_TOTAL
    assert lsr["sections"][0]["section"] == "Кровля"


class _ChosenBook:
    """A ценовая база the operator picked: exact lookup by resource code."""

    def __init__(self, prices: dict[str, float], units: dict[str, str] | None = None) -> None:
        self._prices = dict(prices)
        self._units = dict(units or {})

    def lookup(self, code: str):
        price = self._prices.get(str(code).strip())
        return None if price is None else {"price_current": price, **({"unit": self._units[code]} if code in self._units else {})}


def _cited_book_supply(role):
    code = "SYNTH-SUPPLY-42"
    title = "Изделие модели Ω"
    refs = ["Q1.H1", "Q2.H3"]
    return {
        "source_row": "Поставки Ω!R3",
        "operation_id": "supply-operation",
        "row_role": role,
        "work_row": "Поставки Ω!R2",
        "title": title,
        "resource_code": code,
        "unit": "шт",
        "quantity": 4,
        "evidence_refs": list(refs),
        "evidence_bindings": {ref: f"owned-evidence-{ref}" for ref in refs},
        "_resource_evidence_check": {
            "resource_code": code,
            "cited_resource_codes": [code],
            "code_in_refs": True,
        },
        "supply_resource_match": {
            "source_title": title,
            "resource_code": code,
            "decision": "yes",
            "reason": "Выбранное исполнение соответствует исходной поставке по приведённым карточкам.",
            "source_refs": list(refs),
        },
    }


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("proof_case", [
    "missing",
    "unknown",
    "no",
    "empty_reason",
    "blank_reason",
    "stale_selected_code",
    "stale_title",
    "stale_refset",
    "empty_source_refs",
])
def test_book_priced_supply_without_current_model_match_stays_preliminary(role, proof_case):
    """A cited price keeps known money, but cannot prove the selected product fits."""
    from copy import deepcopy

    row = _cited_book_supply(role)
    if proof_case == "missing":
        row.pop("supply_resource_match")
    else:
        field, value = {
            "unknown": ("decision", "unknown"),
            "no": ("decision", "no"),
            "empty_reason": ("reason", ""),
            "blank_reason": ("reason", "   "),
            "stale_selected_code": ("resource_code", "SYNTH-PREVIOUS-17"),
            "stale_title": ("source_title", "Другое исполнение модели Ω"),
            "stale_refset": ("source_refs", ["Q1.H1", "Q9.H9"]),
            "empty_source_refs": ("source_refs", []),
        }[proof_case]
        row["supply_resource_match"][field] = value
    original = deepcopy(row)

    trace = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
    )
    binding = trace["row_bindings"][0]
    position = trace["sections"][0]["positions"][0]
    summary = position["summary"]
    resource = next(item for item in position["rows"] if item["type"].startswith("resource"))

    assert row == original
    assert binding["source_row"] == original
    assert binding["status"] == "bound"
    assert resource["columns"]["2"] == row["resource_code"]
    assert resource["columns"]["3"] == row["title"]
    assert resource["columns"]["7"] == 4
    assert resource["source"] == "fgis_current"
    assert summary["known_amount"] == 500.0
    assert "supply_resource_match_unsubstantiated" in summary["completeness_reasons"]
    assert not any(str(reason).startswith("price_missing:") for reason in summary["completeness_reasons"])
    assert summary["amount_status"] == "partial"
    assert summary["full_amount"] is None
    assert trace["summary"]["confirmed_amount"] == 0.0
    assert trace["summary"]["preliminary_amount"] == 500.0
    assert trace["summary"]["full_amount"] is None
    assert trace["summary"]["result_status"] == "priced_partial"


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("reverse_refs", [False, True])
def test_book_priced_supply_with_current_model_match_is_complete(role, reverse_refs):
    from copy import deepcopy

    row = _cited_book_supply(role)
    if reverse_refs:
        row["supply_resource_match"]["source_refs"].reverse()
    original = deepcopy(row)
    trace = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
    )
    summary = trace["sections"][0]["positions"][0]["summary"]

    assert row == original
    assert trace["row_bindings"][0]["source_row"] == original
    assert "supply_resource_match_unsubstantiated" not in summary["completeness_reasons"]
    assert summary["amount_status"] == "complete"
    assert summary["full_amount"] == 500.0
    assert trace["summary"]["confirmed_amount"] == 500.0
    assert trace["summary"]["preliminary_amount"] == 0.0
    assert trace["summary"]["result_status"] == "priced_final"


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("overlay", [
    {"quantity": 6},
    {"unit": "шт"},
    {"quantity": 6, "unit": "шт"},
], ids=["quantity-only", "unit-only", "quantity-and-unit"])
def test_quantity_or_unit_confirmation_does_not_confirm_book_supply_selection(role, overlay):
    from copy import deepcopy

    row = _cited_book_supply(role)
    row.pop("supply_resource_match")
    original = deepcopy(row)
    trace = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
        confirmed_data={row["source_row"]: overlay},
    )
    binding = trace["row_bindings"][0]
    position = trace["sections"][0]["positions"][0]
    summary = position["summary"]
    resource = next(item for item in position["rows"] if item["type"].startswith("resource"))
    expected_quantity = overlay.get("quantity", row["quantity"])

    assert row == original
    assert binding["source_row"]["confirmed_by"] == "operator"
    assert "confirmed_price" not in binding["source_row"]
    assert binding["source_row"]["resource_code"] == row["resource_code"]
    assert resource["columns"]["7"] == expected_quantity
    assert resource["source"] == "fgis_current"
    assert summary["known_amount"] == expected_quantity * 125.0
    assert "supply_resource_match_unsubstantiated" in summary["completeness_reasons"]
    assert summary["full_amount"] is None
    assert trace["summary"]["confirmed_amount"] == 0.0
    assert trace["summary"]["preliminary_amount"] == expected_quantity * 125.0
    assert trace["summary"]["result_status"] == "priced_partial"


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("overlay", [
    {"quantity": 6},
    {"unit": "шт"},
    {"quantity": 6, "unit": "шт"},
], ids=["quantity-only", "unit-only", "quantity-and-unit"])
def test_quantity_overlay_cannot_own_a_model_authored_supply_price(role, overlay):
    from copy import deepcopy

    row = _cited_book_supply(role)
    row.pop("supply_resource_match")
    row.update(confirmed_price=80.0, price_source="operator")
    original = deepcopy(row)
    trace = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
        confirmed_data={row["source_row"]: overlay},
    )
    binding = trace["row_bindings"][0]
    summary = trace["sections"][0]["positions"][0]["summary"]

    assert row == original
    assert binding["source_row"]["confirmed_price"] == 80.0
    assert binding["source_row"]["resource_code"] == row["resource_code"]
    assert "supply_resource_match" not in binding["source_row"]
    assert "supply_resource_match_unsubstantiated" in summary["completeness_reasons"]
    assert summary["amount_status"] == "partial"
    assert summary["full_amount"] is None
    assert trace["summary"]["confirmed_amount"] == 0.0
    assert trace["summary"]["preliminary_amount"] > 0.0
    assert trace["summary"]["result_status"] == "priced_partial"


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("unit_price", [80.0, 0.0], ids=["positive-price", "explicit-zero-price"])
def test_actual_operator_unit_price_can_confirm_supply_without_model_match(role, unit_price):
    from copy import deepcopy

    row = _cited_book_supply(role)
    row.pop("supply_resource_match")
    original = deepcopy(row)
    trace = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
        confirmed_data={row["source_row"]: {
            "confirmed_price": unit_price,
            "price_source": "operator",
        }},
    )
    binding = trace["row_bindings"][0]
    summary = trace["sections"][0]["positions"][0]["summary"]
    expected_amount = row["quantity"] * unit_price

    assert row == original
    assert binding["source_row"]["confirmed_by"] == "operator"
    assert binding["source_row"]["confirmed_price"] == unit_price
    assert binding["source_row"]["resource_code"] == row["resource_code"]
    assert "supply_resource_match_unsubstantiated" not in summary["completeness_reasons"]
    assert summary["amount_status"] == "complete"
    assert summary["full_amount"] == expected_amount
    assert trace["summary"]["confirmed_amount"] == expected_amount
    assert trace["summary"]["preliminary_amount"] == 0.0
    assert trace["summary"]["result_status"] == "priced_final"


@pytest.mark.parametrize("role", ["material", "equipment"])
@pytest.mark.parametrize("unit_price", [80.0, 0.0], ids=["positive-price", "explicit-zero-price"])
def test_later_quantity_edit_keeps_an_actual_operator_supply_price(role, unit_price):
    from copy import deepcopy
    from proxy.services.lsr_source_coverage_service import apply_confirmed_data

    row = _cited_book_supply(role)
    row.pop("supply_resource_match")
    trusted_rows = apply_confirmed_data([row], {row["source_row"]: {
        "confirmed_price": unit_price,
        "price_source": "operator",
    }})
    original = deepcopy(trusted_rows)
    trace = build_lsr_trace_from_visible_rows(
        trusted_rows,
        pricebook=_ChosenBook({row["resource_code"]: 125.0},
                              units={row["resource_code"]: "шт"}),
        confirmed_data={row["source_row"]: {"quantity": 6}},
    )
    binding = trace["row_bindings"][0]
    summary = trace["sections"][0]["positions"][0]["summary"]

    assert trusted_rows == original
    assert binding["source_row"]["confirmed_price"] == unit_price
    assert binding["source_row"]["quantity"] == 6
    assert "supply_resource_match_unsubstantiated" not in summary["completeness_reasons"]
    assert summary["amount_status"] == "complete"
    assert summary["full_amount"] == 6 * unit_price
    assert trace["summary"]["confirmed_amount"] == 6 * unit_price
    assert trace["summary"]["preliminary_amount"] == 0.0


@pytest.mark.parametrize("role", ["material", "equipment"])
def test_unsubstantiated_supply_reaches_readiness_retry_without_reopening_trusted_neighbor(role):
    from copy import deepcopy
    from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness
    from proxy.services.lsr_workbook_adapter_service import collect_lsr_workbook_gaps

    supply = _cited_book_supply(role)
    supply.pop("supply_resource_match")
    parent = {
        "source_row": supply["work_row"],
        "operation_id": "pending-parent-work",
        "row_role": "work",
        "title": "Работа по установке изделий",
        "unit": "шт",
        "quantity": 1,
        "norm_code": "",
        "norm_status": "MISSING",
        "open_gap": "Норма работы пока не выбрана.",
    }
    trusted = {
        "source_row": "Поставки Ω!R4",
        "operation_id": "trusted-neighbor",
        "row_role": "material",
        "work_row": parent["source_row"],
        "title": "Поставка с подтверждённой ценой инженера",
        "unit": "шт",
        "quantity": 2,
        "confirmed_price": 80.0,
        "price_source": "operator",
        "confirmed_by": "operator",
    }
    rows = [parent, supply, trusted]
    original = deepcopy(rows)
    basis = {"actual_book": "synthetic_selected_book", "match": "exact",
             "pricebook_sha256": "b" * 64}
    trace = build_lsr_trace_from_visible_rows(
        rows,
        pricebook=_ChosenBook({supply["resource_code"]: 125.0},
                              units={supply["resource_code"]: "шт"}),
    )
    summary = trace["summary"]
    readiness = build_lsr_result_readiness(
        packaged=True,
        rows=rows,
        missing=collect_lsr_workbook_gaps(trace),
        pricing_status=summary["result_status"],
        row_projection=trace["row_projection"],
        calculated_positions=summary["calculated_positions"],
        confirmed_amount=summary["confirmed_amount"],
        preliminary_amount=summary["preliminary_amount"],
        price_basis=basis,
        catalog_work_count=1,
        job_status="complete",
    )
    binding = trace["row_bindings"][1]
    assert "supply_resource_match_unsubstantiated" in binding["readiness"]["open_reasons"]
    assert readiness["estimate_complete"] is False
    assert set(readiness["research_operation_ids"]) == {parent["operation_id"], supply["operation_id"]}
    assert set(readiness["research_locators"]) == {parent["source_row"], supply["source_row"]}
    assert trusted["operation_id"] not in readiness["research_operation_ids"]
    assert trusted["source_row"] not in readiness["research_locators"]
    assert any(action["id"] == "resume" for action in readiness["actions"])
    assert readiness["counts"]["confirmed_amount"] == 160.0
    assert readiness["counts"]["preliminary_amount"] == 500.0
    assert trace["row_bindings"][2]["readiness"]["position_calculated"] is True
    assert binding["source_row"]["resource_code"] == supply["resource_code"]
    assert rows == original


def test_resources_priced_from_the_chosen_book_are_confirmed_final():
    """The priced-final path ADR-16 § 7 requires: every counted rate is owned.

    Each resource resolves through the book the operator chose, so no rate is
    the norm card's own, no card-price flag appears, and the whole amount lands
    in ``confirmed_amount``.  The numbers come from the book, not from a card.
    """
    resources = [
        {"kind": "labor", "code": "1-100-36", "name": "Средний разряд работы 3.6",
         "unit": "чел.-ч", "per_unit": 10.0},
        {"kind": "material", "code": "01.7.15.06-0111", "name": "Гвозди строительные",
         "unit": "т", "per_unit": 0.001},
    ]
    book = _ChosenBook({"1-100-36": 500.0, "01.7.15.06-0111": 90000.0})

    lsr = build_lsr_trace(
        [{"code": CODE, "qty": 1.0, "resources": resources, "nr_pct": 0, "sp_pct": 0}],
        pricebook=book,
        name="База выбрана",
    )
    summary = lsr["summary"]
    position = lsr["sections"][0]["positions"][0]["summary"]

    assert [row["source"] for row in lsr["sections"][0]["positions"][0]["rows"]
            if row["type"].startswith("resource")] == ["fgis_current", "fgis_current"]
    assert position["card_priced_resources"] == []
    assert not [flag for flag in position["flags"]
                if str(flag).startswith("цена взята из карточки нормы")]
    assert position["amount_status"] == "complete"
    assert position["full_amount"] == position["total"]
    assert summary["result_status"] == "priced_final"
    # 10.0 чел.-ч × 500 + 0.001 т × 90000 — the book's numbers, by code
    assert position["total"] == 5090.0
    assert summary["total"] == 5090.0


def test_material_priced_from_the_chosen_book_is_confirmed_money():
    """The confirmed/preliminary split follows price ownership, not the number.

    The same 30 × 100 that the card fallback would leave preliminary counts as
    confirmed when the chosen book supplies the rate.
    """
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "СКС!R9",
                "row_role": "material",
                "work_row": "СКС!R8",
                "title": "Лоток лестничный",
                "resource_code": "01.7.15.06-0111",
                "unit": "м.п",
                "quantity": 30,
                "evidence_refs": ["selected-resource-card"],
                "supply_resource_match": {
                    "source_title": "Лоток лестничный",
                    "resource_code": "01.7.15.06-0111",
                    "decision": "yes",
                    "reason": "Выбранный ресурс соответствует указанному исполнению поставки.",
                    "source_refs": ["selected-resource-card"],
                },
            }
        ],
        pricebook=_ChosenBook({"01.7.15.06-0111": 100.0}, units={"01.7.15.06-0111": "м.п"}),
    )
    binding = lsr["row_bindings"][0]
    assert binding["readiness"]["prices_provided"] is True
    assert binding["readiness"]["price_from_norm_card"] is False
    assert lsr["summary"]["confirmed_amount"] == 3000.0
    assert lsr["summary"]["preliminary_amount"] == 0.0


def test_operator_confirmed_price_makes_the_position_money_confirmed_and_says_so():
    """An operator price is owned money plus its own visible provenance flag."""
    row = {
        "source_row": "PDF!R2",
        "row_role": "work",
        "basis": "ГЭСН 12-01-034-02",
        "title": "Устройство обрешетки",
        "quantity": "61",
        "unit": "м2",
        "norm_status": "EXACT",
    }
    draft = build_lsr_trace_from_visible_rows([row], pricebook=_Book(), name="Видимые строки")
    confirmed = build_lsr_trace_from_visible_rows(
        [row],
        pricebook=_Book(),
        name="Видимые строки",
        confirmed_data={"PDF!R2": {"confirmed_price": 20000.0, "price_source": "operator"}},
    )

    assert draft["summary"]["result_status"] == "priced_partial"
    assert draft["summary"]["confirmed_amount"] == 0.0
    assert draft["summary"]["preliminary_amount"] == GOLD_TOTAL

    binding = confirmed["row_bindings"][0]
    assert binding["readiness"]["prices_provided"] is True
    assert binding["readiness"]["operator_confirmed_price"] is True
    # 0.61 (100 м2) × 20000 — the operator price, not the norm composition
    assert confirmed["summary"]["confirmed_amount"] == 12200.0
    assert confirmed["summary"]["preliminary_amount"] == 0.0
    assert [flag for flag in confirmed["summary"]["flags"]
            if "подтверждённой цене оператора" in str(flag)]


def test_named_needs_research_cipher_is_priced_but_not_final():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "PDF!R18",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": "61",
                "unit": "м2",
                "norm_status": "NEEDS_RESEARCH",
                "quality_message": "аналог, требуется ревизия",
            }
        ],
        pricebook=_Book(),
        name="Видимые строки",
    )

    assert lsr["summary"]["result_status"] == "priced_partial"
    assert lsr["summary"]["norms_status"] == "model_norms_partial"
    assert lsr["summary"]["bound_rows"] == 1
    assert lsr["summary"]["total"] == GOLD_TOTAL
    assert lsr["row_bindings"][0]["status"] == "bound"
    assert lsr["row_bindings"][0]["norm_status"] == "NEEDS_RESEARCH"
    assert lsr["row_bindings"][0]["selection_kind"] == "analogue"


def test_needs_research_without_code_stays_unbound():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "PDF!R21",
                "row_role": "work",
                "title": "Монтаж БАП",
                "quantity": "1",
                "unit": "шт",
                "norm_status": "NEEDS_RESEARCH",
            }
        ]
    )
    assert lsr["summary"]["bound_rows"] == 0
    assert lsr["row_bindings"][0]["status"] == "needs_research"


def test_terminal_exact_status_prices_without_independent_verifier():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": "61",
                "unit": "м2",
                "section": "Кровля",
                "norm_status": "EXACT",
            }
        ],
        pricebook=_Book(),
        name="Видимые строки",
    )

    # A terminal EXACT status settles the norm question only.  The rates behind this
    # total belong to the chosen book, and ADR-16 § 7 keeps a position whose resource
    # list still contains a named price hole out of a final document.
    assert lsr["summary"]["result_status"] == "priced_partial"
    assert lsr["summary"]["norms_status"] == "model_norms_selected"
    assert lsr["summary"]["total"] == GOLD_TOTAL
    assert lsr["row_bindings"][0]["status"] == "bound"
    assert lsr["row_bindings"][0]["quality_gate"] == "not_run"


def test_lsr_trace_from_visible_rows_keeps_unbound_rows_out_of_priced_final():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {"basis": "ГЭСН 12-01-034-02", "title": "Устройство обрешетки", "quantity": "61", "unit": "м2"},
            {"title": "Работа без выбранной нормы", "quantity": "1", "unit": "шт"},
        ]
    )

    assert lsr["summary"]["result_status"] == "priced_partial"
    assert lsr["summary"]["bound_rows"] == 1
    assert lsr["summary"]["unbound_rows"] == 1
    assert lsr["row_bindings"][1]["status"] == "unresolved"


def test_unbound_lsr_placeholder_uses_blank_missing_price_not_zero():
    from proxy.smeta_core.contracts import WorkItem
    from proxy.smeta_core.lsr_renderer import _placeholder

    missing = _placeholder(
        WorkItem(work_id="row-2", title="Работа без нормы", quantity=1, unit="шт"),
        "norm_selection_required",
        "норма не выбрана",
    )

    assert missing["summary"]["total"] is None
    assert missing["rows"][0]["columns"]["12"] is None
    assert missing["rows"][1]["columns"]["12"] is None


def test_lsr_trace_from_visible_rows_does_not_bind_unknown_norm_code():
    lsr = build_lsr_trace_from_visible_rows(
        [{"basis": "ГЭСН 99-99-999-99", "title": "Неизвестная норма", "quantity": "1", "unit": "шт"}]
    )

    assert lsr["summary"]["result_status"] == "norm_selection_required"
    assert lsr["summary"]["bound_rows"] == 0
    assert lsr["summary"]["unbound_rows"] == 1
    assert lsr["row_bindings"][0]["status"] == "norm_not_found"


def test_explicit_norm_cell_does_not_select_code_from_rejected_prose():
    row = {
        "source_row": "Sheet!R34", "row_role": "work",
        "title": "Install tray", "unit": "м2", "quantity": 61,
        "norm_code": "Not suitable: ГЭСН12-01-034-02. Find another norm.",
        "norm_status": "yes",
    }
    trace = build_lsr_trace_from_visible_rows([row])
    assert trace["summary"]["bound_rows"] == 0
    assert trace["row_bindings"][0]["status"] == "norm_selection_required"
    assert trace["row_bindings"][0]["source_row"]["norm_code"] == row["norm_code"]


def test_cited_packaging_converts_to_piece_norm_qty(monkeypatch):
    from proxy.services import gesn_service

    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda *_args, **_kwargs: {
            "code": "ГЭСНм08-03-575-01",
            "name": "Монтаж аппарата",
            "unit": "шт",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    )
    bound = positions_from_visible_lsr_rows([
        {
            "source_row": "СКС!R2",
            "basis": "ГЭСНм:08-03-575-01",
            "title": "Монтаж аппарата",
            "unit": "упак",
            "quantity": 2,
            "conversion_factor": 12,
            "conversion_source": "упаковка 12 шт",
            "row_role": "work",
        }
    ])

    assert bound["row_bindings"][0]["status"] == "bound"
    assert bound["positions"][0]["qty"] == 24
    assert bound["row_bindings"][0]["quantity_trace"]["source_quantity"] == 2.0
    assert bound["row_bindings"][0]["quantity_trace"]["status"] == "cited_conversion_factor"


def test_cited_mass_per_meter_converts_length_to_ton_norm(monkeypatch):
    from proxy.services import gesn_service

    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda *_args, **_kwargs: {
            "code": "ГЭСНм08-02-409-09",
            "name": "Прокладка кабеля",
            "unit": "т",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    )
    without_source = positions_from_visible_lsr_rows([
        {
            "source_row": "СКС!R4",
            "basis": "ГЭСНм08-02-409-09",
            "title": "Кабель",
            "unit": "м",
            "quantity": 100,
            "row_role": "work",
        }
    ])
    assert without_source["row_bindings"][0]["status"] == "needs_input"
    assert without_source["row_bindings"][0]["code"] == "ГЭСНм08-02-409-09"

    cited = positions_from_visible_lsr_rows([
        {
            "source_row": "СКС!R4",
            "basis": "ГЭСНм08-02-409-09",
            "title": "Кабель",
            "unit": "м",
            "quantity": 100,
            "conversion_factor": 0.001,
            "conversion_source": "0,001 т/м по спецификации",
            "row_role": "work",
        }
    ])
    assert cited["row_bindings"][0]["status"] == "bound"
    assert cited["positions"][0]["qty"] == 0.1


def test_confirmed_overlay_fixes_one_row_without_rewriting_others():
    rows = [
        {
            "source_row": "A!R2",
            "basis": "ГЭСН 12-01-034-02",
            "title": "Обрешетка",
            "quantity": "61",
            "unit": "м2",
            "row_role": "work",
        },
        {
            "source_row": "A!R3",
            "basis": "ГЭСН 12-01-034-02",
            "title": "Обрешетка повтор",
            "quantity": "61",
            "unit": "м2",
            "row_role": "work",
        },
    ]
    lsr = build_lsr_trace_from_visible_rows(
        rows,
        confirmed_data={
            "A!R3": {
                "calc_quantity": 100,
                "conversion_source": "оператор: 100 м2",
            }
        },
    )
    by_row = {str(item["row"]): item for item in lsr["row_bindings"]}
    assert by_row["A!R2"]["position"]["qty"] == 0.61
    assert by_row["A!R3"]["position"]["qty"] == 1.0
    assert by_row["A!R2"]["code"] == CODE
    assert by_row["A!R3"]["code"] == CODE


def test_supply_row_preserves_gesn_materials_and_keeps_work_link_with_possible_overlap(monkeypatch):
    from proxy.services import gesn_service

    monkeypatch.setattr(
        gesn_service,
        "get_norm",
        lambda *_args, **_kwargs: {
            "code": "ГЭСНм08-03-575-01",
            "name": "Монтаж",
            "unit": "шт",
            "resources": [
                {"kind": "labor", "name": "Монтажник", "unit": "чел.-ч", "per_unit": 1, "price": 100},
                {"kind": "material", "name": "Аппарат из нормы", "unit": "шт", "per_unit": 1, "price": 50},
            ],
            "nr_pct": 0,
            "sp_pct": 0,
        },
    )
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "СКС!R2",
                "basis": "ГЭСНм08-03-575-01",
                "title": "Монтаж шкафа",
                "unit": "шт",
                "quantity": 1,
                "row_role": "work",
                "nr_pct": 0,
                "sp_pct": 0,
            },
            {
                "source_row": "СКС!R3",
                "title": "Шкаф поставка",
                "unit": "шт",
                "quantity": 1,
                "row_role": "material",
                "work_row": "СКС!R2",
                "supply_kind": "supply",
                "confirmed_price": 1000,
                "price_source": "счёт",
            },
        ]
    )
    work = next(
        position
        for section in lsr["sections"]
        for position in section["positions"]
        if position.get("source_row") == "СКС!R2"
    )
    materials = [row for row in work["rows"] if row.get("type") == "resource_material"]
    assert len(materials) == 1
    assert materials[0]["columns"]["3"] == "Аппарат из нормы"
    assert materials[0]["columns"]["5"] == 1
    assert materials[0]["columns"]["7"] == 1
    assert work["summary"]["full_amount"] is None
    assert work["summary"]["resource_binding_conflicts"][0]["code"] == "separate_supply_possible_overlap"
    assert "resource_bindings" not in work["summary"] and "resource_review" not in work["summary"]
    assert any("Возможное пересечение" in flag for flag in work["summary"]["flags"])
    resource = next(
        item for item in lsr["row_bindings"] if item.get("row") == "СКС!R3"
    )
    assert resource["status"] == "bound"
    assert resource["work_row"] == "СКС!R2"


# --- Stage Б: source quantity/unit are the calculation boundary ----------------


def _tray_source_stamp():
    # Факты исходника для СКС!R47: 30 м.п. лотка.
    return {
        "_catalog_source_values": [
            "1025",
            "Лестничный перфорированный кабельный лоток, длина секции 3000 мм",
            "Ostec",
            "м.п",
            "30",
        ],
        "_catalog_source_numbers": [30.0],
    }


def test_unconfirmed_unit_change_stays_visible_but_uncalculated():
    good = {
        "source_row": "СКС!R2",
        "row_role": "work",
        "title": "Монтаж шкафа",
        "norm_code": CODE,
        "quantity": 61,
        "unit": "м2",
    }
    bad = {
        "source_row": "СКС!R47",
        "row_role": "work",
        "title": "Монтаж лотка",
        "norm_code": CODE,
        "quantity": 30,
        "unit": "шт.",
        **_tray_source_stamp(),
    }
    alone = build_lsr_trace_from_visible_rows([good], pricebook=_Book())
    combined = build_lsr_trace_from_visible_rows([good, bad], pricebook=_Book())
    assert combined["summary"]["total"] == alone["summary"]["total"] > 0
    # only the gated row is missing from the money: one gold position survives
    assert combined["summary"]["total"] == GOLD_TOTAL
    gap = next(b for b in combined["row_bindings"] if b["row"] == "СКС!R47")
    assert gap["status"] == "unconfirmed_quantity"
    payload = gap["quantity_gap"]
    assert payload["field"] == "unit"
    assert payload["authored_value"] == "шт."
    assert "м.п" in payload["source_value"]
    assert "conversion" in payload["needed"]
    # The row stays visible in the source echo and produces no position.
    assert gap["source_row"]["unit"] == "шт."
    assert "СКС!R47" not in {
        str(p.get("source_row"))
        for s in combined["sections"]
        for p in s["positions"]
    }


def test_cited_conversion_basis_allows_the_same_row_to_calculate():
    bad = {
        "source_row": "СКС!R47",
        "row_role": "work",
        "title": "Монтаж лотка",
        "norm_code": CODE,
        "quantity": 30,
        "unit": "шт.",
        **_tray_source_stamp(),
    }
    before = build_lsr_trace_from_visible_rows([bad])
    assert before["summary"]["bound_rows"] == 0
    fixed = dict(bad)
    fixed.update({"calc_quantity": 30, "conversion_source": "длина секции 3 м из исходника"})
    after = build_lsr_trace_from_visible_rows([fixed])
    assert after["summary"]["bound_rows"] == 1
    binding = next(b for b in after["row_bindings"] if b["row"] == "СКС!R47")
    assert binding["status"] == "bound"


def test_unconfirmed_quantity_change_blocks_material_but_matching_passes():
    stamp = {
        "_catalog_source_values": ["990", "Шкаф напольный 42U", "шт", "12"],
        "_catalog_source_numbers": [990.0],
    }
    changed = {
        "source_row": "СКС!R12",
        "row_role": "material",
        "title": "Шкаф",
        "quantity": 12,
        "unit": "шт",
        **stamp,
    }
    trace = build_lsr_trace_from_visible_rows([changed])
    assert trace["summary"]["bound_rows"] == 0
    binding = trace["row_bindings"][0]
    assert binding["status"] == "unconfirmed_quantity"
    assert binding["quantity_gap"]["field"] == "quantity"
    same = dict(changed, quantity=990)
    ok = build_lsr_trace_from_visible_rows([same])
    assert ok["summary"]["bound_rows"] == 1


def test_zero_quantity_is_not_missing_and_no_implicit_one():
    material_no_qty = {"source_row": "СКС!R20", "row_role": "material", "title": "Лоток", "unit": "м.п."}
    trace = build_lsr_trace_from_visible_rows([material_no_qty])
    assert trace["summary"]["bound_rows"] == 0
    assert trace["row_bindings"][0]["status"] == "missing_quantity"
    zero = {"source_row": "СКС!R21", "row_role": "material", "title": "Заглушка", "unit": "шт", "quantity": 0}
    zero_trace = build_lsr_trace_from_visible_rows([zero])
    assert zero_trace["summary"]["bound_rows"] == 1


def test_placeholder_unit_is_not_gated_against_source():
    row = {
        "source_row": "СКС!R4",
        "row_role": "unresolved",
        "title": "Раздел",
        "quantity": None,
        "unit": "-",
        "norm_status": "MISSING",
        "_catalog_source_values": ["982", 'Раздел 6.21 "СКС"'],
        "_catalog_source_numbers": [982.0],
    }
    trace = build_lsr_trace_from_visible_rows([row])
    assert trace["row_bindings"][0]["status"] == "missing_norm"


def test_adapter_stamps_catalog_facts_without_rewriting_model_cells(monkeypatch):
    from proxy.services import attachment_source_catalog_service as catalog_module
    from proxy.services import lsr_workbook_adapter_service as adapter_module

    class _CatRow:
        values = ("м.п", "30")
        numbers = (30.0,)
        excel_row = 47
        locator = "СКС!R47"

    class _Cat:
        def lookup(self, _key):
            return _CatRow()

        def row_origin(self, locator):
            assert locator == "СКС!R47"
            return {}

    monkeypatch.setattr(
        catalog_module, "load_catalog_from_attachment_id", lambda _aid: _Cat()
    )
    decisions = [{"source_row": "СКС!R47", "unit": "шт.", "quantity": 30}]
    snapshot = dict(decisions[0])
    adapter_module._attach_catalog_facts(decisions, "read_x")
    assert decisions[0]["unit"] == snapshot["unit"]
    assert decisions[0]["_catalog_source_values"] == ["м.п", "30"]
    assert decisions[0]["_catalog_source_numbers"] == [30.0]


def test_equivalent_measure_passes_real_catalog_adapter_and_money_trace(tmp_path, monkeypatch):
    from openpyxl import Workbook
    from proxy.services import attachment_source_catalog_service as catalog_module
    from proxy.services import lsr_workbook_adapter_service as adapter_module

    path = tmp_path / "Объёмы с масштабом.xlsx"
    wb = Workbook()
    wb.active.append(["Наименование", "Ед. изм.", "Количество"])
    wb.active.append(["Устройство обрешетки", "м²", 61])
    wb.save(path)
    catalog = catalog_module.load_tabular_source_catalog(path)
    monkeypatch.setattr(catalog_module, "load_catalog_from_attachment_id", lambda _: catalog)
    for qty, unit in [(61, "м2"), (0.61, "100 м2"), (0.061, "1000 м²")]:
        row = {"source_row": "Sheet!R2", "title": "Устройство обрешетки", "row_role": "work",
               "quantity": qty, "unit": unit, "norm_code": CODE, "nr_sp_rule_id": "nrsp-12"}
        assert catalog_module.compare_model_row_to_catalog(row, catalog) == []
        adapter_module._attach_catalog_facts([row], "read_test")
        trace = build_lsr_trace_from_visible_rows([row], pricebook=_Book())
        assert trace["row_bindings"][0]["status"] == "bound"
        assert trace["summary"]["total"] == GOLD_TOTAL
        assert row["quantity"] == qty and row["unit"] == unit

    wrong = {"source_row": "Sheet!R2", "title": "Монтаж обрешетки", "row_role": "work",
             "quantity": 61, "unit": "м", "norm_code": CODE}
    adapter_module._attach_catalog_facts([wrong], "read_test")
    assert build_lsr_trace_from_visible_rows([wrong], pricebook=_Book())["row_bindings"][0]["status"] == "unconfirmed_quantity"


# ---------------------------------------------------------------- stage V split states
def test_stage_v_header_row_is_accounted_without_norm_demand():
    """§6: headings/notes/subtotals are accounted and never demand a mounting norm."""
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "СКС!R1",
                "row_role": "header",
                "title": "Слаботочные системы. Раздел 1",
                "norm_status": "MISSING",
            },
            {
                "source_row": "СКС!R2",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": 61,
                "unit": "м2",
                "norm_status": "EXACT",
            },
        ],
        pricebook=_Book(),
    )
    header = next(item for item in lsr["row_bindings"] if item["row"] == "СКС!R1")
    assert header["status"] == "structure_accounted"
    assert header["readiness"]["header_row"] is True
    assert header["readiness"]["operation_defined"] is True
    assert header["readiness"]["norm_selected"] is None
    assert lsr["summary"]["kind_counts"]["header"] == 1
    # a heading adds no flag; the only flags name the one resource the chosen book
    # does not price (11.1.03.01-0076), one per calculated position
    assert lsr["summary"]["result_status"] == "priced_partial"
    assert not [flag for flag in lsr["summary"]["flags"] if "СКС!R1" in str(flag)]
    assert lsr["summary"]["flags"]
    assert all("11.1.03.01-0076" in str(flag) for flag in lsr["summary"]["flags"])
    from proxy.services.lsr_result_readiness_service import classify_norm_gap

    assert classify_norm_gap(header["source_row"]) == ""


def test_stage_v_bound_material_without_price_is_not_ready():
    """§6: ``bound`` must not read as 'position fully ready / price found'."""
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "СКС!R9",
                "row_role": "material",
                "work_row": "СКС!R8",
                "title": "Лоток лестничный",
                "unit": "м.п",
                "quantity": 30,
            }
        ]
    )
    binding = lsr["row_bindings"][0]
    assert binding["status"] == "bound"
    assert binding["readiness"]["norm_selected"] is None
    # § 4: an unpriced supply line is not a finished calculation.  Its structure is
    # accounted, its known amount is not, so the position answer is «нет» — and the
    # remark that goes into the file says which of the two is missing.
    assert binding["readiness"]["structure_calculated"] is True
    assert binding["readiness"]["known_amount_present"] is False
    assert binding["readiness"]["position_calculated"] is False
    assert binding["readiness"]["prices_provided"] is False
    assert lsr["summary"]["kind_counts"]["material"] == 1
    assert lsr["summary"]["confirmed_amount"] == 0.0
    assert lsr["summary"]["preliminary_amount"] >= 0.0
    assert lsr["summary"]["result_status"] != "priced_final"


def test_stage_v_bound_material_with_confirmed_price_is_confirmed():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "СКС!R9",
                "row_role": "material",
                "work_row": "СКС!R9",
                "title": "Лоток лестничный",
                "unit": "м.п",
                "quantity": 30,
                "confirmed_price": 100,
                "price_source": "operator",
                "confirmed_by": "operator",  # application-owned confirmation overlay
            }
        ]
    )
    binding = lsr["row_bindings"][0]
    assert binding["status"] == "bound"
    assert binding["readiness"]["prices_provided"] is True
    assert lsr["summary"]["confirmed_amount"] == 3000.0
    assert lsr["summary"]["preliminary_amount"] == 0.0


def test_stage_v_needs_research_bound_stays_preliminary_and_keeps_status():
    """§6: analogue/NEEDS_RESEARCH calculation is visible but separated from confirmed."""
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "PDF!R18",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": "61",
                "unit": "м2",
                "norm_status": "NEEDS_RESEARCH",
                "quality_message": "аналог, требуется ревизия",
            }
        ],
        pricebook=_Book(),
    )
    binding = lsr["row_bindings"][0]
    assert binding["status"] == "bound"
    assert binding["norm_status"] == "NEEDS_RESEARCH"
    assert binding["readiness"]["norm_selected"] is True
    assert binding["readiness"]["applicability_substantiated"] is False
    # Directive 2.2: readiness bases are independent. NEEDS_RESEARCH fails
    # norm_applicability_ok, which is what keeps the position preliminary.
    # prices_provided now tracks only the price dimension.
    assert binding["readiness"]["norm_applicability_ok"] is False
    assert lsr["summary"]["confirmed_amount"] == 0.0
    assert lsr["summary"]["preliminary_amount"] == GOLD_TOTAL
    assert lsr["summary"]["result_status"] == "priced_partial"


def test_stage_v_exact_work_without_a_chosen_price_basis_stays_preliminary():
    """Every readiness base of an EXACT work row, and where its money lands.

    The norm and volume questions, and the structure of the calculation, are settled
    by the row itself.  The price question is not: the canonical base card states no
    prices at all, so with no ценовая книга and no confirmation the position carries
    no money whatsoever — every counted resource is a named hole (ADR-16 § 7,
    acceptance § 12 "no false complete status").  The priced counterpart is
    ``test_rim_trace_from_base_norm_card_reproduces_gold_position``, the confirmed
    counterpart is ``test_resources_priced_from_the_chosen_book_are_confirmed_final``.
    """
    lsr = build_lsr_trace_from_visible_rows(
        [
            {
                "source_row": "PDF!R2",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Устройство обрешетки",
                "quantity": "61",
                "unit": "м2",
                "norm_status": "EXACT",
            }
        ]
    )
    binding = lsr["row_bindings"][0]
    readiness = binding["readiness"]
    assert readiness["operation_defined"] is True
    assert readiness["norm_selected"] is True
    assert readiness["applicability_substantiated"] is True
    assert readiness["volume_verified"] is True
    # § 4: the structure question and the money question have separate answers.  The
    # position is built from the card rows, but no rate has an owner, so there is no
    # known amount either — the row is not «полностью рассчитана».
    assert readiness["structure_calculated"] is True
    assert readiness["known_amount_present"] is False
    assert readiness["position_calculated"] is False
    assert readiness["prices_provided"] is False
    assert readiness["price_from_norm_card"] is False
    assert readiness["price_ok"] is False
    # the open questions are the card's seven rows, each named by its шифр
    assert readiness["open_reasons"] == NO_BASIS_REASONS
    assert lsr["summary"]["total"] == 0.0
    assert lsr["summary"]["confirmed_amount"] == 0.0
    assert lsr["summary"]["preliminary_amount"] == 0.0
    assert lsr["summary"]["selected_norms"] == 1
    assert lsr["summary"]["positions_built"] == 1
    assert lsr["summary"]["calculated_positions"] == 0
    assert lsr["summary"]["result_status"] == "priced_partial"
    # the structure survives the missing prices: нормативные quantities are printed
    position = lsr["sections"][0]["positions"][0]["summary"]
    assert position["labor_qty"] == GOLD_LABOR_QTY
    assert position["machinist_qty"] == GOLD_MACHINIST_QTY


def test_stage_v_counts_are_separate_per_kind():
    lsr = build_lsr_trace_from_visible_rows(
        [
            {"source_row": "A!R1", "row_role": "header", "title": "Раздел"},
            {
                "source_row": "A!R2",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Работа",
                "quantity": 61,
                "unit": "м2",
                "norm_status": "EXACT",
            },
            {
                "source_row": "A!R3",
                "row_role": "material",
                "work_row": "A!R2",
                "title": "Материал",
                "unit": "шт",
                "quantity": 2,
            },
            {
                "source_row": "A!R4",
                "row_role": "resource",
                "work_row": "A!R2",
                "title": "Ресурс",
            },
        ]
    )
    counts = lsr["summary"]["kind_counts"]
    assert counts["header"] == 1
    assert counts["work"] == 1
    assert counts["material"] == 1
    assert counts["resource"] == 1
    assert lsr["summary"]["selected_norms"] == 1
    # § 4 / tracker #59: two position rows were built, none of them is fully
    # calculated, and the two numbers are reported separately instead of one
    # standing in for the other.
    assert lsr["summary"]["positions_built"] == 2
    assert lsr["summary"]["calculated_positions"] == 0
    # The projection states one role per locator, so a heading or a resource line is
    # never counted as an unfinished work (§ 4: «заголовки не имеют применимости
    # "нет" как дефекта»).
    assert [item["row_role"] for item in lsr["row_projection"]] == [
        "header", "work", "material", "resource"
    ]
    for role, item in zip(
        ("header", "work", "material", "resource"), lsr["row_projection"], strict=True
    ):
        if role in {"header", "resource"}:
            assert item["norm_selected"] is None
            assert item["position_calculated"] is None
            assert not [
                blocker for blocker in (lsr.get("blockers") or [])
                if blocker["code"] == item["locator"]
            ]


def test_stage_v_coverage_accounts_header_without_gap():
    from proxy.services.lsr_source_coverage_service import build_source_coverage

    coverage = build_source_coverage(
        catalog_locators=["A!R1", "A!R2"],
        rows=[
            {"source_row": "A!R1", "row_role": "заголовок", "title": "Раздел 1"},
            {
                "source_row": "A!R2",
                "row_role": "work",
                "norm_code": "ГЭСН12-01-034-02",
            },
        ],
    )
    assert coverage["complete"] is True
    assert coverage["accounted"] == 2
    assert coverage["gaps"] == []


def test_stage_v_renderer_writes_readiness_column(tmp_path):
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
    from openpyxl import load_workbook

    lsr = build_lsr_trace_from_visible_rows(
        [
            {"source_row": "A!R1", "row_role": "header", "title": "Раздел"},
            {
                "source_row": "A!R2",
                "row_role": "work",
                "basis": "ГЭСН 12-01-034-02",
                "title": "Работа",
                "quantity": 61,
                "unit": "м2",
                "norm_status": "EXACT",
            },
        ],
        pricebook=_Book(),
    )
    out = tmp_path / "lsr.xlsx"
    render_lsr_xlsx(lsr, out)
    review = load_workbook(out)["Проверка"]
    headers = [cell.value for cell in review[1]]
    assert "Готовность строки" in headers
    column = headers.index("Готовность строки") + 1
    rows = list(review.iter_rows(min_row=2))
    text = "\n".join(str(row[column - 1].value) for row in rows if row[column - 1].value)
    # § 4: the sheet prints the seven questions in one wording, and a heading answers
    # «н/п» where it has no question rather than «нет» as if it were a defect.
    # «цены=нет» is the honest wording now: the chosen book prices every card row but
    # one, and no rate is applied from the card, so there is nothing to call «не по базе».
    assert "цены=нет" in text
    assert "шифр выбран=н/п" in text
    assert "применимость обоснована=н/п" in text
    assert "расчёт структуры построен=да" in text
    assert "известная стоимость есть=да" in text
    assert "позиция полностью рассчитана=нет" in text
    # § 4 (tracker #54, #60): the remarks block below the rows carries the open
    # questions of this same revision instead of being an empty section.
    block = "\n".join(
        str(cell.value)
        for row in rows
        for cell in row
        if cell.value not in (None, "")
    )
    assert "Общие замечания и незакрытые позиции" in block
    assert "позиция полностью рассчитана — нет" in block
    assert "price_missing:11.1.03.01-0076" in block
    assert "нужен КАЦ" in block


def test_a_card_resource_identity_reaches_the_operator_without_markup():
    """Installed 0.31.89, case 3 of the installed-trace check.

    The КАЦ question read «нужен КАЦ: Уайт-спирит (14.5.09.<em>11-01</em>02)»,
    because the emphasis the norm card puts around the matched part of a шифр
    travelled with the resource identity into every operator surface.  A code
    with tags in it is a code no ценовая база contains, so the question could
    not be answered.  Only tags are removed here: digits, dots and dashes stay
    exactly as the card has them.
    """
    trace = build_position_trace(
        {
            "code": "ГЭСН00-00-000-00",
            "name": "Окраска поверхности",
            "unit": "м2",
            "qty": 1,
            "nr_pct": 0,
            "sp_pct": 0,
            "resources": [
                {
                    "kind": "material",
                    "code": "14.5.09.<em>11-01</em>02",
                    "name": "Уайт-спирит <em>очищенный</em>",
                    "unit": "<em>л</em>",
                    "per_unit": 1,
                },
            ],
        }
    )

    summary = trace["summary"]
    flags = [str(item) for item in summary.get("flags") or []]
    reasons = [str(item) for item in summary.get("completeness_reasons") or []]
    row = next(item for item in trace["rows"] if item["type"] == "resource_material")

    assert any(item.startswith("нужен КАЦ:") for item in flags)
    assert not any("<" in item or ">" in item for item in flags)
    assert "price_missing:14.5.09.11-0102" in reasons
    assert row["columns"]["2"] == "14.5.09.11-0102"
    assert row["columns"]["3"] == "Уайт-спирит очищенный"
    assert row["columns"]["4"] == "л"



def test_a_norm_card_that_was_never_shown_builds_no_priced_position(monkeypatch, tmp_path):
    """Installed 931 pinned the money rule for an unread card.

    The live БАП run on 930 showed why the verdict had to be total: a row that named a шифр
    with no card body delivered still reached the estimate as money.  Once the row carries the
    verdict, this builder produces no priced position, no known amount and no preliminary
    contribution for it — the amounts stay only where a card was actually shown.
    """
    priced = _active_resource_norm(monkeypatch)
    shown = build_lsr_trace_from_visible_rows([dict(priced)])
    unread = build_lsr_trace_from_visible_rows(
        [dict(priced, _evidence_check={"card_body_shown": False})])

    assert shown["summary"]["positions_built"] == 1
    assert shown["summary"]["preliminary_amount"] > 0
    assert shown["row_projection"][0]["known_amount_present"] is True

    assert unread["summary"]["positions_built"] == 0
    assert unread["summary"]["preliminary_amount"] == 0.0
    assert unread["row_projection"][0]["status"] == "normative_generation_conflict"
    assert unread["row_projection"][0]["known_amount_present"] is not True
