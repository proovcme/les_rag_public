"""Модель видит карточки цен ФССЦ и сама выбирает код изделия.

ЗАЧЕМ. Самостоятельный материал оценивается только по точному `resource_code`
(`pricebook.lookup`). Сметный поиск модели ищет карточки норм, а карточек цен она не видела
вовсе: `PriceBook.browse` — ранжированный поиск «для модели или пользователя» без автопривязки —
в активном пути не вызывался ни разу (17.09). Код ФССЦ модели оставалось угадывать.
При этом в московском справочнике есть и витая пара UTP кат. 5e, и «Покрытие ковровое
(ковролин)»: «цены нет» было бы неправдой, неправдой был бы и подбор чужого товара кодом.

ЧТО ЗДЕСЬ. Запрос с префиксом «цена:» уходит в поиск карточек цен того справочника, по
которому считается смета. Карточки приходят модели как свидетельства; код выбирает модель.
Без настроенного справочника — технический отказ, а не «совпадений нет».
"""
import asyncio
import ast
from decimal import Decimal, ROUND_HALF_UP

import pytest

from proxy.services.model_research_tool_service import (
    ModelResearchToolService,
    model_rag_empty_queries,
    model_rag_failure_detail,
)

@pytest.mark.parametrize('value', [56223.48, 1234567.89, 0.123456789])
def test_price_card_preserves_pricebook_precision(value):
    from proxy.services.model_research_tool_service import price_card_text

    record = dict(code='21.1.04.01-1024', name='Resource', unit='1000 м',
                  price_current_eff=value)
    text = price_card_text(record)
    shown = text.split('Цена текущая: ', 1)[1].split(' руб за ', 1)[0]
    assert Decimal(shown) == Decimal(str(value))
    assert record['price_current_eff'] == value


CARDS = [
    dict(code='21.1.04.01-0003', name='Кабель витая пара UTP 2х2х0,52, категория 5e (внешний)',
         unit='м', price_current_eff=28.4, region='Москва', quarter='3 кв. 2026'),
    dict(code='21.1.04.01-1016', name='Кабель витая пара F/FTP 4х2х0,57, категория 7',
         unit='м', price_current_eff=95.1, region='Москва', quarter='3 кв. 2026'),
]


@pytest.mark.parametrize('price_unit,price_fields,expected', [
    ('1000 м', {'price_base': 28272.17, 'index': 1.44}, '346051.36'),
    ('1000 м', {'price_current': 56223.48}, '477899.58'),
    ('м', {'price_base': 28.27, 'index': 1.44}, '346035.00'),
    ('100000 м', {'price_base': 12345.67, 'index': 1.23456789}, '1295.54'),
])
def test_scaled_price_xlsx_formula_matches_calculation(tmp_path, price_unit, price_fields, expected):
    import openpyxl
    from proxy.services import rim_lsr_trace_service as rim
    from proxy.services.fgis_price_service import PriceBook
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx

    book = PriceBook([dict(code='fixture-cable', name='Synthetic cable', unit=price_unit, **price_fields)])
    trace = rim.build_lsr_trace_from_visible_rows([
        dict(source_row='fixture!1', work_id='supply', title='Synthetic cable', unit='м',
             quantity=8500, norm_code='', decision_kind='материал', resource_code='fixture-cable')
    ], pricebook=book)
    position = trace['sections'][0]['positions'][0]
    target = Decimal(expected)
    assert Decimal(str(position['summary']['mat'])) == target
    path = render_lsr_xlsx(trace, tmp_path / 'scaled.xlsx')
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb['ЛСР РИМ']

    def evaluate(node):
        if isinstance(node, ast.Constant):
            return Decimal(str(node.value))
        if isinstance(node, ast.Name):
            value = ws[node.id].value
            return evaluate(ast.parse(value[1:], mode='eval').body) if isinstance(value, str) and value.startswith('=') else Decimal(str(value))
        if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Mult):
            return evaluate(node.left) * evaluate(node.right)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == 'ROUND':
            return evaluate(node.args[0]).quantize(Decimal(10) ** -int(evaluate(node.args[1])), rounding=ROUND_HALF_UP)
        raise AssertionError(ast.dump(node))

    try:
        row = next(c.row for cells in ws for c in cells if c.column == 2 and c.value == 'fixture-cable')
        assert evaluate(ast.parse(ws.cell(row, 12).value[1:], mode='eval').body) == target
        # Editable quantity must still recalculate from the price, not a frozen receipt total.
        ws.cell(row, 7).value = 4250
        expected_half = (Decimal(str(next(r for r in position['rows'] if r['type'] == 'resource_material')['columns']['10'])) * 4250).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        assert evaluate(ast.parse(ws.cell(row, 12).value[1:], mode='eval').body) == expected_half
    finally:
        wb.close()


def _tool(price_browse=None, norm_calls=None):
    def norm_retrieve(query, limit=6):
        if norm_calls is not None:
            norm_calls.append(query)
        return {'cards': [], 'retrieval_trace': {'rag': {'status': 'ok', 'collection': 'les_smeta_norm_cards'}}}

    async def unused(**_kwargs):
        raise AssertionError('общий RAG не должен вызываться в сметном пути')

    return ModelResearchToolService(retrieve=unused, frozen_dataset_ids=('d',), retrieval_kwargs={},
                                    fallback=lambda *_: {}, smeta_norm_retrieve=norm_retrieve,
                                    price_browse=price_browse)


def _search(tool, query):
    return asyncio.run(tool.execute({'tool': 'search_sources', 'args': {'q': query}}))


def test_price_query_returns_price_cards_for_the_model_to_choose():
    seen = []
    tool = _tool(price_browse=lambda query, limit: (seen.append((query, limit)) or CARDS))
    result = _search(tool, 'цена: кабель витая пара UTP категория 5e')
    assert seen and seen[0][0] == 'кабель витая пара UTP категория 5e', 'префикс снимается'
    assert result.payload['status'] == 'ok' and result.payload['operation'] == 'fgis_price_browse'
    text = result.chunks[0].content
    for part in ('21.1.04.01-0003', 'UTP 2х2х0,52', 'Единица: м', '28.4', 'Москва', '3 кв. 2026'):
        assert part in text, part
    assert result.chunks[0].meta['resource_code'] == '21.1.04.01-0003'
    assert model_rag_failure_detail([result.payload]) is None


def test_price_query_without_matches_is_an_executed_empty_result():
    result = _search(_tool(price_browse=lambda query, limit: []), 'цена: несуществующее изделие')
    assert result.payload['status'] == 'missing'
    assert model_rag_failure_detail([result.payload]) is None
    assert model_rag_empty_queries([result.payload]) == [1]


def test_price_query_without_a_configured_pricebook_is_a_technical_failure():
    result = _search(_tool(price_browse=lambda query, limit: None), 'цена: кабель')
    assert result.payload['status'] == 'blocked'
    detail = model_rag_failure_detail([result.payload])
    assert detail and detail['query_failures'][0]['error_code'] == 'pricebook_not_configured'


def test_ordinary_query_still_searches_norm_cards():
    calls = []
    tool = _tool(price_browse=lambda query, limit: CARDS, norm_calls=calls)
    result = _search(tool, 'прокладка кабеля витая пара в лотке')
    assert calls == ['прокладка кабеля витая пара в лотке']
    assert result.payload['operation'] == 'smeta_norm_native_rrf'


def test_query_instruction_teaches_price_queries_only_for_the_estimator():
    from proxy.services.chat_evidence_application_service import model_rag_query_instruction

    estimator = model_rag_query_instruction({'mode': 'estimator'})
    assert 'цена:' in estimator
    assert 'цена:' not in model_rag_query_instruction({'mode': 'general'})


def test_answer_contract_binds_material_codes_to_shown_price_cards():
    from proxy.services.chat_evidence_application_service import model_answer_request

    text = model_answer_request('Смета', model_driven=True)
    assert 'карточки цены' in text and 'resource_code' in text
    assert 'чужой товар' in text


# --- Единица карточки с множителем (17.09) ------------------------------------------------
# На реальном московском справочнике витая пара стоит «за 1000 м», а в ВОР — 8 500 м.
# Поставка требовала буквального совпадения единиц, и кабель в метрах не мог получить цену
# ни при каком решении модели. «1000 м» и «м» — одна физическая единица с множителем: это
# арифметика, а не физический пересчёт, и её честно делать в коде.

def test_price_unit_scale_is_recognised_only_for_the_same_physical_unit():
    from proxy.services.rim_lsr_trace_service import _price_unit_scale

    assert _price_unit_scale('м', 'м') == 1.0
    assert _price_unit_scale('м', '1000 м') == 1000.0
    assert _price_unit_scale('100 м2', 'м2') == 0.01
    assert _price_unit_scale('шт', '100 шт') == 100.0
    assert _price_unit_scale('м', 'кг') is None
    assert _price_unit_scale('шт', '10 м') is None


def test_standalone_cable_in_metres_is_priced_from_a_card_per_1000_metres(monkeypatch):
    from types import SimpleNamespace
    from proxy.services import chat_evidence_application_service as chat
    from proxy.services import rim_lsr_trace_service as rim
    from proxy.services.fgis_price_service import PriceBook
    from proxy.smeta_core.calculation_receipt import calculation_receipts

    card = dict(code='21.1.04.01-1024', name='Кабель витая пара F/UTP 4х2х0,52, категория 5e',
                unit='1000 м', price_current=56223.48)
    pricebook = PriceBook([card])
    raw = ('| source_row | work_id | title | unit | quantity | norm_code | decision_kind | resource_code | evidence_refs |\n'
           '|---|---|---|---|---|---|---|---|---|\n'
           '| СКС!5.6 | supply | Витая пара кат.5Е 4 пары | м | 200 | — | материал | 21.1.04.01-1024 | Q1.H1 |\n')
    rows = chat.parse_model_rag_result(raw)[1]
    trace = rim.build_lsr_trace_from_visible_rows(rows, pricebook=pricebook)
    supply = next(b for b in trace['row_bindings'] if str(b.get('work_id')) == 'supply')
    assert supply['status'] == 'material_bound', supply
    calculation = calculation_receipts(trace, pricebook=pricebook)[0]['calculation']
    resource = calculation['resources'][0]
    assert abs(resource['known_cost'] - 11244.70) < 0.011, resource
    assert '1000' in str((resource.get('price_card') or {}).get('unit') or '') or '1000' in str(resource)


def test_different_physical_units_remain_a_conflict(monkeypatch):
    from proxy.services import chat_evidence_application_service as chat
    from proxy.services import rim_lsr_trace_service as rim
    from proxy.services.fgis_price_service import PriceBook

    pricebook = PriceBook([dict(code='X-1', name='Кабель', unit='кг', price_current=100)])
    raw = ('| source_row | work_id | title | unit | quantity | norm_code | decision_kind | resource_code | evidence_refs |\n'
           '|---|---|---|---|---|---|---|---|---|\n'
           '| С!1 | supply | Кабель | м | 200 | — | материал | X-1 | |\n')
    trace = rim.build_lsr_trace_from_visible_rows(chat.parse_model_rag_result(raw)[1], pricebook=pricebook)
    assert next(b for b in trace['row_bindings'] if str(b.get('work_id')) == 'supply')['status'] == 'material_price_unit_conflict'


# --- Живая проверка 17.09, проход 1: три причины, почему смета по трём строкам не собралась ---

def test_model_price_intent_is_recognised_in_its_own_wording():
    """Модель написала «Поиск цены на кабель…», а не «цена:», и карточек не получила."""
    seen = []
    tool = _tool(price_browse=lambda query, limit: (seen.append(query) or CARDS))
    for query in ('Поиск цены на кабель витая пара UTP 4 пары категория 5e HF LSZH',
                  'цена на кабель витая пара', 'Цены: кабель витая пара', 'цена: кабель витая пара'):
        assert _search(tool, query).payload['operation'] == 'fgis_price_browse', query
    assert seen[0] == 'кабель витая пара UTP 4 пары категория 5e HF LSZH'
    assert all(item.startswith('кабель витая пара') for item in seen), seen


def test_norm_queries_are_not_mistaken_for_price_queries():
    calls = []
    tool = _tool(price_browse=lambda query, limit: CARDS, norm_calls=calls)
    for query in ('Поиск нормы на демонтаж оповещателя', 'ценообразование монтажных работ'):
        assert _search(tool, query).payload['operation'] == 'smeta_norm_native_rrf', query


def test_base_header_carries_work_id_so_the_model_copies_it():
    """Модель копирует шапку: без work_id две операции одной строки отвергались целиком."""
    from proxy.services.chat_evidence_application_service import model_answer_request

    text = model_answer_request('Смета', model_driven=True)
    # Проход 19: на широкой шапке модель промахивается по ячейкам, поэтому обязательных
    # колонок семь, а section/analogue/coverage/coefficient стали необязательными (18.09).
    assert '| source_row | work_id | title | unit | quantity | norm_code | evidence_refs |' in text
    listing = text.split('К базовой шапке добавь отдельные колонки', 1)[1]
    for name in ('section', 'analogue', 'coverage', 'coefficient'):
        assert name in listing, name
    header = next(line for line in text.splitlines() if line.startswith('| source_row | work_id'))
    separator = text.split('разделитель ', 1)[1].split(',', 1)[0].strip()
    assert header.count('|') == separator.count('|'), (header, separator)


def test_contract_keeps_installation_work_next_to_the_product_supply():
    """Для витой пары модель дала только поставку — монтаж исчез."""
    from proxy.services.chat_evidence_application_service import model_answer_request

    assert 'в дополнение к работе по монтажу' in model_answer_request('Смета', model_driven=True)


# --- Живая проверка 17.09, проход 2: карточки цен не показали модели ковровое покрытие ------
# Запрос «цена: ковровая плитка 500х500 мм Atlas 05 темно серая 4.0 мм» вернул тротуарную
# плитку и противопожарные клапаны 500х500: слово «ковровая» не совпадало с «ковровое», а
# «мм», «серая», «500х500» весили столько же, сколько само изделие. «Покрытие ковровое» в
# московском справочнике есть; модель его не увидела и взяла материал из нормы (тактильная
# ПВХ-плитка). Ранжирование — это порядок свидетельств, а выбор кода остаётся за моделью.

def _book(*names):
    from proxy.services.fgis_price_service import PriceBook

    return PriceBook([dict(code=f'01.1.01.01-{index:04d}', name=name, unit='м2', price_current_eff=100.0)
                      for index, name in enumerate(names, 1)])


def _filler():
    return ([f'Клапан противопожарный стальной квадратный, EI {60 + n}, размеры 500х500 мм' for n in range(20)]
            + [f'Плитка вибропрессованная тротуарная, цвет серая, размеры 240х{n}0 мм' for n in range(10, 20)])


def test_price_browse_matches_word_forms_and_weighs_the_rare_product_word():
    book = _book(*_filler(), 'Покрытие ковровое (ковролин) на джутовой основе')
    names = [row['name'] for row in book.browse('ковровая плитка 500х500 мм Atlas 05 темно серая 4.0 мм', limit=3)]
    assert names[0].startswith('Покрытие ковровое'), names


def test_price_browse_puts_the_named_product_above_shared_attributes():
    book = _book(*[f'Водонагреватель электрический проточный настенный, мощность {n} кВт' for n in range(8)],
                 'Вешалка-крючок металлический настенный',
                 'Оповещатели охранно-пожарные звуковые, IP 42')
    names = [row['name'] for row in book.browse('оповещатель настенный Глагол-Н1-1', limit=2)]
    assert names[0].startswith('Оповещатели'), names


def test_one_family_of_similar_rows_does_not_crowd_out_other_products():
    book = _book(*[f'Плитка стеклянная облицовочная коврово-мозаичная, размеры {n}х{n}х4 мм' for n in range(10, 18)],
                 'Покрытие ковровое из химических волокон безворсовое',
                 *_filler())
    names = [row['name'] for row in book.browse('плитка ковровая', limit=4)]
    assert any(name.startswith('Покрытие ковровое') for name in names), names


def test_price_browse_keeps_exact_code_first_and_hides_haulage_rows():
    from proxy.services.fgis_price_service import PriceBook

    book = PriceBook([dict(code='21.1.04.01-1046', name='Кабель витая пара U/UTP 4х2х0,52, категория 5e', unit='1000 м'),
                      dict(code='21.1.04.01-1044', name='Кабель витая пара U/UTP 2х2х0,52, категория 5e', unit='1000 м'),
                      dict(code='01.01.01.01-01', name='Перевозка кабеля витая пара', unit='т')])
    assert book.browse('21.1.04.01-1046', limit=2)[0]['code'] == '21.1.04.01-1046'
    rows = book.browse('кабель витая пара', limit=5)
    assert [row['code'] for row in rows] == ['21.1.04.01-1044', '21.1.04.01-1046'], rows
    assert rows[0]['match']['schema'] == 'fgis_price_lexical_match_v1'
    assert 'кабель' in rows[0]['match']['matched_tokens']


def test_price_cards_show_the_model_more_candidates_than_norm_cards():
    seen = []
    many = [dict(CARDS[0], code=f'21.1.04.01-{n:04d}') for n in range(12)]
    tool = _tool(price_browse=lambda query, limit: (seen.append(limit) or many))
    result = _search(tool, 'цена: кабель витая пара')
    assert seen == [10] and len(result.chunks) == 10


def test_query_instruction_asks_for_generic_names_without_brand_or_base_prefix():
    from proxy.services.chat_evidence_application_service import model_rag_query_instruction

    text = model_rag_query_instruction({'mode': 'estimator'})
    assert 'родовым наименованием' in text and 'марку' in text
    assert 'регион' in text and 'квартал' in text


# --- Живая проверка 17.09, проход 3: модель подписала запросы номером строки -------------
# «СКС!5.6: цена: кабель витая пара UTP 4 пары категория 5e» ушёл в поиск норм, и карточек
# цен модель снова не получила. Подпись строки — её разметка, а не часть запроса.

def test_price_query_is_recognised_after_the_row_label():
    seen = []
    tool = _tool(price_browse=lambda query, limit: (seen.append(query) or CARDS))
    for query in ('СКС!5.6: цена: кабель витая пара UTP 4 пары категория 5e',
                  'АР!1.31 — цена: плитка напольная ковровая 500х500 мм'):
        assert _search(tool, query).payload['operation'] == 'fgis_price_browse', query
    assert seen == ['кабель витая пара UTP 4 пары категория 5e', 'плитка напольная ковровая 500х500 мм']


def test_norm_query_is_searched_without_the_row_label():
    calls = []
    tool = _tool(price_browse=lambda query, limit: CARDS, norm_calls=calls)
    _search(tool, 'СОУЭ!7.2: Демонтаж оповещателя настенного')
    assert calls == ['Демонтаж оповещателя настенного']


# --- Живая проверка 17.09, проход 4 -------------------------------------------------------
# Карточки цен наконец дошли до модели, но только пять: чатовый путь резал выдачу каждого
# запроса до лимита карточек норм. И среди них не было «Покрытия коврового» — выше стояли
# «Антисептик для ковровых покрытий» и тактильная плитка; модель взяла тактильную плитку
# ради цены. В наименовании ФССЦ изделие названо первыми словами: совпадение там весит больше.

def test_product_named_first_outranks_a_product_merely_for_it():
    book = _book(*_filler(),
                 'Антисептик-антипирен для обработки ковровых покрытий с полиамидным ворсом',
                 'Полоса разметочная нетканая для искусственных ковровых покрытий',
                 'Покрытие ковровое (ковролин) на джутовой основе')
    names = [row['name'] for row in book.browse('ковровая плитка', limit=3)]
    assert names[0].startswith('Покрытие ковровое'), names


def test_price_group_keeps_all_price_cards_when_evidence_is_packed():
    from types import SimpleNamespace
    from proxy.services.chat_evidence_application_service import build_model_rag_evidence_groups

    cards = [SimpleNamespace(content=f'Ресурс ФССЦ: 01.1.01.01-{n:04d}', doc_name='fgis_price_cards.v1',
                             score=1.0, meta={'price_card': True, 'resource_code': f'01.1.01.01-{n:04d}'})
             for n in range(10)]
    norms = [SimpleNamespace(content=f'Шифр: ГЭСН01-01-001-{n:02d}', doc_name='smeta_norm_cards.v1',
                             score=1.0, meta={'norm_code': f'ГЭСН01-01-001-{n:02d}'}) for n in range(10)]
    groups, chunks = build_model_rag_evidence_groups(
        [('цена: плитка', cards), ('норма плитки', norms)], max_chars=100000, hits_per_query=6)
    prices = [c for c in chunks if 'Ресурс ФССЦ' in str(getattr(c, 'content', '') or getattr(c, 'text', ''))]
    assert len(prices) == 10, len(prices)
    assert len(chunks) - len(prices) == 6


def test_price_query_after_a_search_prefix_with_colon():
    """Проход 6: модель написала «поиск: цена: кабель…», и карточек цен снова не было."""
    seen = []
    tool = _tool(price_browse=lambda query, limit: (seen.append(query) or CARDS))
    assert _search(tool, 'поиск: цена: кабель витая пара UTP').payload['operation'] == 'fgis_price_browse'
    assert seen == ['кабель витая пара UTP']
    calls = []
    tool = _tool(price_browse=lambda query, limit: CARDS, norm_calls=calls)
    _search(tool, 'поиск: демонтаж оповещателя настенного')
    assert calls == ['демонтаж оповещателя настенного']


def test_price_query_survives_the_model_own_prefix():
    """Чистая серия, проход 9: цены не дошли из-за приставки самой модели.

    Она пишет «смета: цена ковровая плитка 500х500мм 4.0мм темно серая». Слово «цена»
    стояло не в начале, запрос уходил в поиск норм, карточек цен модель не видела и
    обе поставки остались material_price_missing. Запрос нормы ценовым не становится.
    """
    from proxy.services.model_research_tool_service import price_query_text

    assert price_query_text('смета: цена ковровая плитка 500х500мм 4.0мм темно серая') == (
        'ковровая плитка 500х500мм 4.0мм темно серая')
    assert price_query_text('СКС!5.6: смета: цена кабель витой пары') == 'кабель витой пары'
    assert price_query_text('смета: норма на устройство покрытий из ковровой плитки') is None
    assert price_query_text('устройство покрытий') is None
