"""Коэффициент, предложенный моделью и не принятый расчётом, виден в книге.

Повышающий коэффициент берётся из МДС или приложения к сборнику и требует основания:
применить предложенное число без него значит поменять деньги догадкой. Но и молчать
нельзя — строка уходит в смету по норме, а проверяющий не узнаёт, что её считали нужным
поднять. До 11.09 отметка жила только в разборе и до самой книги не доходила
(ревью Codex 11.09, шаг 3).
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx


COEFFICIENT_COLUMN = 6


def visible_row(**fields: object) -> dict[str, object]:
    row = {
        "source_row": "995",
        "norm_code": "ГЭСНм08-02-398-01",
        "наименование": "Провод в лотках",
        "unit": "100 м",
        "quantity": "3.66",
        "раздел": "СКС",
    }
    row.update(fields)
    return row


def work_rows(trace: dict) -> list[dict]:
    return [
        row
        for section in trace.get("sections") or []
        for position in section.get("positions") or []
        for row in position.get("rows") or []
        if row.get("type") == "work"
    ]


def comments(path: Path) -> list[tuple[int, int, str]]:
    """Coefficient notes only; source/quantity provenance has its own columns."""
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    found = [
        (cell.row, cell.column, cell.comment.text)
        for row in sheet.iter_rows()
        for cell in row
        if cell.comment is not None and cell.column == COEFFICIENT_COLUMN
    ]
    workbook.close()
    return found


def test_the_proposed_coefficient_reaches_the_trace_row() -> None:
    trace = build_lsr_trace_from_visible_rows(
        [visible_row(coefficient="1.15")], pricebook=None, name="Проверка"
    )
    binding = (trace.get("row_bindings") or [{}])[0]
    assert binding.get("coefficient_status") == "needs_basis"
    assert binding.get("proposed_coefficient") == pytest.approx(1.15)

    rows = work_rows(trace)
    assert rows, "позиция не собралась — проверять нечего"
    assert (rows[0].get("meta") or {}).get("coefficient_status") == "needs_basis"


def test_the_workbook_carries_a_note_on_the_coefficient_cell(tmp_path: Path) -> None:
    trace = build_lsr_trace_from_visible_rows(
        [visible_row(coefficient="1.15")], pricebook=None, name="Проверка"
    )
    out = tmp_path / "lsr.xlsx"
    render_lsr_xlsx(trace, out, title="Проверка", meta={})

    found = comments(out)
    assert len(found) == 1
    _row_no, column, text = found[0]
    assert column == COEFFICIENT_COLUMN
    assert "1.15" in text
    assert "не принят без основания" in text


def test_the_calculated_cell_keeps_the_accepted_coefficient(tmp_path: Path) -> None:
    """Значение не подменяется: формулы считали бы по необоснованному числу."""
    trace = build_lsr_trace_from_visible_rows(
        [visible_row(coefficient="1.15")], pricebook=None, name="Проверка"
    )
    out = tmp_path / "lsr.xlsx"
    render_lsr_xlsx(trace, out, title="Проверка", meta={})

    found = comments(out)
    sheet = openpyxl.load_workbook(out).active
    assert sheet.cell(row=found[0][0], column=COEFFICIENT_COLUMN).value == 1


@pytest.mark.parametrize("coefficient", ["", "1", "1.0", "1,0"])
def test_no_note_when_nothing_was_proposed(tmp_path: Path, coefficient: str) -> None:
    """Коэффициент, равный единице, ничего не меняет и пометкой не является."""
    trace = build_lsr_trace_from_visible_rows(
        [visible_row(coefficient=coefficient)], pricebook=None, name="Проверка"
    )
    out = tmp_path / "lsr.xlsx"
    render_lsr_xlsx(trace, out, title="Проверка", meta={})
    assert comments(out) == []
