"""Real source cells retain coordinates, unknowns and the complete denominator."""
import asyncio
import copy
import csv

import fitz
import openpyxl
import pytest

from backend.pdf_layout import extract_layout_markdown
from proxy.services import lsr_workbook_adapter_service as adapter
from proxy.services.portion_plan_service import portion_plan
from proxy.services.tabular_attachment_service import format_tabular_attachment_context
from proxy.smeta_core.vor_rows import row_anchors, rows_from_attachment_text, source_revision
from tests.test_smeta_resource_receipt import norm, work

HEADER=['№','Наименование','Артикул','Производитель','Ед. изм.','Количество','Единица массы','Масса','Стоимость']
SOURCE=[
    [1,'Монтаж устройства (всего 2 комплекта)',None,None,'шт',2,None,None,None],
    [2,'Монтаж с неизвестным количеством',None,None,'шт',None,'кг',7,None],
    [3,'Монтаж с настоящим нулём',None,None,'шт',0,None,None,None],
    [4,'Нестандартная единица',None,None,'нестандартная длинная единица','не задано','кг',8,None],
    [5,'Нумерованная неоднозначная строка',None,None,None,None,'кг',7,None],
    [6,'Количество без единицы',None,None,None,2,'кг',7,None],
    [7,'Исходная запись количества',None,None,'шт','2,50',None,None,None],
    [8,'Монтаж | контроль',r'C:\parts\a|b',None,'шт',2,None,None,None],
]


def source_file(path):
    rows=[HEADER,['Раздел 1. Синтетика']+[None]*8,*SOURCE,['Итого']+[None]*7+[999]]
    if path.suffix=='.csv':
        with path.open('w',encoding='utf-8-sig',newline='') as stream:
            csv.writer(stream,delimiter=';').writerows(rows)
    else:
        book=openpyxl.Workbook()
        sheet=book.active
        sheet.title='S'
        for row in rows:
            sheet.append(row)
        sheet.merge_cells('A2:I2')
        book.save(path)
        book.close()
    return path


@pytest.mark.parametrize('suffix',['.xlsx','.csv'])
def test_actual_source_cells_keep_empty_zero_unknown_and_literal_delimiter(tmp_path,suffix):
    source=source_file(tmp_path/('source'+suffix))
    context,truncated=format_tabular_attachment_context(source,source.name,max_chars=None,max_cell_chars=None)
    assert not truncated
    rows=rows_from_attachment_text(context)
    assert len(rows)==len(SOURCE)
    assert [r['source_row'] for r in rows]==[str(n) for n in range(1,9)]
    assert [r['quantity'] for r in rows]==[2,None,0,None,None,2,2.5,2]
    assert [r['unit'] for r in rows]==['шт','шт','шт','нестандартная длинная единица','','','шт','шт']
    assert [r['quantity_text'] for r in rows]==['2','','0','не задано','','2','2,50','2']
    assert rows[1]['cells'][6:8]==['кг','7']
    assert rows[7]['title']=='Монтаж | контроль' and rows[7]['cells'][2]==r'C:\parts\a|b'
    assert all(len(r['cells'])==len(HEADER) for r in rows)
    assert all(r['section']=='Раздел 1. Синтетика' for r in rows)
    if suffix=='.xlsx':
        direct=row_anchors(source)
        assert [r['cells'] for r in rows]==[r['cells'] for r in direct]
        assert [r['quantity_text'] for r in rows]==[r['quantity_text'] for r in direct]
    plan=portion_plan(rows=rows,limit=2,session_id='synthetic',source_revision=source_revision(rows))
    assert plan['total']==8 and len(plan['rows'])==2 and plan['remaining']==6 and not plan['complete']
    assert plan['rows'][1]['quantity'] is None


def test_full_width_empty_quantity_cannot_be_mass(tmp_path):
    # This bypasses the formatter so its old row squeezing cannot mask the parser bug.
    context='S!R1: '+ ' | '.join(HEADER)+'\nS!R2: 1 | Original work | | | шт | | кг | 7 |'
    row,=rows_from_attachment_text(context)
    assert row['unit']=='шт' and row['quantity'] is None and row['quantity_text']==''


def test_legacy_short_compressed_row_retains_compatible_quantity():
    context=('S!R1: № | Наименование | Артикул | Производитель | Ед. изм. | Количество\n'
        'S!R2: 1 | Legacy work | Producer | шт | 4')
    row,=rows_from_attachment_text(context)
    assert row['unit']=='шт' and row['quantity']==4


def test_partial_authored_reply_keeps_every_source_in_actual_workbook(norm,monkeypatch,tmp_path):
    source=source_file(tmp_path/'source.xlsx')
    decisions=[work(source_row='S!1')]
    original=copy.deepcopy(decisions)
    monkeypatch.setattr(adapter.fgis_price_service,'resolve_pricebook_path',lambda *_:None)
    target=tmp_path/'partial.xlsx'
    result=asyncio.run(adapter.build_lsr_workbook_from_decisions(source,dict(decisions=decisions),target,lambda *_:None))
    assert result['source_rows']==8 and result['bound_rows']==1
    assert result['pricing_status']=='priced_partial'
    assert decisions==original
    for number in range(2,9):
        assert f'source:S!{number}:not_processed' in result['missing']
    book=openpyxl.load_workbook(target,data_only=False)
    records=list(book['Исходные строки'].values)[1:]
    assert [r[0] for r in records]==[f'S!{n}' for n in range(1,9)]
    assert records[0][3]=='Монтаж устройства (всего 2 комплекта)' and records[0][6]=='bound'
    assert records[1][4:7]==('шт',None,'not_processed')
    assert records[2][4:7]==('шт',0,'not_processed')
    assert records[4][4:7]==(None,None,'not_processed')
    assert records[7][3]=='Монтаж | контроль'
    assert all('Итого' not in str(r[3]) for r in records)
    book.close()


def test_actual_sparse_pdf_preserves_quantity_and_mass_columns(tmp_path):
    path=tmp_path/'sparse.pdf'
    headings=['Code','Name','Article','Maker','Unit','Qty','Note','Mass']
    rows=[headings,*[['',f'Device-{n}','','','',str(n),'',str(n*7)] for n in range(1,6)]]
    document=fitz.open()
    page=document.new_page(width=620,height=300)
    x,y,width,height=35,40,68,30
    for ri,row in enumerate(rows):
        for ci,value in enumerate(row):
            if value:
                page.insert_text((x+ci*width+3,y+ri*height+18),value,fontsize=8)
    for ci in range(9):
        page.draw_line((x+ci*width,y),(x+ci*width,y+len(rows)*height))
    for ri in range(len(rows)+1):
        page.draw_line((x,y+ri*height),(x+8*width,y+ri*height))
    document.save(path)
    document.close()
    markdown=extract_layout_markdown(path)
    lines=[line for line in markdown.splitlines() if line.startswith('|') and '---' not in line]
    cells=[[cell.strip() for cell in line.strip('|').split('|')] for line in lines]
    assert cells[0]==headings
    assert cells[1:]==rows[1:]


@pytest.mark.parametrize('quantity',[None,2])
def test_bound_numeric_nonstandard_unit_stays_literal(tmp_path,quantity):
    path=tmp_path/'numeric-unit.xlsx'
    book=openpyxl.Workbook()
    book.active.title='S'
    book.active.append(HEADER)
    book.active.append([1,'Numeric literal unit',None,None,1000,quantity,'кг',7,None])
    book.save(path)
    book.close()
    context,_=format_tabular_attachment_context(path,path.name,max_chars=None,max_cell_chars=None)
    row,=rows_from_attachment_text(context)
    direct,=row_anchors(path)
    assert row['unit']==direct['unit']=='1000'
    assert row['quantity']==quantity


@pytest.mark.parametrize('second_format',['v2','legacy'])
def test_new_file_resets_same_sheet_header_and_section(tmp_path,second_format):
    first=source_file(tmp_path/'first.xlsx')
    first_text,_=format_tabular_attachment_context(first,first.name,max_chars=None,max_cell_chars=None)
    second=tmp_path/'second.xlsx'
    book=openpyxl.Workbook()
    book.active.title='S'
    book.active.append(['№','Ед. изм.','Количество','Наименование'])
    book.active.append([9,'м',3,'Second file work'])
    book.save(second)
    book.close()
    second_text,_=format_tabular_attachment_context(second,second.name,max_chars=None,max_cell_chars=None)
    if second_format=='legacy':
        second_text=second_text.replace('Формат ячеек: les.tabular.cells.v2\n','')
    rows=rows_from_attachment_text(first_text+'\n'+second_text)
    last=rows[-1]
    assert (last['source_row'],last['title'],last['unit'],last['quantity'],last['section'])==('9','Second file work','м',3,'')


def test_format_only_far_tail_does_not_expand_source_or_change_revision(tmp_path):
    from openpyxl.styles import PatternFill
    plain=tmp_path/'plain'/'same.xlsx'
    styled=tmp_path/'styled'/'same.xlsx'
    plain.parent.mkdir()
    styled.parent.mkdir()
    book=openpyxl.Workbook()
    book.active.title='S'
    book.active.append(['№','Наименование','Ед. изм.','Количество'])
    book.active.append([1,'Same source','шт',2])
    book.save(plain)
    book.active['XFD2'].fill=PatternFill(fill_type='solid',fgColor='FFFF00')
    book.save(styled)
    book.close()
    first,_=format_tabular_attachment_context(plain,plain.name,max_chars=None,max_cell_chars=None)
    second,_=format_tabular_attachment_context(styled,styled.name,max_chars=None,max_cell_chars=None)
    assert first==second
    rows_first=rows_from_attachment_text(first)
    rows_second=rows_from_attachment_text(second)
    assert rows_first==rows_second and len(rows_second[0]['cells'])==4
    assert source_revision(rows_first)==source_revision(rows_second)
    assert row_anchors(plain)==row_anchors(styled)


def test_real_far_column_data_keeps_its_original_ordinal(tmp_path):
    path=tmp_path/'far-data.xlsx'
    book=openpyxl.Workbook()
    book.active.title='S'
    book.active.append(['№','Наименование','Ед. изм.','Количество'])
    book.active.append([1,'Far actual data','шт',None])
    book.active['XFD2']='Actual far field | 0'
    book.save(path)
    book.close()
    context,_=format_tabular_attachment_context(path,path.name,max_chars=None,max_cell_chars=None)
    row,=rows_from_attachment_text(context)
    direct,=row_anchors(path)
    assert len(row['cells'])==16384
    assert row['cells'][16383]=='Actual far field | 0'
    assert row['cells']==direct['cells']
    assert row['quantity'] is None and row['quantity_text']=='' and row['unit']=='шт'


@pytest.mark.parametrize('suffix',['.xlsx','.csv'])
def test_trailing_unknown_cells_are_padded_to_header_without_mass_fallback(tmp_path,suffix):
    path=tmp_path/('tail'+suffix)
    rows=[['№','Наименование','Ед. изм.','Количество'],[1,'Unknown trailing quantity','шт']]
    if suffix=='.csv':
        with path.open('w',encoding='utf-8-sig',newline='') as stream:
            csv.writer(stream,delimiter=';').writerows(rows)
    else:
        book=openpyxl.Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
    context,_=format_tabular_attachment_context(path,path.name,max_chars=None,max_cell_chars=None)
    row,=rows_from_attachment_text(context)
    assert row['cells']==['1','Unknown trailing quantity','шт','']
    assert row['quantity'] is None and row['unit']=='шт'
