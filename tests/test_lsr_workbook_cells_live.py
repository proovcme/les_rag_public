"""Решения модели доходят до ячеек скачиваемого XLSX.

ЗАЧЕМ. Проекция для отображения и подменённый исполнитель этого не доказывают: между
решением и файлом стоит расчёт трассы и рендер книги. Ревью Codex 10.09 (пункт 2)
требовало сверки настоящих ячеек — она здесь.

ЖИВАЯ БАЗА. Тест читает активную базу норм и книгу цен, поэтому пропускается там, где
их нет: `data/` под gitignore. Остальные тесты сметного среза герметичны.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest


def _active_base() -> Path | None:
    try:
        from proxy.smeta_core.norm_browser import _base_path

        base = Path(_base_path())
    except Exception:  # noqa: BLE001 — база не настроена, тест не про это
        return None
    return base if base.is_file() else None


BASE = _active_base()
pytestmark = pytest.mark.skipif(
    BASE is None, reason="активная база норм недоступна (data/ в gitignore)"
)


def _real_decisions(limit: int = 2) -> list[dict]:
    import sqlite3

    conn = sqlite3.connect(f"file:{BASE}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute(
            "SELECT display_code, norm_name, norm_unit FROM norms"
            " WHERE display_code LIKE 'ГЭСНм10-01-001-%' ORDER BY display_code LIMIT ?",
            (limit,),
        ).fetchall()
    finally:
        conn.close()
    return [
        {
            "source_row": str(984 + index),
            "title": row["norm_name"],
            "unit": row["norm_unit"],
            "quantity": 12.5 + index,
            "norm_code": row["display_code"],
        }
        for index, row in enumerate(rows)
    ]


@pytest.fixture(scope="module")
def workbook(tmp_path_factory):
    from proxy.services.lsr_workbook_adapter_service import (
        build_lsr_workbook_from_decisions,
    )

    decisions = _real_decisions()
    assert decisions, "в базе не нашлось норм для проверки"
    output = tmp_path_factory.mktemp("lsr") / "lsr.xlsx"
    result = asyncio.run(
        build_lsr_workbook_from_decisions(
            Path("."),
            {"decisions": decisions, "question": "Сверка ячеек"},
            output,
            lambda *_args: None,
        )
    )
    return decisions, output, result


def test_every_decision_becomes_a_position_row(workbook):
    """Шифр, наименование, единица и количество стоят в ячейках, а не только в трассе."""
    import openpyxl

    decisions, output, result = workbook
    assert result["source_rows"] == len(decisions)
    assert not result["blockers"], result["blockers"]

    sheet = openpyxl.load_workbook(output, data_only=True).worksheets[0]
    table = [
        [("" if cell is None else str(cell)) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    for decision in decisions:
        line = next(
            (row for row in table if len(row) > 1 and row[1] == decision["norm_code"]),
            None,
        )
        assert line is not None, f"шифр {decision['norm_code']} не попал в книгу"
        assert decision["title"][:30] in line[2]
        assert line[3] == decision["unit"]
        assert float(line[4]) == pytest.approx(decision["quantity"])


def test_positions_keep_the_order_of_the_decisions(workbook):
    import openpyxl

    decisions, output, _result = workbook
    sheet = openpyxl.load_workbook(output, data_only=True).worksheets[0]
    codes = [
        str(row[1])
        for row in sheet.iter_rows(values_only=True)
        if row and len(row) > 1 and str(row[1] or "").startswith("ГЭСНм")
    ]
    assert codes == [decision["norm_code"] for decision in decisions]


def test_the_source_row_of_the_vor_survives_in_the_workbook(workbook):
    """Связь с исходной ведомостью не должна теряться: книга нумерует позиции заново."""
    import openpyxl

    decisions, output, _result = workbook
    book = openpyxl.load_workbook(output, data_only=True)
    text = "\n".join(
        str(cell)
        for sheet in book.worksheets
        for row in sheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )
    for decision in decisions:
        assert decision["source_row"] in text, (
            "номер строки ВОР пропал из книги — по готовому файлу нельзя будет "
            "проверить, какая позиция ведомости чем закрыта"
        )
