"""Explicit VOR core rows survive text formats without weakening LSR transport."""
from __future__ import annotations

import json

import openpyxl
import pytest

from proxy.services import chat_evidence_application_service as service
from proxy.services.model_lsr_result_contract import decode_model_artifact_kind


CORE_ROW = {'source_row': 'Sheet!R2', 'row_role': 'work', 'work_row': '',
            'title': 'Literal authored installation Ω', 'unit': 'м', 'quantity': 12.5, 'open_gap': ''}


def render(row, fmt, declaration='artifact_kind: vor\n'):
    if fmt == 'markdown':
        header = '| ' + ' | '.join(row) + ' |\n'
        divider = '| ' + ' | '.join('---' for _ in row) + ' |\n'
        values = '| ' + ' | '.join(str(v).replace('|', r'\|') for v in row.values()) + ' |'
        return declaration + header + divider + values
    return declaration + '\n'.join(f'{key}: {value}' for key, value in row.items())


@pytest.mark.parametrize('fmt', ['markdown', 'block_text'])
def test_explicit_vor_seven_core_fields_are_literal_rows_without_norm_fields(fmt):
    raw = render(CORE_ROW, fmt)
    diagnostics = []
    parsed = service.parse_model_rag_result(raw, diagnostics=diagnostics)
    assert decode_model_artifact_kind(raw)['artifact_kind'] == 'vor'
    assert parsed == (raw, [CORE_ROW]) and diagnostics == []
    assert all(key not in parsed[1][0] for key in ('norm_code', 'norm_status', 'evidence_refs'))


@pytest.mark.parametrize('fmt', ['markdown', 'block_text'])
@pytest.mark.parametrize('declaration', [
    '', 'artifact_kind: lsr\n', 'artifact_kind: unknown\n',
    'artifact_kind: vor\nartifact_kind: lsr\n',
    'artifact_kind: vor\nartifact_kind: vor\n',
    'artifact_kind: null\n',
])
def test_undeclared_lsr_invalid_or_conflicting_target_keeps_old_norm_code_guard(fmt, declaration):
    assert service.parse_model_rag_result(render(CORE_ROW, fmt, declaration)) is None


@pytest.mark.parametrize('fmt', ['markdown', 'block_text'])
def test_explicit_empty_core_values_do_not_become_zero_roles_or_norms(fmt):
    row = {**CORE_ROW, 'row_role': 'header', 'unit': '', 'quantity': '', 'open_gap': ''}
    raw = render(row, fmt)
    parsed = service.parse_model_rag_result(raw)
    assert parsed == (raw, [row])
    assert parsed[1][0]['quantity'] == '' and not parsed[1][0].get('norm_code')


def test_duplicate_vor_table_header_is_not_selected_or_merged():
    fields = list(CORE_ROW) + ['quantity']
    raw = ('artifact_kind: vor\n| ' + ' | '.join(fields) + ' |\n| '
           + ' | '.join('---' for _ in fields) + ' |\n| '
           + ' | '.join([*(str(value) for value in CORE_ROW.values()), '99']) + ' |')
    assert service.parse_model_rag_result(raw) is None


def test_wrong_vor_value_width_preserves_diagnostic_without_inventing_cells():
    raw = render(CORE_ROW, 'markdown').replace('| work |', '| work | extra |')
    diagnostics = []
    parsed = service.parse_model_rag_result(raw, diagnostics=diagnostics)
    assert parsed is None or parsed[1] == []
    assert diagnostics[0]['code'] == 'MODEL_RESULT_TABLE_ROW_WIDTH_MISMATCH'
    assert diagnostics[0]['decision']['source_row'] == 'Sheet!R2'


@pytest.mark.asyncio
@pytest.mark.parametrize('fmt', ['markdown', 'block_text'])
@pytest.mark.parametrize('quantity', [12.5, 191])
async def test_seven_core_vor_row_reaches_physical_workbook_and_audit_without_norm_fields(tmp_path, fmt, quantity):
    from proxy.services import workbook_tool_service as workbook
    from tests.test_workbook_tool_service import _source_workbook

    row = {**CORE_ROW, 'quantity': quantity}
    raw = render(row, fmt)
    parsed = service.parse_model_rag_result(raw)
    assert parsed is not None and parsed[1] == [row]
    source = _source_workbook(tmp_path / 'Исходник с пробелом №1.xlsx')
    target = tmp_path / 'ВОР без норм.xlsx'
    result = await workbook._default_vor_adapter(source, {'decisions': parsed[1]}, target, lambda *_: None)
    assert result['source_rows'] == 1
    wb = openpyxl.load_workbook(target, data_only=True)
    try:
        cells = next(wb['ВОР'].iter_rows(min_row=4, max_row=4, values_only=True))
        assert wb.active.title == 'Ведомость работ' and wb.active.max_column == 5
        printed = [values for values in wb.active.iter_rows(min_row=9, values_only=True)
                   if isinstance(values[0], int) and not isinstance(values[0], bool)]
        assert [(values[1], values[2], values[3]) for values in printed] == [(row['title'], row['unit'], quantity)]
        assert json.loads(wb['Решения модели']['C4'].value.removeprefix('JSON: ')) == row
        assert cells[2] == row['title'] and cells[4] == row['unit'] and cells[5] == quantity
        assert cells[9] == row['row_role'] and 'Sheet!R2' in cells[6]
        if quantity == 12.5:
            assert result['requires_decision_count'] == 0 and not result['vor_source_errors']
        else:
            assert result['requires_decision_count'] == 1
            assert result['vor_source_errors'] == ['source_quantity_mismatch:Sheet!R2']
            assert 'исходной строкой' in cells[11]
    finally:
        wb.close()
    # The audit marks the mismatch, never rewrites the model's proposed volume
    # or manufactures norms, flags, references, or installation decisions.
    assert parsed[1] == [row]
    assert all(key not in row for key in ('norm_code', 'norm_status', 'evidence_refs'))
