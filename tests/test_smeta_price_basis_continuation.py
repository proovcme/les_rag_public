"""Document prices are coherent facts, not whichever global book exists now."""
import asyncio
from copy import deepcopy
import hashlib
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from proxy.services import fgis_price_service as prices
from proxy.services import rim_lsr_trace_service as rim
from proxy.services.lsr_workbook_adapter_service import build_lsr_workbook_from_decisions
from tests.test_smeta_resource_identity_contract import setup_rows, isolated


def price_basis(tmp_path, name='A', price=20, unit='шт', *, duplicates=False):
    rows=[dict(code='R1', name='Current synthetic material',unit=unit,price_current=price,
               region=name,quarter='2026Q3',version='synthetic-v1')]
    if duplicates:
        rows.append(dict(rows[0],price_current=price+5))
    path=tmp_path/(name+'.parquet')
    pd.DataFrame(rows).to_parquet(path,index=False)
    identity=dict(schema='les.pricebook-revision.v1',sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        pricebook_id=name,region=name,period='2026Q3',version='synthetic-v1')
    return dict(schema='les.document-price-basis.v1',status='pinned',path=str(path.resolve()),identity=identity)


def package(tmp_path, decisions, basis, filename='book.xlsx', *, review=False):
    return asyncio.run(build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        dict(decisions=decisions,_document_price_basis=basis,review_calculation=review),tmp_path/filename,lambda *_:None))


def test_primary_uses_current_book_not_historical_sql_price(tmp_path, monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis=price_basis(tmp_path)
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:basis['path'])
    result=package(tmp_path,[parent],basis)
    resource,=result['row_calculation'][0]['calculation']['resources']
    assert resource['known_cost']==40
    assert resource['price_source']=='fgis_current'
    assert resource['normative_price']==10
    assert result['row_calculation'][0]['pricebook']['identity']==basis['identity']
    assert resource['price_record_identity']['pricebook_identity']==basis['identity']


def test_default_B_does_not_change_A_document_or_restart(tmp_path, monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis_a=price_basis(tmp_path,'A',20)
    basis_b=price_basis(tmp_path,'B',90)
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:basis_a['path'])
    first=package(tmp_path,[parent],basis_a,'first.xlsx')
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:basis_b['path'])
    prices.get_pricebook.cache_clear()
    later=package(tmp_path,[parent],deepcopy(basis_a),'later.xlsx')
    assert later['row_calculation']==first['row_calculation']
    assert later['row_calculation'][0]['calculation']['resources'][0]['known_cost']==40


def test_equal_amount_different_basis_invalidates_receipt_review(tmp_path, monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    a,b=price_basis(tmp_path,'A',20),price_basis(tmp_path,'B',20)
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:a['path'])
    first=package(tmp_path,[parent],a,'a.xlsx')['row_calculation'][0]
    second=package(tmp_path,[parent],b,'b.xlsx')['row_calculation'][0]
    assert first['calculation']['summary']['known_amount']==second['calculation']['summary']['known_amount']==40
    assert first['calculation_ref']!=second['calculation_ref']
    confirmed=dict(parent,resource_review_ref=first['calculation_ref'],
                   resource_review_status='keep_all_confirmed')
    same=package(tmp_path,[confirmed],a,'a-confirmed.xlsx',review=True)['row_calculation'][0]
    changed=package(tmp_path,[confirmed],b,'b-stale-review.xlsx',review=True)['row_calculation'][0]
    assert same['status']=='bound'
    assert changed['status']=='resource_review_required'
    assert changed['calculation_ref']==second['calculation_ref']


@pytest.mark.parametrize('status',['not_configured','legacy_unpinned'])
def test_missing_current_basis_never_uses_source_price(tmp_path,monkeypatch,status):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:None)
    basis=dict(schema='les.document-price-basis.v1',status=status,path=None,identity=None)
    result=package(tmp_path,[parent],basis)
    fact=result['row_calculation'][0]
    assert fact['calculation']['summary']['full_amount'] is None
    assert fact['calculation']['resources'][0]['known_cost'] is None
    assert fact['calculation']['summary']['known_amount']==0
    book=openpyxl.load_workbook(result['file_path'])
    resource=next(row for row in book.worksheets[0] if len(row)>11 and row[1].value=='R1')
    assert resource[9].value is None and resource[11].value is None
    book.close()


def test_conflicting_price_records_leave_amount_unknown(tmp_path,monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis=price_basis(tmp_path,duplicates=True)
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:basis['path'])
    result=package(tmp_path,[parent],basis)
    item=result['row_calculation'][0]['calculation']['resources'][0]
    assert item['known_cost'] is None
    assert item['price_action']=='price_record_ambiguous'
    assert result['pricing_status']!='priced_final'


def test_primary_unit_thousand_m_and_actual_excel_operands(tmp_path,monkeypatch):
    # Independent supply exercises the actual primary material path and book.
    basis=price_basis(tmp_path,price=56223.48,unit='1000 м')
    monkeypatch.setattr(prices,'resolve_pricebook_path',lambda *_:basis['path'])
    row=dict(source_row='S!1',work_id='supply',title='Synthetic material',unit='м',quantity=8500,
             decision_kind='материал',resource_code='R1')
    result=package(tmp_path,[row],basis)
    item=result['row_calculation'][0]['calculation']['resources'][0]
    assert item['quantity']==8500 and item['known_cost']==477899.58
    book=openpyxl.load_workbook(result['file_path'])
    resource=next(row for row in book.worksheets[0] if len(row)>11 and row[1].value=='R1')
    assert resource[6].value==8500
    assert resource[9].value==pytest.approx(56.22348)
    assert str(resource[11].value).startswith('=ROUND(G')
    comments='\n'.join(c.comment.text for s in book for cells in s for c in cells if c.comment)
    assert basis['identity']['sha256'] in comments
    book.close()


def test_legacy_explicit_manual_api_still_uses_manual_price():
    trace=rim.build_position_trace(dict(qty=2,unit='шт',nr_pct=0,sp_pct=0,
        resources=[dict(kind='material',code='R1',name='Manual',unit='шт',per_unit=1,price=10)]))
    assert trace['summary']['known_amount']==20


def test_primary_incompatible_current_price_unit_stays_unknown_in_book(tmp_path,monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis=price_basis(tmp_path,unit='кг')
    result=package(tmp_path,[parent],basis)
    item=result['row_calculation'][0]['calculation']['resources'][0]
    assert item['known_cost'] is None
    assert item['price_action']=='price_unit_conflict'
    assert result['row_calculation'][0]['calculation']['summary']['full_amount'] is None
    book=openpyxl.load_workbook(result['file_path'])
    resource=next(row for row in book.worksheets[0] if len(row)>11 and row[1].value=='R1')
    assert resource[9].value is None and resource[11].value is None
    book.close()


def test_primary_missing_period_does_not_infer_it_from_filename(tmp_path,monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis=price_basis(tmp_path,name='spb_2kv2026')
    path=Path(basis['path'])
    pd.read_parquet(path).drop(columns=['quarter']).to_parquet(path,index=False)
    basis['identity'].update(period='',sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    result=package(tmp_path,[parent],basis)
    assert result['row_calculation'][0]['pricebook']['identity']['period']==''
    book=openpyxl.load_workbook(result['file_path'])
    headers=[str(cell.value) for row in book.worksheets[0] for cell in row
             if str(cell.value).startswith('Составлен(а) в текущем уровне цен:')]
    assert headers and all('Не указан' in h and 'spb_2kv2026' not in h for h in headers)
    book.close()


def test_missing_price_of_excluded_resource_is_zero_effect_not_fake_price(tmp_path,monkeypatch):
    parent,action,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    basis=price_basis(tmp_path)
    pd.DataFrame([dict(code='UNRELATED',unit='шт',price_current=20,region='A',quarter='2026Q3',version='synthetic-v1')]).to_parquet(basis['path'],index=False)
    basis['identity']['sha256']=hashlib.sha256(Path(basis['path']).read_bytes()).hexdigest()
    result=package(tmp_path,[parent,action],basis)
    item=result['row_calculation'][0]['calculation']['resources'][0]
    assert item['known_cost']==0 and item['excluded_per_unit']==1
    book=openpyxl.load_workbook(result['file_path'])
    resource=next(row for row in book.worksheets[0] if len(row)>11 and row[1].value=='R1')
    assert resource[9].value is None and resource[11].value==0
    book.close()


def test_replaced_pinned_file_after_restart_is_visible_conflict(tmp_path,monkeypatch):
    parent,_,_=setup_rows(tmp_path,monkeypatch,prices=(10,))
    a,b=price_basis(tmp_path,'A',20),price_basis(tmp_path,'B',90)
    Path(a['path']).write_bytes(Path(b['path']).read_bytes())
    prices.get_pricebook.cache_clear()
    result=package(tmp_path,[parent],a)
    assert result['pricing_status']=='price_basis_conflict'
    assert result['price_basis_conflict']['expected']==a['identity']
    assert result['row_calculation'][0]['calculation']['summary']['known_amount']==0
    assert result['row_calculation'][0]['price_basis_conflict']['expected']==a['identity']
