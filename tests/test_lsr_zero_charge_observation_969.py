"""Literal source-confirmed zero work needs no charge, not a confirmed norm."""
from copy import deepcopy
from dataclasses import replace

import openpyxl
import pytest

from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog
from proxy.services.lsr_financial_feedback_service import incomplete_financial_operations
from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness
from proxy.services.lsr_workbook_adapter_service import (
    _attach_catalog_facts, collect_lsr_price_gaps, collect_lsr_workbook_gaps,
)
from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx


def _catalog(tmp_path, quantities=(0,)):
    path = tmp_path / "Нулевой объём с пробелами.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active.append(["Наименование", "Ед. изм.", "Количество"])
    for index, quantity in enumerate(quantities):
        workbook.active.append([f"Literal work {index}", "шт", quantity])
    workbook.save(path)
    workbook.close()
    return load_tabular_source_catalog(path)


def _row(quantity=0, number=2):
    return {"source_row": f"S!R{number}", "row_role": "work", "title": "Literal work",
            "quantity": quantity, "unit": "шт", "operation_id": f"literal-op-{number}",
            "revision": 1, "norm_code": "", "norm_status": None,
            "evidence_refs": [], "coverage": "Авторский комментарий сохранён"}


def _pipeline(rows, catalog):
    decisions = deepcopy(rows)
    if catalog is not None:
        assert _attach_catalog_facts(decisions, "fixture-attachment", catalog_override=catalog) is catalog
    trace = build_lsr_trace_from_visible_rows(decisions, source_catalog=catalog)
    missing = collect_lsr_workbook_gaps(trace)
    prices = collect_lsr_price_gaps(trace)
    ready = build_lsr_result_readiness(
        packaged=True, rows=decisions, source_catalog=catalog,
        catalog_locators=[row["source_row"] for row in decisions],
        catalog_work_count=len(decisions), row_projection=trace["row_projection"],
        missing=missing, price_gaps=prices, pricing_status=trace["summary"]["result_status"],
        calculated_positions=trace["summary"]["calculated_positions"], job_status="complete",
    )
    return trace, ready, missing, prices


@pytest.mark.parametrize("source_quantity,expected", [(0, True), (0.0, True), (-0.0, True),
    (1, False), (-1, False), (None, False), ("0", False), (False, False),
    (float("nan"), False), (float("inf"), False)])
def test_zero_fact_belongs_to_source_catalog_not_a_pricebook(tmp_path, source_quantity, expected):
    from proxy.services.attachment_source_catalog_service import source_confirms_literal_zero_quantity

    catalog = _catalog(tmp_path)
    source = catalog.lookup("S!R2")
    changed = replace(catalog, rows=tuple(
        replace(row, header_quantity=source_quantity) if row is source else row for row in catalog.rows))
    assert source_confirms_literal_zero_quantity(changed, "S!R2") is expected
    assert source_confirms_literal_zero_quantity(changed, "S!R999") is False
    assert source_confirms_literal_zero_quantity(None, "S!R2") is False
    assert source_confirms_literal_zero_quantity(replace(changed, extraction_complete=False), "S!R2") is False


@pytest.mark.parametrize("quantity", [0, 0.0, -0.0])
def test_catalog_zero_has_no_norm_price_calculation_or_retry_debt(tmp_path, quantity):
    catalog = _catalog(tmp_path)
    rows = [_row(quantity)]
    before = deepcopy(rows)
    trace, ready, missing, prices = _pipeline(rows, catalog)
    binding = trace["row_bindings"][0]
    projection = trace["row_projection"][0]
    assert binding["charge_not_required"] is True
    assert binding["source_quantity"] == 0
    assert "начисление не требуется" in binding["message"].casefold()
    assert not binding.get("position") and trace["sections"] == []
    assert projection["charge_not_required"] is True
    assert projection["source_accounted"] is True and projection["volume_verified"] is True
    assert projection["norm_selected"] is None and projection["applicability_substantiated"] is None
    assert projection["position_calculated"] is None and projection["prices_provided"] is None
    assert trace["summary"]["selected_norms"] == trace["summary"]["calculated_positions"] == 0
    assert trace["summary"]["bound_rows"] == trace["summary"]["unbound_rows"] == 0
    assert trace["summary"]["charge_not_required_rows"] == 1
    assert trace["summary"]["confirmed_amount"] == trace["summary"]["preliminary_amount"] == 0
    assert missing == prices == []
    assert ready["dimensions"]["norms"]["gaps"] == ready["dimensions"]["prices"]["gaps"] == []
    assert ready["gap_states"]["calculation_incomplete"]["count"] == 0
    assert ready["research_locators"] == ready["research_operation_ids"] == []
    assert ready["counts"]["works"] == 1 and ready["counts"]["norms_selected"] == 0
    assert ready["coverage"]["accounted"] == 1 and ready["estimate_complete"] is True
    assert incomplete_financial_operations(rows, trace["row_projection"]) == []
    assert rows == before
    assert {key: binding["source_row"].get(key) for key in rows[0]} == rows[0]


@pytest.mark.parametrize("other_quantity", [1, -1])
def test_mixed_source_retries_only_nonzero_operation_without_hiding_its_sign(tmp_path, other_quantity):
    catalog = _catalog(tmp_path, (0, other_quantity))
    rows = [_row(0), _row(other_quantity, 3)]
    before = deepcopy(rows)
    trace, ready, _, _ = _pipeline(rows, catalog)
    assert trace["row_projection"][0]["charge_not_required"] is True
    assert not trace["row_projection"][1].get("charge_not_required")
    assert ready["research_locators"] == ["S!R3"]
    assert ready["research_operation_ids"] == ["literal-op-3"]
    assert ready["gap_states"]["calculation_incomplete"]["count"] == 1
    assert ready["coverage"]["accounted"] == 2
    assert trace["summary"]["charge_not_required_rows"] == trace["summary"]["unbound_rows"] == 1
    assert rows == before and trace["row_bindings"][1]["source_quantity"] == other_quantity


# Exact numeric strings follow the quantity contract and are covered in 972;
# arbitrary text still cannot establish a source-confirmed zero exemption.
@pytest.mark.parametrize("quantity", [None, False, "0шт", "zero", float("nan"), float("inf")])
def test_unknown_boolean_text_and_nonfinite_are_not_a_literal_zero_exemption(tmp_path, quantity):
    trace, ready, _, _ = _pipeline([_row(quantity)], _catalog(tmp_path))
    assert not trace["row_projection"][0].get("charge_not_required")
    assert ready["dimensions"]["norms"]["complete"] is False
    assert ready["research_operation_ids"] == ["literal-op-2"]


@pytest.mark.parametrize("kind", ["missing", "unconfirmed", "unknown_column", "different_quantity",
                                 "unknown_unit", "incomplete_extraction"])
def test_a_model_zero_cannot_confirm_its_own_source(tmp_path, kind):
    catalog = _catalog(tmp_path)
    source = catalog.lookup("S!R2")
    if kind == "missing":
        catalog = None
    elif kind == "incomplete_extraction":
        catalog = replace(catalog, extraction_complete=False)
    else:
        changes = ({"header_quantity": None} if kind == "unconfirmed" else
                   {"quantity_column_known": False} if kind == "unknown_column" else
                   {"header_quantity": 1} if kind == "different_quantity" else
                   {"column_headers": tuple((column, label) for column, label in source.column_headers
                                            if column != 2)})
        catalog = replace(catalog, rows=(replace(source, **changes),))
    trace, ready, _, _ = _pipeline([_row()], catalog)
    assert not trace["row_projection"][0].get("charge_not_required")
    assert ready["estimate_complete"] is False and ready["research_operation_ids"] == ["literal-op-2"]


@pytest.mark.parametrize("fields", [
    {"calc_quantity": 2, "calc_unit": "шт", "conversion_source": "Literal authored basis"},
    {"calc_quantity": 2, "conversion_factor": 3, "conversion_source": "Literal conflicting basis"},
    {"confirmed_price": 100, "price_kind": "position_total", "price_source": "Literal quote"},
    {"resource_bindings": [{"action": "add", "quantity_total": 2, "source_refs": ["Q1.H1"]}]},
    {"norm_components": [{"norm_code": "Literal selected addition", "evidence_refs": ["Q1.H1"]}]},
    {"norm_components": {"malformed": "Literal authored value"}},
])
def test_explicit_derived_charge_or_quote_is_preserved_and_not_silently_exempted(tmp_path, fields):
    row = {**_row(), **deepcopy(fields)}
    before = deepcopy(row)
    trace, ready, _, _ = _pipeline([row], _catalog(tmp_path))
    assert not trace["row_projection"][0].get("charge_not_required")
    assert ready["estimate_complete"] is False
    assert row == before
    assert {key: trace["row_bindings"][0]["source_row"].get(key) for key in row} == row


def test_authored_exemption_flag_is_not_trusted_for_nonzero_work(tmp_path):
    trace, ready, _, _ = _pipeline([{**_row(1), "charge_not_required": True}], _catalog(tmp_path, (1,)))
    assert not trace["row_projection"][0].get("charge_not_required")
    assert ready["research_operation_ids"] == ["literal-op-2"]


def test_rendered_zero_information_row_keeps_quantity_and_explanation(tmp_path):
    trace, _, _, _ = _pipeline([_row()], _catalog(tmp_path))
    output = tmp_path / "Нулевая работа.xlsx"
    render_lsr_xlsx(trace, output)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    main = workbook["ЛСР РИМ"]
    visible = [row for row in main.iter_rows(values_only=True)
               if "начисление не требуется" in str(row[2] or "").casefold()]
    assert len(visible) == 1 and visible[0][6] == 0
    check_rows = list(workbook["Проверка"].iter_rows(values_only=True))
    assert any("literal-op-2" in str(row[3] or "") for row in check_rows)
    assert any("начисление не требуется" in str(row[8] or "").casefold() for row in check_rows)
    workbook.close()


@pytest.mark.parametrize("value", [None, []])
def test_explicit_empty_components_do_not_create_a_charge_for_a_confirmed_zero(tmp_path, value):
    trace, ready, _, _ = _pipeline([{**_row(), "norm_components": value}], _catalog(tmp_path))
    assert trace["row_projection"][0]["charge_not_required"] is True
    assert ready["research_operation_ids"] == []
    assert ready["counts"]["norms_selected"] == 0


@pytest.mark.parametrize("kind", ["zero", "positive", "negative", "selected_norm", "boolean", "unknown_source"])
def test_validator_uses_the_same_confirmed_zero_boundary_and_preserves_other_ref_rules(tmp_path, kind):
    from proxy.services.chat_evidence_application_service import (
        blocking_model_rag_result_errors, validate_model_rag_result_structure,
    )
    quantity = 1 if kind == "positive" else -1 if kind == "negative" else 0
    row = _row(quantity)
    catalog = _catalog(tmp_path, (quantity,))
    if kind == "selected_norm":
        row["norm_code"] = "ГЭСН99-99-999-99"
    if kind in {"positive", "negative"}:
        row["norm_status"] = "MISSING"
    if kind == "boolean":
        row["quantity"] = False
    if kind == "unknown_source":
        catalog = None
    before = deepcopy(row)
    errors = validate_model_rag_result_structure(
        [row], [], expected_source_rows=1, source_catalog=catalog, artifact_kind="lsr")
    assert ("missing_evidence_ref:S!R2" in errors) is (kind != "zero")
    assert blocking_model_rag_result_errors(errors, work_row_count=1, artifact_kind="lsr") == []
    assert row == before


def test_mixed_validator_calculator_renderer_readiness_keeps_only_nonzero_ref_and_retry_debt(tmp_path):
    from proxy.services.chat_evidence_application_service import validate_model_rag_result_structure
    catalog = _catalog(tmp_path, (0, 1))
    rows = [_row(), {**_row(1, 3), "norm_status": "MISSING", "open_gap": "Literal open model decision"}]
    before = deepcopy(rows)
    errors = validate_model_rag_result_structure(
        rows, [], expected_source_rows=2, source_catalog=catalog, artifact_kind="lsr")
    assert "missing_evidence_ref:S!R2" not in errors
    assert "missing_evidence_ref:S!R3" in errors
    trace, ready, _, _ = _pipeline(rows, catalog)
    output = tmp_path / "Mixed zero and open work.xlsx"
    render_lsr_xlsx(trace, output)
    assert output.is_file()
    assert ready["research_locators"] == ["S!R3"]
    assert ready["research_operation_ids"] == ["literal-op-3"]
    assert rows == before
