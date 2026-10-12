import openpyxl

from proxy.services.rim_lsr_trace_service import positions_from_visible_lsr_rows
from tools.smeta_book_audit import read_book, classify, IN_BOOK_NO_PRICE

CODE = 'ГЭСНм08-02-395-01'


def test_resource_reference_is_not_another_installation(monkeypatch):
    from proxy.services import rim_lsr_trace_service as service
    monkeypatch.setattr(service.gesn_service, 'get_norm', lambda code: {'unit': 'шт'})
    row = dict(source_row='12', norm_code=CODE, quantity='8', unit='шт',
               decision_kind='ресурс', included_in='S!11', coverage='Крепёж входит в работу')
    result = positions_from_visible_lsr_rows([row])
    assert result['positions'] == []
    binding = result['row_bindings'][0]
    assert binding['status'] == 'resource_parent_missing'
    assert binding['source_row'] == row


def test_diagnostic_norm_is_not_a_workbook_position(tmp_path):
    path = tmp_path / 'book.xlsx'
    book = openpyxl.Workbook()
    book.active.title = 'ЛСР РИМ'
    check = book.create_sheet('Проверка')
    check.append(['Строка', 'work_id', 'Статус', 'Норма'])
    check.append([12, None, 'unit_conflict', CODE])
    book.save(path)
    assert CODE not in read_book(path)


def test_one_work_position_cannot_confirm_two_source_rows(tmp_path):
    path = tmp_path / 'book.xlsx'
    book = openpyxl.Workbook()
    book.active.append([1, CODE, 'Лоток', 'шт', 2, 1, 2])
    book.save(path)
    positions = read_book(path)
    row = dict(source_row='12', title='Лоток', quantity=2, unit='шт')
    decision = dict(row, norm_code=CODE)
    assert classify(row, decision, positions)['state'] == IN_BOOK_NO_PRICE
    assert classify(dict(row, source_row='13'), decision, positions)['state'] != IN_BOOK_NO_PRICE


def test_named_parent_does_not_prove_resource_calculation():
    row = dict(source_row='12', quantity='8', unit='шт')
    verdict = classify(row, dict(row, decision_kind='ресурс', included_in=CODE,
                                coverage='Входит в монтаж'), {})
    assert verdict['state'] == 'учёт ресурсом не подтверждён'


def test_uncached_total_formula_is_not_reported_as_missing_price(tmp_path):
    path = tmp_path / 'formula.xlsx'
    book = openpyxl.Workbook()
    book.active.append([1, CODE, 'Лоток', 'шт', 2, 1, 2])
    book.active.append([None, None, 'Всего по позиции', *([None] * 8), '=ROUND(2*10,2)'])
    book.save(path)
    row = dict(source_row='12', title='Лоток', quantity=2, unit='шт')
    verdict = classify(row, dict(row, norm_code=CODE), read_book(path))
    assert verdict['state'] == 'в книге с формулой стоимости; кеш отсутствует'


def test_wrong_quantity_is_not_confirmed(tmp_path):
    path = tmp_path / 'quantity.xlsx'
    book = openpyxl.Workbook()
    book.active.append([1, CODE, 'Лоток', '100 м', 200, 1, 200])
    book.save(path)
    row = dict(source_row='12', title='Лоток', quantity=200, unit='м')
    assert classify(row, dict(row, norm_code=CODE), read_book(path))['state'] != IN_BOOK_NO_PRICE


def test_resource_diagnostic_reaches_downloadable_workbook(tmp_path, monkeypatch):
    import asyncio
    from proxy.services import lsr_workbook_adapter_service as adapter
    monkeypatch.setattr(adapter.fgis_price_service, 'resolve_pricebook_path', lambda _: None)
    path = tmp_path / 'resource.xlsx'
    decision = dict(source_row='12', norm_code=CODE, quantity='8', unit='шт',
                    decision_kind='ресурс', included_in='S!11', coverage='Крепёж входит в работу')
    result = asyncio.run(adapter.build_lsr_workbook_from_decisions(
        path, {'decisions': [decision]}, path, lambda *args: None))
    assert result['bound_rows'] == 0
    assert 'row:12:resource_parent_missing' in result['missing']
    book = openpyxl.load_workbook(path)
    try:
        assert any('resource_parent_missing' in row for row in book['Проверка'].values)
    finally:
        book.close()
