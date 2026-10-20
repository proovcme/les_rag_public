"""Actual chat/repack boundaries retain accepted rows without claiming a final document."""
from copy import deepcopy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from proxy.routers import chat
from proxy.services.workbook_research_job_service import WorkbookResearchJobService
from proxy.services.chat_evidence_application_service import workbook_document_issues
from proxy.services.openai_compatible_transport_service import InferenceResponse
from test_model_rag_sheet_evidence_reads import NORM_CODE, NORM_CODE_B, _run_sheet
from tests.test_norm_card_proof_request_930 import SECOND_NORM_ANSWER


def _assert_draft(payload, readiness, root: Path, raw_line: str):
    assert payload["status"] == "partial"
    assert payload["artifact"]["download_url"]
    assert payload["artifact"]["draft"] is True
    assert payload["pricing_status"] == "priced_partial"
    assert payload["confirmed_amount"] == 0
    assert payload["preliminary_amount"] == 600
    assert payload["calculated_positions"] == 1  # the independent accepted position remains calculated
    assert payload["document_issues"]
    assert readiness["technical_complete"] is False
    assert readiness["estimate_complete"] is False
    assert readiness["document_issues"] == payload["document_issues"]
    paths = list(root.rglob(payload["artifact"]["filename"]))
    assert paths
    workbook = load_workbook(paths[0], data_only=False)
    try:
        text = "\n".join(str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value is not None)
        assert "РАСЧЁТ НЕ ЗАВЕРШЁН (PRICED_PARTIAL)" in text
        assert "MODEL_RESULT_TABLE_ROW_WIDTH_MISMATCH" in text
        assert "S!R2" in text and raw_line in text
        assert "ИЗВЕСТНАЯ РАССЧИТАННАЯ ЧАСТЬ" in text
    finally:
        workbook.close()


@pytest.mark.asyncio
async def test_application_delivers_accepted_subset_as_technical_draft(tmp_path, monkeypatch):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, malformed_first_packet=True)
    assert len(run.requests) == 4 and not run.responses
    job = WorkbookResearchJobService().get(run.result["retrieval_trace"]["workbook_research_job"]["job_id"])
    assert job["status"] == "paused" and job["remaining"] == ["S!R2"]
    assert [row["source_row"] for row in job["accepted_rows"]] == ["S!R3"]
    assert len(run.executions) == 1
    call, payload = run.executions[0]
    assert call["args"]["decisions"] == job["accepted_rows"]
    assert "document_issues" not in call["args"]  # trusted execution context only
    _assert_draft(payload, run.result["artifact"]["result_readiness"], tmp_path,
                  job["unapplied_proposals"][0]["decision"]["raw_line"])
    meta = run.result["retrieval_trace"]["model_result"]
    assert meta["status"] == "structural_error" and meta["packaged"] is True


@pytest.mark.asyncio
async def test_repack_does_not_turn_pending_raw_rows_into_final_money(tmp_path, monkeypatch):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, malformed_first_packet=True)
    store = WorkbookResearchJobService()
    parent = store.get(run.result["retrieval_trace"]["workbook_research_job"]["job_id"])
    calls_before = len(run.requests)
    identity = parent["identity"]
    result = await chat.repack_lsr_without_model(chat.LsrRepackRequest(
        job_id=parent["job_id"], session_id=identity["session_id"], attachment_id=identity["attachment_id"],
        question=parent["task_context"]["original_question"], dataset_ids=identity["dataset_ids"],
        profile_revision_id=identity["profile_revision_id"], model_identity=identity["model_identity"]), _user={})
    assert result["model_called"] is False and len(run.requests) == calls_before
    _assert_draft(result["result"], result["harvested"]["artifact"]["result_readiness"], tmp_path,
                  parent["unapplied_proposals"][0]["decision"]["raw_line"])
    updated = store.get(parent["job_id"])
    for field in ("status", "accepted_rows", "remaining", "unapplied_proposals"):
        assert updated[field] == parent[field]
    replay = await chat.repack_lsr_without_model(chat.LsrRepackRequest(
        job_id=parent["job_id"], session_id=identity["session_id"], attachment_id=identity["attachment_id"],
        question=parent["task_context"]["original_question"], dataset_ids=identity["dataset_ids"],
        profile_revision_id=identity["profile_revision_id"], model_identity=identity["model_identity"]), _user={})
    assert replay["result"]["artifact"]["revision_id"] == result["result"]["artifact"]["revision_id"]
    assert replay["result"]["document_issues"] == result["result"]["document_issues"]
    assert len(run.requests) == calls_before


def test_pending_read_stop_and_raw_proposals_are_visible_until_research_resolves():
    raw = {"code": "MODEL_RESULT_JSON_DUPLICATE_FIELD", "decision": {"raw_json": '{"quantity":1,"quantity":2}'}}
    job = {"status": "paused", "remaining": ["Раздел!R14"], "unapplied_proposals": [raw],
           "held_packet": {"read_stop": "evidence_read_repeat"}}
    issues = workbook_document_issues(job)
    assert issues[0]["remaining_locators"] == ["Раздел!R14"]
    assert issues[1]["proposal"] == raw
    assert issues[2]["reason"] == "evidence_read_repeat"
    job.update(status="complete", remaining=[])
    assert workbook_document_issues(job) == []  # resolved raw history is retained by the job, not an open gap


@pytest.mark.asyncio
async def test_model_tool_arguments_cannot_supply_document_integrity_metadata(tmp_path, monkeypatch):
    run = await _run_sheet(tmp_path, monkeypatch, reads=True, malformed_first_packet=True)
    job = WorkbookResearchJobService().get(run.result["retrieval_trace"]["workbook_research_job"]["job_id"])
    call = {"tool": "build_lsr_workbook", "args": {
        **run.executions[0][0]["args"], "document_issues": [{"code": "model-owned-fake"}]}}
    identity = job["identity"]
    rejected = await chat._execute_chat_workbook_tool(call, {
        "session_id": identity["session_id"], "attachment_id": identity["attachment_id"],
        "dataset_ids": identity["dataset_ids"], "profile_revision_id": identity["profile_revision_id"],
        "model_identity": identity["model_identity"]}, None)
    assert rejected["status"] not in {"complete", "partial"}
    assert not rejected.get("artifact")
    assert len(run.requests) == 4


async def _ambiguous_proof_draft_977(tmp_path, monkeypatch):
    # The existing real sheet pipeline receives an authentic unknown-ID proof
    # update. The merge guard must retain it as a proposal, not price its norm.
    proposed = {
        "source_row": "S!R2", "row_role": "work",
        "operation_id": "unissued-proof-operation",
        "title": "Literal unaccepted priced-work proposal",
        "unit": "m", "quantity": 12, "norm_code": NORM_CODE,
        "evidence_refs": ["Q1.H1"], "coverage": "Literal proposed replacement",
    }
    raw = json.dumps({"artifact_kind": "lsr", "rows": [proposed]}, ensure_ascii=False)
    run = await _run_sheet(
        tmp_path, monkeypatch, reads=False, group_char_budget=7000,
        distinct_norms=True, final_answer=SECOND_NORM_ANSWER,
        proof_replies=[InferenceResponse(
            text=raw, tool_calls=(), finish_reason="stop", usage={})],
    )
    return run, proposed, raw


@pytest.mark.asyncio
async def test_identity_rejected_proof_delivers_only_accepted_checkpoint_as_draft_977(
        tmp_path, monkeypatch):
    run, proposed, raw = await _ambiguous_proof_draft_977(tmp_path, monkeypatch)
    trace = run.result["retrieval_trace"]
    store = WorkbookResearchJobService()
    job = store.get(trace["workbook_research_job"]["job_id"])
    before = deepcopy(job)
    assert job["status"] == "paused" and job["remaining"]
    assert len(run.requests) == 3 and not run.responses, "packaging must not buy another model call"
    assert len(job["accepted_rows"]) == 2
    assert all(row["norm_code"] == NORM_CODE_B for row in job["accepted_rows"])
    assert all(row["operation_id"] != proposed["operation_id"] for row in job["accepted_rows"])
    diagnostics = trace["model_result_transport_diagnostics"]
    assert len(diagnostics) == 1
    assert diagnostics[0]["code"] == "MODEL_PROOF_UPDATE_IDENTITY_AMBIGUOUS"
    assert diagnostics[0]["proposal"] == proposed and diagnostics[0]["raw_answer"] == raw
    assert len(job["unapplied_proposals"]) == 1
    assert job["unapplied_proposals"][0]["proposal"] == proposed
    assert job["unapplied_proposals"][0]["raw_answer"] == raw

    assert len(run.executions) == 1, "a rejected new update must not suppress the saved draft"
    call, payload = run.executions[0]
    assert call["args"]["decisions"] == job["accepted_rows"]
    assert call["args"]["draft"] is True
    assert "document_issues" not in call["args"]
    assert payload["status"] == "partial" and payload["artifact"]["draft"] is True
    assert payload["pricing_status"] == "priced_partial" and payload["confirmed_amount"] == 0
    projection = json.dumps(payload["row_projection"], ensure_ascii=False)
    assert proposed["operation_id"] not in projection and NORM_CODE not in projection
    issues = payload["document_issues"]
    assert any(issue["code"] == "MODEL_PROOF_UPDATE_IDENTITY_AMBIGUOUS" for issue in issues)
    assert any(issue["code"] == "MODEL_PROPOSAL_UNAPPLIED"
               and issue["proposal"] == job["unapplied_proposals"][0] for issue in issues)
    readiness = run.result["artifact"]["result_readiness"]
    assert readiness["technical_complete"] is False and readiness["estimate_complete"] is False
    assert readiness["document_issues"][:len(issues)] == issues
    assert readiness["document_issues"][len(issues):] == [
        {"code": "normative_generation_conflict", "locator": locator,
         "message": "поколение процитированной нормы не воспроизведено для расчёта; авторский выбор сохранён"}
        for locator in ("S!R2", "S!R3")
    ]
    meta = trace["model_result"]
    assert meta["status"] == "structural_error" and meta["packaged"] is True
    assert meta["transport_diagnostics"] == diagnostics

    output = next(tmp_path.rglob(payload["artifact"]["filename"]))
    workbook = load_workbook(output, data_only=False)
    try:
        visible = "\n".join(str(cell.value) for sheet in workbook
                            for row in sheet for cell in row if cell.value is not None)
        assert "РАСЧЁТ НЕ ЗАВЕРШЁН" in visible
        assert "MODEL_PROOF_UPDATE_IDENTITY_AMBIGUOUS" in visible
        assert proposed["operation_id"] in visible and proposed["title"] in visible
    finally:
        workbook.close()
    after = store.get(job["job_id"])
    for field in ("status", "accepted_rows", "remaining", "unapplied_proposals", "held_packet", "identity"):
        assert after[field] == before[field], "rendering must not admit or rewrite the pending update"


@pytest.mark.asyncio
async def test_no_accepted_checkpoint_cannot_be_packaged_as_saved_draft_977(tmp_path, monkeypatch):
    rows = [
        {"source_row": source, "row_role": "work", "title": title,
         "unit": "m", "quantity": 999, "norm_code": "MISSING",
         "open_gap": "Literal unaccepted quantity proposal"}
        for source, title in (("S!R2", "Cable A"), ("S!R3", "Cable B"))
    ]
    raw = json.dumps({"artifact_kind": "lsr", "rows": rows}, ensure_ascii=False)
    run = await _run_sheet(tmp_path, monkeypatch, reads=False, final_answer=raw)
    trace = run.result["retrieval_trace"]
    job = WorkbookResearchJobService().get(trace["workbook_research_job"]["job_id"])
    assert job["status"] == "paused" and job["accepted_rows"] == []
    assert len(job["unapplied_proposals"]) == 2 and len(job["remaining"]) == 2
    assert run.executions == [] and not run.result.get("artifact")
    assert trace["model_result"]["status"] == "structural_error"
    assert trace["model_result"]["packaged"] is False
    assert len(run.requests) == 2 and not run.responses


@pytest.mark.asyncio
async def test_saved_draft_cannot_borrow_accepted_rows_from_another_source_identity_977(
        tmp_path, monkeypatch):
    # Corrupt only the trusted snapshot handed back at the paused boundary.
    # The real durable rows/source remain unchanged; no foreign-scope rows may
    # reach the exporter even if an earlier stage returned the wrong snapshot.
    original_get = WorkbookResearchJobService.get
    mismatches = []

    def foreign_snapshot(self, job_id, **kwargs):
        current = original_get(self, job_id, **kwargs)
        if current.get("status") == "paused" and current.get("unapplied_proposals"):
            current = deepcopy(current)
            mismatches.append(current["identity"]["source_sha256"])
            current["identity"]["source_sha256"] = "0" * 64
        return current

    monkeypatch.setattr(WorkbookResearchJobService, "get", foreign_snapshot)
    run, _proposed, _raw = await _ambiguous_proof_draft_977(tmp_path, monkeypatch)
    assert mismatches and all(value != "0" * 64 for value in mismatches)
    trace = run.result["retrieval_trace"]
    durable = original_get(WorkbookResearchJobService(), trace["workbook_research_job"]["job_id"])
    assert durable["accepted_rows"] and durable["unapplied_proposals"]
    assert durable["identity"]["source_sha256"] != "0" * 64
    assert run.executions == [] and not run.result.get("artifact")
    assert trace["model_result"]["status"] == "structural_error"
    assert trace["model_result"]["packaged"] is False
    assert len(run.requests) == 3 and not run.responses
