"""A source-confirmed zero exemption must survive public repack and bounded history."""
from copy import deepcopy
import hashlib
from io import BytesIO

import openpyxl
import pytest

from proxy.routers import chat
from proxy.services import fgis_price_service
from proxy.services.artifact_revision_service import ArtifactRevisionStore
from proxy.services.chat_attachment_service import preserve_read_attachment
from proxy.services.chat_evidence_application_service import (
    harvest_workbook_tool_result, safe_workbook_history_projection,
)
from proxy.services.workbook_research_job_service import WorkbookResearchJobService


@pytest.mark.asyncio
@pytest.mark.parametrize("neighbor_source_quantity,neighbor_authored_quantity", [(1, 1), (None, 0)])
async def test_actual_repack_keeps_source_zero_exemption_through_history_and_public_readiness(
    tmp_path, monkeypatch, neighbor_source_quantity, neighbor_authored_quantity,
):
    attachment_root = tmp_path / "attachments"
    monkeypatch.setenv("LES_CHAT_ATTACHMENT_ROOT", str(attachment_root))
    monkeypatch.setattr(chat, "mutable_path", lambda path: tmp_path / path)
    # No installed user-owned pricebook or model may affect this source-only seam.
    monkeypatch.setattr(fgis_price_service, "available_pricebooks", lambda *a, **kw: [])
    monkeypatch.setattr(chat, "run_chat_evidence_application",
                        lambda *a, **kw: pytest.fail("model called during repack"))

    source = tmp_path / "Source zero and unresolved neighbor.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active.append(["Наименование", "Ед. изм.", "Количество"])
    workbook.active.append(["Literal zero work", "м", 0])
    workbook.active.append(["Literal unresolved neighbor", "м", neighbor_source_quantity])
    workbook.save(source)
    workbook.close()
    source_bytes = source.read_bytes()
    attachment = preserve_read_attachment(source, attachment_id="read_123456abcdef",
        original_name=source.name, root=attachment_root)

    rows = [
        {"source_row": "S!R2", "operation_id": "literal-zero", "revision": 1,
         "row_role": "work", "title": "Literal zero work", "quantity": 0,
         "unit": "м", "norm_code": "", "norm_status": None, "evidence_refs": []},
        {"source_row": "S!R3", "operation_id": "literal-neighbor", "revision": 1,
         "row_role": "work", "title": "Literal unresolved neighbor",
         "quantity": neighbor_authored_quantity, "unit": "м", "norm_code": "",
         "norm_status": "MISSING", "open_gap": "No model selection yet", "evidence_refs": []},
    ]
    store = WorkbookResearchJobService(tmp_path / "owned-jobs.sqlite")
    question = "Собери ЛСР"
    identity = store.identity(session_id="owned-zero-repack-970", attachment_id=attachment["attachment_id"],
        source_sha256=attachment["sha256"], dataset_ids=[], profile_revision_id="fixture-profile:r1",
        model_identity="offline-fixture", question=question)
    job = store.begin_or_resume(identity, task_context={"original_question": question, "artifact_intent": "lsr"})
    store.record_packet(job["job_id"], sheet="S", packet_locators=["S!R2", "S!R3"],
        accepted_rows=rows, evidence=[], remaining=[])
    store.finish(job["job_id"], status="complete")
    monkeypatch.setattr(chat, "WorkbookResearchJobService", lambda: store)
    before = deepcopy(store.get(job["job_id"]))
    operation_ids = {row["source_row"]: row["operation_id"] for row in before["accepted_rows"]}
    assert set(operation_ids) == {"S!R2", "S!R3"}
    assert operation_ids["S!R2"] and operation_ids["S!R3"] != operation_ids["S!R2"]
    request = chat.LsrRepackRequest(session_id=identity.session_id, job_id=job["job_id"])

    first = await chat.repack_lsr_without_model(request, _user={})
    assert first["model_called"] is False
    artifact = first["result"]["artifact"]
    assert artifact["artifact_kind"] == "lsr_workbook" and artifact["download_url"]
    ready = first["harvested"]["artifact"]["result_readiness"]
    projection = first["result"]["row_projection"]
    zero = next(row for row in projection if row["locator"] == "S!R2")
    neighbor = next(row for row in projection if row["locator"] == "S!R3")
    # RED boundary: 969 removes this calculator-owned boolean in history's field whitelist.
    assert zero["charge_not_required"] is True
    assert zero["norm_selected"] is None and zero["position_calculated"] is None
    assert not neighbor.get("charge_not_required")
    assert ready["gap_states"]["norm_missing"]["count"] == 1
    assert ready["research_locators"] == ["S!R3"]
    assert ready["research_operation_ids"] == [operation_ids["S!R3"]]
    assert operation_ids["S!R2"] not in ready["research_operation_ids"]
    assert ready["gap_states"]["calculation_incomplete"]["count"] == 1

    saved = safe_workbook_history_projection(first["result"])
    assert safe_workbook_history_projection(saved) == saved
    twice = harvest_workbook_tool_result(saved)
    assert next(row for row in twice["row_projection"] if row["locator"] == "S!R2")["charge_not_required"] is True

    artifacts = ArtifactRevisionStore(tmp_path / "storage/artifacts/meta.db", tmp_path / "storage/artifacts/files")
    old_bytes = artifacts.read_bytes(artifact["revision_id"])
    assert hashlib.sha256(old_bytes).hexdigest() == artifact["sha256"]
    rendered = openpyxl.load_workbook(BytesIO(old_bytes), read_only=True, data_only=False)
    try:
        visible = [values for values in rendered["Проверка"].iter_rows(values_only=True)
                   if values[0] == "S!R2"]
        assert len(visible) == 1 and "начисление не требуется" in str(visible[0][8]).casefold()
    finally:
        rendered.close()

    replay = await chat.repack_lsr_without_model(request, _user={})
    assert replay["model_called"] is False
    assert replay["result"]["artifact"]["revision_id"] == artifact["revision_id"]
    assert replay["harvested"]["artifact"]["result_readiness"]["research_operation_ids"] == [operation_ids["S!R3"]]
    assert artifacts.read_bytes(artifact["revision_id"]) == old_bytes
    assert source.read_bytes() == source_bytes
    after = store.get(job["job_id"])
    for field in ("identity", "accepted_rows", "query_cache", "evidence", "remaining", "operator_overrides", "held_packet"):
        assert after.get(field) == before.get(field), field
    assert after["accepted_rows"][0]["quantity"] == 0
    assert not any(row.get("charge_not_required") for row in after["accepted_rows"])
