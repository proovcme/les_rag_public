"""Red reproductions for the six confirmed defects in 0.31.75/781.

Each test names the defect from the finalization directive:
1.1 merge keeps two operations on one locator (operation identity).
1.2 trace sums bind by operation identity, not by source_row.
1.3 norms_selected counts real ciphers only ("-", MISSING are not norms).
1.4 price basis checks region and period together; contradictions stay visible.
1.5 double-count conflict is explicit; a confirmed price is addressed.
1.6 stale guard via the public API; attempt identity survives a restart.
"""

from __future__ import annotations

import pytest

from tests.lsr_citation_support import cited

from proxy.services.chat_evidence_application_service import (
    merge_model_rag_rows,
    validate_model_rag_result_structure,
)
from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness
from proxy.services.lsr_source_coverage_service import apply_confirmed_data
from proxy.services.lsr_workbook_adapter_service import summarize_lsr_price_basis
from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
from proxy.services.workbook_research_job_service import WorkbookResearchJobService


def _work_row(locator: str, title: str, **overrides):
    row = {
        "source_row": locator,
        "row_role": "work",
        "title": title,
        "unit": "шт",
        "quantity": 1,
        "norm_code": "",
        "norm_status": "MISSING",
    }
    row.update(overrides)
    # A row that names a cipher carries the citation record the application stores
    # for it; readiness no longer accepts the status token on its own (§ 3).
    return cited(row)


def _supply_row(locator: str, title: str, **overrides):
    row = {
        "source_row": locator,
        "row_role": "material",
        "supply_kind": "supply",
        "title": title,
        "unit": "м",
        "quantity": 1,
        "norm_code": "",
    }
    row.update(overrides)
    return row


def _operator_price(row, price):
    # This fixture uses the same application-owned confirmation overlay as the
    # chat. A model's confirmed_price field alone does not create an operator.
    return apply_confirmed_data([row], {row["source_row"]: {"confirmed_price": price}})[0]


# --------------------------------------------------------------------------- 1.1


def test_merge_keeps_two_works_on_one_locator():
    a = _work_row("СКС!R6", "Установка шкафа")
    b = _work_row("СКС!R6", "Установка сервера")
    merged = merge_model_rag_rows([], [a, b])
    assert len(merged) == 2, "the second work on the same locator was dropped"
    assert {item["title"] for item in merged} == {"Установка шкафа", "Установка сервера"}


def test_merge_still_drops_identical_redelivery():
    a = _work_row("СКС!R6", "Установка шкафа")
    b = _work_row("СКС!R6", "Установка сервера")
    merged = merge_model_rag_rows(merge_model_rag_rows([], [a, b]), [dict(a), dict(b)])
    assert len(merged) == 2, "a redelivered packet must not create new operations"


def test_merge_preserves_operation_identity_from_checkpoint():
    a = dict(_work_row("СКС!R6", "Установка шкафа"), operation_id="op_a", revision=1)
    b = dict(_work_row("СКС!R6", "Установка сервера"), operation_id="op_b", revision=3)
    merged = merge_model_rag_rows([a, b], [])
    assert [item.get("operation_id") for item in merged] == ["op_a", "op_b"]
    # A redelivered checkpoint row (same operation_id) never duplicates itself.
    merged = merge_model_rag_rows(merged, [dict(b)])
    assert len(merged) == 2


def test_validator_allows_distinct_operations_on_one_locator():
    rows = [
        _work_row("СКС!R6", "Установка шкафа"),
        _work_row("СКС!R6", "Установка сервера"),
    ]
    errors = validate_model_rag_result_structure(
        rows, [], expected_source_rows=1, validate_evidence_refs=False
    )
    assert not [item for item in errors if item.startswith("duplicate_source_row")]


def test_validator_still_flags_identical_repeat():
    row = _work_row("СКС!R6", "Установка шкафа")
    errors = validate_model_rag_result_structure(
        [dict(row), dict(row)], [], expected_source_rows=1, validate_evidence_refs=False
    )
    assert "duplicate_source_row:СКС!R6" in errors


# --------------------------------------------------------------------------- 1.2


def _priced_supply_pair(**id_overrides):
    a = _operator_price(_supply_row("СКС!R9", "Кабель А", **id_overrides.get("a", {})), 10)
    b = _operator_price(_supply_row("СКС!R9", "Кабель Б", **id_overrides.get("b", {})), 20)
    return a, b


def test_trace_amounts_bind_by_operation_not_locator():
    a, b = _priced_supply_pair()
    trace = build_lsr_trace_from_visible_rows([a, b])
    summary = trace["summary"]
    assert summary["calculated_positions"] == 2
    assert summary["confirmed_amount"] == pytest.approx(30.0), (
        "the last position's amount was counted for both operations"
    )
    assert summary["preliminary_amount"] == pytest.approx(0.0)
    # sum invariant: confirmed + preliminary == calculated known amount
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )


def test_trace_amounts_bind_by_operation_id():
    a, b = _priced_supply_pair(
        a={"operation_id": "op_a", "revision": 1},
        b={"operation_id": "op_b", "revision": 1},
    )
    trace = build_lsr_trace_from_visible_rows([a, b])
    summary = trace["summary"]
    assert summary["confirmed_amount"] == pytest.approx(30.0)
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )


def test_trace_mixed_confirmed_and_unpriced_positions():
    a = _operator_price(_supply_row("СКС!R9", "Кабель А"), 10)
    b = _supply_row("СКС!R9", "Кабель Б")  # no price -> preliminary
    trace = build_lsr_trace_from_visible_rows([a, b])
    summary = trace["summary"]
    assert summary["confirmed_amount"] == pytest.approx(10.0)
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )


def test_trace_zero_price_is_not_a_missing_norm_counter():
    a = _operator_price(_supply_row("СКС!R9", "Кабель А"), 0)
    trace = build_lsr_trace_from_visible_rows([a])
    summary = trace["summary"]
    # a zero price is a provided price, but must never be mistaken for a total
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )


def test_bare_model_price_does_not_borrow_confirmation_from_a_sibling_operation():
    a = _supply_row("СКС!R9", "Кабель А", confirmed_price=10)
    b = _operator_price(_supply_row("СКС!R9", "Кабель Б"), 20)
    trace = build_lsr_trace_from_visible_rows([a, b])
    assert trace["summary"]["confirmed_amount"] == pytest.approx(20)
    assert trace["summary"]["preliminary_amount"] == pytest.approx(10)
    assert trace["summary"]["known_amount"] == pytest.approx(30)


# --------------------------------------------------------------------------- 1.3


def test_norms_selected_counts_real_ciphers_only():
    rows = [
        _work_row("СКС!R2", "Монтаж сети", norm_status="EXACT", norm_code="10-01-001-01"),
        _supply_row("СКС!R3", "Кабель", norm_code="-"),
        _supply_row("СКС!R4", "Розетка", norm_code="", norm_status="MISSING"),
        _work_row("СКС!R5", "Монтаж поста", norm_status="EXACT", norm_code="10-01-001-01"),
    ]
    payload = build_lsr_result_readiness(packaged=True, rows=rows)
    counts = payload["counts"]
    assert counts["norms_selected"] == 2, (
        "'-' and empty cells must not raise the selected-norm counter"
    )
    assert counts["norms_unique_ciphers"] == 1


# --------------------------------------------------------------------------- 1.4


class _Book:
    region = "Город Санкт-Петербург"
    quarter = "2 квартал 2026 г."


def test_price_basis_region_and_period_checked_together():
    result = summarize_lsr_price_basis(
        question_text="Собери ЛСР по Москве, цены 2 квартала 2026",
        pricebook=_Book(),
        pricebook_name="sankt-peterburg_2kv2026",
    )
    assert result.get("requested_region") == "Москва"
    assert result["actual_region"] == "Город Санкт-Петербург"
    assert result["match"] != "exact", (
        "a matching quarter must not certify a Saint Petersburg book for Moscow"
    )


def test_price_basis_region_is_explicit_context_parameter():
    result = summarize_lsr_price_basis(
        requested_region="Москва",
        requested_region_origin="параметр расчёта",
        question_text="цены 2 квартал 2026",
        pricebook=_Book(),
        pricebook_name="sankt-peterburg_2kv2026",
    )
    assert result["requested_region"] == "Москва"
    assert result["requested_region_origin"] == "параметр расчёта"
    assert result["match"] != "exact"


def test_price_basis_contradiction_between_sources_is_visible():
    result = summarize_lsr_price_basis(
        requested_book="sankt-peterburg_2kv2026",
        question_text="ЛСР, цены IV квартала 2026",
        pricebook=_Book(),
        pricebook_name="sankt-peterburg_2kv2026",
    )
    assert result["match"] != "exact", (
        "conflicting explicit parameters must not be resolved by hidden priority"
    )
    assert "противореч" in result["note"].lower()


def test_price_basis_operator_confirmation_of_alternative_book():
    result = summarize_lsr_price_basis(
        requested_region="Москва",
        question_text="Собери ЛСР по Москве, цены 2 квартала 2026",
        pricebook=_Book(),
        pricebook_name="sankt-peterburg_2kv2026",
        operator_confirmation={
            "pricebook": "sankt-peterburg_2kv2026",
            "confirmed_by": "operator",
        },
    )
    assert result["match"] == "confirmed_alternative"
    assert "подтверждена оператором" in result["note"]


def test_price_basis_unconfirmed_stays_mismatch():
    result = summarize_lsr_price_basis(
        requested_region="Москва",
        question_text="Собери ЛСР по Москве, цены 2 квартала 2026",
        pricebook=_Book(),
        pricebook_name="sankt-peterburg_2kv2026",
        operator_confirmation={"pricebook": "other_book", "confirmed_by": "operator"},
    )
    assert result["match"] == "mismatch"


# --------------------------------------------------------------------------- 1.5


def _double_row_pair():
    eq = _supply_row("СКС!R6", "Шкаф 42U", row_role="equipment", quantity=2,
                     operation_id="op_eq")
    mat = _supply_row("СКС!R6", "Шкаф 42U", quantity=2, operation_id="op_mat")
    return eq, mat


def test_confirmed_price_is_addressed_to_one_operation():
    eq, mat = _double_row_pair()
    patched = apply_confirmed_data(
        [eq, mat],
        {"СКС!R6": {"operation_id": "op_mat", "confirmed_price": 99.5}},
    )
    by_op = {item["operation_id"]: item for item in patched}
    assert by_op["op_mat"].get("confirmed_price") == 99.5
    assert "confirmed_price" not in by_op["op_eq"], (
        "an addressed confirmation must not price the other operation"
    )


def test_unaddressed_price_on_ambiguous_locator_is_flagged():
    eq, mat = _double_row_pair()
    patched = apply_confirmed_data([eq, mat], {"СКС!R6": {"confirmed_price": 99.5}})
    assert all(item.get("confirm_ambiguous") for item in patched), (
        "an unaddressed confirmation on a multi-row locator must stay visible"
    )
    assert all("confirmed_price" not in item for item in patched)


def test_single_row_locator_confirmation_still_applies():
    mat = _supply_row("СКС!R7", "Кабель", operation_id="op_c")
    patched = apply_confirmed_data([mat], {"СКС!R7": {"confirmed_price": 12.0}})
    assert patched[0]["confirmed_price"] == 12.0
    assert not patched[0].get("confirm_ambiguous")


def test_confirmed_price_prices_the_addressed_work_position(monkeypatch, tmp_path):
    """1.5 continuation: an operator-confirmed price must reach the specific
    calculated work position, addressed by operation_id, and be confirmed."""
    from proxy.services import gesn_service

    catalog = {
        "ГЭСНм08-03-575-01": {
            "code": "ГЭСНм08-03-575-01",
            "name": "Монтаж прибора или аппарата",
            "unit": "шт",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        }
    }
    monkeypatch.setattr(
        gesn_service, "get_norm",
        lambda code, *a, **k: catalog.get(str(code or "")),
    )
    work = _operator_price(
        _work_row("ВОР!R3", "Монтаж шкафа", operation_id="op_w", revision=1,
                  norm_status="EXACT", norm_code="ГЭСНм08-03-575-01", quantity=2),
        123.45,
    )
    trace = build_lsr_trace_from_visible_rows([work])
    positions = [
        pos for entry in trace.get("sections") or [] for pos in entry.get("positions") or []
    ]
    assert positions[0]["op_key"] == "op:op_w"
    totals = [r for r in positions[0]["rows"] if r.get("type") == "position_total"]
    assert totals and float(totals[0]["columns"]["12"]) == pytest.approx(246.90), (
        "the addressed confirmed price never reached the work position"
    )
    summary = trace["summary"]
    assert summary["confirmed_amount"] == pytest.approx(246.90)
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )
    # The rendered XLSX must carry the operator's flat cost as a literal value
    # in the "Всего по позиции" column, not a resource-derived SUM formula.
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
    from openpyxl import load_workbook

    out = render_lsr_xlsx(trace, tmp_path / "op_price.xlsx", title="t", meta={})
    wb = load_workbook(out, data_only=False)
    sheet = next(n for n in wb.sheetnames if "\u041b\u0421\u0420" in n)
    ws = wb[sheet]
    total_rows = [
        r for r in range(1, ws.max_row + 1)
        if str(ws.cell(row=r, column=3).value or "").startswith(
            "\u0412\u0441\u0435\u0433\u043e \u043f\u043e \u043f\u043e\u0437\u0438\u0446\u0438\u0438"
        )
    ]
    assert total_rows, "position total row missing in rendered file"
    value = ws.cell(row=total_rows[0], column=12).value
    assert isinstance(value, (int, float)) and abs(float(value) - 246.90) < 0.005, (
        f"renderer must write a literal for operator-confirmed total, got {value!r}"
    )


def test_double_count_conflict_is_detected_and_visible():
    from proxy.services.lsr_source_coverage_service import double_count_conflicts

    eq, mat = _double_row_pair()
    conflicts = double_count_conflicts([eq, mat])
    assert any(item["locator"] == "СКС!R6" for item in conflicts)

    eq_p = _operator_price(dict(eq, quantity=2), 99.5)
    mat_p = _operator_price(dict(mat, quantity=2), 99.5)
    trace = build_lsr_trace_from_visible_rows([eq_p, mat_p])
    flags = " ".join(trace["summary"].get("flags") or [])
    assert "двойн" in flags.lower(), "the double-accounting conflict must stay visible"
    # both rows remain (nothing deleted), but the amount is not silently confirmed
    assert trace["summary"]["confirmed_amount"] != pytest.approx(398.0)


# --------------------------------------------------------------------------- 1.6


def _tmp_job(tmp_path):
    service = WorkbookResearchJobService(tmp_path / "meta.db")
    identity = service.identity(
        session_id="s1",
        attachment_id="att1",
        source_sha256="0" * 64,
        dataset_ids=["d1"],
        profile_revision_id="p1",
        model_identity="m1",
        question="q",
    )
    job = service.begin_or_resume(identity)
    service.record_packet(
        job["job_id"],
        sheet="СКС",
        packet_locators=["СКС!R6"],
        accepted_rows=[_work_row("СКС!R6", "Установка шкафа")],
        evidence=[],
        remaining=[],
    )
    return service, service.get(job["job_id"])


def test_stale_guard_works_through_public_api(tmp_path, monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    service, job = _tmp_job(tmp_path)
    stored = job["accepted_rows"][0]
    op_id = str(stored["operation_id"])
    revision = int(stored.get("revision") or 1)

    from proxy.routers import chat as chat_router
    from proxy.security import require_user

    monkeypatch.setattr(chat_router, "WorkbookResearchJobService", lambda: service)
    app = FastAPI()
    app.include_router(chat_router.router)
    app.dependency_overrides[require_user] = lambda: {"role": "user"}
    client = TestClient(app)

    stale = dict(stored, norm_code="10-01-001-01", norm_status="EXACT",
                 base_revision=revision + 5)
    late = client.post(
        "/api/chat/patch-lsr-rows",
        json={
            "session_id": "s1",
            "job_id": job["job_id"],
            "replace_rows": [stale],
            "repack": False,
        },
    )
    assert late.status_code == 409, (
        "a late refinement must not overwrite a newer revision via the public API"
    )
    current = service.get(job["job_id"])["accepted_rows"][0]
    assert current["operation_id"] == op_id
    assert int(current["revision"]) == revision

    fresh = dict(stored, norm_code="10-01-001-01", norm_status="EXACT",
                 base_revision=revision)
    ok = client.post(
        "/api/chat/patch-lsr-rows",
        json={
            "session_id": "s1",
            "job_id": job["job_id"],
            "replace_rows": [fresh],
            "repack": False,
        },
    )
    assert ok.status_code == 200
    updated = service.get(job["job_id"])["accepted_rows"][0]
    assert int(updated["revision"]) == revision + 1
    assert updated["operation_id"] == op_id


def test_packet_attempts_survive_a_process_restart(tmp_path):
    service, job = _tmp_job(tmp_path)
    service.record_packet(
        job["job_id"],
        sheet="СКС",
        packet_locators=["СКС!R6"],
        accepted_rows=[],
        evidence=[],
        remaining=["СКС!R6"],
    )
    events = service.list_events(job["job_id"], event_type="packet_checkpoint")
    from proxy.services.chat_evidence_application_service import (
        packet_attempt_seed,
    )

    seed = packet_attempt_seed(events)
    assert seed.get(("СКС!R6",)) == 2, (
        "attempt counters must be rebuilt from durable checkpoints, not from zero"
    )
    # A restarted process therefore namespaces evidence as :a3, never :a1 again.
    fresh_service = WorkbookResearchJobService(tmp_path / "meta.db")
    replayed = fresh_service.list_events(job["job_id"], event_type="packet_checkpoint")
    assert packet_attempt_seed(replayed).get(("СКС!R6",)) == 2


# ------------------------------------------------------- 2.2 / 2.3 readiness bases + summary
# Directive 2.2: "Confirmation of a price does not confirm norm applicability."
# A position with norm_status NEEDS_RESEARCH or applicability_substantiated=false
# must remain preliminary even after the operator supplies a price.
# Directive 2.3: the resource-level rows of an operator-replaced position must
# not inflate the estimate's aggregate direct / fot / nr / sp.


def test_confirmed_price_does_not_confirm_norm_applicability(monkeypatch):
    from proxy.services import gesn_service

    catalog = {
        "\u0413\u042d\u0421\u041d\u043c08-03-575-01": {
            "code": "\u0413\u042d\u0421\u041d\u043c08-03-575-01",
            "name": "Mount a device",
            "unit": "\u0448\u0442",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        }
    }
    monkeypatch.setattr(
        gesn_service, "get_norm",
        lambda code, *a, **k: catalog.get(str(code or "")),
    )
    work = _operator_price(
        _work_row(
            "\u0412\u041e\u0420!R5", "Mount a device",
            operation_id="op_needs_research",
            norm_status="NEEDS_RESEARCH",
            norm_code="\u0413\u042d\u0421\u041d\u043c08-03-575-01",
            quantity=1,
        ),
        246.90,
    )
    trace = build_lsr_trace_from_visible_rows([work])
    summary = trace["summary"]
    # The operator set a price; that alone must NOT certify the position.
    assert summary["confirmed_amount"] == pytest.approx(0.0), (
        "a NEEDS_RESEARCH position was counted as confirmed after a price"
    )
    assert summary["preliminary_amount"] == pytest.approx(246.90), (
        "the price-only position should have stayed preliminary"
    )
    assert summary["confirmed_amount"] + summary["preliminary_amount"] == pytest.approx(
        float(summary["known_amount"])
    )


def test_norm_reference_resources_do_not_inflate_summary(monkeypatch):
    """2.3 mechanism: an operator-replaced position is flagged reference_only
    so the summary loop skips its norm-derived resource aggregates."""
    from proxy.services import gesn_service

    catalog = {
        "\u0413\u042d\u0421\u041d\u043c08-03-575-01": {
            "code": "\u0413\u042d\u0421\u041d\u043c08-03-575-01",
            "name": "Mount a device",
            "unit": "\u0448\u0442",
            "resources": [],
            "nr_pct": 0,
            "sp_pct": 0,
        }
    }
    monkeypatch.setattr(
        gesn_service, "get_norm",
        lambda code, *a, **k: catalog.get(str(code or "")),
    )
    work = _operator_price(
        _work_row(
            "\u0412\u041e\u0420!R6", "Mount a device",
            operation_id="op_replaced",
            revision=1,
            norm_status="EXACT",
            norm_code="\u0413\u042d\u0421\u041d\u043c08-03-575-01",
            quantity=2,
        ),
        99.99,
    )
    trace = build_lsr_trace_from_visible_rows([work])
    positions = [
        pos for entry in trace.get("sections") or [] for pos in entry.get("positions") or []
    ]
    assert positions, "position trace missing"
    assert positions[0]["summary"].get("resource_reference_only") is True, (
        "an operator-replaced position must be flagged reference_only so the "
        "summary loop can exclude its norm-derived resource aggregates"
    )
    summary = trace["summary"]
    # 2 x 99.99 = 199.98 literal, carried into the estimate total.
    assert summary["total"] == pytest.approx(199.98), (
        f"operator-replaced position total must equal qty x price, got {summary['total']}"
    )
    assert summary.get("reference_only_positions", 0) >= 1, (
        "the reference_only counter should have been incremented"
    )

# ------------------------------------------------------- 2.1 note / evidence / status
# Directive 2.1: "material 10 x 100 after reprocessing yields 2 000 instead of 1 000,
# without conflict." A note / evidence_refs / status change is not a new paid operation.


def test_note_change_does_not_create_a_new_operation_without_id():
    original = _supply_row("AR!R30", "Carpet tile Atlas 05", quantity=10)
    original["unit_price"] = 100
    original.pop("operation_id", None)
    redelivered = dict(original)
    redelivered["note"] = "from specification, clarify colour"
    merged = merge_model_rag_rows([dict(original)], [redelivered])
    assert len(merged) == 1, (
        "a note change without operation_id produced a duplicate line"
    )
    assert merged[0]["note"] == "from specification, clarify colour"


def test_evidence_refs_change_does_not_create_a_new_operation():
    original = _work_row("SKS!R6", "Mount A")
    original["evidence_refs"] = ["chunk_a"]
    original.pop("operation_id", None)
    redelivered = dict(original)
    redelivered["evidence_refs"] = ["chunk_b"]
    merged = merge_model_rag_rows([dict(original)], [redelivered])
    assert len(merged) == 1, "evidence_refs drift must not create a new operation"


def test_status_change_does_not_create_a_new_operation():
    original = _work_row("SKS!R6", "Mount A")
    original["status"] = "proposed"
    original.pop("operation_id", None)
    redelivered = dict(original)
    redelivered["status"] = "confirmed"
    merged = merge_model_rag_rows([dict(original)], [redelivered])
    assert len(merged) == 1, "status transition must not create a new operation"


def test_two_genuinely_different_operations_still_survive_note_fix():
    """2.1 boundary: the volatile-field exclusion must NOT fuse distinct operations."""
    a = _work_row("SKS!R6", "Mount cabinet")
    b = _work_row("SKS!R6", "Mount server")
    a.pop("operation_id", None)
    b.pop("operation_id", None)
    merged = merge_model_rag_rows([], [a, b])
    assert len(merged) == 2, (
        "different titles on one locator are different decisions; volatile-field "
        "exclusion must not fuse them"
    )


def test_two_materials_on_same_locator_are_two_operations():
    """2.1 boundary: protection must work INSIDE a role as well as across roles."""
    mat1 = _supply_row("AR!R30", "Carpet tile", quantity=10)
    mat1.pop("operation_id", None)
    mat2 = _supply_row("AR!R30", "Adhesive", quantity=5)
    mat2.pop("operation_id", None)
    merged = merge_model_rag_rows([], [mat1, mat2])
    assert len(merged) == 2


def test_price_basis_confirmation_is_durable(tmp_path):
    service, job = _tmp_job(tmp_path)
    updated = service.record_overrides(
        job["job_id"],
        {"_price_basis": {"pricebook": "sankt-peterburg_2kv2026", "note": "иначе нечего"}},
    )
    stored = updated["operator_overrides"].get("_price_basis") or {}
    assert stored.get("pricebook") == "sankt-peterburg_2kv2026"
