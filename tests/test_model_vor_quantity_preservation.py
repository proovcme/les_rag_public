"""VOR must retain authored derived work as an audited draft, without accepting it."""
from __future__ import annotations

import copy
from io import BytesIO
import json

import openpyxl
import pytest

from proxy.services import chat_evidence_application_service as service


def _catalog(tmp_path):
    from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog

    source = tmp_path / "Generic packaged input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active.append(["Наименование", "Ед. изм.", "Количество"])
    workbook.active.append(["Bundle with 24 m per pack", "pack", 8])
    workbook.save(source)
    workbook.close()
    return load_tabular_source_catalog(source)


@pytest.mark.parametrize("artifact_kind", ["", "lsr", "unresolved", "VOR"])
def test_legacy_lsr_quantity_guard_is_unchanged(tmp_path, artifact_kind):
    catalog = _catalog(tmp_path)
    row = {"source_row": "S!R2", "row_role": "work", "title": "Literal work",
           "quantity": 192, "unit": "m", "evidence_refs": [], "norm_code": ""}
    errors = ["source_quantity_mismatch:S!R2", "source_unit_mismatch:S!R2", "missing_evidence_ref:S!R2"]
    assert service.sheet_pass_keep_rows([row], sheet="S", catalog=catalog, errors=errors,
                                        artifact_kind=artifact_kind) == []
    assert service.blocking_model_rag_result_errors(errors, work_row_count=1,
                                                    artifact_kind=artifact_kind) == [errors[0]]


@pytest.mark.parametrize("guard", ["duplicate", "unknown_identity", "wrong_identity", "tail", "foreign_source"])
def test_vor_preservation_retains_exact_existing_ownership_and_tail_guards(tmp_path, guard):
    catalog = _catalog(tmp_path)
    row = {"source_row": "S!R2", "row_role": "work", "title": "Literal proposal",
           "quantity": 192, "unit": "m", "evidence_refs": ["Q777.H2"], "norm_code": "Literal norm"}
    errors, drop = ["source_quantity_mismatch:S!R2"], ""
    if guard == "duplicate":
        errors.append("duplicate_source_row:S!R2")
    elif guard in {"unknown_identity", "wrong_identity"}:
        row["operation_id"] = "authored-operation"
        prefix = "unknown_operation_id" if guard == "unknown_identity" else "operation_source_mismatch"
        errors.append(f"{prefix}:S!R2:authored-operation")
    elif guard == "tail":
        drop = "S!R2"
    else:
        row["source_row"] = "FOREIGN!R2"
    before = copy.deepcopy(row)
    assert service.sheet_pass_keep_rows([row], sheet="S", catalog=catalog, errors=errors,
                                        drop_locator=drop, artifact_kind="vor") == []
    assert row == before


def test_vor_source_derivation_is_never_inferred_or_promoted_by_preservation(tmp_path):
    from proxy.services.attachment_source_catalog_service import compare_model_row_to_catalog

    catalog = _catalog(tmp_path)
    row = {"source_row": "S!R2", "row_role": "work", "title": "Literal proposal",
           "quantity": 191, "unit": "m", "norm_code": "", "evidence_refs": [],
           "conversion_factor": 24, "calc_quantity": 192,
           "conversion_source": "S!R2 title: 24 m per pack", "quantity_formula": "8 * 24"}
    before = copy.deepcopy(row)
    errors = compare_model_row_to_catalog(row, catalog)
    assert errors == ["source_quantity_mismatch:S!R2", "source_unit_mismatch:S!R2"]
    assert service.sheet_pass_keep_rows([row], sheet="S", catalog=catalog, errors=errors,
                                        artifact_kind="vor") == [before]
    assert row == before
    readiness = service.build_vor_result_readiness(rows=[row], has_file=True, remaining_locators=[],
                                                    catalog_work_count=1, model_errors=errors)
    assert readiness["vor_complete"] is False and readiness["technical_complete"] is False
    assert readiness["open_questions"] == errors


def test_optional_vor_norm_citation_checks_remain_visible_without_requiring_norms():
    rows = [{"source_row": "S!R2", "row_role": "work", "norm_code": "Authored norm"},
            {"source_row": "S!R3", "row_role": "work", "norm_code": ""}]
    errors = ["missing_evidence_ref:S!R2", "missing_evidence_ref:S!R3",
              "unknown_evidence_ref:Q777.H2", "norm_code_not_in_referenced_evidence:S!R2"]
    assert service.vor_model_result_audit_errors(rows, errors) == [errors[0], *errors[2:]]
    readiness = service.build_vor_result_readiness(rows=rows, has_file=True, remaining_locators=[],
                                                    catalog_work_count=2, model_errors=errors)
    assert readiness["technical_complete"] is False and readiness["vor_complete"] is False
    plain = service.build_vor_result_readiness(rows=rows[1:], has_file=True, remaining_locators=[],
                                                catalog_work_count=1, model_errors=[errors[1]])
    assert plain["technical_complete"] is True and plain["vor_complete"] is True


@pytest.mark.parametrize("variant", ["B0_DIRECT", "B1_CANDIDATE"])
@pytest.mark.asyncio
async def test_actual_vor_keeps_derived_work_quantity_and_optional_bad_norm_ref(
    tmp_path, monkeypatch, variant,
):
    from tests import test_model_rag_sheet_evidence_reads as fixture
    from backend.runtime_paths import mutable_path
    from proxy.services.artifact_revision_service import ArtifactRevisionStore
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService

    original_preserve = fixture.preserve_read_attachment

    def preserve_packaged_source(source, **kwargs):
        workbook = openpyxl.load_workbook(source)
        sheet = workbook.active
        for row, values in ((2, ["Bundle with 24 m per pack", "pack", 8]),
                            (3, ["Manufactured lengths of 3 m", "piece", 7])):
            for column, value in enumerate(values, 1):
                sheet.cell(row, column, value)
        workbook.save(source)
        workbook.close()
        return original_preserve(source, **kwargs)

    authored = [
        {"source_row": "S!R2", "row_role": "work", "title": "Authored laying operation",
         "unit": "m", "quantity": 192, "norm_code": "AUTHORED-UNVERIFIED-NORM",
         "norm_status": "ANALOGUE_CANDIDATE", "evidence_refs": ["Q999.H1"],
         "coverage": "8 packs x 24 m = 192 m; installed extent is unverified", "open_gap": ""},
        {"source_row": "S!R2", "row_role": "material", "title": "Authored bundle supply",
         "unit": "pack", "quantity": 8, "norm_code": "", "evidence_refs": [], "work_row": "S!R2"},
        {"source_row": "S!R3", "row_role": "work", "title": "Authored section installation",
         "unit": "m", "quantity": 21, "norm_code": "", "evidence_refs": [],
         "coverage": "7 sections x 3 m = 21 m", "open_gap": ""},
        {"source_row": "S!R3", "row_role": "material", "title": "Authored section supply",
         "unit": "piece", "quantity": 7, "norm_code": "", "evidence_refs": [], "work_row": "S!R3"},
    ]
    before = copy.deepcopy(authored)
    monkeypatch.setattr(fixture, "preserve_read_attachment", preserve_packaged_source)
    monkeypatch.setattr(fixture, "_answer", lambda **kwargs: json.dumps(
        {"artifact_kind": "vor", "rows": authored}, ensure_ascii=False))
    run = await fixture._run_sheet(tmp_path, monkeypatch, reads=False, variant=variant,
                                   query_declaration='{"artifact_kind":"vor","queries":[]}')
    trace = run.result["retrieval_trace"]
    assert len(run.requests) == 2 and run.queries == [] and len(run.executions) == 1
    call, result = run.executions[0]
    assert call["tool"] == "build_vor_workbook"
    exported = call["args"]["decisions"]
    assert len(exported) == 4, "the source audit must not delete model-authored installation work"
    for original, actual in zip(authored, exported, strict=True):
        for field, value in original.items():
            assert actual[field] == value
    assert authored == before
    codes = {issue["code"] for issue in result["document_issues"]}
    assert {"source_quantity_mismatch:S!R2", "source_quantity_mismatch:S!R3",
            "source_unit_mismatch:S!R2", "source_unit_mismatch:S!R3",
            "unknown_evidence_ref:Q999.H1"}.issubset(codes)
    assert "missing_evidence_ref:S!R3" not in codes, "a VOR work without a norm needs no normative citation"
    assert result["status"] == "partial"
    readiness = trace["result_readiness"]
    assert readiness["technical_complete"] is False and readiness["vor_complete"] is False
    assert "unknown_evidence_ref:Q999.H1" in readiness["open_questions"]
    assert "source_quantity_mismatch:S!R2" in readiness["open_questions"]

    store = ArtifactRevisionStore(mutable_path("storage/artifacts/meta.db"),
                                  mutable_path("storage/artifacts/files"))
    raw = store.read_bytes(result["artifact"]["revision_id"])
    workbook = openpyxl.load_workbook(BytesIO(raw), data_only=True)
    sheet = workbook["ВОР"]
    assert "НЕ ЗАВЕРШЁН" in sheet["A2"].value
    assert [sheet.cell(row, 3).value for row in range(4, 8)] == [row["title"] for row in authored]
    assert [sheet.cell(row, 5).value for row in range(4, 8)] == ["m", "pack", "m", "piece"]
    assert [sheet.cell(row, 6).value for row in range(4, 8)] == [192, 8, 21, 7]
    assert sheet["I4"].value == "ТРЕБУЕТ РЕШЕНИЯ" and sheet["I6"].value == "ТРЕБУЕТ РЕШЕНИЯ"
    assert "исходной строкой" in sheet["L4"].value and "исходной строкой" in sheet["L6"].value
    checks = "\n".join(str(cell.value or "") for row in workbook["Проверка"] for cell in row)
    assert "source_quantity_mismatch:S!R2" in checks and "unknown_evidence_ref:Q999.H1" in checks
    workbook.close()
    if variant == "B1_CANDIDATE":
        from types import SimpleNamespace
        from proxy.routers import chat

        job = WorkbookResearchJobService().get(trace["workbook_research_job"]["job_id"])
        assert job["remaining"] == []
        assert len(job["accepted_rows"]) == 4
        accepted_before = copy.deepcopy(job["accepted_rows"])
        resumed = await run.resume_same_request()
        assert len(run.requests) == 2, "resume must not invent another professional model pass"
        assert resumed["retrieval_trace"]["result_readiness"]["vor_complete"] is False
        assert WorkbookResearchJobService().get(job["job_id"])["accepted_rows"] == accepted_before
        repacked = await chat.repack_lsr_without_model(chat.LsrRepackRequest(
            session_id="sheet-session", job_id=job["job_id"]), _user=SimpleNamespace())
        assert repacked["model_called"] is False and len(run.requests) == 2
        repack_codes = {issue["code"] for issue in repacked["result"]["document_issues"]}
        assert codes.issubset(repack_codes), "the saved-workbook path must retain its application-owned audits"
        repack_ready = repacked["harvested"]["artifact"]["result_readiness"]
        assert repack_ready["vor_complete"] is False
        assert "unknown_evidence_ref:Q999.H1" in repack_ready["open_questions"]
        repack_raw = store.read_bytes(repacked["result"]["artifact"]["revision_id"])
        checkbook = openpyxl.load_workbook(BytesIO(repack_raw), data_only=True)
        assert [checkbook["ВОР"].cell(row, 6).value for row in range(4, 8)] == [192, 8, 21, 7]
        recheck = "\n".join(str(cell.value or "") for row in checkbook["Проверка"] for cell in row)
        assert "unknown_evidence_ref:Q999.H1" in recheck
        checkbook.close()
        replay = await chat.repack_lsr_without_model(chat.LsrRepackRequest(
            session_id="sheet-session", job_id=job["job_id"]), _user=SimpleNamespace())
        assert replay["model_called"] is False and len(run.requests) == 2
        assert replay["result"]["artifact"]["revision_id"] == repacked["result"]["artifact"]["revision_id"]
        assert store.read_bytes(replay["result"]["artifact"]["revision_id"]) == repack_raw
        assert replay["result"]["document_issues"] == repacked["result"]["document_issues"]
        assert WorkbookResearchJobService().get(job["job_id"])["accepted_rows"] == accepted_before
