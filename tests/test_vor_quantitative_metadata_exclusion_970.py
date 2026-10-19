"""Unexplained quantitative-source exclusion is audited without choosing work roles."""
from copy import deepcopy
from dataclasses import replace
import json

import openpyxl
import pytest

from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog
from proxy.services.chat_attachment_service import preserve_read_attachment
from proxy.services.chat_evidence_application_service import build_vor_result_readiness
from proxy.services.vor_result_integrity_service import audit_vor_result_integrity
from proxy.services.workbook_tool_service import build_vor_workbook
from test_workbook_tool_service import _context


CODE = "unexplained_quantitative_source_exclusion"


def _source(tmp_path):
    path = tmp_path / "Measured source with spaces.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Literal source item", "м²", 17.25])
    sheet.append(["Literal other item", "шт", 2])
    workbook.save(path)
    workbook.close()
    return path, load_tabular_source_catalog(path)


def _rows(role="header", **extra):
    return [
        {"source_row": "S!R2", "operation_id": "literal-metadata", "row_role": role,
         "title": "Literal authored metadata", "quantity": None, "unit": "", **extra},
        {"source_row": "S!R3", "operation_id": "literal-other", "row_role": "work",
         "title": "Literal modeled work", "quantity": 2, "unit": "шт"},
    ]


def _specific_issues(rows, catalog):
    return [issue for issue in audit_vor_result_integrity(rows, source_catalog=catalog)
            if issue["code"] == CODE]


@pytest.mark.parametrize("role", ["header", "note", "subtotal"])
@pytest.mark.parametrize("quantity", [17.25, 0, -2])
def test_metadata_exclusion_of_a_literal_measured_source_requires_only_an_explanation(tmp_path, role, quantity):
    _, catalog = _source(tmp_path)
    catalog = replace(catalog, rows=tuple(replace(row, header_quantity=quantity)
        if row.locator == "S!R2" else row for row in catalog.rows))
    decisions = _rows(role)
    before = deepcopy(decisions)
    issue, = _specific_issues(decisions, catalog)
    assert issue["source_row"] == issue["source_locator"] == "S!R2"
    assert issue["source_value"] == quantity and issue["source_unit"] == "м²"
    assert issue["row_index"] == 0 and issue["operation_id"] == "literal-metadata"
    assert "Поясните исключение" in issue["reason"]
    assert decisions == before


@pytest.mark.parametrize("boundary", ["source_header", "source_subtotal", "source_section", "blank_unit",
    "unknown_quantity", "unbound_quantity", "boolean", "nan", "inf", "incomplete_catalog"])
def test_existing_metadata_and_unconfirmed_measurements_do_not_gain_a_blind_warning(tmp_path, boundary):
    _, catalog = _source(tmp_path)
    source = catalog.lookup("S!R2")
    if boundary.startswith("source_"):
        source = replace(source, row_classification=boundary.removeprefix("source_"))
    elif boundary == "blank_unit":
        source = replace(source, cells=tuple(replace(cell, raw_value="", display_value="")
            if cell.column == 2 else cell for cell in source.cells))
    elif boundary == "unbound_quantity":
        source = replace(source, quantity_column_known=False)
    elif boundary != "incomplete_catalog":
        value = {"unknown_quantity": None, "boolean": False, "nan": float("nan"), "inf": float("inf")}[boundary]
        source = replace(source, header_quantity=value)
    catalog = replace(catalog, rows=tuple(source if row.locator == "S!R2" else row for row in catalog.rows),
                      extraction_complete=boundary != "incomplete_catalog")
    decisions = _rows()
    before = deepcopy(decisions)
    assert _specific_issues(decisions, catalog) == []
    assert decisions == before


def test_a_real_column_header_with_blank_unit_and_no_quantity_remains_a_header(tmp_path):
    _, catalog = _source(tmp_path)
    decisions = _rows()
    decisions[0]["source_row"] = "S!R1"
    before = deepcopy(decisions)
    assert _specific_issues(decisions, catalog) == []
    assert decisions == before


@pytest.mark.parametrize("explanation", [{"coverage": "Literal explained exclusion"},
    {"open_gap": "Engineer should decide whether this scope belongs here"}])
def test_a_literal_explanation_is_preserved_without_being_professionally_approved(tmp_path, explanation):
    _, catalog = _source(tmp_path)
    decisions = _rows(**explanation)
    before = deepcopy(decisions)
    assert _specific_issues(decisions, catalog) == []
    assert decisions == before


@pytest.mark.parametrize("explanation", [{"coverage": "-"}, {"coverage": "—"}, {"open_gap": "  "}, {"coverage": True}])
def test_placeholder_or_boolean_is_not_an_explanation(tmp_path, explanation):
    _, catalog = _source(tmp_path)
    decisions = _rows(**explanation)
    before = deepcopy(decisions)
    assert len(_specific_issues(decisions, catalog)) == 1
    assert decisions == before


@pytest.mark.parametrize("sibling_role", ["work", "material", "unresolved"])
def test_an_existing_content_operation_at_the_exact_source_prevents_the_all_metadata_warning(tmp_path, sibling_role):
    _, catalog = _source(tmp_path)
    decisions = _rows()
    decisions.append({"source_row": "S!R2", "operation_id": "literal-sibling", "row_role": sibling_role,
        "title": "Literal source disposition", "quantity": 17.25, "unit": "м²",
        "work_row": "literal-other", "open_gap": "Literal open disposition" if sibling_role == "unresolved" else ""})
    before = deepcopy(decisions)
    assert _specific_issues(decisions, catalog) == []
    assert decisions == before


def test_duplicate_metadata_aliases_are_one_question_for_one_exact_source(tmp_path):
    _, catalog = _source(tmp_path)
    decisions = _rows()
    decisions.insert(1, {**decisions[0], "source_row": catalog.source_file_id + "!R2",
        "operation_id": "literal-alias", "title": "Literal second metadata"})
    before = deepcopy(decisions)
    issue, = _specific_issues(decisions, catalog)
    assert issue["source_locator"] == "S!R2"
    assert decisions == before


def _assert_visible_exclusion_question(ready):
    # Pure presentation assertion, separate from invoking the workbook builder.
    # Keeping these apart also preserves the architecture guard against a real
    # language matcher dispatching a workbook tool.
    assert any("Поясните исключение" in question for question in ready["open_questions"])


@pytest.mark.asyncio
@pytest.mark.parametrize("role", ["header", "note"])
async def test_real_vor_public_builder_xlsx_and_readiness_show_the_question_without_adding_work(tmp_path, role):
    source, catalog = _source(tmp_path)
    attachment = preserve_read_attachment(source, attachment_id="read_123456abcdef",
        original_name=source.name, root=tmp_path / "attachments")
    decisions = _rows(role)
    before = deepcopy(decisions)
    context = _context(tmp_path)
    result = await build_vor_workbook({"attachment_id": attachment["attachment_id"], "decisions": decisions}, context)
    assert result["status"] == "partial" and result["source"]["rows"] == len(decisions)
    issue, = [issue for issue in result["document_issues"] if issue["code"] == CODE]
    assert issue["source_locator"] == "S!R2" and issue["source_value"] == 17.25
    workbook = openpyxl.load_workbook(context.artifacts.resolve_path(result["artifact"]["revision_id"]), read_only=True)
    try:
        authored = [json.loads(values[2].removeprefix("JSON: "))
            for values in workbook["Решения модели"].iter_rows(min_row=4, values_only=True)]
        assert authored == before
        assert workbook["ВОР"]["F4"].value is None and workbook["ВОР"]["E4"].value is None
        assert "Поясните исключение" in workbook["ВОР"]["L4"].value
        assert any(values[0] == CODE for values in workbook["Проверка"].iter_rows(min_row=4, values_only=True))
        printed = [values[1] for values in workbook.active.iter_rows(min_row=9, values_only=True)
                   if isinstance(values[0], int) and not isinstance(values[0], bool)]
        assert printed == [before[1]["title"]]
    finally:
        workbook.close()
    ready = build_vor_result_readiness(rows=decisions, has_file=True, remaining_locators=[],
        catalog_work_count=len(catalog.work_locators()), model_errors=[], source_catalog=catalog)
    assert ready["technical_complete"] is True and ready["vor_complete"] is False
    assert len([issue for issue in ready["row_integrity_issues"] if issue["code"] == CODE]) == 1
    _assert_visible_exclusion_question(ready)
    assert decisions == before
