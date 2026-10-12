"""Actual intake/export regressions; no inference or live services are involved."""
from __future__ import annotations

import copy
import json
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import pytest

from proxy.services import attachment_source_catalog_service as catalog_service
from proxy.services import workbook_tool_service as workbook_service
from proxy.services.artifact_revision_service import ArtifactRevisionStore
from proxy.services.chat_attachment_service import preserve_read_attachment
from proxy.services.workflow_checkpoint_service import WorkflowCheckpointService


def _source_workbook(path: Path, quantity=12, *, cached_quantity=None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Исходный предмет", "м", quantity])
    notes = workbook.create_sheet("Условия")
    notes.append(["Дословные условия исходного документа"])
    workbook.save(path)
    workbook.close()
    if cached_quantity is not None:
        with zipfile.ZipFile(path) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        sheet_xml = ET.fromstring(parts["xl/worksheets/sheet1.xml"])
        value = sheet_xml.find(".//s:c[@r='C2']/s:v", namespace)
        assert value is not None
        value.text = str(cached_quantity)
        parts["xl/worksheets/sheet1.xml"] = ET.tostring(sheet_xml, encoding="utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, data in parts.items():
                archive.writestr(name, data)
    return path


@pytest.mark.parametrize("quantity,cached,expected", [
    (None, None, None), ("=6*2", None, None), ("=6*2", 12, 12),
])
def test_subject_without_literal_number_remains_in_required_source_scope(
    tmp_path, quantity, cached, expected,
):
    source = _source_workbook(tmp_path / "Источник 'Ω'.xlsx", quantity, cached_quantity=cached)
    catalog = catalog_service.load_tabular_source_catalog(source)
    row = catalog.lookup("S!R2")

    assert row.row_classification == "ambiguous"
    assert row.is_work_row()
    assert row.quantity_column_known and row.header_quantity == expected
    assert catalog.work_locators() == ["S!R2"]
    assert catalog_service.catalog_sheet_packet_facts(catalog, "S")["requested_source_rows"] == ["S!R2"]
    assert catalog_service.coverage_errors([], catalog) == ["source_coverage_incomplete:0/1"]
    assert catalog.lookup("Условия!R1").row_classification == "note"
    quantity_cell = next((cell for cell in row.cells if cell.column == 3), None)
    if quantity is None:
        assert quantity_cell is None
    else:
        assert quantity_cell.raw_value == quantity_cell.formula == quantity
        assert quantity_cell.cached_formula_value == cached
    assert catalog_service.quantity_candidates(row) == (() if expected is None else (expected,))


def test_pdf_escaped_pipes_keep_physical_columns_and_source_quantity(tmp_path, monkeypatch):
    source = tmp_path / "Literal table.pdf"
    source.write_bytes(b"%PDF-fixture")
    markdown = (
        "| Наименование | Ед. изм. | Количество |\n|---|---|---|\n"
        "| Cabinet \\| variant | м | 12 |\n"
    )
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = catalog_service.load_tabular_source_catalog(source)
    row = catalog.lookup("PDF!R2")

    assert row.values == ("Cabinet | variant", "м", "12")
    assert catalog_service.header_bound_unit(row) == "м"
    assert row.header_quantity == 12
    assert catalog_service.compare_model_row_to_catalog(
        {"source_row": "PDF!R2", "quantity": 12, "unit": "м"}, catalog,
    ) == []
    assert markdown[row.source_start:row.source_end].strip() == "| Cabinet \\| variant | м | 12 |"


@pytest.mark.parametrize("literal,expected", [
    (r"| Cabinet \\| м | 12 |", "Cabinet \\"),
    (r"| Cabinet \\\| variant | м | 12 |", r"Cabinet \| variant"),
])
def test_markdown_escaped_backslash_keeps_the_next_pipe_boundary(literal, expected):
    assert catalog_service._markdown_table_rows(literal) == [(expected, "м", "12")]


@pytest.mark.parametrize("title", ["Решения модели", "Проверка"])
@pytest.mark.parametrize("partial_marker", [False, True])
def test_audit_sheet_name_or_partial_marker_never_excludes_arbitrary_work(
    tmp_path, title, partial_marker,
):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = title
    if partial_marker:
        sheet.append([catalog_service.VOR_MODEL_AUDIT_SCHEMA])
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Произвольный исходный предмет", "м", 12])
    source_row = sheet.max_row
    source = tmp_path / "Произвольный исходник.xlsx"
    workbook.save(source)
    workbook.close()

    catalog = catalog_service.load_tabular_source_catalog(source)

    assert catalog.work_locators() == [f"{title}!R{source_row}"]
    assert catalog_service.coverage_errors([], catalog) == ["source_coverage_incomplete:0/1"]


def _context(tmp_path, adapter):
    return workbook_service.WorkbookExecutionContext(
        session_id="intake-928", idempotency_key="intake-928", model_decision_revision="authored",
        profile_revision_id="profile", model_identity="fixture", model_preset="fixture",
        attachment_root=tmp_path / "attachments", work_dir=tmp_path / "work",
        checkpoints=WorkflowCheckpointService(tmp_path / "checkpoints.db"),
        artifacts=ArtifactRevisionStore(tmp_path / "artifacts.db", tmp_path / "artifacts"),
        lsr_adapter=adapter,
    )


@pytest.mark.asyncio
@pytest.mark.parametrize("suffix", [".docx", ".csv", ".rtf"])
async def test_supported_literal_sources_reach_authored_lsr_packaging(
    tmp_path, monkeypatch, suffix,
):
    source = tmp_path / f"Источник 'Ω'{suffix}"
    if suffix == ".docx":
        from docx import Document
        document = Document()
        document.add_paragraph("Буквальный исходник")
        document.save(source)
    elif suffix == ".csv":
        source.write_text("Наименование;Ед. изм.;Количество\nПредмет;м;12\n", encoding="utf-8")
    else:
        source.write_bytes(b"{\\rtf1\\ansi Literal source}")
    metadata = preserve_read_attachment(
        source, attachment_id="read_123456abcdef", original_name=source.name,
        root=tmp_path / "attachments",
    )
    decisions = [{"source_row": "SOURCE!R2", "title": "Авторская операция", "quantity": 12, "unit": "м",
                  "resource_bindings": [{"action": "reuse", "reason": "Literal authored reason"}]}]
    before = copy.deepcopy(decisions)
    calls = []

    def adapter(path, args, output, _progress):
        calls.append((path.suffix, copy.deepcopy(args["decisions"])))
        workbook = openpyxl.Workbook()
        workbook.active.append([args["decisions"][0]["title"], args["decisions"][0]["quantity"]])
        workbook.save(output)
        workbook.close()
        return {"file_path": output, "source_rows": len(args["decisions"]), "missing": [], "blockers": []}

    context = _context(tmp_path, adapter)
    monkeypatch.setattr(workbook_service, "packaging_fault", lambda: None)
    result = await workbook_service.build_lsr_workbook(
        {"attachment_id": metadata["attachment_id"], "decisions": decisions}, context,
    )

    assert result["status"] == "complete"
    assert calls == [(suffix, before)]
    assert decisions == before
    assert result["source"]["sha256"] == metadata["sha256"]
    assert context.artifacts.resolve_path(result["artifact"]["revision_id"]).is_file()


@pytest.mark.asyncio
@pytest.mark.parametrize("long_metadata", [False, True])
async def test_vor_physical_roundtrip_preserves_complete_authored_operation_metadata(
    tmp_path, monkeypatch, long_metadata,
):
    source = _source_workbook(tmp_path / "Исходник.xlsx")
    decisions = [
        {"source_row": "S!R2", "operation_id": "op-parent", "revision": 2,
         "row_role": "work", "work_row": "S!R2", "title": "Главная операция",
         "unit": "м", "quantity": 12, "open_gap": "", "evidence_refs": ["S!R2"],
         "conversion_source": "Авторское основание", "quantity_formula": "6 * 2",
         "calc_quantity": 12, "conversion_factor": 2,
         "custom_metadata": {"literal": "Точное поле\nс пробелами  и кавычкой \"Ω\" и\u00a0NBSP"}},
        {"source_row": "S!R2", "operation_id": "op-covered", "revision": 3,
         "row_role": "work", "work_row": "S!R2", "title": "Учтённая операция",
         "unit": "м", "quantity": 12, "open_gap": "", "evidence_refs": ["S!R2"],
         "covered_by_work_id": "op-parent", "coverage": "Авторское подтверждение покрытия"},
        {"source_row": "S!R2", "operation_id": "op-material", "revision": 1,
         "row_role": "material", "work_row": "S!R2", "title": "Авторский материал",
         "unit": "м", "quantity": 12, "open_gap": "", "supply_kind": "supply",
         "resource_bindings": [{"action": "reuse", "target_resource_code": "R-01",
                                "reason": "Авторская связь", "source_refs": ["S!R2"]}],
         "resource_review": [{"literal": "Дословное решение модели"}]},
    ]
    if long_metadata:
        decisions[0]["custom_metadata"]["long_literal"] = "x\"\\\nΩ" * 14000
    before = copy.deepcopy(decisions)
    monkeypatch.setattr("proxy.services.spec_to_bor_service.derive_work_for_item",
                        lambda *_a, **_k: pytest.fail("authored VOR cannot invoke legacy professional choice"))
    target = tmp_path / "Авторская ВОР.xlsx"

    result = await workbook_service._default_vor_adapter(
        source, {"decisions": decisions}, target, lambda *_a: None,
    )

    assert result["source_rows"] == 3
    assert decisions == before
    workbook = openpyxl.load_workbook(target, data_only=False)
    audit = workbook["Решения модели"]
    assert audit.sheet_state == "visible"
    main = workbook["ВОР"]
    headers = {cell.value: cell.column for cell in main[3]}
    assert [main.cell(row, headers["ID операции"]).value for row in range(4, 7)] == [
        item["operation_id"] for item in decisions
    ]
    covered = main.cell(5, headers["Учтено работой"])
    assert covered.value == "op-parent" and covered.hyperlink.target == "#'ВОР'!M4"
    for row in range(4, 7):
        link = main.cell(row, headers["Полное решение модели"])
        assert link.hyperlink is not None and "Решения модели" in link.hyperlink.target
    assert [main.cell(row, 6).value for row in range(4, 7)] == [12, 12, 12]
    workbook.close()

    catalog = catalog_service.load_tabular_source_catalog(target)
    fragments = {}
    for row in catalog.rows:
        if row.sheet != "Решения модели" or not row.values[0].startswith("ВОР!R"):
            continue
        assert not row.is_work_row(), "the audit must be reference context, not duplicate operations"
        literal = next(cell.raw_value for cell in row.cells if cell.column == 3)
        assert row.values[2] == literal, "the model's displayed JSON must retain literal whitespace"
        assert literal.startswith("JSON: ")
        fragments.setdefault(row.values[0], []).append(literal[len("JSON: "):])
    restored = [json.loads("".join(fragments[f"ВОР!R{row}"])) for row in range(4, 7)]
    assert restored == before
    assert catalog.work_locators() == ["ВОР!R4", "ВОР!R5", "ВОР!R6"]
    packet = catalog_service.catalog_sheet_pack_lines(catalog, "ВОР", ["ВОР!R4", "ВОР!R5", "ВОР!R6"])
    assert any("op-covered" in line and "covered_by_work_id" in line for line in packet)
    assert any("quantity_formula" in line and "6 * 2" in line for line in packet)
    assert any("resource_bindings" in line and "R-01" in line for line in packet)
    scoped_packet = catalog_service.catalog_sheet_pack_lines(catalog, "ВОР", ["ВОР!R4"])
    assert any("quantity_formula" in line and "6 * 2" in line for line in scoped_packet)
    assert not any("op-covered" in line or "op-material" in line for line in scoped_packet)
