"""Explicit separate-supply decisions and exact financial retry targets."""
from copy import deepcopy

import pytest

from proxy.services import model_resource_binding_service as binding
from proxy.services.lsr_result_readiness_service import build_lsr_result_readiness
from tests.lsr_citation_support import cited_all


NORM = "ГЭСНм08-02-147-03"
RESOURCES = [{"kind": "material", "code": "FIX", "name": "Norm fastener", "unit": "kg", "per_unit": 0.1}]
SUPPLY = {"source_row": "S!R3", "resource_code": "CABLE"}


def _review(**changes):
    return {"source_row": "S!R3", "relation": "separate", "reason": "Literal model reason",
            "evidence_refs": ["Q1.H1"], **changes}


def _row(reviews=None):
    return {"source_row": "S!R2", "norm_code": NORM, "evidence_refs": ["Q1.H1"],
            "evidence_bindings": {"Q1.H1": "immutable-norm"},
            "resource_review": {"status": "keep_all_confirmed", "reason": "Norm resources retained",
                                "separate_supply_reviews": [_review()] if reviews is None else reviews}}


def _checked(row, *, shown=True, code=NORM, bindings=None):
    row = deepcopy(row)
    row["_separate_supply_review_evidence_checks"] = binding.separate_supply_review_evidence_checks(
        row, bindings if bindings is not None else {"Q1.H1": "immutable-norm"},
        {"Q1.H1": code}, {"immutable-norm": shown})
    return row


def _apply(row, supplies=(SUPPLY,), ambiguous=False):
    return binding.resource_binding_position_fields(row, RESOURCES, work_id="W", work_qty=1,
        physical_quantity=1, physical_unit="m", linked_supply_rows=supplies,
        supply_parent_ambiguous=ambiguous)


def test_explicit_checked_separate_supply_preserves_norm_and_authored_review():
    row = _checked(_row())
    before = deepcopy((row, RESOURCES, SUPPLY))
    result = _apply(row)
    assert result["resource_binding_conflicts"] == []
    assert result["resource_review_status"] == "keep_all_confirmed"
    assert result["resource_review"] == row["resource_review"]
    assert result["separate_supply_review_audit"][0]["status"] == "confirmed"
    assert (row, RESOURCES, SUPPLY) == before
    assert result.get("resources", RESOURCES) == RESOURCES


@pytest.mark.parametrize("review", [
    _review(source_row="S!R99"), _review(relation="different"), _review(reason=""),
    _review(evidence_refs=[]), _review(evidence_refs=["Q1.H2"]), _review(extra="unknown"),
    _review(evidence_refs="Q1.H1"),
])
def test_invalid_or_foreign_relationship_never_resolves_overlap(review):
    row = _checked(_row([review]))
    result = _apply(row)
    assert any(item["code"] == "separate_supply_possible_overlap" for item in result["resource_binding_conflicts"])
    assert result["resource_review"] == row["resource_review"]


@pytest.mark.parametrize("changes", [
    {"shown": False}, {"code": "ГЭСН11-01-011-12"}, {"bindings": {}},
])
def test_separate_review_requires_pinned_full_body_of_its_chosen_norm(changes):
    result = _apply(_checked(_row(), **changes))
    assert any(item["code"] == "separate_supply_possible_overlap" for item in result["resource_binding_conflicts"])


def test_other_delivered_card_does_not_certify_an_unread_selected_reference():
    row = _row([_review(evidence_refs=["Q1.H2"])])
    checks = binding.separate_supply_review_evidence_checks(row,
        {"Q1.H1": "shown", "Q1.H2": "deferred"}, {"Q1.H1": NORM, "Q1.H2": NORM},
        {"shown": True, "deferred": False})
    assert checks[0]["selected_norm_body_shown"] is False


def test_nested_relation_can_pin_its_own_shown_selected_norm_reference():
    row = _row([_review(evidence_refs=["Q2.H1"])])
    checks = binding.separate_supply_review_evidence_checks(row, {"Q1.H1": "top"}, {"Q1.H1": NORM},
        {"top": True, "nested": True}, first_bind=True,
        issued_alias_bindings={"Q1.H1": "top", "Q2.H1": "nested"}, global_cards={"top": NORM, "nested": NORM})
    assert checks[0]["selected_norm_body_shown"] is True
    assert checks[0]["evidence_bindings"] == {"Q2.H1": "nested"}
    assert "_separate_supply_review_evidence_bindings" not in row


def test_nested_replay_never_rebinds_reused_packet_alias_to_another_card():
    row = _row([_review(evidence_refs=["Q2.H1"])])
    row["_separate_supply_review_evidence_bindings"] = {"0": {"Q2.H1": "original"}}
    before = deepcopy(row)
    checks = binding.separate_supply_review_evidence_checks(row, {"Q1.H1": "top"}, {"Q1.H1": NORM},
        {"original": True, "other": True}, first_bind=False,
        issued_alias_bindings={"Q2.H1": "other"},
        global_cards={"original": NORM, "other": "ГЭСН11-01-011-12"})
    assert checks[0]["selected_norm_body_shown"] is True
    assert checks[0]["evidence_bindings"] == {"Q2.H1": "original"}
    assert row == before
    missing_original = binding.separate_supply_review_evidence_checks(row, {}, {}, {"other": True}, first_bind=False,
        issued_alias_bindings={"Q2.H1": "other"}, global_cards={"other": NORM})
    assert missing_original[0]["selected_norm_body_shown"] is False
    assert missing_original[0]["evidence_bindings"] == {"Q2.H1": "original"}


@pytest.mark.parametrize("case", ["duplicate", "ambiguous_parent", "duplicate_source", "no_checks", "unresolved"])
def test_ambiguous_or_unchecked_statements_remain_conflicts(case):
    row = _row([_review(), _review()] if case == "duplicate" else None)
    if case != "no_checks":
        row = _checked(row)
    if case == "unresolved":
        row["resource_review"]["status"] = "unresolved"
    result = _apply(row, supplies=(SUPPLY, SUPPLY) if case == "duplicate_source" else (SUPPLY,),
                    ambiguous=case == "ambiguous_parent")
    assert any(item["code"] == "separate_supply_possible_overlap" for item in result["resource_binding_conflicts"])


def test_one_confirmed_relationship_does_not_close_a_different_supply():
    result = _apply(_checked(_row()), supplies=(SUPPLY, {"source_row": "S!R4", "resource_code": "OTHER"}))
    overlap = next(item for item in result["resource_binding_conflicts"] if item["code"] == "separate_supply_possible_overlap")
    assert overlap["linked_supply_rows"] == [{"source_row": "S!R4", "resource_code": "OTHER"}]


def test_general_keep_all_and_different_resource_codes_are_not_relationship_decisions():
    row = _row([])
    result = _apply(row)
    assert any(item["code"] == "separate_supply_possible_overlap" for item in result["resource_binding_conflicts"])


@pytest.mark.parametrize("index", [[], {}, True, "0", -1, None])
def test_invalid_private_evidence_index_is_unverified_not_an_exception(index):
    row = _row()
    row["_separate_supply_review_evidence_checks"] = [{"index": index, "selected_norm_body_shown": True}]
    before = deepcopy(row)
    result = _apply(row)
    assert any(item["code"] == "separate_supply_possible_overlap" for item in result["resource_binding_conflicts"])
    assert row == before


def test_financial_retry_targets_exact_incomplete_operations_not_complete_siblings():
    rows = cited_all([
        {"source_row": "S!R2", "operation_id": "done", "row_role": "work", "norm_code": NORM, "norm_status": "EXACT"},
        {"source_row": "S!R2", "operation_id": "open", "row_role": "work", "norm_code": NORM, "norm_status": "EXACT"},
        {"source_row": "S!R3", "operation_id": "supply", "row_role": "material", "work_row": "S!R2"},
        {"source_row": "S!R4", "operation_id": "covered", "row_role": "work", "covered_by_work_id": "done"},
        {"source_row": "S!R5", "operation_id": "header", "row_role": "header"},
    ])
    projection = [
        {"locator": "S!R2", "operation_id": "op:done", "row_role": "work", "position_calculated": True},
        {"locator": "S!R2", "operation_id": "op:open", "row_role": "work", "position_calculated": False,
         "open_reasons": ["separate_supply_possible_overlap"]},
        {"locator": "S!R3", "operation_id": "op:supply", "row_role": "material", "position_calculated": False,
         "open_reasons": ["needs_price"]},
        {"locator": "S!R4", "operation_id": "op:covered", "row_role": "work", "status": "covered",
         "covered_by_operation_id": "op:done", "position_calculated": False},
        {"locator": "S!R5", "operation_id": "op:header", "row_role": "header", "position_calculated": False},
        {"locator": "S!R2", "operation_id": "op:foreign", "row_role": "work", "position_calculated": False},
    ]
    before = deepcopy((rows, projection))
    result = build_lsr_result_readiness(packaged=True, rows=rows, row_projection=projection, pricing_status="priced_partial")
    assert result["research_locators"] == ["S!R2", "S!R3"]
    assert result["research_operation_ids"] == ["open", "supply"]
    assert (rows, projection) == before


def test_missing_financial_projection_does_not_invent_retry_targets():
    rows = cited_all([{"source_row": "S!R2", "operation_id": "done", "row_role": "work", "norm_code": NORM, "norm_status": "EXACT"}])
    assert build_lsr_result_readiness(packaged=True, rows=rows)["research_operation_ids"] == []


def _feedback_fixture():
    rows = [{"source_row": "S!R2", "operation_id": "open", "revision": 2, "row_role": "work", "quantity": 0,
             "norm_code": NORM}, {"source_row": "S!R3", "operation_id": "done", "revision": 1, "row_role": "work"}]
    projection = [{"locator": "S!R2", "operation_id": "op:open", "row_role": "work", "position_calculated": False,
                   "open_reasons": ["separate_supply_possible_overlap", "needs_price"], "known_amount": 0},
                  {"locator": "S!R3", "operation_id": "op:done", "row_role": "work", "position_calculated": True}]
    basis = {"pricebook": "spb_3kv2026", "region": "SPb", "quarter": "3:2026", "pricebook_sha256": "p" * 64}
    identity = {"source_sha256": "s" * 64, "source_frame_contract_sha256": "f" * 64, "dataset_ids": ["b", "a"]}
    return rows, projection, basis, identity


def test_financial_feedback_is_exact_immutable_projection_not_recalculation():
    from proxy.services.lsr_financial_feedback_service import financial_feedback_snapshot, financial_feedback_for_rows
    rows, projection, basis, identity = _feedback_fixture()
    before = deepcopy((rows, projection, basis, identity))
    snapshot = financial_feedback_snapshot(rows, projection, revision_id="issued", price_basis=basis, identity=identity)
    facts = financial_feedback_for_rows(snapshot, rows[:1], price_basis=basis, identity=identity)
    assert facts == [{"operation_id": "open", "source_row": "S!R2", "accepted_revision": 2,
                      "document_revision_id": "issued", "projection": projection[0]}]
    assert financial_feedback_for_rows(snapshot, rows[1:], price_basis=basis, identity=identity) == []
    assert (rows, projection, basis, identity) == before


@pytest.mark.parametrize("change", ["revision", "quantity", "book_hash", "book", "source", "frame", "datasets", "legacy", "unknown_book_hash"])
def test_financial_feedback_never_presents_stale_scope_as_current(change):
    from proxy.services.lsr_financial_feedback_service import financial_feedback_snapshot, financial_feedback_for_rows
    rows, projection, basis, identity = _feedback_fixture()
    snapshot = financial_feedback_snapshot(rows, projection, revision_id="issued", price_basis=basis, identity=identity)
    if change == "revision": rows[0]["revision"] += 1
    if change == "quantity": rows[0]["quantity"] = 12
    if change == "book_hash": basis["pricebook_sha256"] = "changed"
    if change == "book": basis["pricebook"] = "other"
    if change == "source": identity["source_sha256"] = "changed"
    if change == "frame": identity["source_frame_contract_sha256"] = "changed"
    if change == "datasets": identity["dataset_ids"] = ["new"]
    if change == "legacy": snapshot = {}
    if change == "unknown_book_hash":
        basis.pop("pricebook_sha256")
        snapshot["price_basis"]["pricebook_sha256"] = ""
    assert financial_feedback_for_rows(snapshot, rows, price_basis=basis, identity=identity) == []


def test_packaged_feedback_survives_requested_retry_restart_without_reopening_complete_sibling(tmp_path):
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    from proxy.services.lsr_financial_feedback_service import financial_feedback_for_rows
    service = WorkbookResearchJobService(tmp_path / "retry.db")
    identity = service.identity(session_id="same", attachment_id="literal", source_sha256="a" * 64,
        dataset_ids=["norms"], profile_revision_id="profile", model_identity="model", question="LSR")
    job = service.begin_or_resume(identity, source_locators=["S!R2", "S!R3"])
    service.record_packet(job["job_id"], sheet="S", packet_locators=["S!R2", "S!R3"], evidence=[], remaining=[],
        accepted_rows=[{"source_row": "S!R2", "row_role": "work", "title": "Open", "quantity": 5, "norm_code": NORM},
                       {"source_row": "S!R3", "row_role": "work", "title": "Complete", "quantity": 10, "norm_code": NORM}])
    accepted = service.get(job["job_id"])["accepted_rows"]
    open_row, complete = accepted
    basis = {"pricebook": "spb", "pricebook_sha256": "p" * 64}
    projection = [{"locator": "S!R2", "operation_id": "op:" + open_row["operation_id"], "row_role": "work",
                   "position_calculated": False, "open_reasons": ["needs_price"]},
                  {"locator": "S!R3", "operation_id": "op:" + complete["operation_id"], "row_role": "work",
                   "position_calculated": True}]
    service.record_packaging(job["job_id"], tool="build_lsr_workbook", revision_id="issued", draft=True,
                             price_basis=basis, financial_projection=projection)
    reopened = service.reopen_locators(job["job_id"], ["S!R2"], preserve_accepted=True,
                                      operation_ids=[open_row["operation_id"]])
    context = WorkbookResearchJobService(service.db_path).chat_review_context(reopened["job_id"])
    assert context["financial_feedback"]["revision_id"] == "issued"
    assert context["financial_feedback"]["records"][0]["projection"] == projection[0]
    assert reopened["accepted_rows"] == accepted
    assert service.pending_research_operation_ids(reopened) == {open_row["operation_id"]}
    assert financial_feedback_for_rows(context["financial_feedback"], accepted,
        price_basis=reopened["task_context"]["price_basis"], identity=reopened["identity"])[0]["operation_id"] == open_row["operation_id"]
    updated = {**open_row, "expected_revision": open_row["revision"], "coverage": "Literal engineer clarification"}
    service.record_packet(job["job_id"], sheet="S", packet_locators=["S!R2"], accepted_rows=[updated], evidence=[], remaining=[])
    after = service.get(job["job_id"])["accepted_rows"]
    assert next(row for row in after if row["operation_id"] == complete["operation_id"]) == complete
    assert next(row for row in after if row["operation_id"] == open_row["operation_id"])["revision"] == open_row["revision"] + 1


def test_new_packaging_without_projection_does_not_replay_an_older_snapshot(tmp_path):
    from proxy.services.workbook_research_job_service import WorkbookResearchJobService
    service = WorkbookResearchJobService(tmp_path / "legacy.db")
    identity = service.identity(session_id="same", attachment_id="literal", source_sha256="a" * 64,
        dataset_ids=[], profile_revision_id="p", model_identity="m", question="LSR")
    job = service.begin_or_resume(identity)
    service.record_packaging(job["job_id"], tool="build_lsr_workbook", revision_id="old", draft=True, financial_projection=[])
    service.record_packaging(job["job_id"], tool="build_lsr_workbook", revision_id="new", draft=True)
    assert service.latest_financial_feedback(job["job_id"]) == {}


def test_confirmed_relation_audit_reaches_position_trace_and_downloaded_workbook(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from proxy.services import gesn_service
    from proxy.services.rim_lsr_trace_service import build_lsr_trace_from_visible_rows
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx
    from tests.lsr_citation_support import issued_norm_metadata

    norm = {"code": NORM, "name": "Literal work", "unit": "1 m", "resources": deepcopy(RESOURCES)}
    norm["_source_sha256"] = issued_norm_metadata(norm, synthetic=True)["normative_base_sha256"]
    monkeypatch.setattr(gesn_service, "get_norm", lambda *args, **kwargs: norm)
    row = _checked(_row())
    row.update(title="Literal work", quantity=1, unit="m", row_role="work", norm_status="EXACT", operation_id="parent")
    supply = {**SUPPLY, "row_role": "material", "work_row": "S!R2", "title": "Separate literal supply", "quantity": 1, "unit": "m"}
    before = deepcopy((row, supply, norm))
    trace = build_lsr_trace_from_visible_rows([row, supply])
    summary = trace["sections"][0]["positions"][0]["summary"]
    assert summary["separate_supply_review_audit"][0]["status"] == "confirmed"
    assert trace["row_bindings"][0]["separate_supply_review_audit"] == summary["separate_supply_review_audit"]
    assert trace["row_bindings"][0]["resource_review"] == row["resource_review"]
    assert not summary["resource_binding_conflicts"]
    assert [(item["code"], item["qty"]) for item in summary["resource_decisions"]] == [("FIX", 0.1)]
    output = tmp_path / "relations.xlsx"
    render_lsr_xlsx(trace, output)
    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        text = "\n".join(str(value) for values in workbook["Проверка"].iter_rows(values_only=True) for value in values if value is not None)
    finally:
        workbook.close()
    assert "Отдельная поставка S!R3" in text
    assert "Literal model reason" in text and "Q1.H1" in text
    assert (row, supply, norm) == before


def test_separate_relation_conflict_has_accurate_russian_operator_message():
    message = binding.resource_binding_conflict_note({"code": "separate_supply_review_conflict", "reason": "technical raw reason"})
    assert "отдельной поставки" in message and "карточку нормы" in message
    assert "technical raw reason" not in message
