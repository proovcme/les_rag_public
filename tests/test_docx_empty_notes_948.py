from docx import Document
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
import pytest

from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _notes_xml(kind, *, extra="", separator_content=""):
    note = "footnote" if kind == "footnotes" else "endnote"
    return (
        f'<w:{kind} xmlns:w="{WORD_NS}">\n'
        f'  <w:{note} w:type="separator" w:id="-1"><w:p><w:pPr>'
        '<w:spacing w:after="0" w:line="240"/></w:pPr>'
        '<w:r><w:rPr><w:rFonts w:ascii="Arial"/><w:sz w:val="20"/></w:rPr>'
        f'<w:separator/>{separator_content}</w:r></w:p></w:{note}>\n'
        f'  <w:{note} w:type="continuationSeparator" w:id="0">'
        f'<w:p><w:r><w:continuationSeparator/><w:t> \n\t </w:t></w:r></w:p></w:{note}>'
        f'{extra}\n</w:{kind}>'
    ).encode("utf-8")


def _write_docx(tmp_path, notes, *, body_marker=None):
    path = tmp_path / "Источник с примечаниями Ω.docx"
    document = Document()
    table = document.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ("Наименование", "Ед. изм.", "Количество")):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("Изделие", "шт", "3")):
        cell.text = text
    if body_marker:
        document.add_paragraph()._p.append(OxmlElement("w:" + body_marker))
    for kind, blob in notes.items():
        content_type = f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}+xml"
        part = Part(PackURI(f"/word/{kind}.xml"), content_type, blob, document.part.package)
        relation = RELATIONSHIP_TYPE.FOOTNOTES if kind == "footnotes" else RELATIONSHIP_TYPE.ENDNOTES
        document.part.relate_to(part, relation)
    document.save(path)
    return path


@pytest.mark.parametrize("kinds", [("footnotes",), ("endnotes",), ("footnotes", "endnotes")])
def test_empty_word_separator_parts_do_not_report_unread_document_content(tmp_path, kinds):
    path = _write_docx(tmp_path, {kind: _notes_xml(kind) for kind in kinds})
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete is True
    assert catalog.extraction_warnings == ()
    assert catalog.lookup("DOCX_T1!R2").header_quantity == 3


@pytest.mark.parametrize("kind", ["footnotes", "endnotes"])
@pytest.mark.parametrize("content", [
    "<w:p><w:r><w:t>Условие производства работ</w:t></w:r></w:p>",
    "<w:tbl><w:tr><w:tc><w:p><w:r><w:t>Размер</w:t></w:r></w:p></w:tc></w:tr></w:tbl>",
    "<w:p><w:r><w:drawing/></w:r></w:p>",
    "<w:p><w:r><w:pict/></w:r></w:p>",
    "<w:p><w:r><w:object/></w:r></w:p>",
])
def test_actual_note_text_table_or_media_remains_an_unread_part(tmp_path, kind, content):
    note = "footnote" if kind == "footnotes" else "endnote"
    path = _write_docx(tmp_path, {kind: _notes_xml(kind, extra=f'<w:{note} w:id="5">{content}</w:{note}>')})
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete is False
    assert catalog.extraction_warnings == (f"docx_unread_part:{kind}.xml",)


@pytest.mark.parametrize("kind", ["footnotes", "endnotes"])
@pytest.mark.parametrize("content", [
    "<w:t>Не пустой текст разделителя</w:t>",
    "<w:drawing/>", "<w:pict/>", "<w:object/>", "<w:txbxContent/>",
    "<w:tbl/>", "<w:futurePayload/>",
])
def test_separator_type_cannot_hide_real_or_unsupported_content(tmp_path, kind, content):
    path = _write_docx(tmp_path, {kind: _notes_xml(kind, separator_content=content)})
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete is False
    assert catalog.extraction_warnings == (f"docx_unread_part:{kind}.xml",)


@pytest.mark.parametrize("kind", ["footnotes", "endnotes"])
@pytest.mark.parametrize("blob", [b"<broken", b"not xml", b"<unexpected/>"])
def test_broken_or_unrecognized_notes_xml_is_still_an_explicit_gap(tmp_path, kind, blob):
    path = _write_docx(tmp_path, {kind: blob})
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete is False
    assert catalog.extraction_warnings == (f"docx_unread_part:{kind}.xml",)


@pytest.mark.parametrize("body_marker", ["drawing", "pict", "object", "txbxContent"])
def test_empty_notes_never_suppress_unread_images_or_objects_in_the_document(tmp_path, body_marker):
    path = _write_docx(
        tmp_path, {"footnotes": _notes_xml("footnotes"), "endnotes": _notes_xml("endnotes")},
        body_marker=body_marker,
    )
    catalog = load_tabular_source_catalog(path)
    assert catalog.extraction_complete is False
    assert catalog.extraction_warnings == (f"docx_unread_content:DOCX:{body_marker}",)
