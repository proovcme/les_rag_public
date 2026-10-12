"""Реестр строк ведомости: чем именно учтена каждая из 70 позиций.

ЗАЧЕМ. «Решено 12, остаток 0» означает конец очереди попыток, а не готовую смету
(ревью Codex 11.09). Критерий готовности — подтверждённый способ учёта по каждой
строке, и отдельная норма не единственный законный исход: крепёж входит в норму на
монтаж основного изделия, оборудование идёт самостоятельной стоимостью.

Здесь проверяется, что реестр не выдаёт заявление за доказательство.
"""

from __future__ import annotations

import json
import sqlite3

from tools.smeta_row_registry import (
    CALCULATED,
    CLAIMED_IN_WORK,
    NOT_IN_BOOK,
    NOT_REVIEWED,
    NO_NORM_FOUND,
    UNCLEAR,
    registry,
)


VOR = (
    "Файл: СКС.xlsx\n"
    "СКС!R1: № П.П | НАИМЕНОВАНИЕ РАБОТ И ЗАТРАТ | ЕД. ИЗМ | ОБЪЕМ\n"
    "СКС!R2: 984 | Шкаф напольный 42U | шт. | 2\n"
    "СКС!R3: 985 | Гайка М8 | шт. | 40\n"
    "СКС!R4: 986 | Заглушка-редукция | шт. | 12\n"
)


def _revision():
    """Ревизия той же ведомости — реестр берёт записи только по ней."""
    from proxy.smeta_core.vor_rows import rows_from_attachment_text, source_revision

    return source_revision(rows_from_attachment_text(VOR))


def _store(tmp_path, rows):
    path = tmp_path / "cursor.db"
    conn = sqlite3.connect(path)
    conn.execute(
        "CREATE TABLE smeta_portion_done (session_id TEXT, source_revision TEXT,"
        " anchor TEXT, decision_json TEXT, row_state TEXT, attempts INTEGER,"
        " PRIMARY KEY (session_id, source_revision, anchor))"
    )
    revision = _revision()
    for anchor, decision, state in rows:
        conn.execute(
            "INSERT INTO smeta_portion_done VALUES ('s1',?,?,?,?,1)",
            (revision, anchor, json.dumps(decision, ensure_ascii=False), state),
        )
    conn.commit()
    conn.close()
    return path


def _book(tmp_path, codes):
    import openpyxl

    path = tmp_path / "book.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["№ п/п", "Обоснование", "Наименование", "Ед. изм.", "Кол-во на ед.", "коэф.", "Кол-во всего"])
    for index, code in enumerate(codes, 1):
        sheet.append([index, code, "Работа", "шт", 2, 1, 2])
    book.save(path)
    return path


def test_a_row_counts_as_accounted_only_when_it_reaches_the_book(tmp_path):
    """Норма в решении и норма в книге — разные вещи. Считается вторая."""
    store = _store(tmp_path, [
        ("СКС!984", {"source_row": "984", "norm_code": "ГЭСНм10-06-037-01"}, "decided"),
        ("СКС!985", {"source_row": "985", "norm_code": "ГЭСНм99-99-999-99"}, "decided"),
    ])
    rows = registry(
        session="s1", attachment_text=VOR,
        book=_book(tmp_path, ["ГЭСНм10-06-037-01"]), db=store,
    )
    by_row = {row["source_row"]: row for row in rows}
    assert by_row["984"]["accounted_by"] == CALCULATED
    assert by_row["984"]["book_cells"] is not None
    assert by_row["985"]["accounted_by"] == NOT_IN_BOOK, (
        "норма выбрана, но в книгу не попала — учтённой строка не считается"
    )


def test_a_written_claim_is_not_proof_of_accounting(tmp_path):
    """«Учитывается в норме на монтаж» — заявление, а не доказательство.

    Ревью Codex 11.09: такая строка не связана ни с конкретной работой, ни с ресурсом,
    и расчётчик помечает её как «норма не выбрана». Реестр называет это отдельным
    состоянием, а не успехом.
    """
    store = _store(tmp_path, [
        ("СКС!985", {
            "source_row": "985", "norm_code": "",
            "coverage": "Нет отдельной нормы на гайку; учитывается в норме на монтаж",
        }, "exhausted"),
    ])
    rows = registry(session="s1", attachment_text=VOR, book=None, db=store)
    row = next(r for r in rows if r["source_row"] == "985")
    assert row["accounted_by"] == CLAIMED_IN_WORK
    assert row["accounted_by"] != CALCULATED
    assert row["norm_code"] is None


def test_a_row_the_model_never_saw_is_not_accounted(tmp_path):
    """Строка без записи в курсоре не может считаться учтённой ни в каком смысле."""
    rows = registry(session="s1", attachment_text=VOR, book=None, db=_store(tmp_path, []))
    assert len(rows) == 3
    assert {row["accounted_by"] for row in rows} == {NOT_REVIEWED}
    assert all(row["cursor_state"] == "нет записи" for row in rows)


def test_the_registry_keeps_the_order_of_the_vor(tmp_path):
    rows = registry(session="s1", attachment_text=VOR, book=None, db=_store(tmp_path, []))
    assert [row["source_row"] for row in rows] == ["984", "985", "986"]


def test_one_book_position_is_not_credited_to_two_rows(tmp_path):
    """Позиция книги расходуется один раз — иначе учтённой считается и чужая строка.

    Ревью Codex 11.09: сопоставление шло по одному шифру, и две строки с общей нормой
    обе числились учтёнными, даже когда вторую расчёт отверг.
    """
    store = _store(tmp_path, [
        ("СКС!984", {"source_row": "984", "norm_code": "ГЭСНм10-01-052-01"}, "decided"),
        ("СКС!985", {"source_row": "985", "norm_code": "ГЭСНм10-01-052-01"}, "decided"),
    ])
    # В книге одна позиция на 2 шт — она может принадлежать только строке 984.
    import openpyxl

    book = tmp_path / "one.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["№", "Обоснование", "Наименование", "Ед. изм.", "Кол-во", "коэф.", "Всего"])
    sheet.append([1, "ГЭСНм10-01-052-01", "Кроссировка", "шт", 2, 1, 2])
    workbook.save(book)

    rows = registry(session="s1", attachment_text=VOR, book=book, db=store)
    states = [row["accounted_by"] for row in rows[:2]]
    assert states.count(CALCULATED) == 1, "одна позиция — одна учтённая строка"
    assert CALCULATED not in states[1:], "вторая строка чужой расчёт не получает"


def test_a_gap_is_not_a_claim_of_accounting(tmp_path):
    """«Нет подходящей нормы в базе» — признание пробела, а не заявление об учёте.

    Ревью Codex 11.09: любой комментарий без нормы попадал в «заявлено в монтаже»,
    и 44 строки с признанием пробела считались объяснёнными.
    """
    store = _store(tmp_path, [
        ("СКС!984", {"source_row": "984", "norm_code": "",
                     "coverage": "Нет подходящей нормы в базе"}, "exhausted"),
        ("СКС!985", {"source_row": "985", "norm_code": "",
                     "coverage": "Нет отдельной нормы на гайку; учитывается в составе работ"},
         "exhausted"),
        ("СКС!986", {"source_row": "986", "norm_code": "", "coverage": ""}, "exhausted"),
    ])
    rows = registry(session="s1", attachment_text=VOR, book=None, db=store)
    by_row = {row["source_row"]: row["accounted_by"] for row in rows}
    assert by_row["984"] == NO_NORM_FOUND
    assert by_row["985"] == CLAIMED_IN_WORK, "пробел объяснён включением в работу"
    assert by_row["986"] == UNCLEAR


def test_rows_of_another_vor_do_not_leak_in(tmp_path):
    """Одна сессия может пройти несколько вложений — реестр берёт только своё."""
    path = tmp_path / "cursor.db"
    conn = sqlite3.connect(path)
    conn.execute(
        "CREATE TABLE smeta_portion_done (session_id TEXT, source_revision TEXT,"
        " anchor TEXT, decision_json TEXT, row_state TEXT, attempts INTEGER,"
        " PRIMARY KEY (session_id, source_revision, anchor))"
    )
    conn.execute(
        "INSERT INTO smeta_portion_done VALUES ('s1','чужая-ревизия','СКС!984',?, 'decided',1)",
        (json.dumps({"source_row": "984", "norm_code": "ГЭСН-ЧУЖАЯ"}),),
    )
    conn.commit()
    conn.close()

    rows = registry(session="s1", attachment_text=VOR, book=None, db=path)
    assert rows[0]["norm_code"] is None, "решение другой ведомости не должно попасть"
    assert rows[0]["accounted_by"] == NOT_REVIEWED


def test_a_hundredfold_quantity_error_is_not_a_match(tmp_path):
    """120 шт при измерителе «100 шт» — это 1,2 в книге, и никак не 120.

    Ревью Codex 11.09: сопоставление принимало и «как есть», и после деления, поэтому
    ошибка в сто раз проходила за совпадение, а «расхождений ноль» ничего не доказывало.
    """
    from tools.smeta_row_registry import _take_position

    wrong = [{"code": "A", "no": "1", "unit": "100 шт", "per_unit": "120",
              "total": "120", "taken": False}]
    assert _take_position(wrong, "A", 120) is None

    right = [{"code": "A", "no": "1", "unit": "100 шт", "per_unit": "1.2",
              "total": "1.2", "taken": False}]
    assert _take_position(right, "A", 120) is not None


def test_a_quantity_that_matches_nothing_is_not_called_accounted(tmp_path):
    """Позиция с тем же шифром, но чужим количеством, строке не принадлежит.

    Ревью Codex 11.09: реестр забирал первую попавшуюся позицию, помечал расхождение
    и всё равно считал строку «норма в книге».
    """
    from tools.smeta_row_registry import QUANTITY_UNMATCHED

    store = _store(tmp_path, [
        ("СКС!984", {"source_row": "984", "norm_code": "ГЭСНм10-06-037-01"}, "decided"),
    ])
    import openpyxl

    book = tmp_path / "book.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["№", "Обоснование", "Наименование", "Ед. изм.", "Кол-во", "коэф.", "Всего"])
    sheet.append([1, "ГЭСНм10-06-037-01", "Работа", "шт", 99, 1, 99])
    workbook.save(book)

    rows = registry(session="s1", attachment_text=VOR, book=book, db=store)
    assert rows[0]["accounted_by"] == QUANTITY_UNMATCHED


def test_a_negated_claim_is_a_gap(tmp_path):
    """«Не входит в состав работ» — противоположность заявлению об учёте."""
    from tools.smeta_row_registry import _classify_without_norm

    assert _classify_without_norm("Не входит в состав работ; требуется поиск") == NO_NORM_FOUND
    assert _classify_without_norm("Не учитывается в норме — нужен поиск") == NO_NORM_FOUND
    assert _classify_without_norm("Стоимость включена в стоимость оборудования") == CLAIMED_IN_WORK
    assert _classify_without_norm("Учитывается в норме на монтаж шкафа") == CLAIMED_IN_WORK
