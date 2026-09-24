# -*- coding: utf-8
"""§1 release run (25.09): a real customer estimate names its quantity graph
«Ко-во» — one letter short of «Кол-во».  The ВОР writer bound no quantity column
at all and every delivered row lost its volume while the units survived.  The
variant spelling must bind, and no other column may be mistaken for it."""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from proxy.services.bor_service import source_rows_to_vor_xlsx
from proxy.services.spec_to_bor_service import _map_headers, rows_from_spec_xlsx


ESTIMATE_HEADER = ["№ п/п", "Обоснование", "Наименование работ и затрат",
                   "Ед. изм.", "Ко-во", "Цена за ед.  руб.", "", "Стоимость, руб."]


def test_variant_quantity_header_binds_the_quantity_column():
    mapped = _map_headers(ESTIMATE_HEADER)
    assert mapped["name"] == 2
    assert mapped["unit"] == 3
    assert mapped.get("qty") == 4


def test_a_word_starting_like_the_variant_is_not_bound_as_quantity():
    # «Ковровые…» must never be read as a quantity graph.
    mapped = _map_headers(["Наименование", "Ковровые покрытия", "Цена"])
    assert "qty" not in mapped


def test_alias_priority_over_a_bare_volume_column_is_unchanged():
    mapped = _map_headers(["Наименование", "Объем", "Кол-во ФАКТ"])
    assert mapped["qty"] == 2


def test_rows_from_spec_xlsx_keeps_the_source_quantities(tmp_path: Path):
    path = tmp_path / "smeta.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(ESTIMATE_HEADER)
    sheet.append([1, "Договорная цена", "Кабельный ввод M40", "шт.", 3, 623.76, 0, 1871.28])
    sheet.append([2, '- " -', "DIN-рейка, 35мм", "м.", 5, 518.71, 0, 2593.55])
    workbook.save(path)

    rows = rows_from_spec_xlsx(path, source_label=path.name)

    assert [(r["name"], r["unit"], r["qty"]) for r in rows] == [
        ("Кабельный ввод M40", "шт.", 3.0),
        ("DIN-рейка, 35мм", "м.", 5.0),
    ]
    # the delivered sheet must carry the value in the Кол-во graph
    out = tmp_path / "vor.xlsx"
    source_rows_to_vor_xlsx(rows, out, title="ВОР")
    exported = list(openpyxl.load_workbook(out).active.iter_rows(min_row=4,
                                                                values_only=True))
    assert exported[0][5] == pytest.approx(3.0)
    assert exported[1][5] == pytest.approx(5.0)
