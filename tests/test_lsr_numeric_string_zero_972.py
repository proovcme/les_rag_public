"""Numeric-string zeros obey the quantity contract without float underflow.

An exact zero string may observe the same authoritative numeric source zero.
Unknown source facts and authored charge, norm or coverage actions keep their
existing path, and the model's quantity remains unchanged.
"""
from __future__ import annotations

from copy import deepcopy
from dataclasses import replace
import hashlib

import openpyxl
import pytest


def _source(tmp_path, quantity=0):
    from proxy.services.attachment_source_catalog_service import load_tabular_source_catalog

    path = tmp_path / "Numeric zero source.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Literal source work", "шт", quantity])
    workbook.save(path)
    workbook.close()
    catalog = load_tabular_source_catalog(path)
    assert catalog is not None and catalog.extraction_complete
    return path, catalog


def _row(quantity):
    return {
        "source_row": "S!R2", "row_role": "work", "title": "Literal source work",
        "quantity": quantity, "unit": "шт", "operation_id": "literal-zero-operation",
        "revision": 1, "norm_code": "", "norm_status": None, "evidence_refs": [],
        "coverage": "Literal authored explanation remains unchanged",
    }


@pytest.mark.parametrize("quantity", ["0", "0.00", "-0", "0,0", "0e0"])
def test_exact_numeric_string_zero_observes_known_source_zero_without_rewriting_qty(tmp_path, quantity):
    from proxy.services.rim_lsr_trace_service import source_confirmed_zero_work_charge_not_required

    path, catalog = _source(tmp_path)
    row = _row(quantity)
    before = deepcopy(row)
    source_sha = hashlib.sha256(path.read_bytes()).hexdigest()

    assert source_confirmed_zero_work_charge_not_required(row, catalog) is True
    assert row == before and type(row["quantity"]) is str
    assert row["quantity"] == quantity
    assert hashlib.sha256(path.read_bytes()).hexdigest() == source_sha


@pytest.mark.parametrize(
    "quantity",
    ["1e-999999", "-1e-999999", "1", "-1", "NaN", "nan", "inf", "Infinity",
     "-Infinity", "false", "0 шт", "zero", "", "   ", None, False],
    ids=["positive-underflow", "negative-underflow", "positive", "negative", "NaN", "nan",
         "inf", "Infinity", "negative-Infinity", "text-false", "unit-in-number", "nonnumeric",
         "empty", "whitespace", "null", "boolean-false"],
)
def test_nonzero_or_nonnumeric_quantity_cannot_observe_a_zero_charge_exemption(tmp_path, quantity):
    from proxy.services.rim_lsr_trace_service import source_confirmed_zero_work_charge_not_required

    _, catalog = _source(tmp_path)
    row = _row(quantity)
    before = deepcopy(row)

    assert source_confirmed_zero_work_charge_not_required(row, catalog) is False
    assert row == before


@pytest.mark.parametrize("source_quantity", [1, -1, None])
def test_zero_numeric_string_cannot_replace_positive_negative_or_blank_source(tmp_path, source_quantity):
    from proxy.services.rim_lsr_trace_service import source_confirmed_zero_work_charge_not_required

    _, catalog = _source(tmp_path, source_quantity)
    row = _row("0.00")
    before = deepcopy(row)

    assert source_confirmed_zero_work_charge_not_required(row, catalog) is False
    assert row == before


@pytest.mark.parametrize("kind", ["missing-catalog", "unknown-column", "unknown-unit", "incomplete-extraction"])
def test_zero_numeric_string_requires_complete_known_source_quantity_and_unit(tmp_path, kind):
    from proxy.services.rim_lsr_trace_service import source_confirmed_zero_work_charge_not_required

    _, catalog = _source(tmp_path)
    source = catalog.lookup("S!R2")
    if kind == "missing-catalog":
        catalog = None
    elif kind == "unknown-column":
        catalog = replace(catalog, rows=tuple(
            replace(item, quantity_column_known=False) if item is source else item
            for item in catalog.rows
        ))
    elif kind == "unknown-unit":
        catalog = replace(catalog, rows=tuple(
            replace(item, column_headers=tuple(
                (column, "Unspecified label" if column == 2 else label)
                for column, label in item.column_headers
            )) if item is source else item
            for item in catalog.rows
        ))
    else:
        catalog = replace(catalog, extraction_complete=False)
    row = _row("0")
    before = deepcopy(row)

    assert source_confirmed_zero_work_charge_not_required(row, catalog) is False
    assert row == before


@pytest.mark.parametrize(
    "authored",
    [
        {"norm_code": "ГЭСН11-01-011-01"},
        {"row_role": "material"},
        {"calc_quantity": 2},
        {"conversion_factor": 2},
        {"calc_unit": "м"},
        {"conversion_source": "Literal derived basis"},
        {"quantity_source": "Literal authored quantity source"},
        {"confirmed_price": 0},
        {"resource_bindings": [{"action": "add", "resource_code": "literal-resource"}]},
        {"resource_review": {"reason": "Literal authored review"}},
        {"covered_by_work_id": "literal-parent-work"},
        {"covered_by_operation_id": "literal-parent-operation"},
        {"norm_components": [{"norm_code": "literal-component"}]},
    ],
    ids=["selected-norm", "supply-role", "derived-quantity", "conversion-factor", "derived-unit",
         "conversion-basis", "quantity-basis", "explicit-zero-price", "resource-action",
         "resource-review", "covered-work", "covered-operation", "norm-components"],
)
def test_numeric_string_zero_does_not_erase_any_authored_charge_or_coverage_action(tmp_path, authored):
    from proxy.services.rim_lsr_trace_service import source_confirmed_zero_work_charge_not_required

    _, catalog = _source(tmp_path)
    row = {**_row("0e0"), **authored}
    before = deepcopy(row)

    assert source_confirmed_zero_work_charge_not_required(row, catalog) is False
    assert row == before


@pytest.mark.parametrize("quantity", ["0", "0.00"])
def test_numeric_string_zero_has_no_norm_price_retry_debt_and_renders_as_numeric_observation(tmp_path, quantity):
    from tests.test_lsr_zero_charge_observation_969 import _catalog, _pipeline
    from proxy.services.rim_trace_xlsx_service import render_lsr_xlsx

    rows = [_row(quantity)]
    before = deepcopy(rows)
    trace, ready, missing, prices = _pipeline(rows, _catalog(tmp_path))

    assert trace["row_bindings"][0]["charge_not_required"] is True
    assert trace["row_bindings"][0]["source_row"]["quantity"] == quantity
    assert trace["row_projection"][0]["charge_not_required"] is True
    assert trace["sections"] == []
    assert trace["summary"]["selected_norms"] == 0
    assert trace["summary"]["calculated_positions"] == 0
    assert trace["summary"]["confirmed_amount"] == 0
    assert trace["summary"]["preliminary_amount"] == 0
    assert ready["dimensions"]["norms"]["complete"] is True
    assert ready["dimensions"]["prices"]["complete"] is True
    assert ready["gap_states"]["calculation_incomplete"]["count"] == 0
    assert ready["research_locators"] == ready["research_operation_ids"] == []
    assert ready["estimate_complete"] is True
    assert missing == prices == []
    assert rows == before and type(rows[0]["quantity"]) is str

    output = tmp_path / "Zero string charge observation.xlsx"
    render_lsr_xlsx(trace, output)
    workbook = openpyxl.load_workbook(output, read_only=True, data_only=False)
    try:
        observations = [row for row in workbook["ЛСР РИМ"].iter_rows()
                        if "начисление не требуется" in str(row[2].value or "").casefold()]
        assert len(observations) == 1
        observation = observations[0]
        assert observation[6].value == 0 and observation[6].data_type == "n"
        # This is an informational source observation, not a zero-priced charge.
        assert all(cell.value is None for cell in observation[7:12])
    finally:
        workbook.close()
    assert rows == before and rows[0]["quantity"] == quantity


def test_numeric_string_zero_keeps_only_a_positive_neighbor_in_research(tmp_path):
    from tests.test_lsr_zero_charge_observation_969 import _catalog, _pipeline

    rows = [_row("0.00"), {
        **_row(1), "source_row": "S!R3", "operation_id": "positive-neighbor",
        "title": "Literal positive work", "norm_status": "MISSING",
        "open_gap": "No model norm selection yet",
    }]
    before = deepcopy(rows)
    trace, ready, _, _ = _pipeline(rows, _catalog(tmp_path, (0, 1)))

    projections = {row["locator"]: row for row in trace["row_projection"]}
    assert projections["S!R2"]["charge_not_required"] is True
    assert not projections["S!R3"].get("charge_not_required")
    assert ready["research_locators"] == ["S!R3"]
    assert ready["research_operation_ids"] == ["positive-neighbor"]
    assert ready["gap_states"]["calculation_incomplete"]["count"] == 1
    assert ready["estimate_complete"] is False
    assert rows == before and rows[0]["quantity"] == "0.00"
