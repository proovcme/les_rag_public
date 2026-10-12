import asyncio
from types import SimpleNamespace

from proxy.services import chat_evidence_application_service as chat
from proxy.services import fgis_price_service as prices
from proxy.services.portion_cursor_service import PortionCursorService, decisions_revision


def basis(letter='a'):
    return dict(schema='les.document-price-basis.v1',status='pinned',path='/fixture/'+letter+'.parquet',
        identity=dict(schema='les.pricebook-revision.v1',sha256=letter*64,pricebook_id=letter,
                      region='fixture',period='fixture',version=None))


def test_document_pin_survives_default_change_and_cursor_restart(tmp_path):
    db=tmp_path/'cursor.sqlite'
    store=PortionCursorService(db)
    assert store.document_price_basis(session_id='s',source_revision='r',capture=lambda:basis())==basis()
    store=PortionCursorService(db)
    assert store.document_price_basis(session_id='s',source_revision='r',capture=lambda:basis('b'))==basis()
    assert store.document_price_basis(session_id='s',source_revision='new',capture=lambda:basis('b'))==basis('b')


def test_unparsed_real_attachments_do_not_share_empty_price_identity(tmp_path):
    import hashlib
    import openpyxl
    from proxy.services.tabular_attachment_service import format_tabular_attachment_context
    from proxy.smeta_core.vor_rows import rows_from_attachment_text
    store=PortionCursorService(tmp_path/'cursor.sqlite')
    revisions=[]
    for letter in ('a','b'):
        path=tmp_path/(letter+'.xlsx')
        book=openpyxl.Workbook(); book.active.append(['Unnumbered source '+letter]); book.save(path); book.close()
        text,_=format_tabular_attachment_context(path,path.name,max_chars=None,max_cell_chars=None)
        assert rows_from_attachment_text(text)==[]
        sha=hashlib.sha256(path.read_bytes()).hexdigest()
        revision=chat.document_price_source_revision(text,sha)
        revisions.append(revision)
        captured=store.document_price_basis(session_id='same',source_revision=revision,capture=lambda:basis(letter))
        assert captured==basis(letter)
        assert chat.document_price_source_revision('Renamed preview',sha)==revision
    assert revisions[0]!=revisions[1]


def test_unparsed_pasted_context_has_own_price_identity():
    assert chat.document_price_source_revision('Unnumbered A')!=chat.document_price_source_revision('Unnumbered B')


def test_old_document_without_pin_never_claims_current_default(tmp_path):
    store=PortionCursorService(tmp_path/'cursor.sqlite')
    store.record_pass(session_id='s',source_revision='r',completed=[],decisions={},pass_id='old')
    def forbidden():
        raise AssertionError('legacy document must not resolve current default')
    result=store.document_price_basis(session_id='s',source_revision='r',capture=forbidden)
    assert result['status']=='legacy_unpinned' and result['identity'] is None


def test_explicit_trusted_change_invalidates_artifact_freshness(tmp_path):
    store=PortionCursorService(tmp_path/'cursor.sqlite')
    store.document_price_basis(session_id='s',source_revision='r',capture=lambda:basis())
    decision=dict(source_row='S!1',work_id='w',norm_code='ГЭСН01-01-001-01')
    old=decisions_revision([decision],document_price_basis=basis())
    store.replace_document_price_basis(session_id='s',source_revision='r',basis=basis('b'))
    current=store.document_price_basis(session_id='s',source_revision='r',capture=lambda:basis())
    assert decisions_revision([decision],document_price_basis=current)!=old
    assert chat.stored_decisions_need_a_workbook(dict(prior_decisions={'S!1':decision},
        artifact_revision_id='existing',artifact_decisions_revision=old,document_price_basis=current))


def test_request_local_price_callback_is_bound_to_document_not_default(monkeypatch):
    seen=[]
    monkeypatch.setattr(prices,'pricebook_for_basis',lambda value:
        SimpleNamespace(browse=lambda query,limit:seen.append((value,query,limit)) or [dict(code='M')]),raising=False)
    browse=chat.bind_document_price_browse(basis())
    assert browse('model query',6)==[dict(code='M')]
    assert seen==[(basis(),'model query',6)]


def test_price_cards_preserve_reader_book_and_record_identity():
    from proxy.services.model_research_tool_service import ModelResearchToolService
    async def retrieve(*args,**kwargs):
        return []
    record=dict(code='01.1.01.01-0001',name='Fixture',unit='м',price_current_eff=10,
        pricebook_identity=basis()['identity'],price_record_identity=dict(record_id='row:1'))
    service=ModelResearchToolService(retrieve=retrieve,frozen_dataset_ids=(),retrieval_kwargs={},
        fallback=retrieve,price_browse=lambda q,n:[record])
    result=asyncio.run(service._price_cards('fixture'))
    assert result.chunks[0].meta['pricebook_identity']==record['pricebook_identity']
    assert result.chunks[0].meta['price_record_identity']==record['price_record_identity']


def test_cached_answer_a_cannot_bypass_changed_document_basis_b():
    cached=SimpleNamespace(answer='old A answer')
    calls=[]
    cache=SimpleNamespace(lookup_session_unvalidated=lambda *a:calls.append(a) or cached)
    assert chat.document_price_safe_cache_lookup(cache,'continue','scope','fp','session',primary_document=True) is None
    assert calls==[]
    assert chat.document_price_safe_cache_lookup(cache,'continue','scope','fp','session',primary_document=False) is cached


def test_price_revision_failure_is_explicit_research_blocker(monkeypatch):
    from proxy.services.model_research_tool_service import ModelResearchToolService
    class Conflict(Exception):
        def __init__(self):
            self.expected=basis()['identity']; self.actual=basis('b')['identity']; self.reason='replaced'
    monkeypatch.setattr(prices,'PriceBookRevisionConflict',Conflict,raising=False)
    def browse(q,n):
        raise Conflict()
    async def retrieve(*args,**kwargs):
        return []
    service=ModelResearchToolService(retrieve=retrieve,frozen_dataset_ids=(),retrieval_kwargs={},fallback=retrieve,price_browse=browse)
    result=asyncio.run(service._price_cards('fixture'))
    assert result.payload['status']=='blocked'
    assert result.retrieval_trace['error_code']=='pricebook_revision_conflict'
    assert result.retrieval_trace['expected']==basis()['identity']


def test_initial_price_capture_conflict_is_visible_422(monkeypatch,tmp_path):
    import pytest
    from fastapi import HTTPException
    class Conflict(Exception):
        def __init__(self):
            self.expected=None; self.actual={'region':'mixed'}; self.reason='mixed_region'
    monkeypatch.setattr(prices,'PriceBookRevisionConflict',Conflict,raising=False)
    def capture():
        raise Conflict()
    with pytest.raises(HTTPException) as raised:
        chat.document_price_basis_or_error(PortionCursorService(tmp_path/'cursor.sqlite'),
            session_id='session',source_revision='source',capture=capture,http_exception=HTTPException)
    assert raised.value.status_code==422
    assert raised.value.detail['code']=='pricebook_revision_conflict'
    assert raised.value.detail['reason']=='mixed_region'


def test_real_capture_default_change_restart_keeps_same_price_cards(tmp_path,monkeypatch):
    import pandas as pd
    from proxy.services.model_research_tool_service import ModelResearchToolService
    a,b=tmp_path/'book_a.parquet',tmp_path/'book_b.parquet'
    for path,amount in ((a,10),(b,90)):
        pd.DataFrame([dict(code='01.1.01.01-0001',name='Synthetic cable',unit='м',
            price_current=amount,price_current_eff=amount,region='Fixture',quarter='2026Q1')]).to_parquet(path,index=False)
    monkeypatch.setattr(prices,'available_pricebooks',lambda *a,**k:[str(a_path),str(b_path)])
    a_path,b_path=a,b
    monkeypatch.setattr(prices,'pricebook_manifest',lambda:{})
    monkeypatch.setenv('LES_DEFAULT_PRICEBOOK','book_a')
    db=tmp_path/'cursor.sqlite'
    pin=PortionCursorService(db).document_price_basis(session_id='s',source_revision='source',capture=prices.capture_price_basis)
    monkeypatch.setenv('LES_DEFAULT_PRICEBOOK','book_b')
    assert prices.capture_price_basis()['identity']['pricebook_id']=='book_b'
    prices.get_pricebook.cache_clear()
    restored=PortionCursorService(db).document_price_basis(session_id='s',source_revision='source',capture=prices.capture_price_basis)
    assert restored==pin
    async def retrieve(*args,**kwargs):
        return []
    service=ModelResearchToolService(retrieve=retrieve,frozen_dataset_ids=(),retrieval_kwargs={},fallback=retrieve,
        price_browse=chat.bind_document_price_browse(restored))
    result=asyncio.run(service._price_cards('Synthetic cable'))
    assert result.chunks[0].meta['price_current_eff']==10
    assert result.chunks[0].meta['pricebook_identity']==pin['identity']
    assert result.chunks[0].meta['price_record_identity']['pricebook_identity']==pin['identity']
    prices.get_pricebook.cache_clear()


def test_trusted_workbook_context_routes_basis_and_separates_checkpoints(tmp_path):
    import openpyxl
    from dataclasses import replace
    from proxy.services.chat_attachment_service import preserve_read_attachment
    from proxy.services.workflow_checkpoint_service import WorkflowCheckpointService
    from proxy.services.artifact_revision_service import ArtifactRevisionStore
    from proxy.services.workbook_tool_service import WorkbookExecutionContext, build_lsr_workbook
    source=tmp_path/'source.xlsx'
    book=openpyxl.Workbook(); book.active.append(['Fixture']); book.save(source); book.close()
    attachment=preserve_read_attachment(source,attachment_id='read_123456abcdef',original_name='source.xlsx',root=tmp_path/'attachments')
    seen=[]
    def adapter(source_path,args,output_path,progress):
        seen.append(dict(args))
        book=openpyxl.Workbook(); book.active.append(['Synthetic artifact']); book.save(output_path); book.close()
        return dict(file_path=output_path)
    ctx=WorkbookExecutionContext(session_id='session',idempotency_key='basis-a',model_decision_revision='decision',
        profile_revision_id='profile',model_identity='model',model_preset='preset',attachment_root=tmp_path/'attachments',
        work_dir=tmp_path/'work',checkpoints=WorkflowCheckpointService(tmp_path/'checkpoints.sqlite'),
        artifacts=ArtifactRevisionStore(tmp_path/'artifacts.sqlite',tmp_path/'artifacts'),lsr_adapter=adapter,
        document_price_basis=basis(),source_revision='source',source_attachment_sha256=attachment['sha256'])
    args=dict(attachment_id=attachment['attachment_id'],decisions=[dict(source_row=1,title='Synthetic',quantity=1,unit='шт',norm_code='')])
    first=asyncio.run(build_lsr_workbook(args,ctx))
    resumed=asyncio.run(build_lsr_workbook(args,ctx))
    second=asyncio.run(build_lsr_workbook(args,replace(ctx,idempotency_key='basis-b',document_price_basis=basis('b'))))
    assert first['artifact']['revision_id']==resumed['artifact']['revision_id']
    assert first['artifact']['revision_id']!=second['artifact']['revision_id']
    assert [s['_document_price_basis'] for s in seen]==[basis(),basis('b')]
    assert all(s['_source_revision']=='source' for s in seen)
    # A repeat upload of the same bytes is valid; another file must not receive A's pin.
    alias=preserve_read_attachment(source,attachment_id='read_abcdef123456',original_name='alias.xlsx',root=tmp_path/'attachments')
    alias_result=asyncio.run(build_lsr_workbook(dict(args,attachment_id=alias['attachment_id']),replace(ctx,idempotency_key='alias')))
    assert alias_result['status']=='complete'
    other=tmp_path/'other.xlsx'
    book=openpyxl.Workbook(); book.active.append(['Other source']); book.save(other); book.close()
    other_attachment=preserve_read_attachment(other,attachment_id='read_abcdefabcdef',original_name='other.xlsx',root=tmp_path/'attachments')
    rejected=asyncio.run(build_lsr_workbook(dict(args,attachment_id=other_attachment['attachment_id']),replace(ctx,idempotency_key='wrong-source')))
    assert rejected['code']=='DOCUMENT_PRICE_BASIS_SOURCE_MISMATCH'
    assert len(seen)==3
