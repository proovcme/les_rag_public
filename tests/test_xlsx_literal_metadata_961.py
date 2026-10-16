"""Printed metadata and classic comments are literal references, never works."""
import hashlib
import json
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl
from openpyxl.comments import Comment
import pytest

from proxy.services.attachment_source_catalog_service import (
    catalog_sheet_pack_lines, compare_model_row_to_catalog, load_tabular_source_catalog,
)
from proxy.services.model_source_frame_service import build_model_source_frame, render_model_source_frame

_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def _source(tmp_path, *, comment=False):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Лист Ω"
    sheet.append(["Наименование", "Ед. изм.", "Количество"])
    sheet.append(["Исходная операция", "м2", 12])
    if comment:
        sheet["A2"].comment = Comment("Полный текст\n30 мм & <слой> Ω", "Автор 'Ω'")
    return workbook, sheet, tmp_path / "Исходник «Ω» и 'quote'.xlsx"


def _save(workbook, source):
    workbook.save(source)
    workbook.close()
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    catalog = load_tabular_source_catalog(source)
    assert catalog is not None
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    return catalog


def _metadata(catalog):
    return [row for row in catalog.rows if row.classification_reason == "xlsx_literal_metadata"]


def _replace_xml(source, part, mutate):
    with ZipFile(source) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    target = next(data for item, data in members if item.filename == part)
    node = ET.fromstring(target)
    mutate(node)
    with ZipFile(source, "w") as archive:
        for item, data in members:
            archive.writestr(item, ET.tostring(node, encoding="utf-8", xml_declaration=True)
                if item.filename == part else data)


@pytest.mark.parametrize("kind", ["oddHeader", "evenHeader", "firstHeader", "oddFooter", "evenFooter", "firstFooter"])
@pytest.mark.parametrize("position", ["left", "center", "right"])
def test_all_print_parts_and_sections_reach_every_packet_as_literal_reference(tmp_path, kind, position):
    workbook, sheet, source = _source(tmp_path)
    text = 'Слой <Ω> & 30 мм\n&"Arial,Bold"&12run1 &"Arial,Italic"run2 && &P / &N &F &A &Z &D &T'
    getattr(getattr(sheet, kind), position).text = text
    catalog = _save(workbook, source)
    assert catalog.extraction_complete and catalog.extraction_warnings == ()
    assert catalog.work_locators() == ["Лист Ω!R2"]
    metadata = _metadata(catalog)
    assert len(metadata) == 1 and metadata[0].excel_row == 0
    assert metadata[0].row_classification == "note" and not metadata[0].is_work_row()
    body = json.loads(metadata[0].values[0])
    with ZipFile(source) as archive:
        raw = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    original = raw.find(f"{{{_NS}}}headerFooter/{{{_NS}}}{kind}")
    assert body["parts"] == [{"kind": child.tag.rsplit("}", 1)[-1], "literal": child.text or ""}
        for child in raw.find(f"{{{_NS}}}headerFooter")]
    assert ET.fromstring(body["literal_xml"]).find(f"{{{_NS}}}{kind}").text == original.text
    assert metadata[0].line in catalog_sheet_pack_lines(catalog, "Лист Ω", ["Лист Ω!R2"])
    frame = build_model_source_frame(catalog, ["Лист Ω!R2"])
    assert metadata[0].locator in frame.reference_locators
    assert metadata[0].locator not in frame.requested_locators
    shown = json.loads(render_model_source_frame(frame).payload_json)["shown_rows"]
    assert next(row["source_text"] for row in shown if row["source_row"] == metadata[0].locator) == metadata[0].line
    assert catalog.lookup(metadata[0].locator) == metadata[0]
    assert compare_model_row_to_catalog({"source_row": "Лист Ω!R2", "quantity": 12, "unit": "м2"}, catalog) == []


@pytest.mark.parametrize("coordinate", ["A2", "D8"])
@pytest.mark.parametrize("hidden", [False, True])
def test_classic_comment_text_author_and_coordinate_survive_without_work_candidate(tmp_path, coordinate, hidden):
    workbook, sheet, source = _source(tmp_path)
    sheet[coordinate].comment = Comment("Условия:\nдва слоя Ω <x> & 30 мм", "Инженер Ω")
    sheet.row_dimensions[8].hidden = hidden
    if hidden:
        other = workbook.create_sheet("Видимый")
        other["A1"] = "Контекст"
        sheet.sheet_state = "hidden"
    catalog = _save(workbook, source)
    assert catalog.extraction_complete and catalog.extraction_warnings == ()
    assert catalog.work_locators() == ["Лист Ω!R2"]
    metadata = _metadata(catalog)
    assert len(metadata) == 1 and metadata[0].hidden_sheet is hidden
    body = json.loads(metadata[0].values[0])
    assert body["coordinate"] == coordinate and body["author"] == "Инженер Ω"
    assert body["text"] == "Условия:\nдва слоя Ω <x> & 30 мм"
    frame = build_model_source_frame(catalog, ["Лист Ω!R2"])
    assert metadata[0].locator in frame.reference_locators
    assert catalog.lookup(metadata[0].locator).line == metadata[0].line


def test_comment_rich_text_runs_preserve_exact_text_and_full_xml(tmp_path):
    workbook, sheet, source = _source(tmp_path, comment=True)
    _save(workbook, source)
    def rich_text(node):
        text = node.find(f"{{{_NS}}}commentList/{{{_NS}}}comment/{{{_NS}}}text")
        text.clear()
        for value in ("Первая ", "Ω <&>\n", " последняя"):
            run = ET.SubElement(text, f"{{{_NS}}}r")
            ET.SubElement(run, f"{{{_NS}}}rPr")
            ET.SubElement(run, f"{{{_NS}}}t").text = value
    with ZipFile(source) as archive:
        comment_part = next(name for name in archive.namelist() if name.startswith("xl/comments/") and name.endswith(".xml"))
    _replace_xml(source, comment_part, rich_text)
    catalog = load_tabular_source_catalog(source)
    body = json.loads(_metadata(catalog)[0].values[0])
    assert catalog.extraction_complete and body["text"] == "Первая Ω <&>\n последняя"
    assert len(ET.fromstring(body["literal_xml"]).findall(f"{{{_NS}}}text/{{{_NS}}}r")) == 3


@pytest.mark.parametrize("code,unread", [("&G", True), ("&&G", False), ("&&&G", True)])
def test_picture_token_stays_unread_but_literal_escaped_ampersand_does_not(tmp_path, code, unread):
    workbook, sheet, source = _source(tmp_path)
    sheet.oddHeader.left.text = "Источник " + code
    catalog = _save(workbook, source)
    assert catalog.extraction_complete is not unread
    assert any(item.endswith(":picture") for item in catalog.extraction_warnings) is unread
    assert code in _metadata(catalog)[0].line
    assert catalog.work_locators() == ["Лист Ω!R2"]


@pytest.mark.parametrize("unknown", ["child", "attribute", "direct_text", "tail"])
def test_unknown_print_structure_remains_visible_and_unread(tmp_path, unknown):
    workbook, sheet, source = _source(tmp_path)
    sheet.oddFooter.right.text = "Полный footer &P"
    _save(workbook, source)
    def mutate(node):
        container = node.find(f"{{{_NS}}}headerFooter")
        if unknown == "child":
            ET.SubElement(container, "{urn:unhandled}unknown").text = "Неразобранный текст"
        elif unknown == "attribute":
            container.set("{urn:unhandled}unknown", "Неразобранный атрибут")
        elif unknown == "direct_text":
            container.text = "Неразобранный прямой текст"
        else:
            list(container)[0].tail = "Неразобранный хвост"
    _replace_xml(source, "xl/worksheets/sheet1.xml", mutate)
    catalog = load_tabular_source_catalog(source)
    assert not catalog.extraction_complete
    assert "xlsx_unread_content:Лист Ω:headerFooter_unsupported" in catalog.extraction_warnings
    assert "Неразобранный" in _metadata(catalog)[0].line


def test_unknown_comment_extension_keeps_warning_and_full_literal_text(tmp_path):
    workbook, sheet, source = _source(tmp_path, comment=True)
    _save(workbook, source)
    with ZipFile(source) as archive:
        comment_part = next(name for name in archive.namelist() if name.startswith("xl/comments/") and name.endswith(".xml"))
    _replace_xml(source, comment_part, lambda node: ET.SubElement(node, "{urn:unsupported}extension"))
    catalog = load_tabular_source_catalog(source)
    assert not catalog.extraction_complete
    assert "xlsx_unread_content:Лист Ω:cell_comments_unsupported" in catalog.extraction_warnings
    body = json.loads(_metadata(catalog)[0].values[0])
    assert body["text"] == "Полный текст\n30 мм & <слой> Ω"


@pytest.mark.parametrize("placement", ["root_text", "authors_text", "author_tail", "comment_tail"])
def test_unknown_comment_container_text_or_tail_is_preserved_and_never_complete(tmp_path, placement):
    workbook, sheet, source = _source(tmp_path, comment=True)
    _save(workbook, source)
    with ZipFile(source) as archive:
        part = next(name for name in archive.namelist() if name.startswith("xl/comments/") and name.endswith(".xml"))
    marker = "UNREAD_LITERAL_Ω"
    def mutate(node):
        if placement == "root_text":
            node.text = marker
        elif placement == "authors_text":
            node.find(f"{{{_NS}}}authors").text = marker
        elif placement == "author_tail":
            node.find(f"{{{_NS}}}authors/{{{_NS}}}author").tail = marker
        else:
            node.find(f"{{{_NS}}}commentList/{{{_NS}}}comment").tail = marker
    _replace_xml(source, part, mutate)
    catalog = load_tabular_source_catalog(source)
    assert catalog is not None and not catalog.extraction_complete
    assert "xlsx_unread_content:Лист Ω:cell_comments_unsupported" in catalog.extraction_warnings
    bodies = [json.loads(row.values[0]) for row in _metadata(catalog)]
    unsupported = next(body for body in bodies if body["kind"] == "xlsx_comments_unsupported")
    assert marker in unsupported["literal_xml"]
    assert next(body["text"] for body in bodies if body["kind"] == "xlsx_cell_comment") == "Полный текст\n30 мм & <слой> Ω"
    assert catalog.work_locators() == ["Лист Ω!R2"]


def test_threaded_comment_relationship_stays_unread_without_classic_comments(tmp_path):
    workbook, sheet, source = _source(tmp_path)
    sheet.oddFooter.right.text = "Литерал &P"
    _save(workbook, source)
    with ZipFile(source) as archive:
        members = [(item, archive.read(item.filename)) for item in archive.infolist()]
    relation_xml = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="thread1" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" '
        'Target="../threadedComments/threadedComment1.xml"/></Relationships>')
    with ZipFile(source, "w") as archive:
        for item, data in members:
            archive.writestr(item, data)
        archive.writestr("xl/worksheets/_rels/sheet1.xml.rels", relation_xml)
        archive.writestr("xl/threadedComments/threadedComment1.xml", "<unread>Контекст Ω</unread>")
    catalog = load_tabular_source_catalog(source)
    assert catalog is not None and not catalog.extraction_complete
    assert "xlsx_unread_content:Лист Ω:threaded_comments" in catalog.extraction_warnings
    assert catalog.work_locators() == ["Лист Ω!R2"]
    assert not any(json.loads(row.values[0])["kind"] == "xlsx_cell_comment" for row in _metadata(catalog))


def test_comment_only_hidden_sheet_context_reaches_existing_main_sheet_packet(tmp_path):
    workbook, sheet, source = _source(tmp_path)
    notes = workbook.create_sheet("Примечания Ω")
    notes.sheet_state = "hidden"
    notes["D8"].comment = Comment("Отдельное исходное условие", "Автор")
    catalog = _save(workbook, source)
    metadata = _metadata(catalog)
    assert len(metadata) == 1 and metadata[0].sheet == "Примечания Ω"
    assert metadata[0].line in catalog_sheet_pack_lines(catalog, "Лист Ω", ["Лист Ω!R2"])
    frame = build_model_source_frame(catalog, ["Лист Ω!R2"])
    assert metadata[0].locator in frame.reference_locators
    assert metadata[0].locator not in frame.requested_locators
    assert catalog.work_locators() == ["Лист Ω!R2"]


def test_header_picture_relationship_without_text_is_still_unread(tmp_path):
    workbook, sheet, source = _source(tmp_path)
    _save(workbook, source)
    def picture(node):
        drawing = ET.SubElement(node, f"{{{_NS}}}legacyDrawingHF")
        drawing.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "hfPicture")
    _replace_xml(source, "xl/worksheets/sheet1.xml", picture)
    catalog = load_tabular_source_catalog(source)
    assert not catalog.extraction_complete
    assert "xlsx_unread_content:Лист Ω:header_footer_pictures" in catalog.extraction_warnings


def test_empty_print_formatting_keeps_the_old_cell_only_catalog(tmp_path):
    workbook, sheet, source = _source(tmp_path)
    sheet.oddHeader.center.font = "Arial,Bold"
    sheet.oddHeader.center.size = 12
    sheet.firstFooter.left.text = ""
    catalog = _save(workbook, source)
    assert catalog.extraction_complete and not catalog.extraction_warnings
    assert not _metadata(catalog)
    assert [row.locator for row in catalog.rows] == ["Лист Ω!R1", "Лист Ω!R2"]


def test_existing_typed_reader_returns_the_complete_metadata_body_and_source_identity(tmp_path):
    from proxy.services.model_research_tool_service import ModelResearchToolService
    workbook, sheet, source = _source(tmp_path, comment=True)
    sheet.oddFooter.right.text = "Литерал Ω &P / &N"
    catalog = _save(workbook, source)
    def forbidden(*args, **kwargs):
        raise AssertionError("No retrieval, fallback, or external call expected")
    reader = ModelResearchToolService(retrieve=forbidden, frozen_dataset_ids=(), retrieval_kwargs={},
        fallback=forbidden, attachment_catalog=catalog, expected_attachment_sha256=catalog.source_file_sha256)
    metadata = _metadata(catalog)
    assert len(metadata) == 2
    for row in metadata:
        result = reader._read_attachment_source({"source_row": row.locator})
        assert result.payload["status"] == "ok"
        assert result.payload["result"]["source_sha256"] == catalog.source_file_sha256
        assert result.payload["result"]["body_complete"] is True
        assert result.chunks[0].content == row.line
        assert result.payload["warnings"] == []
    assert reader._read_attachment_source({"source_row": "Лист Ω!META:missing"}).payload["status"] != "ok"
