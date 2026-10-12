"""Unknown source consumption survives the actual SQLite and workbook path."""
import asyncio
import json
import sqlite3
from contextlib import closing

import openpyxl
import pandas as pd
import pytest

from proxy.services import gesn_service as gs, fgis_price_service, rim_lsr_trace_service as rim
from proxy.services import fsem_machinist_service as fsem
from proxy.services.lsr_workbook_adapter_service import build_lsr_workbook_from_decisions
from proxy.services.rim_trace_xlsx_service import render_trace_xlsx
from proxy.smeta_core.calculation_receipt import calculation_receipts
from proxy.smeta_core.norm_browser import _code_listing
from tools.build_smeta_structured_base import build_structured_base
from tools.gesn_import import RESOURCE_FIELDS


@pytest.fixture(autouse=True)
def isolated_rates(monkeypatch):
    monkeypatch.setattr(rim.nr_sp_service, 'resolve', lambda *a, **k:
        dict(status='resolved_fixture', nr_pct=0, sp_pct=0))
    for loader in (gs.load_base_norms, gs.load_structured_base_norms, gs.load_norms):
        loader.cache_clear()
    yield
    for loader in (gs.load_base_norms, gs.load_structured_base_norms, gs.load_norms):
        loader.cache_clear()


def make_base(tmp_path, monkeypatch, value):
    records = []
    for code, amount in [('UNKNOWN', value), ('KNOWN', 1)]:
        row = dict.fromkeys(RESOURCE_FIELDS)
        row.update(norm_code='ГЭСН01-01-001-01', base_type='ГЭСН', norm_key='ГЭСН:01-01-001-01',
            norm_name='Synthetic operation', norm_unit='шт', kind='material',
            resource_code=code, resource_name='Synthetic ' + code, resource_unit='шт',
            per_unit=amount, price=10, work_steps=json.dumps(['Synthetic source step']),
            source_doc='synthetic source', source_guid='synthetic-guid')
        records.append(row)
    # Object strings preserve invalid source text through the real parquet boundary.
    for row in records:
        if row['per_unit'] is not None:
            row['per_unit'] = str(row['per_unit'])
    source = tmp_path / 'source.parquet'
    pd.DataFrame(records, columns=list(RESOURCE_FIELDS)).to_parquet(source, index=False)
    base = tmp_path / 'base.sqlite'
    build_structured_base(source=source, out=base, manifest_out=tmp_path / 'base_manifest.json', edition='synthetic')
    monkeypatch.setenv('LES_SMETA_STRUCTURED_BASE', str(base))
    return base


@pytest.mark.parametrize('value,expected', [(None,None), ('',None), (float('nan'),None),
    ('invalid',None), ('inf',None), (0,0), ('0.00',0), ('0,00',0), (1,1)])
def test_source_consumption_survives_builder_reader_receipt_adapter(tmp_path, monkeypatch, value, expected):
    base = make_base(tmp_path, monkeypatch, value)
    with sqlite3.connect(base) as conn:
        stored, = conn.execute("SELECT per_unit FROM resources WHERE resource_code='UNKNOWN'").fetchone()
    assert stored == expected
    card, = _code_listing('ГЭСН', '01-01-001-01', base_path=base)
    assert next(r for r in card['resources'] if r['code']=='UNKNOWN')['per_unit'] == expected
    assert next(r for r in gs.get_norm('ГЭСН01-01-001-01')['resources'] if r['code']=='UNKNOWN')['per_unit'] == expected
    decision = dict(source_row='Synthetic!1', work_id='mount', title='Synthetic operation',
        norm_code='ГЭСН01-01-001-01', quantity=2, unit='шт')
    trace = rim.build_lsr_trace_from_visible_rows([decision])
    receipt, = calculation_receipts(trace)
    (tmp_path/'calculation-receipt.json').write_text(json.dumps(receipt, ensure_ascii=False,
        indent=2, allow_nan=False), encoding='utf-8')
    resource = next(r for r in receipt['calculation']['resources'] if r['code']=='UNKNOWN')
    assert resource['per_unit'] == expected
    assert resource['quantity'] == (None if expected is None else expected*2)
    assert resource['known_cost'] == (None if expected is None else expected*20)
    summary = trace['summary']
    assert summary['known_amount'] == 20 + (0 if expected is None else expected*20)
    assert summary['full_amount'] == (None if expected is None else summary['known_amount'])
    assert summary['result_status'] == ('priced_partial' if expected is None else 'priced_final')
    monkeypatch.setattr(fgis_price_service, 'resolve_pricebook_path', lambda *_: None)
    target = tmp_path / 'actual-adapter.xlsx'
    result = asyncio.run(build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        dict(decisions=[decision]), target, lambda *_: None))
    book = openpyxl.load_workbook(target, data_only=False)
    row = next(r for r in book.worksheets[0] if r[1].value=='UNKNOWN')
    assert row[4].value == expected
    assert row[6].value == (None if expected is None else expected*2)
    if expected is None:
        assert row[11].value is None
        assert any('UNKNOWN' in flag for flag in result['missing'])
        assert 'расход' in ' '.join(str(c.value) for ws in book for rr in ws for c in rr if c.value).lower()
    else:
        assert str(row[11].value).startswith('=ROUND(')
    book.close()


@pytest.mark.parametrize('value', [None, '', 'invalid', float('nan'), float('inf')])
def test_manual_unknown_consumption_is_not_a_zero_cost(value, tmp_path):
    trace = rim.build_position_trace(dict(name='Synthetic', qty=2, nr_pct=0, sp_pct=0,
        resources=[dict(kind='material', code='UNKNOWN', name='Synthetic resource', unit='шт', per_unit=value, price=10)]))
    row = next(r for r in trace['rows'] if r['type']=='resource_material')
    assert row['columns'].get('5') is None
    assert row['columns'].get('7') is None
    assert row['columns'].get('12') is None
    assert trace['summary']['full_amount'] is None
    out = render_trace_xlsx(trace, tmp_path/'manual.xlsx')
    book = openpyxl.load_workbook(out)
    resource = next(r for r in book.active if r[1].value=='UNKNOWN')
    assert resource[11].value is None
    book.close()


def test_fsem_unknown_hours_cannot_be_reconciled_as_zero(monkeypatch):
    monkeypatch.setattr(fsem, 'lookup', lambda *a: dict(crew_hours=1, driver_code='DRIVER', driver_grade=3))
    source = [dict(kind='machine', code='M', per_unit=None), dict(kind='machinist', code='', per_unit=0)]
    resources, trace = fsem.enrich_machinists(source, quantity_field='per_unit')
    assert trace['status']=='unresolved'
    assert resources==source
    assert any('unknown' in reason for reason in trace['reasons'])


@pytest.mark.parametrize('source', [
    [dict(kind='machine', code='M', per_unit=0), dict(kind='machinist', code='', per_unit=None)],
    [dict(kind='machine', code='M', per_unit=0, excluded=True, excluded_per_unit=None),
        dict(kind='machinist', code='', per_unit=0)],
])
def test_fsem_unknown_aggregate_or_excluded_hours_are_unresolved(monkeypatch, source):
    monkeypatch.setattr(fsem, 'lookup', lambda *a: dict(crew_hours=1, driver_code='DRIVER', driver_grade=3))
    resources, trace = fsem.enrich_machinists(source, quantity_field='per_unit')
    assert resources==source and trace['status']=='unresolved'
    assert any('unknown' in reason for reason in trace['reasons'])


def test_null_and_zero_same_identity_remain_separate_source_resources(tmp_path, monkeypatch):
    base = make_base(tmp_path, monkeypatch, None)
    source = tmp_path/'source.parquet'
    records = pd.read_parquet(source)
    records.loc[1, ['resource_code', 'resource_name']] = records.loc[0, ['resource_code', 'resource_name']].values
    records.loc[1, 'per_unit'] = '0'
    records.to_parquet(source, index=False)
    build_structured_base(source=source, out=base, manifest_out=tmp_path/'base_manifest.json', edition='synthetic')
    with sqlite3.connect(base) as conn:
        rows = conn.execute('SELECT resource_code,per_unit FROM resources ORDER BY id').fetchall()
    assert rows==[('UNKNOWN',None),('UNKNOWN',0)]


@pytest.mark.parametrize('price', ['invalid', float('nan'), float('inf')])
def test_invalid_manual_price_cannot_be_a_known_zero(price):
    trace = rim.build_position_trace(dict(name='Synthetic', qty=2, nr_pct=0, sp_pct=0,
        resources=[dict(kind='material', code='UNKNOWN_PRICE', name='Synthetic resource',
            unit='шт', per_unit=1, price=price)]))
    row = next(r for r in trace['rows'] if r['type']=='resource_material')
    assert row['columns'].get('10') is None
    assert row['columns'].get('12') is None
    assert trace['summary']['full_amount'] is None


@pytest.mark.parametrize('kind,summary_key,group', [('labor','labor_qty','group_labor'),
    ('machinist','machinist_qty','group_machinist')])
def test_unknown_hours_cannot_be_a_known_group_zero(kind, summary_key, group, tmp_path):
    position = dict(name='Synthetic', qty=2, nr_pct=0, sp_pct=0,
        resources=[dict(kind=kind, code='HOURS', name='Synthetic hours', unit='чел.-ч', per_unit=None, price=10)])
    trace = rim.build_position_trace(position)
    assert trace['summary'][summary_key] is None
    assert rim.build_lsr_trace([position])['summary'][summary_key] is None
    assert next(row for row in trace['rows'] if row['type']==group)['columns'].get('7') is None
    out = render_trace_xlsx(trace, tmp_path/'unknown-hours.xlsx')
    book = openpyxl.load_workbook(out)
    header = next(row for row in book.active if str(row[0].value or '').strip().startswith(
        'нормативные затраты труда ' + ('рабочих' if kind=='labor' else 'машинистов')))
    assert header[3].value is None
    group_row = next(row for row in book.active if row[2].value==('ОТ(ЗТ)' if kind=='labor' else 'ОТм(ЗТм)'))
    assert group_row[6].value is None
    book.close()


def test_legacy_not_null_base_zero_is_not_reinterpreted(tmp_path, monkeypatch):
    base = make_base(tmp_path, monkeypatch, 0)
    with closing(sqlite3.connect(base)) as conn, conn:
        schema, = conn.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='resources'").fetchone()
        schema = schema.replace('per_unit REAL,', 'per_unit REAL NOT NULL,')
        conn.execute('ALTER TABLE resources RENAME TO old_resources')
        conn.execute(schema)
        conn.execute('INSERT INTO resources SELECT * FROM old_resources')
        assert next(row for row in conn.execute('PRAGMA table_info(resources)') if row[1]=='per_unit')[3]==1
    norm = gs.get_norm('ГЭСН01-01-001-01')
    assert next(r for r in norm['resources'] if r['code']=='UNKNOWN')['per_unit']==0
    trace = rim.build_lsr_trace([dict(code='ГЭСН01-01-001-01',qty=2,nr_pct=0,sp_pct=0)])
    assert trace['summary']['known_amount']==20
    card, = _code_listing('ГЭСН', '01-01-001-01', base_path=base)
    assert next(r for r in card['resources'] if r['code']=='UNKNOWN')['per_unit']==0


@pytest.mark.parametrize('value', [None, 0])
def test_model_exclusion_preserves_original_unknown_quantity(tmp_path, monkeypatch, value):
    make_base(tmp_path, monkeypatch, value)
    work = dict(source_row='Synthetic!1', work_id='mount', title='Synthetic operation',
        norm_code='ГЭСН01-01-001-01', quantity=2, unit='шт')
    removal = dict(source_row='Synthetic!1', work_id='remove', title='Model removal',
        decision_kind='ресурс', included_in='mount', resource_action='excluded',
        target_resource_code='UNKNOWN', quantity=2, unit='шт', replaced_by='not_required',
        coverage='Model explicitly confirms this material is not required')
    trace = rim.build_lsr_trace_from_visible_rows([work, removal])
    binding = trace['row_bindings'][1]
    assert binding['status']=='resource_excluded'
    assert binding['resource_bindings'][0]['excluded_quantity']==value
    receipts = calculation_receipts(trace)
    assert receipts[1]['resource_bindings'][0]['excluded_quantity']==value
    material = next(row for row in receipts[0]['calculation']['resources'] if row['code']=='UNKNOWN')
    assert material['quantity']==0 and material['known_cost']==0
    monkeypatch.setattr(fgis_price_service, 'resolve_pricebook_path', lambda *_: None)
    out = tmp_path/'model-exclusion.xlsx'
    asyncio.run(build_lsr_workbook_from_decisions(tmp_path/'source.xlsx',
        dict(decisions=[work, removal]), out, lambda *_: None))
    book = openpyxl.load_workbook(out)
    review = ' '.join(str(c.value) for row in book['Проверка'] for c in row if c.value)
    assert ('исходное количество снятия: неизвестно' if value is None else 'исходное количество снятия: 0') in review
    (tmp_path/'exclusion-receipts.json').write_text(json.dumps(receipts, ensure_ascii=False,
        indent=2, allow_nan=False), encoding='utf-8')
    book.close()


@pytest.mark.parametrize('value,quantity,status', [(None,1,'resource_quantity_unknown'),
    (None,0,'resource_quantity_unknown'), (0,1,'resource_quantity_exceeded'), (0,0,'resource_bound')])
def test_unknown_available_resource_is_not_a_known_zero(tmp_path, monkeypatch, value, quantity, status):
    make_base(tmp_path, monkeypatch, value)
    work = dict(source_row='Synthetic!1', work_id='mount', title='Synthetic operation',
        norm_code='ГЭСН01-01-001-01', quantity=2, unit='шт')
    allocation = dict(source_row='Synthetic!2', work_id='allocation', title='Model allocation',
        decision_kind='ресурс', included_in='mount', resource_action='included',
        target_resource_code='UNKNOWN', quantity=quantity, unit='шт')
    trace = rim.build_lsr_trace_from_visible_rows([work, allocation])
    assert trace['row_bindings'][1]['status']==status
    if value is None:
        assert 'расход' in trace['row_bindings'][1]['reason']
