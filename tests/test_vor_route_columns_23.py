# -*- coding: utf-8
"""§1 (23.09): the real «спецификация → ВОР» path must emit an actual ВОР sheet that
keeps every source position (one output row per input row, no grouping/loss) and marks
any installation work merely *implied* by an item as requiring a decision — a material
or equipment item is not proof of a монтаж."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from proxy.services.bor_service import source_rows_to_vor_xlsx
from proxy.services.spec_to_bor_service import derive_work_for_item


def test_derive_work_for_item_flags_implied_work_only():
    # an item that is not already a work gets a proposal and must be decided
    work, needs = derive_work_for_item("Светильник LED")
    assert needs is True
    assert work and "Светильник LED" in work
    # an explicit work is carried as-is, needs no decision
    work, needs = derive_work_for_item("Монтаж шкафа")
    assert needs is False
    assert work == "Монтаж шкафа"
    # noise names (short / numbered section headers) yield nothing
    assert derive_work_for_item("1. Раздел") == ("", False)
    assert derive_work_for_item("шт") == ("", False)
    assert derive_work_for_item("") == ("", False)


def test_vor_sheet_preserves_positions_and_marks_decisions(tmp_path: Path):
    rows = [
        {"name": "Светильник LED", "unit": "шт", "qty": 24.0,
         "source_file": "spec.xlsx", "pos": "1",
         "derived_work": "Установка Светильник LED", "requires_decision": True},
        {"name": "Монтаж шкафа", "unit": "шт", "qty": 4.0,
         "source_file": "spec.xlsx", "pos": "2",
         "derived_work": "Монтаж шкафа", "requires_decision": False},
        {"name": "Кабель ВВГнг", "unit": "м", "qty": 100.0,
         "source_file": "spec.xlsx", "pos": "3"},
    ]
    out = tmp_path / "vor.xlsx"
    n = source_rows_to_vor_xlsx(rows, out, title="ВОР")
    assert n == len(rows)

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["ВОР"]
    sheet = wb["ВОР"]
    header = [c.value for c in sheet[3]]
    assert "Предполагаемая работа" in header and "Статус решения" in header
    i_status = header.index("Статус решения")
    i_name = header.index("Наименование")
    i_qty = header.index("Кол-во")

    data = [[c.value for c in r] for r in sheet.iter_rows(min_row=4)
            if any(c.value is not None for c in r)]
    # one output row per input row — no source position is lost
    assert len(data) == 3
    assert [row[i_name] for row in data] == ["Светильник LED", "Монтаж шкафа", "Кабель ВВГнг"]
    assert [row[i_qty] for row in data] == [24.0, 4.0, 100.0]
    # implied work flagged; explicit work confirmed; un-annotated row left blank (not invented)
    assert data[0][i_status] == "ТРЕБУЕТ РЕШЕНИЯ"
    assert data[1][i_status] == "подтверждено"
    assert data[2][i_status] in (None, "")
    wb.close()


def test_explicit_empty_vor_decisions_do_not_select_the_historical_spec_reader(tmp_path: Path):
    import asyncio

    from proxy.services.workbook_tool_service import _default_vor_adapter
    from proxy.services.spec_to_bor_service import rows_from_spec_xlsx

    spec = tmp_path / "spec.xlsx"
    wb = openpyxl.Workbook()
    sh = wb.active
    sh.title = "Спецификация"
    sh.append(["№", "Наименование", "Ед. изм.", "Кол-во"])
    sh.append([1, "Светильник LED", "шт", 24])
    sh.append([2, "Монтаж шкафа", "шт", 4])
    wb.save(spec)
    wb.close()

    out = tmp_path / "vor.xlsx"

    async def _run():
        return await _default_vor_adapter(
            spec, {"decisions": []}, out, lambda *a, **k: None,
        )

    result = asyncio.run(_run())
    assert result["source_rows"] == 0
    assert result["requires_decision_count"] == 0
    assert result["blockers"] == ["MODEL_WORK_ROWS_EMPTY"]

    wb2 = openpyxl.load_workbook(out)
    assert wb2.sheetnames == ["ВОР", "Проверка"]
    assert wb2["ВОР"].max_row == 3
    wb2.close()
    # Historical specification utilities remain available when deliberately
    # called; their rows must not replace the model's explicit empty decision.
    historical = rows_from_spec_xlsx(spec)
    assert [row["name"] for row in historical] == ["Светильник LED", "Монтаж шкафа"]
    assert [row["qty"] for row in historical] == [24, 4]
