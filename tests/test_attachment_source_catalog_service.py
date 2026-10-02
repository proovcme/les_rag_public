from pathlib import Path

import pytest

from proxy.services import attachment_source_catalog_service as service


def _write_sheet(tmp_path, rows, *, title="Лист1"):
    from openpyxl import Workbook

    path = tmp_path / "ВОР header strip.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    return path


@pytest.mark.parametrize("source_qty,source_unit,model_qty,model_unit,expected", [
    (800, "м", 0.8, "км", []),
    (0.8, "км", 800, "м", []),
    (250, "кг", 0.25, "т", []),
    (800, "м", 0.9, "км", ["source_quantity_mismatch:Лист1!R2"]),
])
def test_source_comparison_accepts_only_equal_metric_quantities(
    tmp_path, source_qty, source_unit, model_qty, model_unit, expected,
):
    source = _write_sheet(tmp_path, [["Наименование", "Ед. изм.", "Кол-во"],
                                    ["Ресурс", source_unit, source_qty]])
    catalog = service.load_tabular_source_catalog(source)
    row = {"source_row": "Лист1!R2", "quantity": model_qty, "unit": model_unit}
    assert service.compare_model_row_to_catalog(row, catalog) == expected
    assert row["quantity"] == model_qty and row["unit"] == model_unit
    assert catalog.lookup("Лист1!R2").header_quantity == source_qty


def test_header_strip_naming_a_total_column_binds_the_quantity(tmp_path):
    """A сметная header row whose money column is named «(итого)» is a header.

    Regression for the customer's ВОР: the column labels
    «Материал (итого), руб. без НДС» / «Стоимость (итого)…» made the *header row
    itself* match the subtotal marker, and the subtotal rule ran first.  The whole
    workbook then had no header row at all, so nothing bound: the quantity column
    stayed unknown and every work row fell back to "any number in the row" -- the
    item number («1.2»), which turned the model's correct quantity into
    ``source_quantity_mismatch`` on every row.
    """

    catalog = service.load_tabular_source_catalog(
        _write_sheet(
            tmp_path,
            [
                ["№ п/п", "Наименование работ", "Ед.изм.", "Кол-во",
                 "Материал (итого), руб. без НДС", "Примечание"],
                ["1.2", "Демонтаж перегородок из газоблоков", "м2", 98.6,
                 "=E11*D11", None],
            ],
        )
    )
    header = catalog.lookup("Лист1!R1")
    assert header.row_classification == "header"
    assert header.classification_reason == "column_header_marker"

    row = catalog.lookup("Лист1!R2")
    assert row.quantity_column_known is True
    assert row.header_quantity == 98.6
    # the item number is no longer offered as the quantity of the work row
    assert service.quantity_candidates(row) == (98.6,)
    # the labels travel as data, so a consumer can tell «Кол-во» from a money column
    assert "Кол-во" in [label for _column, label in row.column_headers]
    assert "98.6 [Кол-во]" in row.line

    # the model's own quantity against that source: accepted, and a wrong one refused
    assert service.compare_model_row_to_catalog(
        {"source_row": "Лист1!R2", "quantity": 98.6, "unit": "м2"}, catalog
    ) == []
    assert service.compare_model_row_to_catalog(
        {"source_row": "Лист1!R2", "quantity": 1.2, "unit": "м2"}, catalog
    ) == ["source_quantity_mismatch:Лист1!R2"]


def test_subtotal_rows_keep_their_class(tmp_path):
    """Looking for the header strip first does not re-label a real total line.

    A total either carries a resolved number, or its money cells are formulas --
    which the shape test rejects, because a column strip is a row of *names*.  The
    last assertion is the one that matters: a total caption that mentions a column
    name must not become the governing header for the rows below it, or those rows
    lose their quantity column again.
    """

    catalog = service.load_tabular_source_catalog(
        _write_sheet(
            tmp_path,
            [
                ["№ п/п", "Наименование работ", "Ед.изм.", "Кол-во"],
                ["1.1", "Монтаж перегородки", "м2", 41],
                ["ИТОГО по смете, руб. с НДС", 125000, None, None],
                ["Итого по разделу «Наименование работ»", "=SUBTOTAL(9,D9:D40)", None, None],
                ["1.2", "Монтаж облицовки", "м2", 87.6, None, None],
            ],
        )
    )
    assert catalog.lookup("Лист1!R3").row_classification == "subtotal"
    assert catalog.lookup("Лист1!R4").row_classification == "subtotal"
    below = catalog.lookup("Лист1!R5")
    assert below.row_classification == "work"
    assert below.quantity_column_known is True
    assert below.header_quantity == 87.6
    assert service.quantity_candidates(below) == (87.6,)


def test_numberless_prose_row_is_not_read_as_a_header_strip(tmp_path):
    """A single text cell that happens to contain a column word stays a note."""

    catalog = service.load_tabular_source_catalog(
        _write_sheet(
            tmp_path,
            [
                ["№ п/п", "Наименование работ", "Ед.изм.", "Кол-во"],
                ["Работа выполнена по наименованию из проекта"],
            ],
        )
    )
    assert catalog.lookup("Лист1!R1").row_classification == "header"
    assert catalog.lookup("Лист1!R2").row_classification == "note"


def test_pdf_markdown_catalog_has_stable_locators_and_only_work_rows(monkeypatch, tmp_path):
    source = tmp_path / "ВОР монтаж БАП.pdf"
    source.write_bytes(b"%PDF-fixture")
    markdown = """Ведомость объемов работ

| № пп | Наименование | Ед. изм. | Кол-во | Примечание |
| --- | --- | --- | --- | --- |
| 1 | 2 | 3 | 4 | 5 |
| Раздел 1. Демонтажные работы |  |  |  |  |
| 1 | Защитное укрытие | м² | 116 | Пленка |
| Раздел 2. Монтажные работы |  |  |  |  |
| 1 | Монтаж блока аварийного питания | шт. | 16 |  |
"""
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)

    assert catalog is not None
    assert catalog.sheets == (("PDF", 6), ("PDF_TEXT", 1))
    assert catalog.work_locators() == ["PDF!R4", "PDF!R6", "PDF_TEXT!R1"]
    assert catalog.lookup("PDF_TEXT!R1").row_classification == "ambiguous"
    assert catalog.lookup("PDF!R6").section_context.startswith("Раздел 2")
    assert catalog.lookup("PDF!R2").row_classification == "header"


def test_markdown_table_parser_ignores_separator_and_prose():
    assert service._markdown_table_rows(
        "prose\n| A | B |\n| :--- | ---: |\n| x | 2 |\n"
    ) == [("A", "B"), ("x", "2")]


def test_mixed_pdf_keeps_non_table_conditions_after_the_chat_preview(monkeypatch, tmp_path):
    source = tmp_path / "Mixed source.pdf"
    source.write_bytes(b"%PDF-fixture")
    before = "Literal source introduction\n" * 800
    after = "Late literal condition: keep and reinstall existing equipment."
    markdown = before + "\n| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n| Ресурс | м | 7 |\n\n" + after
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)

    assert catalog.lookup("PDF!R2").header_quantity == 7
    text_rows = [row for row in catalog.rows if row.sheet == "PDF_TEXT"]
    assert text_rows
    assert any(after in row.values[0] for row in text_rows)
    assert "Literal source introduction" in "\n".join(row.values[0] for row in text_rows)
    assert all(row.row_classification == "ambiguous" for row in text_rows)
    assert all(row.classification_reason == "pdf_text_requires_model_role" for row in text_rows)
    assert all(row.locator in catalog.work_locators() for row in text_rows)
    assert catalog.lookup("PDF!R2").numbers == (7.0,)
    assert any(error.startswith("source_coverage_incomplete") for error in service.coverage_errors(
        [{"source_row": "PDF!R2"}], catalog,
    ))


def test_pdf_without_tables_preserves_model_owned_text_scope(monkeypatch, tmp_path):
    source = tmp_path / "Text source.pdf"
    source.write_bytes(b"%PDF-fixture")
    markdown = "Literal first condition.\n\nLiteral unquantified operation.\n\n123\n"
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)

    assert catalog is not None
    assert all(row.sheet == "PDF_TEXT" for row in catalog.rows)
    assert "Literal unquantified operation." in "\n".join(row.values[0] for row in catalog.rows)
    assert all(row.row_classification == "ambiguous" for row in catalog.rows)
    assert all(service.quantity_candidates(row) == () for row in catalog.rows)
    assert all(row.header_quantity is None for row in catalog.rows)


def test_pdf_text_blocks_are_literal_bounded_and_old_table_ids_are_stable(monkeypatch, tmp_path):
    source = tmp_path / "Long literal source.pdf"
    source.write_bytes(b"%PDF-fixture")
    before = "  First literal line.  \n" + "X" * 20000 + "\n"
    after = "  Final condition with original spacing.  \n"
    markdown = before + "| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n| Ресурс | м | 7 |\n" + after
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)
    again = service.load_tabular_source_catalog(source)
    text_rows = [row for row in catalog.rows if row.sheet == "PDF_TEXT"]

    assert "".join(row.values[0] for row in text_rows) == before + after
    assert all(len(row.values[0]) <= 4096 for row in text_rows)
    assert [row.locator for row in text_rows] == [f"PDF_TEXT!R{index}" for index in range(1, len(text_rows) + 1)]
    assert [row.row_id for row in catalog.rows] == [row.row_id for row in again.rows]
    assert catalog.lookup("PDF!R1").row_classification == "header"
    assert catalog.lookup("PDF!R2").header_quantity == 7
    assert catalog.lookup("PDF!R3") is None


def test_pdf_repeated_tables_retain_original_order_and_literal_note_offsets(monkeypatch, tmp_path):
    source = tmp_path / "Ordered repeated tables.pdf"
    source.write_bytes(b"%PDF-fixture")
    header = "| Наименование | Ед. изм. | Кол-во |\n"
    work = "| Ресурс | м | 7 |\n"
    table = header + "|---|---|---|\n" + work
    first_note = "First table: condition A.\n"
    second_note = "Second table: condition B.\n"
    markdown = table + first_note + table + second_note
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)
    ordered = ["PDF!R1", "PDF!R2", "PDF_TEXT!R1", "PDF!R3", "PDF!R4", "PDF_TEXT!R2"]
    assert catalog.interleaved_row_lines() == [catalog.lookup(locator).line for locator in ordered]
    for locator, raw_line, start in (
        ("PDF!R1", header, 0),
        ("PDF!R2", work, markdown.index(work)),
        ("PDF_TEXT!R1", first_note, markdown.index(first_note)),
        ("PDF!R3", header, markdown.rindex(header)),
        ("PDF!R4", work, markdown.rindex(work)),
        ("PDF_TEXT!R2", second_note, markdown.index(second_note)),
    ):
        row = catalog.lookup(locator)
        assert row.source_start == start
        assert row.source_end == start + len(raw_line)
        assert markdown[row.source_start:row.source_end] == raw_line
    assert catalog.lookup("PDF!R2").header_quantity == catalog.lookup("PDF!R4").header_quantity == 7
    facts = service.catalog_sheet_packet_facts(catalog, "PDF", ["PDF!R2", "PDF!R4"])
    assert facts["source_order_basis"] == "converter_markdown_char_offsets"
    assert [item["source_row"] for item in facts["source_order"]] == ordered
    assert service.catalog_sheet_pack_lines(catalog, "PDF", ["PDF!R2", "PDF!R4"]) == catalog.interleaved_row_lines()


def test_pdf_late_text_blocks_keep_raw_offsets_after_table_and_bounded_splits(monkeypatch, tmp_path):
    source = tmp_path / "Late offset source.pdf"
    source.write_bytes(b"%PDF-fixture")
    table = "| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n| Ресурс | м | 7 |\n"
    prose = "Before the table.\n"
    late = "X" * 20000 + "\nFinal literal condition.\n"
    markdown = prose + table + late
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)
    text_rows = [row for row in catalog.rows if row.sheet == "PDF_TEXT"]
    assert "".join(markdown[row.source_start:row.source_end] for row in text_rows) == prose + late
    assert all(markdown[row.source_start:row.source_end] == row.values[0] for row in text_rows)
    assert all(len(row.values[0]) <= 4096 for row in text_rows)
    assert text_rows[-1].source_end == len(markdown)


def test_non_pdf_catalog_retains_sheet_interleave_without_pdf_order_fields(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / "Two sheets.xlsx"
    book = Workbook()
    book.active.title = "A"
    book.active.append(["Resource A", "m", 7])
    book.create_sheet("B").append(["Resource B", "m", 9])
    book.save(source)
    catalog = service.load_tabular_source_catalog(source)
    assert catalog.interleaved_row_lines() == [catalog.lookup("A!R1").line, catalog.lookup("B!R1").line]
    facts = service.catalog_sheet_packet_facts(catalog, "A", ["A!R1"])
    assert set(facts) == {
        "source_extraction_complete", "source_extraction_warnings", "requested_source_rows",
        "reference_source_rows", "reference_only_source_rows", "source_quantity_cells", "source_scope_instruction",
    }


@pytest.mark.parametrize("large", [False, True])
def test_pdf_cross_sheet_context_is_bounded_and_unshown_text_remains_its_own_packet(
    monkeypatch, tmp_path, large,
):
    source = tmp_path / "Packet source.pdf"
    source.write_bytes(b"%PDF-fixture")
    prose = ("Literal condition.\n" * (1000 if large else 2)) + "Late exact condition."
    markdown = "| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n| Ресурс | м | 7 |\n" + prose
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)
    text_rows = [row for row in catalog.rows if row.sheet == "PDF_TEXT"]
    lines = service.catalog_sheet_pack_lines(catalog, "PDF", ["PDF!R2"])
    facts = service.catalog_sheet_packet_facts(catalog, "PDF", ["PDF!R2"])

    assert facts["requested_source_rows"] == ["PDF!R2"]
    assert facts["source_extraction_complete"] is True
    assert facts["source_context_complete"] is (not large)
    if large:
        assert len("\n".join(lines)) < 8192
        assert all(row.line not in lines for row in text_rows)
        assert facts["reference_only_source_rows"] == []
        assert facts["reference_omitted_source_count"] == len(text_rows)
        assert facts["reference_omitted_source_locator_range"] == {
            "sheet": "PDF_TEXT", "first": text_rows[0].locator,
            "last": text_rows[-1].locator, "count": len(text_rows), "content_status": "not_shown",
        }
        assert "model_context_partial" in "\n".join(lines)
        assert text_rows[-1].locator in "\n".join(lines)
    else:
        assert all(row.line in lines for row in text_rows)
        assert facts["reference_only_source_rows"] == [row.locator for row in text_rows]
        assert facts["reference_omitted_source_count"] == 0
        assert facts["reference_omitted_source_locator_range"] is None
    own_lines = service.catalog_sheet_pack_lines(catalog, "PDF_TEXT", [text_rows[-1].locator])
    own_facts = service.catalog_sheet_packet_facts(catalog, "PDF_TEXT", [text_rows[-1].locator])
    assert text_rows[-1].line in own_lines
    assert "Late exact condition." in text_rows[-1].values[0]
    assert own_facts["requested_source_rows"] == [text_rows[-1].locator]
    assert own_facts["source_context_complete"] is True
    assert own_facts["reference_omitted_source_count"] == 0


@pytest.mark.parametrize("unread", [
    "![Unread diagram](source.png)",
    "![Unread diagram][source]",
    "<svg><text>Diagram annotation</text></svg>",
    "[WARN] Source.pdf: текст не извлечён (сканированный PDF?)",
    "[Ошибка распознавания страницы: provider unavailable]",
])
def test_pdf_converter_unread_signals_do_not_claim_complete_extraction(monkeypatch, tmp_path, unread):
    source = tmp_path / "Unread content.pdf"
    source.write_bytes(b"%PDF-fixture")
    markdown = "| Наименование | Ед. изм. | Кол-во |\n|---|---|---|\n| Ресурс | м | 7 |\n\n" + unread
    monkeypatch.setattr("backend.converter.convert_to_markdown", lambda _path, **_kwargs: markdown)

    catalog = service.load_tabular_source_catalog(source)

    assert catalog.extraction_complete is False
    assert catalog.extraction_warnings
    assert "source_extraction_incomplete" in service.coverage_errors(
        [{"source_row": row.locator} for row in catalog.rows], catalog,
    )
    assert unread in "\n".join(row.values[0] for row in catalog.rows if row.sheet == "PDF_TEXT")


def test_packet_keeps_unquantified_group_description_beside_variant(tmp_path):
    catalog = service.load_tabular_source_catalog(_write_sheet(tmp_path, [
        ["№ п/п", "Наименование", "Ед. изм.", "Кол-во"],
        ["5.1", "Труба гибкая гофрированная из ПВХ", "", ""],
        ["", "- Ø20 мм", "м", 200],
        ["", "- Ø25 мм", "м", 30],
        ["5.2", "Труба стальная", "", ""],
        ["", "- Ø40 мм", "м", 50],
    ]))
    parent = catalog.lookup("Лист1!R2")
    assert parent.is_work_row() and parent.quantity_column_known and parent.header_quantity is None
    packet = service.catalog_sheet_pack_lines(catalog, "Лист1", remaining_work=["Лист1!R4"])
    assert parent.line in packet
    assert catalog.lookup("Лист1!R4").line in packet
    assert catalog.lookup("Лист1!R3").line not in packet
    assert catalog.lookup("Лист1!R6").line not in packet
    # Source context does not invent a quantity or change the coverage denominator.
    assert service.quantity_candidates(parent) == ()
    assert "Лист1!R2" in catalog.work_locators()
    facts = service.catalog_sheet_packet_facts(catalog, "Лист1", ["Лист1!R4"])
    assert facts["requested_source_rows"] == ["Лист1!R4"]
    assert "Лист1!R2" in facts["reference_source_rows"]
    values = {item["source_row"]: item for item in facts["source_quantity_cells"]}
    assert values["Лист1!R4"]["quantity"] == 30
    assert values["Лист1!R4"]["unit"] == "м"
    assert values["Лист1!R2"]["quantity"] is None


def test_unknown_quantity_layout_does_not_expose_every_other_work_as_context(tmp_path):
    catalog = service.load_tabular_source_catalog(_write_sheet(tmp_path, [
        ["1.1", "Монтаж первого изделия", "шт", 10],
        ["1.2", "Монтаж другого изделия", "шт", 20],
    ]))
    assert not catalog.lookup("Лист1!R1").quantity_column_known
    packet = service.catalog_sheet_pack_lines(catalog, "Лист1", remaining_work=["Лист1!R2"])
    assert catalog.lookup("Лист1!R1").line not in packet
    assert catalog.lookup("Лист1!R2").line in packet
    facts = service.catalog_sheet_packet_facts(catalog, "Лист1", ["Лист1!R2"])
    assert facts["source_quantity_cells"][0]["quantity_column_known"] is False
    assert "quantity" not in facts["source_quantity_cells"][0]


def test_csv_empty_cells_keep_quantity_under_its_header(tmp_path):
    import csv

    for delimiter in (";", ","):
        source = tmp_path / ("Разреженная спецификация " + str(ord(delimiter)) + ".csv")
        with source.open("w", encoding="utf-8", newline="") as stream:
            writer = csv.writer(stream, delimiter=delimiter)
            writer.writerow(["№", "Наименование", "Тип", "Ед. изм.", "Количество", "Масса, кг"])
            writer.writerow([1, "Монтаж изделия", "", "шт.", 25, 0.35])
            writer.writerow(["", "", "", "", "", ""])
            writer.writerow([2, "Монтаж изделия Б", "", "шт.", 14, ""])
        catalog = service.load_tabular_source_catalog(source)
        first = catalog.lookup("CSV!R2")
        assert service.quantity_candidates(first) == (25.0,)
        assert service.quantity_candidates(catalog.lookup("CSV!R4")) == (14.0,)
        assert catalog.lookup("CSV!R3") is None
        assert service.compare_model_row_to_catalog(
            {"source_row": "CSV!R2", "quantity": 25, "unit": "шт."}, catalog
        ) == []
        assert "source_quantity_mismatch:CSV!R2" in service.compare_model_row_to_catalog(
            {"source_row": "CSV!R2", "quantity": 0.35, "unit": "шт."}, catalog
        )
