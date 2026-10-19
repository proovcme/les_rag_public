"""Formal book measures prorate a source quote, never source/model quantities."""
from copy import deepcopy
from decimal import Decimal
import math

from openpyxl import load_workbook
import pytest

from proxy.services import lsr_assembly_service as assembly
from proxy.services import rim_lsr_trace_service as rim
from proxy.services.rim_trace_xlsx_service import render_trace_xlsx, render_lsr_xlsx


class Book:
    region = "Owned region"
    quarter = "Owned period"
    source_name = "owned.parquet"
    source_sha256 = "owned-snapshot-sha"

    def __init__(self, record):
        self.record = record

    def lookup(self, code):
        return self.record if code == "owned-resource" else None


def both(record, resource_unit, quantity, *, work_qty=1, **fields):
    resource = {"kind": "material", "code": "owned-resource", "name": "Owned resource",
                "unit": resource_unit, **fields}
    position = {"name": "Authored work", "qty": work_qty, "nr_pct": 0, "sp_pct": 0}
    inputs = deepcopy((record, resource, position))
    traced = rim.build_position_trace(
        {**position, "resources": [{**resource, "per_unit": quantity}]}, pricebook=Book(record))
    assembled = assembly.compute_position(
        {**position, "resources": [{**resource, "qty": quantity * work_qty}]}, pricebook=Book(record))
    row = next(row for row in traced["rows"] if row["type"] == "resource_material")
    assert (record, resource, position) == inputs
    assert row["columns"]["4"] == resource_unit
    assert row["columns"]["7"] == pytest.approx(quantity * work_qty)
    assert assembled["resources"][0]["unit"] == resource_unit
    assert assembled["resources"][0]["qty"] == pytest.approx(quantity * work_qty)
    return traced, row, assembled


@pytest.mark.parametrize("quantity", [84, 1_000_000, -84, 0])
def test_authoritative_effective_quote_is_prorated_before_final_amount_rounding(quantity):
    record = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49,
              "price_current_eff": 134243.44, "source_ref": "owned-book#row"}
    trace, row, assembled = both(record, "м", quantity)
    expected = float((Decimal("134243.44") * Decimal(quantity) / Decimal(1000)).quantize(Decimal("0.01")))
    assert row["columns"]["10"] == pytest.approx(134.24344)
    assert row["columns"]["12"] == expected
    assert trace["summary"]["total"] == assembled["total"] == expected
    assert assembled["resources"][0]["price_used"] == pytest.approx(134.24344)
    for conversion in (row["meta"]["price_unit_conversion"],
                       assembled["resources"][0]["price_unit_conversion"]):
        assert conversion["book_current_price"] == 134243.44
        assert conversion["quote_field"] == "price_current_eff"
        assert conversion["conversion_factor"] == .001
        assert conversion["source_price_fields"]["price_base"] == 90096.27
        assert conversion["price_owner"] == "chosen_pricebook"
        assert conversion["book_source_sha256"] == "owned-snapshot-sha"
        assert conversion["source_ref"] == "owned-book#row"


@pytest.mark.parametrize("book_unit, resource_unit, quantity, quoted_price, expected", [
    ("100 м3", "м3", 5, 123.45, 6.17),
    ("м3", "100 м3", .05, 123.45, 617.25),
    ("кг", "100 т", .00002, 10, 20),
    ("1000 м", "км", .084, 134243.44, 11276.45),
    ("0,001 т", "кг", 2, 10, 20),
    ("100 м²", "м2", 2, 37.5, .75),
])
def test_formal_measures_and_si_are_applied_once(book_unit, resource_unit, quantity, quoted_price, expected):
    trace, row, assembled = both({"unit": book_unit, "price_current_eff": quoted_price}, resource_unit, quantity)
    assert row["columns"]["12"] == trace["summary"]["total"] == assembled["total"] == expected
    assert row["meta"]["price_unit_conversion"]["conversion_kind"] == "numeric_measure"
    assert assembled["price_requirements"] == []


def test_equivalent_physical_quantities_and_existing_norm_scale_keep_one_amount():
    record = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49, "price_current_eff": 134243.44}
    metres, metre_row, _ = both(record, "м", 84)
    scaled, scaled_row, _ = both(record, "1000 м", .084)
    normalized, normalized_row, _ = both(record, "м", 42, work_qty=2)
    assert metres["summary"]["total"] == scaled["summary"]["total"] == normalized["summary"]["total"] == 11276.45
    assert metre_row["columns"]["7"] == normalized_row["columns"]["7"] == 84
    assert scaled_row["columns"]["7"] == .084


def test_scalar_conversion_is_reversible_without_reconstructing_book_quote():
    original = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49, "price_current_eff": 134243.44}
    before = deepcopy(original)
    forward, _ = rim.convert_book_price_unit(original, "м", pricebook=Book(original))
    backward, trace = rim.convert_book_price_unit(forward, "1000 м", pricebook=Book(forward))
    assert backward["price_current_eff"] == pytest.approx(original["price_current_eff"])
    assert trace["quote_field"] == "price_current_eff"
    assert original == before


def test_missing_current_quote_rounds_book_base_index_before_conversion():
    record = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49}
    _, row, assembled = both(record, "м", 1_000_000)
    assert row["meta"]["price_unit_conversion"]["quote_field"] == "round_book_base_x_index_2"
    assert row["meta"]["price_unit_conversion"]["book_current_price"] == 134243.44
    assert row["columns"]["12"] == assembled["total"] == 134243440.0


@pytest.mark.parametrize("book_unit,resource_unit", [
    ("1000 кг", "л"), ("100 м3", "м2"), ("1000 м", "м кабеля"),
    ("1000 м", "упаковка"), ("0 м", "м"), ("-100 м", "м"), ("1e3 м", "м"),
])
def test_incompatible_noisy_and_nonpositive_measures_remain_gaps(book_unit, resource_unit):
    _, row, assembled = both({"unit": book_unit, "price_current_eff": 10}, resource_unit, 1)
    assert row["columns"].get("10") is None
    assert row["columns"].get("12") is None
    assert assembled["resources"][0]["price_used"] is None
    assert "price_unit_conversion" not in row["meta"]


@pytest.mark.parametrize("record,target_unit", [
    ({"unit": "м", "price_current_eff": 1e308}, "1000 м"),
    ({"unit": "1000 м", "price_current_eff": 1e-323}, "м"),
    ({"unit": "1000 м", "price_current_eff": float("inf")}, "м"),
    ({"unit": "1000 м", "price_current_eff": 10, "price_conflict": {"code": "AMBIGUOUS_PRICE_RECORD"}}, "м"),
])
def test_overflow_underflow_and_source_conflict_do_not_acquire_conversion(record, target_unit):
    before = deepcopy(record)
    unchanged, conversion = rim.convert_book_price_unit(record, target_unit, pricebook=Book(record))
    assert unchanged is record
    assert conversion is None
    assert record == before


def test_scalar_zero_quote_is_not_a_missing_price():
    _, row, assembled = both({"unit": "1000 м", "price_current_eff": 0}, "м", 84)
    assert row["columns"]["10"] == row["columns"]["12"] == assembled["total"] == 0
    assert assembled["price_requirements"] == []


@pytest.mark.parametrize("price, fields", [(5, {"price_confirmed": True}),
    (0, {"price_source_ref": "engineer#reuse", "resource_binding": {"action": "reuse"}})])
def test_manual_and_reuse_money_remain_owned_and_unconverted(price, fields):
    _, row, assembled = both({"unit": "1000 м", "price_current_eff": 10}, "м", 84, price=price, **fields)
    assert row["columns"]["10"] == assembled["resources"][0]["price_used"] == price
    assert "price_unit_conversion" not in row["meta"]


def test_legacy_single_si_unit_rounding_is_unchanged():
    _, row, assembled = both({"unit": "кг", "price_current_eff": 2.34567}, "т", 1)
    assert row["columns"]["10"] == assembled["resources"][0]["price_used"] == 2345.67
    assert "conversion_kind" not in row["meta"]["price_unit_conversion"]


@pytest.mark.parametrize("quantity,expected", [(84, 11276.45), (1_000_000, 134243440.0)])
def test_scalar_price_xlsx_formula_uses_source_quote_not_unrounded_base_product(tmp_path, quantity, expected):
    record = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49, "price_current_eff": 134243.44}
    trace, _, _ = both(record, "м", quantity)
    before = deepcopy(trace)
    output = render_trace_xlsx(trace, tmp_path / "scalar quote.xlsx")
    formulas = load_workbook(output, data_only=False).worksheets[0]
    values = load_workbook(output, data_only=True).worksheets[0]
    row = next(line[0].row for line in formulas if line[1].value == "owned-resource")
    assert formulas.cell(row, 10).value == "=134243.44*0.001"
    assert formulas.cell(row, 12).value == f"=ROUND(G{row}*J{row}*K{row},2)"
    assert values.cell(row, 10).value == pytest.approx(134.24344)
    assert values.cell(row, 12).value == expected
    assert values.cell(row, 7).value == quantity and values.cell(row, 4).value == "м"
    assert formulas.cell(row, 10).number_format != "#,##0.00"
    assert "134243.44" in str(values.cell(row, 3).value)
    assert "1000 м" in str(values.cell(row, 3).value)
    assert "134.24344" in str(values.cell(row, 3).value)
    assert values.row_dimensions[row].height >= 30
    assert trace == before


def test_scalar_renderer_rejects_non_numeric_formula_inputs(tmp_path):
    trace, row, _ = both({"unit": "1000 м", "price_current_eff": 10}, "м", 1)
    row["meta"]["price_unit_conversion"]["book_current_price"] = "=HYPERLINK(1)"
    with pytest.raises(ValueError, match="numeric-measure"):
        render_trace_xlsx(trace, tmp_path / "invalid quote.xlsx")


def test_identical_numeric_measures_keep_the_legacy_quote_priority():
    record = {"unit": "1000 м", "price_base": 2, "index": 1.2, "price_current_eff": 999}
    _, row, assembled = both(record, "1000 м", 1)
    # An inconsistent synthetic record exposes each pre-existing priority:
    # RIM uses base x index; the direct assembly API uses the effective quote.
    # An identical measure must not enter the new conversion branch at all.
    assert row["columns"]["10"] == 2.4
    assert assembled["resources"][0]["price_used"] == 999
    assert "price_unit_conversion" not in row["meta"]
    assert "price_unit_conversion" not in assembled["resources"][0]


def test_current_quote_takes_precedence_over_base_product_and_effective_fallback():
    record = {"unit": "1000 м", "price_current": 100000.01,
              "price_base": 90096.27, "index": 1.49, "price_current_eff": 134243.44}
    _, row, assembled = both(record, "м", 1_000_000)
    assert row["meta"]["price_unit_conversion"]["quote_field"] == "price_current"
    assert row["columns"]["10"] == pytest.approx(100.00001)
    assert row["columns"]["12"] == assembled["total"] == 100000010.0


def test_unresolved_family_is_not_converted_or_replaced_by_a_member(monkeypatch):
    group = {"state": "unresolved_resource_group", "group": "owned-resource", "total": 1,
             "members": [{"code": "must-not-select", "unit": "1000 м", "price_current_eff": 10}]}
    monkeypatch.setattr(Book, "resource_group", lambda self, code: deepcopy(group), raising=False)
    _, row, assembled = both(None, "м", 84)
    assert row["meta"]["price_action"] == "unresolved_resource_group"
    assert row["meta"]["price_resource_candidates"] == group["members"]
    assert row["columns"]["2"] == assembled["resources"][0]["code"] == "owned-resource"
    assert row["columns"].get("10") is None and row["columns"].get("12") is None
    assert assembled["resources"][0]["price_used"] is None
    assert assembled["price_requirements"][0]["action"] == "unresolved_resource_group"
    assert "price_unit_conversion" not in row["meta"]


def test_actual_visible_supply_route_preserves_authored_fields_and_prices_without_model_calc_unit(tmp_path):
    record = {"unit": "1000 м", "price_base": 90096.27, "index": 1.49, "price_current_eff": 134243.44}
    authored = {"source_row": "S!R7", "row_role": "material", "title": "Authored supply",
                "resource_code": "owned-resource", "norm_code": "", "unit": "м", "quantity": 84,
                "evidence_refs": ["Q1.H1"], "evidence_bindings": {"Q1.H1": "owned-card"},
                "_resource_evidence_check": {"code_in_refs": True}}
    before = deepcopy(authored)
    lsr = rim.build_lsr_trace_from_visible_rows([authored], pricebook=Book(record))
    position = lsr["sections"][0]["positions"][0]
    row = next(row for row in position["rows"] if row["type"] == "resource_material")
    binding = lsr["row_bindings"][0]
    assert authored == before and "calc_unit" not in authored
    assert binding["position"]["source_quantity"] == 84
    assert binding["position"]["source_unit"] == "м"
    assert binding["position"]["code"] == "" and binding["row_role"] == "material"
    assert row["columns"]["2"] == "owned-resource"
    assert row["columns"]["4"] == "м" and row["columns"]["7"] == 84
    assert row["columns"]["10"] == pytest.approx(134.24344)
    assert row["columns"]["12"] == lsr["summary"]["total"] == 11276.45
    assert binding["resource_decisions"][0]["price_unit_conversion"]["book_current_price"] == 134243.44
    snapshot = deepcopy(lsr)
    output = render_lsr_xlsx(lsr, tmp_path / "visible scalar supply.xlsx")
    formulas = load_workbook(output, data_only=False).worksheets[0]
    values = load_workbook(output, data_only=True).worksheets[0]
    number = next(line[0].row for line in formulas if line[1].value == "owned-resource")
    assert formulas.cell(number, 10).value == "=134243.44*0.001"
    assert values.cell(number, 10).value == pytest.approx(134.24344)
    assert values.cell(number, 12).value == 11276.45
    assert lsr == snapshot and authored == before



@pytest.mark.parametrize("quantity,expected", [(1, 0.02), (-1, -0.02)])
def test_numeric_measure_cent_tie_matches_excel_round_in_both_calculators_and_xlsx(tmp_path, quantity, expected):
    record = {"unit": "1000 м", "price_current_eff": 15}
    trace, row, assembled = both(record, "м", quantity)
    assert row["columns"]["10"] == assembled["resources"][0]["price_used"] == 0.015
    assert row["columns"]["12"] == trace["summary"]["total"] == assembled["total"] == expected
    assert assembled["resources"][0]["cost"] == expected
    assert row["columns"]["7"] == quantity
    before = deepcopy(trace)
    output = render_trace_xlsx(trace, tmp_path / ("cent-tie-" + str(quantity) + ".xlsx"))
    formulas = load_workbook(output, read_only=True, data_only=False)
    values = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet, cached = formulas.worksheets[0], values.worksheets[0]
        number = next(line[0].row for line in sheet if line[1].value == "owned-resource")
        assert sheet.cell(number, 10).value == "=15.0*0.001"
        assert sheet.cell(number, 12).value == f"=ROUND(G{number}*J{number}*K{number},2)"
        assert cached.cell(number, 10).value == 0.015
        assert cached.cell(number, 12).value == expected
        assert cached.cell(number, 7).value == quantity
    finally:
        formulas.close()
        values.close()
    assert trace == before


@pytest.mark.parametrize("quoted,expected", [(0.015, 0.02), (-0.015, -0.02)])
def test_numeric_measure_missing_quote_rounds_book_base_index_half_away_from_zero(quoted, expected):
    trace, row, assembled = both({"unit": "1000 м", "price_base": quoted, "index": 1}, "м", 1000)
    conversion = row["meta"]["price_unit_conversion"]
    assert conversion["quote_field"] == "round_book_base_x_index_2"
    assert conversion["book_current_price"] == expected
    assert row["columns"]["12"] == trace["summary"]["total"] == assembled["total"] == expected


@pytest.mark.parametrize("invalid", [float("nan"), float("inf"), -float("inf")])
def test_numeric_measure_rounding_helper_rejects_nonfinite_values(invalid):
    with pytest.raises(ValueError, match="finite"):
        rim.round_numeric_measure_amount(invalid)


def test_numeric_measure_rounding_handles_finite_float_boundary_without_decimal_overflow():
    assert math.isfinite(rim.round_numeric_measure_amount(1e308))
    assert rim.round_numeric_measure_amount(-1e308) == -1e308
